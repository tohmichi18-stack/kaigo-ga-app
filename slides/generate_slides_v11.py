"""主任ケアマネ研修スライド生成スクリプト 第11弾

テーマ: ヤングケアラー・ビジネスケアラー支援とケアマネの役割
重点: 改正子ども・若者育成支援推進法を踏まえた多機関連携の実務
トーン: 実務型(気づき・連携・場面対応)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(家族・世代をつなぐ:明るいブルー基調 + 温かいコーラル) =====
BLUE        = RGBColor(0x20, 0x68, 0xA4)   # メイン
BLUE_MID    = RGBColor(0x3E, 0x83, 0xBA)
BLUE_LT     = RGBColor(0x76, 0xA5, 0xC8)
CORAL       = RGBColor(0xD9, 0x6E, 0x4C)   # アクセント(人・行動)
GREEN       = RGBColor(0x3C, 0x80, 0x5A)   # 支援・OK
AMBER       = RGBColor(0xC9, 0x97, 0x33)   # 注意
PLUM        = RGBColor(0x7A, 0x4A, 0x6E)   # 子ども・若者
RED         = RGBColor(0xB3, 0x44, 0x39)   # リスク・NG
LIGHT_BG    = RGBColor(0xEF, 0xF3, 0xF6)
MID_GRAY    = RGBColor(0xD3, 0xDA, 0xE0)
DARK_GRAY   = RGBColor(0x2D, 0x32, 0x37)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xE6, 0xED, 0xF1)
RED_BG      = RGBColor(0xF8, 0xE6, 0xE4)
AMBER_BG    = RGBColor(0xFA, 0xF1, 0xDB)
GREEN_BG    = RGBColor(0xE4, 0xED, 0xE7)
CORAL_BG    = RGBColor(0xFA, 0xEA, 0xE2)
PLUM_BG     = RGBColor(0xEF, 0xE6, 0xEC)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), BLUE)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=BLUE)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), CORAL)


def add_footer(slide, page_num, total=27):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(27), Cm(0.5),
             "主任CM研修 / ヤングケアラー・ビジネスケアラー支援とケアマネの役割",
             size=9, color=BLUE_LT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=BLUE_LT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=BLUE_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=BLUE, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=CORAL):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, BLUE)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=accent_color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=33, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color, fontsize=15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(fontsize)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, BLUE)
add_rect(s, 0, Cm(8), SW, Cm(0.15), CORAL)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), BLUE_LT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.4), SW - Cm(2), Cm(2.5),
         "ヤングケアラー・ビジネスケアラー\n支援とケアマネの役割",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
add_text(s, Cm(1.5), Cm(9.7), SW - Cm(3), Cm(2),
         "ー 改正子ども・若者育成支援推進法を踏まえた多機関連携の実務 ー",
         size=15, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(13), SW - Cm(3), Cm(2),
         "■ ヤングケアラー・ビジネスケアラーとは ― 用語と背景\n"
         "■ 改正子ども・若者育成支援推進法のポイント\n"
         "■ 気づき、つなぎ、両立を支える ― 多機関連携の実務",
         size=14, color=AMBER_BG, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "ヤングケアラー・ビジネスケアラーという言葉と、その背景を理解する",
    "改正子ども・若者育成支援推進法のポイントと、支援の方向性を押さえる",
    "ヤングケアラーの「見えにくさ」と、CMが気づける立場を理解する",
    "ビジネスケアラー(働く介護者)の両立を、ケアプランで支える視点を持つ",
    "高齢者本人の担当であるCMが、世帯の中の家族介護者をどう支えるかを考える",
    "教育・福祉・労働など多機関と連携する実務と、その留意点を持ち帰る",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), CORAL)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "ヤングケアラー・ビジネスケアラーとは",  "10分"),
    ("2", "改正子ども・若者育成支援推進法",          "15分"),
    ("3", "ヤングケアラーへの気づきと支援",          "20分"),
    ("4", "ビジネスケアラー(働く介護者)への支援",  "20分"),
    ("5", "多機関連携の実務",                        "15分"),
    ("6", "まとめ ― 主任CMが持ち帰るもの",           "10分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, BLUE_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=BLUE)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=CORAL, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "ヤングケアラー・\nビジネスケアラーとは",
              "Young Carers and Business (Working) Carers")


# =============================================================
#  SLIDE 5: 用語の整理
# =============================================================
s = add_slide()
add_header(s, "1. 用語と背景", "2つの「ケアラー」 ― 言葉を整理する")
add_footer(s, 5)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(7),
        "ヤングケアラー",
        "■ 本来は大人が担うようなケア責任を引き受け、\n   家事や家族の世話・介護などを\n   日常的に行っている子ども・若者\n\n"
        "■ 高齢の祖父母の介護、障害・疾患のある\n   家族の世話、きょうだいの世話、家事 など\n\n"
        "■ 学業・友人関係・進路・心身の発達に\n   影響が及ぶことがある\n\n"
        "▶ 「子どもらしい時間」が奪われうる",
        title_bg=PLUM, body_bg=PLUM_BG, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(7),
        "ビジネスケアラー(働く介護者)",
        "■ 仕事をしながら、家族の介護を\n   担っている人\n\n"
        "■ 主に「働き盛り」の世代\n   (子育てと重なる「ダブルケア」も)\n\n"
        "■ 仕事と介護の両立が難しく、\n   介護離職に至るリスク\n\n"
        "▶ 本人のキャリア・収入、企業の人材、\n   社会全体に関わる課題",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11.7), Cm(31.4), Cm(5.3),
        "共通するのは「家族介護者」への支援という視点",
        "■ どちらも「家族を介護・ケアする人(ケアラー)」 ― 介護を受ける本人とは別に、支援が要る存在\n\n"
        "■ 居宅CMは「高齢者本人」を担当するが、その世帯には必ず「ケアを担う家族」がいる\n\n"
        "■ 「介護される人」だけでなく「介護する人」にも目を向けることが、これからのケアマネジメント",
        title_bg=BLUE, body_size=13)


# =============================================================
#  SLIDE 6: なぜ今注目されるのか
# =============================================================
s = add_slide()
add_header(s, "1. 用語と背景", "なぜ今、家族介護者の支援が注目されるのか")
add_footer(s, 6)

reasons = [
    ("家族のかたちの変化",
     "・少子化・核家族化\n・ひとり親世帯・共働き世帯\n・介護の担い手が限られる"),
    ("介護の長期化",
     "・高齢化で介護期間が長くなる\n・「老老介護」「多重介護」\n・1人に負担が集中しやすい"),
    ("ケアラーの不可視化",
     "・家族介護は「家のこと」とされ\n  外から見えにくい\n・本人も声をあげにくい"),
    ("社会・経済への影響",
     "・ヤングケアラーの学び・将来\n・介護離職による労働力の損失\n・社会全体で支える必要性"),
]
col_w = Cm(7.7)
col_h = Cm(5.8)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(4.5)
colors = [PLUM, BLUE_MID, AMBER, CORAL]
for i, (title, body) in enumerate(reasons):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.5), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(1.5),
             title, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.5), col_w, col_h - Cm(1.5), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(1.8), col_w - Cm(0.6), col_h - Cm(2.1),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(11.5), Cm(31.4), Cm(5.5),
        "「介護は家族が担うもの」という前提が、限界を迎えている",
        "■ かつては「家族がいれば家族が看る」が当たり前とされてきた\n\n"
        "■ しかし今、その家族自身が ― 子どもであったり、働く現役世代であったり ― 支えを必要としている\n\n"
        "■ 「介護を社会で支える」という介護保険の理念は、「介護する家族」も視野に入れて初めて完成する\n\n"
        "■ 国も、ヤングケアラー・ビジネスケアラーの支援を、重要な政策課題として位置づけ始めている",
        title_bg=BLUE, body_size=13)


# =============================================================
#  SLIDE 7: 第2部
# =============================================================
section_title(2, "改正子ども・若者\n育成支援推進法",
              "The Amended Act on Youth Development Support")


# =============================================================
#  SLIDE 8: 法改正の概要
# =============================================================
s = add_slide()
add_header(s, "2. 改正子ども・若者育成支援推進法", "法改正のポイント ― ヤングケアラーの明記")
add_footer(s, 8)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3.3),
        "ヤングケアラーが、法律上の支援対象として位置づけられた",
        "■ 子ども・若者育成支援推進法の改正により、ヤングケアラーが\n"
        "    国・地方公共団体が支援すべき対象として、法律上明確に位置づけられた\n"
        "■ これまで「定義のない」存在だったヤングケアラーが、施策の対象として明文化された意義は大きい",
        title_bg=BLUE, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(4.8),
        "改正で明確になったこと",
        "■ 「家族の介護その他の日常生活上の世話を\n   過度に行うと認められる子ども・若者」\n   が支援対象として位置づけられた\n\n"
        "■ 国・自治体が、その支援に努めることが\n   求められる",
        title_bg=BLUE_MID, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(4.8),
        "押さえておきたいこと",
        "■ 「ヤングケアラー」は、年齢で一律に\n   区切られるものではない\n\n"
        "■ 子どもだけでなく「若者」も対象\n\n"
        "■ 法律名・条文・施行時期等の詳細は\n   一次情報で確認すること",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.5),
         "▶ 「法律に位置づけられた」=「気づいて、つなぐべき対象」が明確になった、ということ。\n"
         "▶ CMにとっても、ヤングケアラーへの気づきが、より明確に求められるようになった。",
         size=13, color=BLUE, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 9: 国・自治体の動き
# =============================================================
s = add_slide()
add_header(s, "2. 改正子ども・若者育成支援推進法", "支援の方向性 ― 国・自治体の動き")
add_footer(s, 9)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「気づく・つなぐ・支える」体制づくりが進む",
        "ヤングケアラーは「自覚しにくく」「声をあげにくい」存在。\n"
        "だからこそ、社会の側が「気づき、見つけ、支援につなぐ」しくみが求められている。",
        title_bg=BLUE, body_size=12)

items = [
    ("実態把握・\n認知度の向上",
     "・実態調査の実施\n・社会全体への周知・啓発\n・「ヤングケアラー」の認知向上"),
    ("相談・支援体制",
     "・相談窓口の整備\n・コーディネーターの配置\n・ピアサポート・交流の場"),
    ("関係機関の連携",
     "・教育・福祉・医療・介護の連携\n・自治体での連携体制づくり\n・支援につなぐしくみ"),
]
col_w = Cm(10.2)
col_h = Cm(5.5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(7)
colors = [PLUM, GREEN, BLUE_MID]
for i, (title, body) in enumerate(items):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.5), colors[i])
    add_text(s, x, y, col_w, Cm(1.5),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.5), col_w, col_h - Cm(1.5), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.8), col_w - Cm(0.8), col_h - Cm(2.1),
             body, size=11.5, color=DARK_GRAY, line_spacing=1.45)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "CMが知っておきたいこと",
        "■ 多くの自治体に、ヤングケアラーの相談窓口・コーディネーターが置かれつつある\n"
        "■ 「自分の市町村ではどこが窓口か」を把握しておくことが、つなぐための第一歩\n"
        "■ 介護保険分野(CM)も、その連携ネットワークの一員として期待されている",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 10: 第3部
# =============================================================
section_title(3, "ヤングケアラーへの\n気づきと支援",
              "Noticing and Supporting Young Carers")


# =============================================================
#  SLIDE 11: ヤングケアラーの実態
# =============================================================
s = add_slide()
add_header(s, "3. ヤングケアラーへの支援", "ヤングケアラーが担っていること")
add_footer(s, 11)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「お手伝い」と「ヤングケアラー」の境目",
        "家事の手伝いそのものが問題なのではない。\n"
        "「過度な」「日常的な」ケアにより、子ども本人の生活・発達・将来が脅かされることが課題。",
        title_bg=BLUE, body_size=12)

cares = [
    ("家事を担う",
     "・買い物・料理・掃除・洗濯\n・家計の管理"),
    ("家族の世話・介護",
     "・祖父母の介護\n・障害・疾患のある家族の世話\n・通院の付き添い"),
    ("きょうだいの世話",
     "・幼いきょうだいの世話\n・送り迎え・身の回りの世話"),
    ("感情面の支え",
     "・家族の気持ちを支える\n・通訳(日本語・手話等)\n・見守り・声かけ"),
]
col_w = Cm(7.7)
col_h = Cm(5.5)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.2)
for i, (title, body) in enumerate(cares):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), PLUM)
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(1.3),
             title, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(1.6), col_w - Cm(0.6), col_h - Cm(1.9),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "本人が「ケアラー」と気づいていないことが多い",
        "■ 「家族だから当たり前」「ずっとそうだった」 ― 本人にとっては日常で、問題と思っていない\n"
        "■ 「助けて」と言えない/言わない ― 家族を悪く言いたくない、知られたくない\n"
        "▶ だからこそ、周りの大人が「気づく」ことが決定的に重要になる",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 12: サインの見つけ方
# =============================================================
s = add_slide()
add_header(s, "3. ヤングケアラーへの支援", "サインに気づく ― CMが出会う場面")
add_footer(s, 12)

signs = [
    ("訪問時に出会う様子",
     "・利用者宅で、子どもが介護や家事を\n  している場面に出会う\n・学校に行っている時間に家にいる\n・幼い子が高齢者の世話をしている"),
    ("家族・利用者の語り",
     "・「孫が手伝ってくれている」\n・「子どもがいるから大丈夫」\n・家族が子どもをあてにしている様子"),
    ("世帯の状況から",
     "・ほかに介護を担える大人がいない\n・ひとり親・共働き・家族の疾患\n・経済的な厳しさ"),
    ("子ども本人の様子",
     "・疲れている・元気がない\n・遅刻・欠席・学校の話を避ける\n・年齢に比して「しっかりしすぎ」"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.9)
colors = [BLUE_MID, GREEN, AMBER, PLUM]
for i, (title, body) in enumerate(signs):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=11.5, color=DARK_GRAY, line_spacing=1.4)

add_text(s, Cm(1.2), Cm(16.6), SW - Cm(2.4), Cm(0.8),
         "▶ CMは利用者宅に入る ― ヤングケアラーの存在に「最初に気づける」専門職の一人。",
         size=12, color=BLUE, bold=True)


# =============================================================
#  SLIDE 13: 気づいたときの対応
# =============================================================
s = add_slide()
add_header(s, "3. ヤングケアラーへの支援", "気づいたときの対応 ― CMの動き方")
add_footer(s, 13)

steps = [
    ("子どもを\n責めない・否定しない",
     "・「えらいね」だけで終わらせない\n・「家族思い」を否定もしない\n・本人の気持ちをそのまま受けとめる"),
    ("ケアプランで\n軽減を試みる",
     "・高齢者本人へのサービスを見直し、\n  子どもが担う負担を減らせないか検討\n・「家族の介護力」をあてにしすぎない"),
    ("一人で抱えず\n相談する",
     "・主任CM・地域包括に相談\n・子ども本人への支援はCMの範囲を超える\n・つなぐ判断を早めに"),
    ("適切な窓口に\nつなぐ",
     "・自治体のヤングケアラー相談窓口\n・学校・スクールソーシャルワーカー\n・子ども家庭・福祉の相談機関"),
]
y = Cm(4.4)
colors = [PLUM, BLUE_MID, AMBER, GREEN]
for i, (title, body) in enumerate(steps):
    add_numbered_circle(s, Cm(1.5), y + Cm(0.4), Cm(1.2), i + 1, colors[i])
    add_rect(s, Cm(3.2), y, Cm(7), Cm(2.2), colors[i])
    add_text(s, Cm(3.3), y, Cm(6.8), Cm(2.2),
             title, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(10.2), y, SW - Cm(11.4), Cm(2.2), TABLE_BG)
    add_text(s, Cm(10.6), y + Cm(0.2), SW - Cm(12), Cm(1.9),
             body, size=11.5, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
    y += Cm(2.45)

add_box(s, Cm(1.2), Cm(14.6), Cm(31.4), Cm(2.5),
        "CMだからできる、独自の貢献",
        "■ 多くの支援者は「子ども」から入るが、CMは「高齢者本人のケア」を通じて世帯に関われる\n"
        "■ 高齢者へのサービスを適切に組むこと自体が、子どもの負担を直接軽くする ― これがCMの強み",
        title_bg=BLUE, body_size=12)


# =============================================================
#  SLIDE 14: 第4部
# =============================================================
section_title(4, "ビジネスケアラー\n(働く介護者)への支援",
              "Supporting Working Carers")


# =============================================================
#  SLIDE 15: ビジネスケアラーと介護離職
# =============================================================
s = add_slide()
add_header(s, "4. ビジネスケアラーへの支援", "働きながら介護する ― 介護離職という課題")
add_footer(s, 15)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "「仕事を辞めて介護に専念」は、最善ではないことが多い",
        "介護のために仕事を辞めると、収入が途絶え、社会とのつながりも失われ、\n"
        "本人の心身の負担はかえって増し、「介護後」の生活再建も難しくなる。\n"
        "「介護離職を防ぎ、両立を支える」ことが、本人にとっても重要。",
        title_bg=BLUE, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(5),
        "介護離職がもたらすもの",
        "■ 収入の喪失・経済的な不安定\n\n"
        "■ キャリアの中断・再就職の困難\n\n"
        "■ 社会とのつながり・気分転換の喪失\n\n"
        "■ 介護に専念することによる閉塞・孤立\n\n"
        "■ 「介護後」の生活設計の困難",
        title_bg=RED, body_bg=RED_BG, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(5),
        "両立を支えることの意味",
        "■ 介護者本人の生活・人生を守る\n\n"
        "■ 介護を「終わり」でなく「続けられる」ものに\n\n"
        "■ 企業にとっても貴重な人材の損失を防ぐ\n\n"
        "■ 社会全体の労働力・活力を支える\n\n"
        "▶ 「両立支援」は本人・企業・社会の共通利益",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ CMは「介護を組み立てる専門職」 ― 家族の「働き続ける」を支える力を持っている。",
         size=13, color=BLUE, bold=True)


# =============================================================
#  SLIDE 16: 両立支援の制度
# =============================================================
s = add_slide()
add_header(s, "4. ビジネスケアラーへの支援", "仕事と介護の両立を支える制度を知る")
add_footer(s, 16)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "育児・介護休業法による両立支援のしくみ",
        "働きながら介護する人を支えるため、法律でさまざまな制度が定められている。\n"
        "CMが概要を知っておくと、家族に「制度の存在」を伝え、相談につなぐことができる。",
        title_bg=BLUE, body_size=12)

col_widths = [Cm(8.5), Cm(13.5), Cm(9.4)]
headers = ["制度", "概要(一般的な内容)", "ポイント"]
rows = [
    ["介護休業",
     "対象家族の介護のため、一定期間\n仕事を休める制度",
     "「介護に専念」でなく\n体制を整える期間に"],
    ["介護休暇",
     "介護のために、年単位で\n取得できる休暇",
     "通院付き添い等\n短時間の用事に"],
    ["勤務時間等の\n配慮",
     "短時間勤務・時差出勤、所定外・\n時間外労働の制限 等の措置",
     "「働き方」を\n調整して両立"],
    ["相談・情報提供",
     "介護に直面した労働者への\n情報提供・相談対応の充実",
     "近年、両立支援が\n強化されている"],
]
add_table(s, Cm(1.2), Cm(7.3), col_widths, headers, rows,
          header_bg=BLUE, first_col_color=BLUE,
          row_height=Cm(1.85), body_size=11)

add_box(s, Cm(1.2), Cm(15.3), Cm(31.4), Cm(2.2),
        "重要 ― 制度の詳細は専門窓口へ",
        "■ 制度の対象・日数・要件・手続きは法令で定められ、改正もあります ― CMは「概要を伝える」役\n"
        "■ 詳細は勤務先・労働局・専門の相談窓口へ ― CMが正確な制度説明を抱え込まないこと",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 17: CMができること
# =============================================================
s = add_slide()
add_header(s, "4. ビジネスケアラーへの支援", "CMができること ― ケアプランで両立を支える")
add_footer(s, 17)

ways = [
    ("働き方を前提に\nプランを組む",
     "・家族の勤務時間・働き方を把握する\n・「日中独居」を前提にサービスを組む\n・家族の介護負担を当て込みすぎない"),
    ("使える制度・\n資源を伝える",
     "・両立支援制度の存在を家族に伝える\n・「専門窓口に相談を」とつなぐ\n・地域のサービス・社会資源を紹介"),
    ("急変・緊急時に\n備える",
     "・家族が仕事中の緊急時の体制\n・連絡・対応のルールを決めておく\n・「仕事を抜けなくて済む」体制を"),
    ("家族の気持ちを\n支える",
     "・「両立のつらさ」を受けとめる\n・「辞めるしかない」と思い詰める前に相談を\n・罪悪感を一人で抱えさせない"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [BLUE_MID, GREEN, AMBER, CORAL]
for i, (title, body) in enumerate(ways):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "アセスメントで「家族の仕事」を必ず聴く",
        "■ 利用者本人の状態だけでなく「介護する家族が働いているか/どんな働き方か」を必ず把握する\n"
        "■ 「家族がいるから大丈夫」と安易に判断しない ― 家族の生活・仕事を犠牲にした計画は続かない\n"
        "■ 良いケアプランとは「本人を支え、かつ家族の人生も守る」もの",
        title_bg=BLUE, body_size=13)


# =============================================================
#  SLIDE 18: 第5部
# =============================================================
section_title(5, "多機関連携の実務",
              "Multi-Agency Collaboration in Practice")


# =============================================================
#  SLIDE 19: CMの立場と限界
# =============================================================
s = add_slide()
add_header(s, "5. 多機関連携の実務", "CMの立場と限界 ― だから「つなぐ」")
add_footer(s, 19)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "CMの担当は「高齢者本人」 ― しかし出会うのは「世帯」",
        "居宅CMの制度上の担当は、要介護高齢者本人。\n"
        "ヤングケアラー(子ども)やビジネスケアラー(家族)への直接的・継続的な支援は、\n"
        "CMの守備範囲を超える ― だからこそ「気づき、つなぐ」ことが役割になる。",
        title_bg=BLUE, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(4.8),
        "CMにできること",
        "■ 訪問・関わりを通じて気づく\n\n"
        "■ 高齢者へのケアを通じて家族の負担を軽減\n\n"
        "■ 適切な窓口・機関につなぐ\n\n"
        "■ 多機関連携の一員として関わり続ける",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(4.8),
        "CMだけでは担えないこと",
        "■ 子どもへの教育・福祉的な支援\n\n"
        "■ 家族の就労・労働問題への対応\n\n"
        "■ 世帯全体の課題の総合的な調整\n\n"
        "▶ 専門の機関と「役割を分担」して支える",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.5),
         "▶ 「自分の役割ではない」と素通りもせず、「全部抱え込む」こともせず。\n"
         "▶ 気づき、つなぎ、連携の一員であり続ける ― これがCMの現実的な貢献。",
         size=13, color=BLUE, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 20: 連携先
# =============================================================
s = add_slide()
add_header(s, "5. 多機関連携の実務", "連携先を知る ― つなぐ先のマップ")
add_footer(s, 20)

col_widths = [Cm(8), Cm(13), Cm(10.4)]
headers = ["分野", "主な連携先", "こんなときに"]
rows = [
    ["子ども・若者",
     "自治体のヤングケアラー相談窓口・\nコーディネーター、子ども家庭の相談機関",
     "ヤングケアラーに\n気づいたとき"],
    ["教育",
     "学校、スクールソーシャルワーカー、\nスクールカウンセラー",
     "子どもの学校生活に\n影響が見えるとき"],
    ["介護・高齢",
     "地域包括支援センター、\n市町村の介護保険担当",
     "高齢者支援・世帯の\n調整が必要なとき"],
    ["労働・就労",
     "勤務先(人事等)、労働局、\n仕事と介護の両立相談の窓口",
     "両立支援制度の\n相談が必要なとき"],
    ["福祉・横断",
     "重層的支援体制整備事業の窓口、\n生活困窮・障害福祉の相談支援",
     "世帯に複合課題が\nあるとき"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=BLUE, first_col_color=BLUE,
          row_height=Cm(2.0), body_size=10.5, header_size=12)

add_box(s, Cm(1.2), Cm(16), Cm(31.4), Cm(1.5),
        "",
        "▶ 「つなぐ先のリスト」を、平時から自分の地域版でつくっておく ― 困ってから探さない。",
        title_bg=WHITE, body_bg=CORAL_BG, title_size=1, body_size=12)


# =============================================================
#  SLIDE 21: 連携の進め方
# =============================================================
s = add_slide()
add_header(s, "5. 多機関連携の実務", "連携の進め方と、留意すること")
add_footer(s, 21)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.5),
        "連携を進めるステップ",
        "■ 気づいたことを「事実」として記録する\n\n"
        "■ 一人で抱えず、主任CM・包括に相談する\n\n"
        "■ 緊急性・リスクの程度を見立てる\n\n"
        "■ 適切な窓口・機関につなぐ\n\n"
        "■ つないだ後も、連携の一員として関わる\n\n"
        "■ 会議の場(支援会議等)を活用する",
        title_bg=BLUE_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.5),
        "個人情報・同意への留意",
        "■ 家族の情報を他機関と共有するときは、\n   本人・家族の同意を得るのが原則\n\n"
        "■ 「誰に・何のために・何を」共有するかを\n   丁寧に説明する\n\n"
        "■ ただし、子どもの安全に関わる緊急時は、\n   同意を待たず関係機関と連携する判断も\n\n"
        "■ 迷うときは一人で判断せず、組織・包括と",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "連携を「つなぎっぱなし」にしない",
        "■ 「つないだら終わり」ではない ― CMは高齢者支援を通じて、その世帯に関わり続ける\n\n"
        "■ 連携先と「誰が・何を担うか」を確認し、CMの役割を明確にする\n\n"
        "■ 高齢者の状態が変われば、家族の負担も変わる ― 変化を連携先と共有し続ける\n\n"
        "■ 「気づき → つなぐ → 協働する」の協働の部分まで、CMは責任を持つ",
        title_bg=BLUE, body_size=13)


# =============================================================
#  SLIDE 22: 第6部
# =============================================================
section_title(6, "まとめ", "Wrap Up")


# =============================================================
#  SLIDE 23: 主任CMの役割
# =============================================================
s = add_slide()
add_header(s, "Closing", "家族介護者支援における主任CMの役割")
add_footer(s, 23)

roles = [
    ("気づく目を育てる",
     "・後輩CMに、世帯・家族介護者を\n  見る視点を伝える\n・「介護する人」への気づきを促す"),
    ("つなぎ先を持つ",
     "・地域のヤングケアラー窓口・\n  両立支援の相談先を把握\n・「つなぐ先リスト」を整備・共有"),
    ("アセスメントを変える",
     "・「家族の仕事・生活」を必ず聴く\n・家族の介護力をあてにしすぎない\n・本人と家族、両方を支える計画に"),
    ("連携をリードする",
     "・多機関連携の一員として動く\n・地域ケア会議等で課題を共有\n・地域の連携体制づくりに関わる"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [PLUM, GREEN, BLUE_MID, CORAL]
for i, (title, body) in enumerate(roles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "「介護される人」と「介護する人」の、両方を視野に",
        "■ ケアマネジメントは、これまで「介護される高齢者本人」を中心に組み立てられてきた\n"
        "■ これからは「介護する家族」 ― 子ども・働く世代 ― も視野に入れることが求められる\n"
        "■ 主任CMは、この新しい視点を、事業所・地域の実践に広げていく担い手になる",
        title_bg=BLUE, body_size=13)


# =============================================================
#  SLIDE 24: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 24)

actions = [
    ("Action 1", "アセスメントに「家族介護者」の視点を加える",
     "・担当ケースで「誰が・どう介護を担っているか」「家族は働いているか」を確認する\n"
     "・世帯に子ども・若者がいないか、その役割に目を向ける"),
    ("Action 2", "地域の「つなぎ先」を調べておく",
     "・自治体のヤングケアラー相談窓口、仕事と介護の両立相談の窓口を調べる\n"
     "・自分の地域版「つなぎ先リスト」をつくり、事業所内で共有する"),
    ("Action 3", "後輩CMと「気づく視点」を共有する",
     "・事例検討等で、ヤングケアラー・ビジネスケアラーの視点を取り上げる\n"
     "・「介護する人」も支援の対象、という視点を事業所内に広げる"),
]
y = Cm(4.5)
colors = [PLUM, BLUE_MID, GREEN]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=24, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=15, bold=True, color=BLUE)
    add_text(s, Cm(6), y + Cm(1.3), SW - Cm(7.5), Cm(2.3),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 25: クロージング
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, BLUE)
add_rect(s, 0, Cm(7), SW, Cm(0.1), CORAL)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=CORAL, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "介護を「受ける人」の隣には、\nいつも介護を「担う人」がいる。",
         size=26, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "学校に行けない子ども、仕事を辞めようとしている家族。\n"
         "その姿は、利用者宅を訪れるあなたの目に、きっと映っています。\n"
         "気づき、声をかけ、つなぐこと ― それが、もう一人の誰かの人生を守ります。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 26: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 26)

refs = [
    "■ 子ども・若者育成支援推進法(改正法を含む)― ヤングケアラー支援の根拠",
    "■ こども家庭庁・厚生労働省「ヤングケアラー支援」関連の資料・通知",
    "■ 各自治体のヤングケアラー相談窓口・コーディネーター・支援施策の情報",
    "■ 育児・介護休業法 ― 仕事と介護の両立支援制度の根拠",
    "■ 厚生労働省・経済産業省「仕事と介護の両立支援」関連の資料",
    "■ 各都道府県労働局・「仕事と介護の両立」相談窓口の情報",
    "■ 重層的支援体制整備事業など、世帯まるごと支援の関連情報",
]
y = Cm(4.4)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=13, color=DARK_GRAY)
    y += Cm(1.2)

add_box(s, Cm(1.2), Cm(13), Cm(31.4), Cm(4),
        "ご注意 ・ 情報の確認について",
        "■ 法律名・条文・施行時期、両立支援制度の要件等は、改正により変わります\n"
        "■ ヤングケアラーの相談窓口・施策は自治体ごとに異なります ― 自地域の最新情報をご確認ください\n"
        "■ 制度の詳細な説明・個別の判断は、各分野の専門機関・窓口にご確認ください\n"
        "■ 本資料は 2026年5月時点の一般的な情報を基に構成しています",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 27: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, BLUE)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), CORAL)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 事例検討",
         size=18, color=CORAL, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "現場で出会った家族介護者の事例を持ち寄り、気づき方・つなぎ方を考えましょう。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/ヤングケアラー・ビジネスケアラー支援_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
