"""主任ケアマネ研修スライド生成スクリプト 第5弾

テーマ: ケアプランデータ連携システムと介護DXの推進
重点: 導入事例 / 既存システム(楽にネット等)との併用戦略 / 業務効率化の実際
トーン: 実務型(導入手順・判断基準・効果検証)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(DX・テクノロジートーン:インディゴ基調) =====
INDIGO      = RGBColor(0x23, 0x2E, 0x5C)   # メイン
INDIGO_MID  = RGBColor(0x3C, 0x4B, 0x8A)
INDIGO_LT   = RGBColor(0x60, 0x6E, 0xAE)
CYAN        = RGBColor(0x1F, 0x8F, 0xA8)   # アクセント(連携・効率)
GREEN       = RGBColor(0x2F, 0x8C, 0x5A)   # 効果・OK
AMBER       = RGBColor(0xD4, 0x97, 0x2A)   # 注意・要確認
RED         = RGBColor(0xC0, 0x42, 0x3A)   # 課題・NG
LIGHT_BG    = RGBColor(0xF0, 0xF2, 0xF7)
MID_GRAY    = RGBColor(0xD3, 0xD8, 0xE2)
DARK_GRAY   = RGBColor(0x2F, 0x33, 0x3C)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xE7, 0xEA, 0xF2)
RED_BG      = RGBColor(0xFA, 0xE7, 0xE5)
AMBER_BG    = RGBColor(0xFA, 0xF0, 0xDB)
GREEN_BG    = RGBColor(0xE4, 0xEE, 0xE8)
CYAN_BG     = RGBColor(0xDF, 0xED, 0xF0)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), INDIGO)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=INDIGO)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), CYAN)


def add_footer(slide, page_num, total=33):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(26), Cm(0.5),
             "主任CM研修 / ケアプランデータ連携システムと介護DXの推進  ※費用・補助制度は最新情報をご確認ください",
             size=9, color=INDIGO_LT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=INDIGO_LT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=INDIGO_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=INDIGO, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=CYAN):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, INDIGO)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=accent_color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color, fontsize=15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(fontsize)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


def add_arrow(slide, left, top, width, height, color):
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    arr.shadow.inherit = False
    return arr


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, INDIGO)
add_rect(s, 0, Cm(8), SW, Cm(0.15), CYAN)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), INDIGO_LT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.4), SW - Cm(2), Cm(2.5),
         "ケアプランデータ連携システムと\n介護DXの推進",
         size=31, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
add_text(s, Cm(1.5), Cm(9.7), SW - Cm(3), Cm(2),
         "ー 導入事例から学ぶ、現実的な業務効率化の進め方 ー",
         size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(13), SW - Cm(3), Cm(2),
         "■ ケアプランデータ連携システムの仕組みと導入手順\n"
         "■ 既存の介護ソフト(楽にネット等)との併用戦略\n"
         "■ 業務効率化の実際 ― 導入事例とつまずきポイント",
         size=14, color=AMBER, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "介護DXがなぜ今求められているのか、その全体像を理解する",
    "ケアプランデータ連携システムの仕組み・費用・導入手順を把握する",
    "既存の介護ソフト(楽にネット等)との関係・併用の考え方を整理する",
    "導入事例から、業務効率化の「効果が出る条件」と「つまずき」を学ぶ",
    "自事業所での段階的な導入計画を描けるようになる",
    "主任CMとして、事業所内・連携先のDX推進をリードする視点を持つ",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), CYAN)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "介護DXの全体像 ― なぜ今DXなのか",            "10分"),
    ("2", "ケアプランデータ連携システムの基礎",          "20分"),
    ("3", "既存システムとの併用戦略",                    "20分"),
    ("4", "業務効率化の実際 ― 導入事例に学ぶ",           "20分"),
    ("5", "導入を成功させるために ― 進め方",             "15分"),
    ("6", "これからの介護DXと主任CMの役割",              " 5分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, INDIGO_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=INDIGO)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=CYAN, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "介護DXの全体像", "The Big Picture of Care DX")


# =============================================================
#  SLIDE 5: なぜ今DXか
# =============================================================
s = add_slide()
add_header(s, "1. 介護DXの全体像", "なぜ今、介護DXなのか")
add_footer(s, 5)

reasons = [
    ("人材不足", RED,
     "・生産年齢人口の減少\n・介護職員の確保が年々困難に\n・限られた人手で質を維持する必要"),
    ("業務負担", AMBER,
     "・記録・書類・FAX等の事務作業が膨大\n・対人援助の時間が圧迫される\n・残業・持ち帰り業務の常態化"),
    ("国の方針", CYAN,
     "・「介護現場の生産性向上」が政策の柱\n・文書負担の軽減・標準化を推進\n・補助金・報酬上の評価も整備"),
]
col_w = Cm(10.2)
col_h = Cm(7)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(4.5)
for i, (title, color, body) in enumerate(reasons):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.5), color)
    add_text(s, x, y, col_w, Cm(1.5),
             title, size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.5), col_w, col_h - Cm(1.5), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.9), col_w - Cm(0.8), col_h - Cm(2.3),
             body, size=12, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(12.3), Cm(31.4), Cm(4.7),
        "DXは「目的」ではなく「手段」",
        "■ DX(デジタル・トランスフォーメーション)= デジタル技術で業務やしくみを作り変えること\n\n"
        "■ 「システムを入れること」がゴールではない\n"
        "    → 浮いた時間を 利用者と向き合う時間 ・ ケアの質 に振り向けることが本当の目的\n\n"
        "■ 主任CMの問い:「この技術で、誰の・どの作業が・どれだけ楽になるか?」を常に持つ",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 6: 介護DXの全体像マップ
# =============================================================
s = add_slide()
add_header(s, "1. 介護DXの全体像", "介護DXの全体像 ― 居宅介護支援が関わる領域")
add_footer(s, 6)

areas = [
    ("情報連携の\nデジタル化",
     CYAN,
     "・ケアプランデータ連携システム\n・医療介護連携(MCS等)\n・入退院時の情報共有"),
    ("記録・文書の\n効率化",
     INDIGO_MID,
     "・介護ソフトでの記録一元化\n・音声入力・タブレット記録\n・帳票の自動作成"),
    ("データの\n活用",
     GREEN,
     "・LIFE(科学的介護)\n・データに基づくケアの検証\n・AIケアプラン作成支援"),
    ("見守り・\nケアの高度化",
     AMBER,
     "・見守りセンサー・ICT機器\n・オンラインモニタリング\n・(主にサービス事業所)"),
]
col_w = Cm(7.7)
col_h = Cm(7.5)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(4.5)
for i, (title, color, body) in enumerate(areas):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2), color)
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(2),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2), col_w, col_h - Cm(2), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(2.3), col_w - Cm(0.6), col_h - Cm(2.6),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13), Cm(31.4), Cm(4),
        "本日の焦点 ― ケアプランデータ連携システム",
        "■ 4領域のうち、居宅介護支援事業所が「今すぐ・比較的低コストで」取り組めるのが\n"
        "    左端の「情報連携のデジタル化」 ― 中でもケアプランデータ連携システム\n\n"
        "■ FAX・郵送・手渡しという「最も身近で・最も非効率な作業」を変える、DXの入口",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 7: 第2部
# =============================================================
section_title(2, "ケアプランデータ\n連携システムの基礎",
              "Care Plan Data Linkage System")


# =============================================================
#  SLIDE 8: 従来のやり取りの課題
# =============================================================
s = add_slide()
add_header(s, "2. データ連携システム", "従来のケアプラン授受 ― 何が問題だったか")
add_footer(s, 8)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "毎月くり返される作業",
        "居宅介護支援事業所とサービス事業所の間で、提供票・実績などを\n"
        "FAX・郵送・持参 でやり取りしている ― これが毎月、全利用者分発生する。",
        title_bg=INDIGO, body_size=13)

problems = [
    ("時間がかかる",
     "印刷・封入・FAX送信・電話確認。\n1件ずつの手作業が積み上がる"),
    ("転記ミス",
     "受け取った紙を見て介護ソフトに\n手入力 → 入力ミス・転記漏れ"),
    ("コスト",
     "紙・インク・郵送費・FAX回線。\n保管スペースも必要"),
    ("確認の手間",
     "「届いた?」「読めた?」の\n電話連絡。再送のやり取り"),
    ("月末の集中",
     "提供票・実績の授受が月初・月末に\n集中し、残業の原因に"),
    ("情報漏えい\nリスク",
     "FAX誤送信・郵送事故。\n個人情報を含む書類の取扱い"),
]
col_w = Cm(10.2)
gap_x = Cm(0.4)
y_start = Cm(7.3)
row_h = Cm(2.8)
for i, (title, body) in enumerate(problems):
    col = i % 3
    row = i // 3
    x = (SW - (col_w * 3 + gap_x * 2)) // 2 + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(0.9), RED)
    add_text(s, x + Cm(0.3), y + Cm(0.05), col_w - Cm(0.6), Cm(0.8),
             title, size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(0.9), col_w, row_h - Cm(0.9), RED_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.1), col_w - Cm(0.8), row_h - Cm(1.3),
             body, size=11, color=DARK_GRAY, line_spacing=1.35)

add_text(s, Cm(1.2), Cm(16.7), SW - Cm(2.4), Cm(0.8),
         "▶ 「昔からこうだから」で続けてきた作業に、初めてメスを入れるのがデータ連携システム。",
         size=12, color=CYAN, bold=True)


# =============================================================
#  SLIDE 9: システムとは
# =============================================================
s = add_slide()
add_header(s, "2. データ連携システム", "ケアプランデータ連携システムとは")
add_footer(s, 9)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3.3),
        "概要",
        "■ 国民健康保険中央会(国保中央会)が運営する、ケアプラン関連情報を\n"
        "    事業所間で電子的にやり取りする全国共通のしくみ(令和5年4月から本格稼働)\n"
        "■ 居宅介護支援事業所 ⇔ サービス事業所 の間で、提供票・実績等をデータで授受",
        title_bg=INDIGO, body_size=13)

# やり取りのイメージ
add_rect(s, Cm(2), Cm(8.2), Cm(8), Cm(2.2), INDIGO_MID)
add_text(s, Cm(2), Cm(8.2), Cm(8), Cm(2.2),
         "居宅介護支援\n事業所",
         size=14, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, Cm(12.9), Cm(8.0), Cm(8), Cm(2.6), CYAN)
add_text(s, Cm(12.9), Cm(8.0), Cm(8), Cm(2.6),
         "データ連携\nシステム",
         size=14, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, Cm(23.8), Cm(8.2), Cm(8), Cm(2.2), INDIGO_MID)
add_text(s, Cm(23.8), Cm(8.2), Cm(8), Cm(2.2),
         "サービス\n事業所",
         size=14, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_arrow(s, Cm(10.2), Cm(8.9), Cm(2.5), Cm(0.9), GREEN)
add_arrow(s, Cm(21.1), Cm(8.9), Cm(2.5), Cm(0.9), GREEN)
add_text(s, Cm(2), Cm(10.8), Cm(29.8), Cm(0.8),
         "提供票・提供票別表・実績情報 などを、紙でなくデータで送受信",
         size=12, color=DARK_GRAY, align=PP_ALIGN.CENTER)

add_box(s, Cm(1.2), Cm(12.2), Cm(15.5), Cm(4.8),
        "やり取りできる主な情報",
        "■ サービス利用票・利用票別表(提供票)\n\n"
        "■ サービス提供実績\n\n"
        "■ 厚労省の「ケアプラン標準仕様」に\n   準拠したデータ形式でやり取り",
        title_bg=CYAN, body_bg=CYAN_BG, body_size=12)

add_box(s, Cm(17.5), Cm(12.2), Cm(15), Cm(4.8),
        "前提となる環境",
        "■ 介護ソフト(データ連携対応のもの)\n\n"
        "■ 国保連合会の電子請求で使う\n   電子証明書\n\n"
        "■ インターネット接続環境\n\n"
        "■ 連携専用のクライアントソフト",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 10: メリット
# =============================================================
s = add_slide()
add_header(s, "2. データ連携システム", "導入のメリット ― 何が変わるか")
add_footer(s, 10)

col_widths = [Cm(8), Cm(11.7), Cm(11.7)]
headers = ["項目", "従来(紙・FAX)", "データ連携システム"]
rows = [
    ["授受の手間", "印刷・封入・送信・電話確認", "クリック操作でデータ送受信"],
    ["転記作業", "紙を見て介護ソフトへ手入力", "データを取り込み(転記不要)"],
    ["ミス", "転記ミス・送信ミスが起こる", "転記が減りミスの機会が減る"],
    ["コスト", "紙・インク・郵送費・FAX代", "通信費中心(印刷物が大幅減)"],
    ["時間帯", "相手の営業時間に依存", "24時間いつでも送受信可"],
    ["保管", "紙の保管スペースが必要", "データで管理(省スペース)"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=INDIGO, first_col_color=INDIGO,
          row_height=Cm(1.45), body_size=12)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "効果の試算例(国保中央会)",
        "■ 国保中央会は、1事業所あたり 年間で相応の業務コスト削減効果 があるとの試算を公表\n"
        "■ ただし効果は「連携できる相手先がどれだけあるか」に大きく左右される(後述)\n"
        "▶ 具体的な金額・最新の試算は、国保中央会の公表資料で必ず確認すること",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)


# =============================================================
#  SLIDE 11: 費用とライセンス
# =============================================================
s = add_slide()
add_header(s, "2. データ連携システム", "費用 ・ ライセンス ・ 補助制度")
add_footer(s, 11)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(6),
        "ライセンス費用",
        "■ ケアプランデータ連携システムの利用には、\n   1事業所単位の年間ライセンスが必要\n\n"
        "■ 比較的低額な年額ライセンス制\n   (金額は国保中央会の最新情報を要確認)\n\n"
        "■ このほか、対応介護ソフトの利用料\n   電子証明書の取得費用 等がかかる\n\n"
        "▶ 「何にいくらかかるか」を一覧化して試算する",
        title_bg=INDIGO_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(6),
        "補助制度(ICT導入支援事業)",
        "■ 地域医療介護総合確保基金を活用した\n   ICT導入支援事業 がある\n\n"
        "■ 介護ソフト・タブレット・連携システム等の\n   導入費用の一部を補助\n\n"
        "■ 補助率・上限額・要件は 都道府県ごと\n   に異なり、年度ごとに変わる\n\n"
        "▶ 自治体の介護保険担当課・国保連に確認",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "費用対効果の考え方 ― 主任CMの試算ポイント",
        "■ コスト側:ライセンス料 + ソフト利用料 + 証明書 + 導入時の研修・設定の手間\n\n"
        "■ 効果側:削減できる 紙・郵送・FAXコスト + 作業時間(時給換算)+ 残業削減\n\n"
        "■ 「連携相手がゼロ」なら効果もゼロ ― 周囲のサービス事業所の対応状況とセットで判断\n\n"
        "■ 補助金が使える年度に導入すれば、初期負担を大きく抑えられる ― タイミングが重要",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 12: 導入の流れ
# =============================================================
s = add_slide()
add_header(s, "2. データ連携システム", "導入の流れ ― 5ステップ")
add_footer(s, 12)

steps = [
    ("1", "現状確認",
     "・使っている介護ソフトがデータ連携に対応しているか確認\n・電子証明書(電子請求用)の有無を確認"),
    ("2", "情報収集・申込",
     "・国保中央会のサイトで仕様・費用を確認\n・ライセンスの申込手続き"),
    ("3", "環境準備",
     "・連携用クライアントソフトの導入・設定\n・介護ソフト側の連携設定・動作確認"),
    ("4", "連携先の調整",
     "・取引のあるサービス事業所に対応状況を確認\n・連携を始める相手と運用ルールを取り決め"),
    ("5", "運用開始・拡大",
     "・一部の利用者・事業所から試行\n・問題なければ対象を段階的に拡大"),
]
y = Cm(4.5)
colors = [INDIGO_MID, INDIGO_LT, CYAN, GREEN, AMBER]
for i, (num, title, body) in enumerate(steps):
    add_numbered_circle(s, Cm(1.5), y + Cm(0.35), Cm(1.3), num, colors[i])
    add_rect(s, Cm(3.4), y, Cm(6.5), Cm(2), colors[i])
    add_text(s, Cm(3.4), y, Cm(6.5), Cm(2),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(9.9), y, SW - Cm(11.1), Cm(2), TABLE_BG)
    add_text(s, Cm(10.3), y + Cm(0.2), SW - Cm(11.9), Cm(1.7),
             body, size=12, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
    y += Cm(2.35)

add_text(s, Cm(1.2), Cm(16.5), SW - Cm(2.4), Cm(1),
         "▶ ステップ1・4 がカギ。「自分のソフトが対応しているか」「連携できる相手がいるか」を最初に確認。",
         size=12, color=CYAN, bold=True)


# =============================================================
#  SLIDE 13: 第3部
# =============================================================
section_title(3, "既存システムとの\n併用戦略",
              "Coexistence Strategy with Existing Systems")


# =============================================================
#  SLIDE 14: 介護ソフトとの関係
# =============================================================
s = add_slide()
add_header(s, "3. 併用戦略", "ケアプランデータ連携システムと介護ソフトの関係")
add_footer(s, 14)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3),
        "よくある誤解",
        "「データ連携システムを入れたら、今の介護ソフトを買い替えないといけない」\n"
        "→ 誤解。データ連携システムは 既存の介護ソフトと組み合わせて使う ものです。",
        title_bg=RED, body_bg=RED_BG, body_size=13)

# 関係図
add_rect(s, Cm(3), Cm(8), Cm(11), Cm(3), INDIGO_MID)
add_text(s, Cm(3), Cm(8), Cm(11), Cm(3),
         "介護ソフト\n(楽にネット 等)\n日々の記録・帳票・請求",
         size=13, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_arrow(s, Cm(14.3), Cm(9), Cm(5), Cm(1), GREEN)
add_text(s, Cm(14.3), Cm(8.2), Cm(5), Cm(0.8),
         "データ受渡し", size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_rect(s, Cm(19.8), Cm(8), Cm(11), Cm(3), CYAN)
add_text(s, Cm(19.8), Cm(8), Cm(11), Cm(3),
         "データ連携システム\n事業所間でのデータ送受信",
         size=13, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

add_box(s, Cm(1.2), Cm(11.7), Cm(15.5), Cm(5.3),
        "それぞれの役割",
        "■ 介護ソフト:事業所内の業務システム\n   記録・ケアプラン作成・帳票・国保連請求\n\n"
        "■ データ連携システム:事業所間の通り道\n   作ったデータを「外に送る・受け取る」\n\n"
        "▶ 役割が違うので、両方が必要",
        title_bg=INDIGO_MID, body_size=12)

add_box(s, Cm(17.5), Cm(11.7), Cm(15), Cm(5.3),
        "カギは「連携対応」かどうか",
        "■ 介護ソフトが データ連携システムに対応\n   していれば、そのまま使い続けられる\n\n"
        "■ 多くの主要な介護ソフトは対応済み\n   または対応を進めている\n\n"
        "▶ まず自社ソフトのベンダーに\n   「連携対応の有無」を確認する",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 15: 併用パターンの判断
# =============================================================
s = add_slide()
add_header(s, "3. 併用戦略", "自事業所はどのパターン? ― 4つの状況別判断")
add_footer(s, 15)

col_widths = [Cm(9.5), Cm(11), Cm(11.4)]
headers = ["現在の状況", "とるべき方針", "ポイント"]
rows = [
    ["介護ソフトあり\n+ 連携対応済み",
     "そのまま連携システムを\n追加導入する",
     "最もスムーズ。\nソフト買い替え不要"],
    ["介護ソフトあり\n+ 連携未対応",
     "ベンダーに対応予定を確認\n対応版へのアップデートを待つ\n/ 検討",
     "買い替えありきにしない。\nまず対応予定を聞く"],
    ["紙・Excelで\n運用している",
     "連携対応の介護ソフト導入と\nセットで計画する",
     "補助金活用のチャンス。\n業務全体を見直す好機"],
    ["複数ソフトが\n混在している",
     "連携を機にソフトを\n整理・統一を検討",
     "現場の負担を考え\n段階的に進める"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=INDIGO, first_col_color=CYAN,
          row_height=Cm(2.3), body_size=11, header_size=12)

add_box(s, Cm(1.2), Cm(15.3), Cm(31.4), Cm(2.2),
        "「楽にネット等」既存ソフトを使っている事業所へ",
        "■ 慣れたソフトを無理に変えない ― まずは「連携対応しているか」をベンダーに確認するのが第一歩\n"
        "■ 対応していれば、操作を覚え直す負担なく、連携の効果だけを得られる",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)


# =============================================================
#  SLIDE 16: 移行時の注意
# =============================================================
s = add_slide()
add_header(s, "3. 併用戦略", "ソフト選定 ・ 移行時のチェックポイント")
add_footer(s, 16)

checks = [
    ("ケアプランデータ連携システムに対応しているか",
     "対応していなければ連携の効果が得られない。最重要確認項目"),
    ("国保連の電子請求機能と連動しているか",
     "請求業務まで一気通貫だと、二重入力が発生しない"),
    ("LIFE・他システムとの連携に対応しているか",
     "将来の科学的介護・データ活用を見据えた拡張性"),
    ("操作性・サポート体制は十分か",
     "現場のCMが無理なく使えるか。困ったときのサポートの手厚さ"),
    ("費用(初期・月額)が事業規模に見合うか",
     "クラウド型/オンプレ型、利用人数による課金体系を確認"),
    ("データの移行・バックアップが安全に行えるか",
     "ソフト乗り換え時の過去データの引き継ぎ。情報セキュリティ"),
]
y = Cm(4.4)
for i, (title, body) in enumerate(checks):
    bg = TABLE_BG if i % 2 == 0 else WHITE
    add_rect(s, Cm(1.2), y, Cm(31.4), Cm(1.9), bg, line_color=MID_GRAY)
    add_rect(s, Cm(1.6), y + Cm(0.55), Cm(0.8), Cm(0.8), WHITE, line_color=INDIGO)
    add_text(s, Cm(2.8), y + Cm(0.2), SW - Cm(4.5), Cm(0.9),
             title, size=14, bold=True, color=INDIGO)
    add_text(s, Cm(2.8), y + Cm(1.0), SW - Cm(4.5), Cm(0.8),
             body, size=11, color=DARK_GRAY)
    y += Cm(2.05)


# =============================================================
#  SLIDE 17: 第4部
# =============================================================
section_title(4, "業務効率化の実際",
              "Real-World Efficiency Gains")


# =============================================================
#  SLIDE 18: 導入事例A
# =============================================================
s = add_slide()
add_header(s, "4. 業務効率化の実際", "導入事例① ― 中規模・居宅介護支援事業所")
add_footer(s, 18)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "事業所プロフィール(モデル例)",
        "常勤CM 5名 / 担当 約180名 / 周辺のサービス事業所と多くの取引あり\n"
        "→ 提供票・実績の授受がFAX・郵送中心で、月初・月末に事務作業が集中していた",
        title_bg=INDIGO, body_size=12)

add_box(s, Cm(1.2), Cm(7), Cm(15.5), Cm(5),
        "導入前の課題",
        "■ 月初:提供票の送付に丸2日かかる\n\n"
        "■ 月末:実績の回収・転記に追われる\n\n"
        "■ FAXの送信エラー・再送が頻発\n\n"
        "■ 紙・郵送費が月数万円規模",
        title_bg=RED, body_bg=RED_BG, body_size=12)

add_box(s, Cm(17.5), Cm(7), Cm(15), Cm(5),
        "導入後の変化",
        "■ 連携対応のサービス事業所とは\n   データで一括送受信\n\n"
        "■ 提供票送付の作業時間が大幅短縮\n\n"
        "■ 転記作業が減り、ミスも減少\n\n"
        "■ 紙・郵送コストが目に見えて減少",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(1.2), Cm(12.5), Cm(31.4), Cm(4.5),
        "成功のポイント",
        "■ 取引の多い大手サービス事業所が連携に対応していた → 効果が出やすかった\n"
        "■ 「全件いきなり」ではなく、対応事業所から段階的に切り替えた\n"
        "■ 浮いた時間を、モニタリング訪問・記録の充実にあてた(=DXの目的を果たした)",
        title_bg=CYAN, body_bg=CYAN_BG, body_size=13)


# =============================================================
#  SLIDE 19: 導入事例B(うまくいかなかった例)
# =============================================================
s = add_slide()
add_header(s, "4. 業務効率化の実際", "導入事例② ― 効果が出にくかったケース")
add_footer(s, 19)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "事業所プロフィール(モデル例)",
        "常勤CM 2名 / 担当 約60名 / 取引先は地域の小規模サービス事業所が中心\n"
        "→ 「効率化したい」と考え、データ連携システムを導入した",
        title_bg=INDIGO, body_size=12)

add_box(s, Cm(1.2), Cm(7), Cm(15.5), Cm(5),
        "起きたこと",
        "■ 取引先の多くがシステム未対応\n\n"
        "■ 結局、FAX・郵送も続けることに\n\n"
        "■ 「紙」と「データ」の二重運用に\n\n"
        "■ かえって手間が増えた感覚",
        title_bg=RED, body_bg=RED_BG, body_size=12)

add_box(s, Cm(17.5), Cm(7), Cm(15), Cm(5),
        "原因",
        "■ 導入前に連携相手の対応状況を\n   確認していなかった\n\n"
        "■ 「入れれば効率化する」と思い込んだ\n\n"
        "■ 連携先を増やす働きかけをしなかった\n\n"
        "■ 効果検証の指標を持っていなかった",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_box(s, Cm(1.2), Cm(12.5), Cm(31.4), Cm(4.5),
        "ここから学べること",
        "■ データ連携の効果は 「連携できる相手の数」 に正比例する ― 相手がいなければ効果はゼロ\n"
        "■ 導入の前に必ず:取引先のサービス事業所の対応状況を調べる\n"
        "■ 「二重運用」の期間は必ず生じる ― それを織り込み、徐々に連携先を増やす計画を持つ",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 20: 効果が出る条件
# =============================================================
s = add_slide()
add_header(s, "4. 業務効率化の実際", "効果が出る条件 ― 2つの事例の対比から")
add_footer(s, 20)

col_widths = [Cm(9), Cm(11.7), Cm(11.7)]
headers = ["条件", "効果が出やすい", "効果が出にくい"]
rows = [
    ["連携先の対応", "取引先の多くが対応済み", "取引先がほぼ未対応"],
    ["導入前の調査", "相手の状況を事前に確認", "確認せず見切り発車"],
    ["切り替え方", "対応先から段階的に", "全件いきなり / 計画なし"],
    ["連携先拡大", "未対応先に働きかける", "待ちの姿勢"],
    ["効果検証", "時間・コストを数値で把握", "なんとなくの感覚のみ"],
    ["目的の共有", "浮いた時間の使い道が明確", "「導入」自体が目的化"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=INDIGO, first_col_color=INDIGO,
          row_height=Cm(1.45), body_size=12)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "結論",
        "■ システムは「入れれば効率化する魔法の箱」ではない\n"
        "■ 効果 = システムの性能 × 連携先の数 × 運用の工夫 ― どれかがゼロなら効果もゼロ\n"
        "■ 主任CMの仕事は、この「掛け算」のすべてを高めること",
        title_bg=CYAN, body_bg=CYAN_BG, body_size=13)


# =============================================================
#  SLIDE 21: 連携先を増やす
# =============================================================
s = add_slide()
add_header(s, "4. 業務効率化の実際", "連携先を増やす ― 一事業所だけでは完結しない")
add_footer(s, 21)

add_box(s, Cm(1.2), Cm(4.5), Cm(31.4), Cm(2.5),
        "データ連携は「相手あってのDX」",
        "■ 自分の事業所だけ導入しても、相手が紙のままなら連携は成立しない\n"
        "■ 地域全体で対応事業所が増えるほど、すべての事業所の効果が高まる",
        title_bg=INDIGO, body_size=13)

actions = [
    ("取引先に\n働きかける",
     "・よく取引するサービス事業所に\n  対応状況を聞く\n・「一緒に連携しませんか」と提案"),
    ("地域で\n声をあげる",
     "・ケアマネ連絡会・事業者連絡会で\n  データ連携を話題にする\n・地域ケア会議でも共有"),
    ("情報を\n共有する",
     "・補助金情報・導入事例を\n  地域の事業所に伝える\n・つまずきポイントも共有"),
    ("行政・国保連\nと連携",
     "・自治体の生産性向上の\n  取り組みと歩調を合わせる\n・説明会・研修を活用"),
]
col_w = Cm(7.7)
col_h = Cm(6.5)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.7)
colors = [CYAN, INDIGO_MID, GREEN, AMBER]
for i, (title, body) in enumerate(actions):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(2),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2), col_w, col_h - Cm(2), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(2.3), col_w - Cm(0.6), col_h - Cm(2.6),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_text(s, Cm(1.2), Cm(16.5), SW - Cm(2.4), Cm(1),
         "▶ 主任CMは、自事業所を超えて「地域のDX」を後押しできる立場にある。",
         size=13, color=CYAN, bold=True)


# =============================================================
#  SLIDE 22: つまずきポイント
# =============================================================
s = add_slide()
add_header(s, "4. 業務効率化の実際", "よくあるつまずきと対処法")
add_footer(s, 22)

col_widths = [Cm(9.5), Cm(11), Cm(11.4)]
headers = ["つまずき", "なぜ起こるか", "対処法"]
rows = [
    ["二重運用が\n負担になる",
     "連携先と未連携先が\n混在する移行期",
     "移行期は必ずあると割り切る\n連携先を計画的に増やす"],
    ["現場が\n操作に不慣れ",
     "新しい手順への\n戸惑い・抵抗感",
     "少人数・少件数から試行\n操作担当者を決めて教え合う"],
    ["効果が\n実感できない",
     "連携先が少ない\n効果を測っていない",
     "導入前後で時間・コストを記録\n連携率を見える化する"],
    ["設定で\nつまずく",
     "証明書・ソフト設定\nが分かりにくい",
     "ベンダー・国保連のサポートを\n遠慮なく活用する"],
    ["導入が\n目的化する",
     "「入れて満足」\nしてしまう",
     "「浮いた時間で何をするか」を\n最初に決めておく"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=INDIGO, first_col_color=RED,
          row_height=Cm(2.0), body_size=11, header_size=12)

add_text(s, Cm(1.2), Cm(16), SW - Cm(2.4), Cm(1),
         "▶ つまずきの多くは「想定内」。事前に知っておけば、慌てず対処できる。",
         size=12, color=CYAN, bold=True)


# =============================================================
#  SLIDE 23: 第5部
# =============================================================
section_title(5, "導入を成功させるために",
              "Making the Implementation Succeed")


# =============================================================
#  SLIDE 24: 段階的導入ステップ
# =============================================================
s = add_slide()
add_header(s, "5. 導入の進め方", "段階的導入ロードマップ")
add_footer(s, 24)

phases = [
    ("STEP 1\n調べる",
     "0〜1ヶ月",
     "・自社ソフトの連携対応を確認\n・取引先の対応状況を調査\n・費用・補助金を試算"),
    ("STEP 2\n決める",
     "1〜2ヶ月",
     "・導入の可否・時期を判断\n・事業所内で合意形成\n・補助金申請(該当年度)"),
    ("STEP 3\n準備する",
     "2〜3ヶ月",
     "・ライセンス申込・環境設定\n・操作担当者の決定・研修\n・連携先と運用ルール調整"),
    ("STEP 4\n試す",
     "3〜4ヶ月",
     "・一部利用者・事業所で試行\n・問題点の洗い出し・改善\n・効果を記録し始める"),
    ("STEP 5\n広げる",
     "4ヶ月〜",
     "・対象を段階的に拡大\n・連携先を増やす働きかけ\n・効果を検証・共有"),
]
col_w = Cm(6.2)
col_h = Cm(10)
gap = Cm(0.2)
start_x = (SW - (col_w * 5 + gap * 4)) // 2
y = Cm(4.5)
colors = [INDIGO_MID, INDIGO_LT, CYAN, GREEN, AMBER]
for i, (label, time, body) in enumerate(phases):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2.3), colors[i])
    add_text(s, x + Cm(0.1), y + Cm(0.2), col_w - Cm(0.2), Cm(1.4),
             label, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Cm(0.1), y + Cm(1.6), col_w - Cm(0.2), Cm(0.6),
             time, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, x, y + Cm(2.3), col_w, col_h - Cm(2.3), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(2.6), col_w - Cm(0.6), col_h - Cm(2.9),
             body, size=10, color=DARK_GRAY, line_spacing=1.35)

add_box(s, Cm(1.2), Cm(15), Cm(31.4), Cm(2.2),
        "ロードマップの考え方",
        "■ 期間はあくまで目安。事業所の規模・体制に合わせて無理のないペースで\n"
        "■ 「調べる」を飛ばさない ― ここを省くと事例②のような失敗につながる",
        title_bg=INDIGO, body_size=12)


# =============================================================
#  SLIDE 25: 事業所内の合意形成
# =============================================================
s = add_slide()
add_header(s, "5. 導入の進め方", "事業所内の合意形成 ― 「やらされ感」をなくす")
add_footer(s, 25)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(6.5),
        "現場が抵抗を感じる理由",
        "■ 「今のやり方で困っていない」\n\n"
        "■ 「新しい操作を覚えるのが面倒」\n\n"
        "■ 「忙しいのに、また仕事が増える」\n\n"
        "■ 「うまくいかなかったらどうする」\n\n"
        "■ 「結局、上が決めたことでしょう」\n\n"
        "▶ 不安・負担感は当然のもの。否定しない",
        title_bg=RED, body_bg=RED_BG, body_size=13)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(6.5),
        "主任CMの進め方",
        "■ 目的を共有する\n   「あなたの○○の作業が楽になる」と具体的に\n\n"
        "■ 不安を聞く\n   反対意見こそ、改善のヒント\n\n"
        "■ 小さく始める\n   一部から試し、成功体験を共有\n\n"
        "■ 役割を与える\n   操作に強い職員を「推進担当」に\n\n"
        "■ 効果を見せる\n   「これだけ楽になった」を数字で",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11.5), Cm(31.4), Cm(5.5),
        "DX推進は「人」のマネジメント",
        "■ 技術の問題より、人の気持ちの問題でつまずくことが多い\n\n"
        "■ 「効率化のため」ではなく「あなたの負担を減らすため」と、一人ひとりの言葉に翻訳する\n\n"
        "■ 第4弾で学んだ スーパービジョン・ファシリテーションの技術 が、ここでも生きる\n\n"
        "■ 一度の説明で終わらせず、移行期は繰り返し声をかけ、フォローする",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 26: 導入チェックリスト
# =============================================================
s = add_slide()
add_header(s, "5. 導入の進め方", "導入判断チェックリスト")
add_footer(s, 26)

checks = [
    "現在使っている介護ソフトがデータ連携システムに対応しているか確認した",
    "対応していない場合、ベンダーに対応予定を問い合わせた",
    "電子請求で使う電子証明書の有無・有効期限を確認した",
    "よく取引するサービス事業所の対応状況を調べた",
    "ライセンス料・ソフト費用・証明書費用などの総コストを試算した",
    "活用できる補助金(ICT導入支援事業)の有無を自治体に確認した",
    "導入後に「浮いた時間で何をするか」を事業所内で話し合った",
    "段階的な導入スケジュールと、効果を測る指標を決めた",
]
y = Cm(4.3)
for i, c in enumerate(checks):
    bg = TABLE_BG if i % 2 == 0 else WHITE
    add_rect(s, Cm(1.2), y, Cm(31.4), Cm(1.4), bg, line_color=MID_GRAY)
    add_numbered_circle(s, Cm(1.5), y + Cm(0.25), Cm(0.9), i + 1, INDIGO_MID, fontsize=13)
    add_rect(s, Cm(2.8), y + Cm(0.3), Cm(0.7), Cm(0.7), WHITE, line_color=INDIGO)
    add_text(s, Cm(3.8), y + Cm(0.2), SW - Cm(5), Cm(1),
             c, size=12, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE)
    y += Cm(1.45)

add_text(s, Cm(1.2), Cm(16.3), SW - Cm(2.4), Cm(1),
         "▶ 8項目すべてに「はい」と言えるようになってから、本格導入に進むのが安全。",
         size=12, color=CYAN, bold=True)


# =============================================================
#  SLIDE 27: 第6部
# =============================================================
section_title(6, "これからの介護DXと\n主任CMの役割",
              "The Future of Care DX")


# =============================================================
#  SLIDE 28: これからの介護DX
# =============================================================
s = add_slide()
add_header(s, "6. これからの介護DX", "データ連携の先に広がるもの")
add_footer(s, 28)

futures = [
    ("LIFEとの連動",
     "・サービス事業所のLIFEデータを\n  ケアプランに反映\n・データに基づくケアの検証\n・科学的介護の本格化"),
    ("ケアプラン作成\nのAI支援",
     "・アセスメントからの\n  プラン原案づくりの支援\n・あくまでCMの判断を補助\n・最終判断は専門職が担う"),
    ("医療介護連携\nの深化",
     "・入退院時の情報連携の電子化\n・医療機関とのデータ共有\n・多職種のシームレスな連携"),
    ("ケアプラン\nデータの活用",
     "・地域全体のケアの傾向把握\n・政策・事業計画への反映\n・R9改定の論点とも連動"),
]
col_w = Cm(7.7)
col_h = Cm(7.5)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(4.5)
colors = [GREEN, CYAN, INDIGO_MID, AMBER]
for i, (title, body) in enumerate(futures):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(2),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2), col_w, col_h - Cm(2), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(2.3), col_w - Cm(0.6), col_h - Cm(2.6),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13), Cm(31.4), Cm(4),
        "データ連携システムは「入口」",
        "■ ケアプランデータ連携システムは、介護DXの最初の一歩\n\n"
        "■ ここで「データでやり取りする」経験を積むことが、LIFE活用・AI活用への土台になる\n\n"
        "■ 「まだ早い」ではなく「今できることから」 ― 小さな一歩が、次の大きな変化につながる",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 29: 主任CMの役割
# =============================================================
s = add_slide()
add_header(s, "6. これからの介護DX", "DX推進における主任CMの役割")
add_footer(s, 29)

roles = [
    ("見極める",
     "・「何のためのDXか」を問い続ける\n・流行や売り込みに振り回されない\n・自事業所に必要なものを選ぶ"),
    ("リードする",
     "・導入の旗振り役になる\n・現場の不安に寄り添い、合意を作る\n・段階的に・無理なく進める"),
    ("つなげる",
     "・連携先・地域に働きかける\n・事例・情報を共有する\n・地域全体のDXを後押しする"),
    ("活かす",
     "・浮いた時間をケアの質に振り向ける\n・データを読み、ケアに活かす\n・学び続け、後輩にも伝える"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [INDIGO_MID, CYAN, GREEN, AMBER]
for i, (title, body) in enumerate(roles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13), Cm(31.4), Cm(4),
        "DXの主役は「技術」ではなく「人」",
        "■ どんなに優れたシステムも、使いこなす人がいなければ意味がない\n\n"
        "■ 主任CMは、技術と現場の「翻訳者」 ― 難しい話を、現場の言葉に置き換えて伝える\n\n"
        "■ DXのゴールは、利用者と向き合う時間を取り戻し、ケアの質を高めること",
        title_bg=INDIGO, body_size=13)


# =============================================================
#  SLIDE 30: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 30)

actions = [
    ("Action 1", "自社の介護ソフトの「連携対応」を確認する",
     "・使っている介護ソフトのベンダーに、データ連携システム対応の有無を問い合わせる\n"
     "・未対応なら、対応予定の時期も確認する"),
    ("Action 2", "取引先サービス事業所の対応状況を調べる",
     "・よく取引する事業所が何件、データ連携に対応しているかをリスト化する\n"
     "・「効果が出るだけの連携先があるか」を見極める"),
    ("Action 3", "補助金と費用を試算する",
     "・自治体にICT導入支援事業(補助金)の有無・要件を確認する\n"
     "・ライセンス料等の総コストと、削減できるコスト・時間を試算する"),
]
y = Cm(4.5)
colors = [INDIGO_MID, CYAN, GREEN]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=24, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=16, bold=True, color=INDIGO)
    add_text(s, Cm(6), y + Cm(1.4), SW - Cm(7.5), Cm(2.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 31: クロージング
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, INDIGO)
add_rect(s, 0, Cm(7), SW, Cm(0.1), CYAN)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=CYAN, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "DXの目的は、効率化ではなく、\n「人と向き合う時間」を取り戻すこと。",
         size=25, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "FAXを送る時間、紙を探す時間、転記する時間。\n"
         "その一つひとつを減らした先に、利用者ともう一度きちんと向き合える時間があります。\n"
         "完璧でなくていい。まずは「自社のソフトは対応しているか」の一本の電話から。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 32: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 32)

refs = [
    "■ 国民健康保険中央会「ケアプランデータ連携システム」公式情報",
    "■ 厚生労働省「介護分野におけるICTの活用」「ケアプラン標準仕様」関連資料",
    "■ 厚生労働省「介護現場における生産性向上」関連ガイドライン・手引き",
    "■ 各都道府県「ICT導入支援事業(地域医療介護総合確保基金)」の募集要項",
    "■ 利用中の介護ソフトベンダーの提供する連携対応情報・サポート窓口",
    "■ 市町村・国保連合会の説明会・研修資料",
]
y = Cm(4.5)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=14, color=DARK_GRAY)
    y += Cm(1.3)

add_box(s, Cm(1.2), Cm(13), Cm(31.4), Cm(4),
        "ご注意 ・ 情報の確認について",
        "■ ライセンス費用・補助金の要件・対応ソフトの状況は、年度や地域により変わります\n"
        "■ 導入を検討する際は、必ず 国保中央会・自治体・ソフトベンダーの最新情報 をご確認ください\n"
        "■ 本資料は 2026年5月時点の一般的な情報を基に、特定の製品を推奨しない中立的な立場で構成しています\n"
        "■ 「楽にネット等」は、既存の介護ソフトの一例として記載したものです",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 33: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, INDIGO)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), CYAN)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 情報交換",
         size=18, color=AMBER, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "自事業所・連携先の状況を持ち寄り、地域での取り組みにつなげましょう。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/ケアプランデータ連携システムと介護DX_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
