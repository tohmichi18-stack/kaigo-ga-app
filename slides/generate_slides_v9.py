"""主任ケアマネ研修スライド生成スクリプト 第9弾

テーマ: 困難事例・複合課題への対応
重点: 8050問題 / セルフネグレクト / 精神疾患を抱える利用者 /
      重層的支援体制整備事業との連携
トーン: 実務型(視点・手順・場面対応)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(困難事例:鋼(スチール)基調 + 温かいコーラル) =====
STEEL       = RGBColor(0x2B, 0x3D, 0x4F)   # メイン
STEEL_MID   = RGBColor(0x44, 0x5A, 0x6E)
STEEL_LT    = RGBColor(0x73, 0x88, 0x9C)
CORAL       = RGBColor(0xD0, 0x6A, 0x47)   # アクセント(人・行動)
TEAL        = RGBColor(0x2C, 0x76, 0x73)   # 連携
SAGE        = RGBColor(0x53, 0x7E, 0x5A)   # OK・推奨
AMBER       = RGBColor(0xCB, 0x96, 0x33)   # 注意
RED         = RGBColor(0xB2, 0x42, 0x39)   # リスク・NG
LIGHT_BG    = RGBColor(0xF1, 0xF3, 0xF4)
MID_GRAY    = RGBColor(0xD4, 0xDA, 0xDE)
DARK_GRAY   = RGBColor(0x2E, 0x33, 0x36)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xE7, 0xEB, 0xEE)
RED_BG      = RGBColor(0xF8, 0xE6, 0xE4)
AMBER_BG    = RGBColor(0xFA, 0xF1, 0xDB)
SAGE_BG     = RGBColor(0xE6, 0xED, 0xE6)
CORAL_BG    = RGBColor(0xFA, 0xEA, 0xE2)
TEAL_BG     = RGBColor(0xDF, 0xEB, 0xEA)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), STEEL)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=STEEL)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), CORAL)


def add_footer(slide, page_num, total=32):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(26), Cm(0.5),
             "主任CM研修 / 困難事例・複合課題への対応",
             size=9, color=STEEL_LT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=STEEL_LT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=STEEL_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=STEEL, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=CORAL):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, STEEL)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=accent_color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=33, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color, fontsize=15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(fontsize)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, STEEL)
add_rect(s, 0, Cm(8), SW, Cm(0.15), CORAL)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), STEEL_LT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.6), SW - Cm(2), Cm(2),
         "困難事例・複合課題への対応",
         size=33, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1.5), Cm(9.3), SW - Cm(3), Cm(2),
         "ー 介護保険の枠を超えて、世帯まるごとを支える ー",
         size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(12.8), SW - Cm(3), Cm(2.5),
         "■ 8050問題 ・ セルフネグレクト ・ 精神疾患を抱える利用者\n"
         "■ 重層的支援体制整備事業との連携\n"
         "■ 抱え込まず、つなぎ、チームで支える技術",
         size=14, color=AMBER_BG, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "「困難事例」を捉え直し、複合課題・多問題世帯への視点を持つ",
    "重層的支援体制整備事業のしくみと、CMから見た連携の活かし方を理解する",
    "8050問題 ― 親の介護を入口に、世帯全体を見る視点と連携先を知る",
    "セルフネグレクト ― 支援拒否への向き合い方と、信頼関係づくりを学ぶ",
    "精神疾患を抱える利用者への関わりの基本と、精神科医療との連携を整理する",
    "抱え込まず・チームで・つなぐ ― 主任CMとして困難事例に向き合う技術を持ち帰る",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), CORAL)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "困難事例・複合課題とは ― 捉え直す",      "10分"),
    ("2", "重層的支援体制整備事業との連携",          "15分"),
    ("3", "8050問題への対応",                        "15分"),
    ("4", "セルフネグレクトへの対応",                "15分"),
    ("5", "精神疾患を抱える利用者への対応",          "20分"),
    ("6", "困難事例に向き合う技術 ・ まとめ",        "15分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, STEEL_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=STEEL)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=CORAL, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "困難事例・複合課題とは",
              "Rethinking Complex Cases")


# =============================================================
#  SLIDE 5: 困難事例とは
# =============================================================
s = add_slide()
add_header(s, "1. 困難事例・複合課題とは", "「困難事例」とは誰にとっての困難か")
add_footer(s, 5)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "「困難事例」という言葉を捉え直す",
        "「困難事例」と呼ぶとき、それは多くの場合「支援者にとっての困難」を指している。\n"
        "本人にとっての「困りごと」と、支援者にとっての「困難」は、同じではない。",
        title_bg=STEEL, body_size=13)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "支援者にとっての「困難」",
        "■ 支援を拒否される\n\n"
        "■ 制度・サービスにあてはまらない\n\n"
        "■ 関わる機関が多く、調整が大変\n\n"
        "■ 解決の糸口が見えない\n\n"
        "■ 対応に時間と労力がかかる",
        title_bg=STEEL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "本人にとっての「困りごと」",
        "■ これまでの生き方・価値観がある\n\n"
        "■ 助けを求めることへの葛藤・不信\n\n"
        "■ 過去の傷つき体験\n\n"
        "■ 「困っている」自覚がないこともある\n\n"
        "▶ 本人の世界から見ることが出発点",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(2),
         "▶ 「困難事例」とラベルを貼った瞬間、本人が「困らせる人」に見えてしまう。\n"
         "▶ 「困難を抱えた、その人」として捉え直すことが、支援の第一歩。",
         size=13, color=STEEL, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 6: 複合課題の広がり
# =============================================================
s = add_slide()
add_header(s, "1. 困難事例・複合課題とは", "複合課題・多問題世帯の広がり")
add_footer(s, 6)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "1つの世帯に、いくつもの課題が重なる",
        "高齢の親の介護課題だけでなく、同居の子の課題、経済、住まい、孤立 ―\n"
        "複数の課題が絡み合い、1つの制度だけでは支えきれない世帯が増えている。",
        title_bg=STEEL, body_size=12)

issues = [
    ("8050問題",
     "高齢の親と、ひきこもり等\nの中高年の子の同居世帯"),
    ("ダブルケア",
     "介護と育児が\n同時期に重なる"),
    ("セルフネグレクト",
     "自分自身のケアを\n放棄してしまう状態"),
    ("経済的困窮",
     "生活困窮・滞納・\n債務・年金頼みの家計"),
    ("精神疾患・障害",
     "本人や家族が精神疾患・\n障害を抱える"),
    ("社会的孤立",
     "地域・親族とのつながり\nが断たれている"),
]
col_w = Cm(10.2)
gap_x = Cm(0.4)
y_start = Cm(7.2)
row_h = Cm(2.7)
colors = [CORAL, TEAL, STEEL_MID, AMBER, RED, STEEL_LT]
for i, (title, body) in enumerate(issues):
    col = i % 3
    row = i // 3
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(0.95), colors[i])
    add_text(s, x + Cm(0.3), y + Cm(0.05), col_w - Cm(0.6), Cm(0.85),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(0.95), col_w, row_h - Cm(0.95), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.15), col_w - Cm(0.8), row_h - Cm(1.3),
             body, size=11, color=DARK_GRAY, line_spacing=1.35)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "なぜ「複合課題」が支えにくいのか",
        "■ 介護保険は「高齢者本人」を対象とする制度 ― 同居家族の課題は守備範囲外になりがち\n"
        "■ 制度ごとに窓口・担当が分かれ、世帯全体を見る人がいない「制度の縦割り」が壁になる",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 7: CMが世帯を見る
# =============================================================
s = add_slide()
add_header(s, "1. 困難事例・複合課題とは", "CMは「世帯まるごと」に出会っている")
add_footer(s, 7)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3),
        "CMだからこそ気づける立場",
        "居宅CMは、利用者宅を訪問し、生活の場に入る数少ない専門職。\n"
        "高齢者本人だけでなく、同居家族・住まい・暮らしぶり ―\n"
        "「世帯まるごと」の状況に、最初に気づける立場にいる。",
        title_bg=STEEL, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(4.8),
        "CMにできること",
        "■ 訪問を通じて世帯の異変に気づく\n\n"
        "■ 高齢者支援を入口に、世帯全体を見る\n\n"
        "■ 必要な機関・制度につなぐ\n\n"
        "■ 多機関協働の一員として関わり続ける",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(4.8),
        "CMだけでは担えないこと",
        "■ 高齢者以外の家族への直接の支援\n\n"
        "■ 制度横断的な調整の主導\n\n"
        "■ 長期的・専門的な伴走\n\n"
        "▶ だからこそ「つなぐ」「協働する」が要",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(2),
         "▶ 「介護保険の範囲外だから」と見て見ぬふりをしない。\n"
         "▶ かといって一人で抱え込まない ― 「気づき、つなぐ」のがCMの役割。",
         size=13, color=STEEL, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 8: 第2部
# =============================================================
section_title(2, "重層的支援体制整備事業\nとの連携",
              "Working with Community-Based Integrated Support")


# =============================================================
#  SLIDE 9: 重層事業とは
# =============================================================
s = add_slide()
add_header(s, "2. 重層的支援体制整備事業", "重層的支援体制整備事業とは")
add_footer(s, 9)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3.3),
        "制度の「はざま」を埋めるしくみ",
        "■ 社会福祉法に基づき、市町村が実施する事業(任意事業)\n"
        "■ 介護・障害・子ども・生活困窮 ― 分野ごとの相談支援を「一体的」に行い、\n"
        "    属性や世代を問わず、複合課題を抱える人・世帯を丸ごと受けとめる",
        title_bg=STEEL, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(4.8),
        "なぜ生まれたのか",
        "■ 「縦割り」では複合課題に対応できない\n\n"
        "■ 「どの窓口にも当てはまらない」人がいる\n\n"
        "■ 8050・ダブルケア・孤立など\n   制度のはざまの課題の顕在化\n\n"
        "■ 「断らない相談支援」をめざして",
        title_bg=STEEL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(4.8),
        "知っておきたいこと",
        "■ 「任意事業」のため、実施状況は\n   市町村によって異なる\n\n"
        "■ 自分の市町村が実施しているか、\n   窓口はどこかを確認しておく\n\n"
        "■ 未実施の地域でも、各分野の相談\n   機関の連携で同様の対応をめざす",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ まず「自分の市町村ではどうなっているか」を知ることが、連携の出発点。",
         size=13, color=STEEL, bold=True)


# =============================================================
#  SLIDE 10: 重層事業の3つの支援
# =============================================================
s = add_slide()
add_header(s, "2. 重層的支援体制整備事業", "重層事業を支える支援のかたち")
add_footer(s, 10)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(1.8),
        "3つの支援 + それを支える2つの機能",
        "属性を問わない相談 / 参加支援 / 地域づくり ― を、多機関協働とアウトリーチが支える。",
        title_bg=STEEL, body_size=12)

supports = [
    ("包括的相談支援",
     "属性・世代を問わず\n相談を受けとめる\n「断らない相談支援」"),
    ("参加支援",
     "社会的なつながり・\n社会参加・就労等の\n「出番」をつくる支援"),
    ("地域づくりに\n向けた支援",
     "住民同士のつながり・\n交流・支え合いの\n場をつくる"),
]
col_w = Cm(10.2)
col_h = Cm(5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(6.7)
colors = [STEEL_MID, CORAL, SAGE]
for i, (title, body) in enumerate(supports):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.5), colors[i])
    add_text(s, x, y, col_w, Cm(1.5),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.5), col_w, col_h - Cm(1.5), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.8), col_w - Cm(0.8), col_h - Cm(2.1),
             body, size=12, color=DARK_GRAY, line_spacing=1.5,
             align=PP_ALIGN.CENTER)

add_box(s, Cm(1.2), Cm(12.2), Cm(15.5), Cm(4.8),
        "多機関協働",
        "■ 各分野の支援者をつなぎ、世帯全体の\n   支援の方向性を調整する\n\n"
        "■ 「重層的支援会議」「支援会議」等で\n   関係機関が情報を共有・検討する\n\n"
        "■ CMも、その協働のメンバーになる",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_box(s, Cm(17.5), Cm(12.2), Cm(15), Cm(4.8),
        "アウトリーチ等を通じた継続的支援",
        "■ 自ら相談に来られない人へ、支援者から\n   出向いて関わる(アウトリーチ)\n\n"
        "■ すぐに解決しなくても、つながり続ける\n\n"
        "■ 「待つ」のではなく「届ける」支援",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 11: CMから見た連携
# =============================================================
s = add_slide()
add_header(s, "2. 重層的支援体制整備事業", "CMから見た ― 重層事業の活かし方")
add_footer(s, 11)

col_widths = [Cm(10.5), Cm(10.5), Cm(10.4)]
headers = ["こんなとき", "重層事業を通じて", "CMの動き方"]
rows = [
    ["同居の子の課題に\n気づいた",
     "高齢者以外の家族の\n相談先につながる",
     "市町村の相談窓口・\n包括に相談・情報共有"],
    ["どこに相談すれば\nよいか分からない",
     "「断らない相談支援」が\n受けとめてくれる",
     "迷ったら、まず\n包括的相談支援につなぐ"],
    ["関わる機関が多く\n調整が難しい",
     "多機関協働で支援の\n方向性が整理される",
     "重層的支援会議等に\n参加し情報を持ち寄る"],
    ["本人が孤立し\n出番がない",
     "参加支援で社会との\nつながりをつくる",
     "本人の希望を伝え\n参加支援につなぐ"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=STEEL, first_col_color=STEEL,
          row_height=Cm(2.0), body_size=11)

add_box(s, Cm(1.2), Cm(13.5), Cm(31.4), Cm(3.7),
        "連携を活かすために、主任CMがしておくこと",
        "■ 自分の市町村の重層事業の実施状況・相談窓口・キーパーソンを把握しておく\n\n"
        "■ 「介護の課題」として閉じず、世帯の課題として包括・市町村に共有する習慣をつくる\n\n"
        "■ 後輩CMに「一人で抱え込まず、つなぐ」ことを、具体的なつなぎ先とともに伝える",
        title_bg=STEEL, body_size=13)


# =============================================================
#  SLIDE 12: 第3部
# =============================================================
section_title(3, "8050問題への対応",
              "Responding to the 8050 Problem")


# =============================================================
#  SLIDE 13: 8050問題とは
# =============================================================
s = add_slide()
add_header(s, "3. 8050問題", "8050問題とは")
add_footer(s, 13)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3),
        "「8050問題」が指すもの",
        "■ 80代の高齢の親と、50代の中高年の子(ひきこもり等)が同居する世帯の課題\n"
        "■ 親の介護・年金に依存して生活が成り立ってきたが、親の高齢化・要介護化・\n"
        "    死亡により、世帯全体が立ち行かなくなる ― そのリスクが顕在化する",
        title_bg=STEEL, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(5),
        "背景にあるもの",
        "■ ひきこもりの長期化・高年齢化\n\n"
        "■ 親子ともに社会的に孤立している\n\n"
        "■ 親の年金が世帯の唯一の収入\n\n"
        "■ 「家のことは家で」という意識\n\n"
        "■ 子の課題(障害・疾患・離職等)",
        title_bg=STEEL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(5),
        "なぜ気づかれにくいのか",
        "■ 子の存在自体が「見えない」\n\n"
        "■ 親が子のことを語りたがらない\n\n"
        "■ 「恥」「世間体」の意識\n\n"
        "■ 介護の相談に、子の課題は出てこない\n\n"
        "▶ 介護をきっかけに「見えてくる」ことが多い",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.5), SW - Cm(2.4), Cm(1.2),
         "▶ 親の介護に入ったCMが、世帯の中の「もう一つの課題」に最初に気づくことが多い。",
         size=13, color=STEEL, bold=True)


# =============================================================
#  SLIDE 14: 8050へのCMの関わり
# =============================================================
s = add_slide()
add_header(s, "3. 8050問題", "CMの関わり ― 親の介護を入口に、世帯を見る")
add_footer(s, 14)

steps = [
    ("気づく",
     "・同居の子の存在・状況に気を配る\n・「日中、お子さんは?」とさりげなく\n・世帯の経済状況・将来の見通し"),
    ("否定せず受けとめる",
     "・親・子を責めない、せかさない\n・「親亡き後」の不安を受けとめる\n・まず親(利用者)との信頼を大切に"),
    ("一人で抱えない",
     "・主任CM・地域包括に相談\n・子の課題はCMの守備範囲を超える\n・「つなぐ」判断を早めに"),
    ("つなぐ・協働する",
     "・市町村の相談窓口・重層事業へ\n・ひきこもり支援・生活困窮・障害\n・多機関で世帯を支える形をつくる"),
]
y = Cm(4.4)
colors = [STEEL_MID, CORAL, AMBER, SAGE]
for i, (title, body) in enumerate(steps):
    add_numbered_circle(s, Cm(1.5), y + Cm(0.4), Cm(1.2), i + 1, colors[i])
    add_rect(s, Cm(3.2), y, Cm(6.5), Cm(2.2), colors[i])
    add_text(s, Cm(3.3), y, Cm(6.3), Cm(2.2),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(9.7), y, SW - Cm(10.9), Cm(2.2), TABLE_BG)
    add_text(s, Cm(10.1), y + Cm(0.2), SW - Cm(11.5), Cm(1.9),
             body, size=11.5, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
    y += Cm(2.45)

add_box(s, Cm(1.2), Cm(14.6), Cm(31.4), Cm(2.5),
        "連携先を知っておく",
        "■ ひきこもり地域支援センター ■ 生活困窮者自立支援の相談窓口 ■ 障害福祉の相談支援\n"
        "■ 市町村の重層的支援の窓口・地域包括支援センター ― 「つなぐ先のリスト」を持っておく",
        title_bg=STEEL, body_size=12)


# =============================================================
#  SLIDE 15: 第4部
# =============================================================
section_title(4, "セルフネグレクトへの対応",
              "Responding to Self-Neglect")


# =============================================================
#  SLIDE 16: セルフネグレクトとは
# =============================================================
s = add_slide()
add_header(s, "4. セルフネグレクト", "セルフネグレクトとは")
add_footer(s, 16)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "セルフネグレクト = 自己放任",
        "健康・安全・生活に関わる、自分自身のケアを行えない・行わない状態。\n"
        "高齢者虐待防止法の「虐待」の定義には含まれないが、生命・健康に関わる重大な課題。",
        title_bg=STEEL, body_size=12)

signs = [
    ("住環境のサイン",
     "・ごみ・物が大量にたまる\n・不衛生・異臭\n・ライフラインが止まっている"),
    ("身体・健康のサイン",
     "・極端なやせ・脱水\n・必要な受診・服薬をしない\n・けが・病気を放置"),
    ("生活のサイン",
     "・食事をとらない\n・入浴・着替えをしない\n・金銭管理ができていない"),
    ("関わりのサイン",
     "・支援・受診を強く拒む\n・近隣との関係が断たれている\n・「困っていない」と言う"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(7.3)
row_h = Cm(3.2)
colors = [STEEL_MID, RED, AMBER, CORAL]
for i, (title, body) in enumerate(signs):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=11.5, color=DARK_GRAY, line_spacing=1.35)

add_text(s, Cm(1.2), Cm(16.3), SW - Cm(2.4), Cm(1),
         "▶ 背景に、認知症・精神疾患・喪失体験・あきらめ・プライドなどが隠れていることが多い。",
         size=12, color=STEEL, bold=True)


# =============================================================
#  SLIDE 17: 支援拒否への向き合い方
# =============================================================
s = add_slide()
add_header(s, "4. セルフネグレクト", "支援拒否への向き合い方")
add_footer(s, 17)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "「拒否」の奥にあるものを考える",
        "支援を拒むのには理由がある ― 過去の傷つき、不信、プライド、あきらめ、\n"
        "「困っている自覚がない」こと。「拒否」を「NO」と即断せず、その背景を見る。",
        title_bg=STEEL, body_size=12)

col_widths = [Cm(15.7), Cm(15.7)]
headers = ["関係を遠ざける関わり", "信頼に近づく関わり"]
rows = [
    ["「このままでは危険です」と説得する",
     "本人の暮らし方・思いをまず否定せず聴く"],
    ["すぐに片づけ・受診をさせようとする",
     "本人が受け入れられる小さな一歩から"],
    ["訪問のたびに支援を勧める",
     "支援の話を急がず、まず関係をつくる"],
    ["「困っていますよね」と決めつける",
     "本人なりの「困りごと」に焦点を当てる"],
    ["一度断られたら関わりをやめる",
     "断られても、つながり続ける(見守る)"],
]
add_table(s, Cm(1.2), Cm(7.3), col_widths, headers, rows,
          header_bg=STEEL, first_col_color=RED,
          row_height=Cm(1.5), body_size=12, first_col_bold=False)

add_box(s, Cm(1.2), Cm(15.3), Cm(31.4), Cm(2.2),
        "支援の原則",
        "■ ゴールは「すぐに解決」ではなく「つながり続ける」こと ― 関係こそが支援の土台\n"
        "■ ただし、生命の危険が切迫している場合は、本人の同意を待たず関係機関と緊急対応する",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 18: 見守りとチーム対応
# =============================================================
s = add_slide()
add_header(s, "4. セルフネグレクト", "つながり続ける ― 見守りとチーム対応")
add_footer(s, 18)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.3),
        "関わり続けるための工夫",
        "■ 訪問の理由を「支援」でなく自然な形で\n   (「近くまで来たので」等)\n\n"
        "■ 本人の関心・得意なことを糸口にする\n\n"
        "■ 小さな「できた」「受け入れた」を大切に\n\n"
        "■ 焦らず、本人のペースに合わせる\n\n"
        "■ 変化が見えなくても、関わりを記録する",
        title_bg=STEEL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.3),
        "チーム・地域で見守る",
        "■ CM一人で抱えない ― 地域包括と共有\n\n"
        "■ 民生委員・近隣・ライフライン事業者等\n   地域の「見守りの目」と連携\n\n"
        "■ 緊急時の連絡体制を決めておく\n\n"
        "■ 重層事業のアウトリーチ支援と連携\n\n"
        "■ 医療(認知症・精神)の視点も入れる",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "主任CM・支援者として大切にしたいこと",
        "■ 「説得できないこと」を支援者の失敗と捉えない ― つながり続けること自体が成果\n\n"
        "■ 後輩CMが「拒否されてつらい」と感じたとき、その気持ちを支える(支持的SV)\n\n"
        "■ 支援者の安全も確保する ― 不衛生な環境・予測しにくい状況では複数人で訪問する\n\n"
        "■ 「いつ・誰が・どう動くか」を地域包括・関係機関と取り決め、CMだけに集中させない",
        title_bg=STEEL, body_size=13)


# =============================================================
#  SLIDE 19: 第5部
# =============================================================
section_title(5, "精神疾患を抱える\n利用者への対応",
              "Supporting People with Mental Illness")


# =============================================================
#  SLIDE 20: 精神疾患の理解
# =============================================================
s = add_slide()
add_header(s, "5. 精神疾患を抱える利用者", "精神疾患を抱える利用者を理解する")
add_footer(s, 20)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "「分かりにくさ」を、まず受けとめる",
        "精神疾患は外から見えにくく、症状や状態に波がある。\n"
        "「気分」や「性格」と捉えず、疾患による状態として理解することが、関わりの出発点。",
        title_bg=STEEL, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "理解しておきたいこと",
        "■ 症状・状態には波がある(良い時/悪い時)\n\n"
        "■ 服薬の継続が安定の鍵になることが多い\n\n"
        "■ 不安・不信・被害的な思いを抱きやすい\n\n"
        "■ 長年の生きづらさ・孤立を抱えてきた\n\n"
        "■ 認知症とは異なる ― 安易に混同しない",
        title_bg=STEEL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "高齢期に出会う精神疾患",
        "■ 若い頃からの疾患を抱えて高齢化した人\n\n"
        "■ 高齢期に発症・顕在化する場合\n\n"
        "■ うつ・不安、アルコール関連の課題\n\n"
        "■ 認知症との併存\n\n"
        "▶ 「介護」と「精神保健福祉」が交わる領域",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ 診断名で人を見ない。「疾患を抱えながら生きてきた、その人」として理解する。",
         size=13, color=STEEL, bold=True)


# =============================================================
#  SLIDE 21: 関わりの基本
# =============================================================
s = add_slide()
add_header(s, "5. 精神疾患を抱える利用者", "関わりの基本 ― 安心できる関係をつくる")
add_footer(s, 21)

col_widths = [Cm(15.7), Cm(15.7)]
headers = ["不安・不信を強める関わり", "安心できる関わり"]
rows = [
    ["約束・予定を急に変える",
     "予定を一定に保ち、変更は早めに丁寧に伝える"],
    ["大勢で訪問する・距離が近すぎる",
     "落ち着いた人数・適度な距離感を保つ"],
    ["本人の訴え・不安を否定する",
     "訴えをまず受けとめ、否定から入らない"],
    ["一度に多くを求める・急かす",
     "一つずつ、本人のペースで進める"],
    ["症状の波を「わがまま」と捉える",
     "波があることを前提に、調子に合わせる"],
    ["できないことばかりを指摘する",
     "できていること・本人の強みに目を向ける"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=STEEL, first_col_color=RED,
          row_height=Cm(1.55), body_size=12, first_col_bold=False)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "土台となる姿勢",
        "■ 「安心」「一定」「予測できる」関わりが、不安を和らげ、信頼の土台になる\n"
        "■ 本人の訴えの「内容の正誤」より、その背後にある「不安・つらさ」に応える",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 22: 精神科医療との連携
# =============================================================
s = add_slide()
add_header(s, "5. 精神疾患を抱える利用者", "精神科医療・精神保健福祉との連携")
add_footer(s, 22)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "介護だけで抱えず、医療・精神保健福祉と組む",
        "精神疾患を抱える利用者の支援は、医療(精神科)との連携が不可欠。\n"
        "CMは「介護の専門職」として、医療チームと役割を分担し、協働する。",
        title_bg=STEEL, body_size=12)

partners = [
    ("精神科医療機関\n・主治医",
     "病状の管理・服薬\n医学的な見立て・助言"),
    ("精神保健福祉士\n(PSW)",
     "生活・社会復帰の支援\n医療と地域をつなぐ"),
    ("保健所・精神保健\n福祉センター",
     "精神保健の相談\n対応に迷うときの相談先"),
    ("障害福祉サービス",
     "自立生活・就労等の支援\n相談支援専門員と連携"),
]
col_w = Cm(7.7)
col_h = Cm(5.5)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7)
colors = [STEEL_MID, TEAL, AMBER, SAGE]
for i, (title, body) in enumerate(partners):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.8), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(1.8),
             title, size=12.5, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.8), col_w, col_h - Cm(1.8), TABLE_BG)
    add_text(s, x + Cm(0.25), y + Cm(2.1), col_w - Cm(0.5), col_h - Cm(2.4),
             body, size=11, color=DARK_GRAY, line_spacing=1.4,
             align=PP_ALIGN.CENTER)

add_box(s, Cm(1.2), Cm(13.7), Cm(31.4), Cm(3.3),
        "連携を機能させるために",
        "■ 「介護保険サービス」と「障害福祉サービス」の両方を視野に入れる(併用・移行)\n"
        "■ 服薬・通院が途切れていないか、生活の中で見える変化を医療チームに共有する\n"
        "■ 対応に迷うとき・緊急時の相談先(保健所・精神保健福祉センター等)を把握しておく",
        title_bg=STEEL, body_size=12)


# =============================================================
#  SLIDE 23: 家族支援
# =============================================================
s = add_slide()
add_header(s, "5. 精神疾患を抱える利用者", "家族への支援も視野に入れる")
add_footer(s, 23)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "家族もまた、長く支え、疲れてきた",
        "精神疾患を抱える人の家族は、長年にわたり対応を続け、孤立し、\n"
        "疲弊していることが多い。家族もまた「支えられる対象」である。",
        title_bg=STEEL, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5),
        "家族が抱えやすいもの",
        "■ 長年の対応による心身の疲弊\n\n"
        "■ 「自分のせいでは」という自責感\n\n"
        "■ 周囲に相談できず孤立\n\n"
        "■ 親亡き後・自分亡き後への不安\n\n"
        "■ 高齢の親が高齢の子を支える構図",
        title_bg=STEEL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5),
        "CMができること",
        "■ 家族の労をねぎらい、話を聴く\n\n"
        "■ 家族を責めない・追い詰めない\n\n"
        "■ 家族会・ピアの集まりにつなぐ\n\n"
        "■ 家族の課題は、家族支援の窓口へ\n\n"
        "■ 「親亡き後」を見据えた連携を早めに",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13), SW - Cm(2.4), Cm(1.2),
         "▶ 8050問題とも重なる ― 「世帯まるごと」「親亡き後」の視点を、ここでも忘れない。",
         size=13, color=STEEL, bold=True)


# =============================================================
#  SLIDE 24: 第6部
# =============================================================
section_title(6, "困難事例に向き合う技術",
              "Skills & Wrap Up")


# =============================================================
#  SLIDE 25: アセスメントの視点
# =============================================================
s = add_slide()
add_header(s, "6. 向き合う技術", "困難事例のアセスメント ― 3つの視点")
add_footer(s, 25)

views = [
    ("世帯まるごと\n見る",
     "・利用者本人だけでなく、同居家族・\n  世帯全体の状況を捉える\n・収入・住まい・つながりも視野に\n・「見えていない人」に気を配る"),
    ("ストレングス\n(強み)を見る",
     "・課題・できないことだけを見ない\n・本人・世帯の強み・資源・願いを探す\n・これまで生き抜いてきた力に着目\n・本人の「したい」を支援の起点に"),
    ("時間の流れで\n見る",
     "・なぜ今の状態になったのか(経過)\n・本人の人生史・関係史を理解する\n・「親亡き後」など将来も見据える\n・すぐ解決を求めず長い目で"),
]
col_w = Cm(10.2)
col_h = Cm(9)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(4.5)
colors = [STEEL_MID, SAGE, TEAL]
for i, (title, body) in enumerate(views):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.8), colors[i])
    add_text(s, x + Cm(0.2), y, col_w - Cm(0.4), Cm(1.8),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.8), col_w, col_h - Cm(1.8), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(2.2), col_w - Cm(0.8), col_h - Cm(2.6),
             body, size=12, color=DARK_GRAY, line_spacing=1.55)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "アセスメントが支援の質を決める",
        "■ 「困難」に見える事例ほど、丁寧なアセスメントで「糸口」が見えてくる\n"
        "■ 課題の山を前に立ちすくむより、「小さくても動かせる一点」を一緒に探す",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 26: 抱え込まない
# =============================================================
s = add_slide()
add_header(s, "6. 向き合う技術", "抱え込まない ― チームで・つなぐ")
add_footer(s, 26)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "困難事例こそ、一人で支えてはいけない",
        "複合課題は、一人のCM・一つの機関では支えきれない。\n"
        "「抱え込まない」ことは、力不足ではなく、適切な支援判断そのもの。",
        title_bg=STEEL, body_size=12)

points = [
    ("早めに共有する",
     "・「困難かも」と感じた段階で\n  主任CM・包括に共有\n・こじれてからでは連携が難しい"),
    ("会議の場を使う",
     "・地域ケア会議・支援会議等で\n  多機関で検討する\n・一人の見立てを複数の目で"),
    ("役割を分担する",
     "・「誰が・何を担うか」を明確に\n・CMの守備範囲を超える課題は\n  専門機関につなぐ"),
    ("つながり続ける",
     "・つないだら終わり、ではない\n・多機関協働の一員として\n  関わり続ける"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(7.3)
row_h = Cm(3.5)
colors = [STEEL_MID, TEAL, AMBER, SAGE]
for i, (title, body) in enumerate(points):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_text(s, Cm(1.2), Cm(16.4), SW - Cm(2.4), Cm(0.9),
         "▶ 主任CMは、後輩が「抱え込まず相談できる」空気と仕組みをつくる。",
         size=12, color=STEEL, bold=True)


# =============================================================
#  SLIDE 27: 支援者の安全と燃え尽き予防
# =============================================================
s = add_slide()
add_header(s, "6. 向き合う技術", "支援者を守る ― 安全と燃え尽きの予防")
add_footer(s, 27)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.3),
        "支援者の安全を守る",
        "■ リスクのある訪問は複数人で行う\n\n"
        "■ 訪問の予定・行き先を事業所で共有\n\n"
        "■ 不衛生な環境での健康・衛生に配慮\n\n"
        "■ 暴言・暴力等のリスクには組織で対応\n\n"
        "■ 「一人で行かせない」を主任CMが徹底",
        title_bg=RED, body_bg=RED_BG, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.3),
        "燃え尽きを防ぐ",
        "■ 困難事例は、成果が見えにくく消耗する\n\n"
        "■ 「解決できない」を失敗と捉えさせない\n\n"
        "■ 感情を言葉にして吐き出せる場をつくる\n\n"
        "■ 困難事例の負担を一人に集中させない\n\n"
        "■ 「つながり続けたこと」を成果と認める",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "主任CMの役割 ― 支える人を支える",
        "■ 困難事例を担当するCMを孤立させない ― 定期的に状況を聴き、ねぎらう\n\n"
        "■ 「つながり続けられたこと」「小さな変化」を、一緒に成果として確認する\n\n"
        "■ 支持的スーパービジョン(第4弾)で、CMの感情・つらさを受けとめる\n\n"
        "■ 困難事例を特定のCMに集中させず、チーム・事業所全体で分かち合う",
        title_bg=STEEL, body_size=13)


# =============================================================
#  SLIDE 28: 主任CMの役割
# =============================================================
s = add_slide()
add_header(s, "6. 向き合う技術", "困難事例対応における主任CMの役割")
add_footer(s, 28)

roles = [
    ("見立てを支える",
     "・後輩の事例を一緒にアセスメント\n・世帯まるごと・強みの視点を伝える\n・「困難」のラベルを問い直す"),
    ("つなぎ先を持つ",
     "・重層事業・専門機関のリストを整備\n・地域の窓口・キーパーソンを把握\n・「どこにつなぐか」を即答できる"),
    ("会議・連携を動かす",
     "・地域ケア会議・支援会議を活用\n・多機関協働のハブになる\n・地域の連携体制づくりに関わる"),
    ("人を支える",
     "・担当CMを孤立させない\n・安全を確保し、燃え尽きを防ぐ\n・困難事例を一人に集中させない"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [STEEL_MID, TEAL, CORAL, SAGE]
for i, (title, body) in enumerate(roles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "困難事例は「地域の力」で支える",
        "■ 困難事例は、CM個人の力量の問題ではなく、地域全体で支えるべき課題\n"
        "■ 主任CMは、事業所と地域・多機関をつなぎ、「支える網」をつくる役割を担う\n"
        "■ 本シリーズで学んだSV・ファシリテーション・意思決定支援が、ここでも一貫して生きる",
        title_bg=STEEL, body_size=13)


# =============================================================
#  SLIDE 29: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 29)

actions = [
    ("Action 1", "自分の市町村の「つなぎ先」を確認する",
     "・重層的支援体制整備事業の実施状況・相談窓口を調べる\n"
     "・ひきこもり支援・生活困窮・精神保健等の連携先をリスト化しておく"),
    ("Action 2", "担当ケースを「世帯まるごと」で見直す",
     "・気になるケースを、同居家族・経済・つながりも含めて捉え直す\n"
     "・「介護の課題」に閉じず、世帯の課題として包括と共有する"),
    ("Action 3", "困難事例を「一人で抱えない」体制を点検する",
     "・後輩CMが抱える困難事例を把握し、相談できる場をつくる\n"
     "・リスクのある訪問の安全確保(複数訪問・行き先共有)を徹底する"),
]
y = Cm(4.5)
colors = [STEEL_MID, TEAL, CORAL]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=24, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=16, bold=True, color=STEEL)
    add_text(s, Cm(6), y + Cm(1.4), SW - Cm(7.5), Cm(2.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 30: クロージング
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, STEEL)
add_rect(s, 0, Cm(7), SW, Cm(0.1), CORAL)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=CORAL, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "「困難事例」とは、\n一人では支えきれない、という合図。",
         size=25, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "解決できないことに立ちすくむより、つながり続けること。\n"
         "一人で抱えるより、つなぎ、分かち合うこと。\n"
         "その積み重ねが、制度のはざまで孤立する誰かを、地域の網で受けとめます。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 31: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 31)

refs = [
    "■ 社会福祉法(重層的支援体制整備事業の根拠法)",
    "■ 厚生労働省「重層的支援体制整備事業」関連の通知・手引き",
    "■ 厚生労働省「地域共生社会」関連資料",
    "■ 8050問題・ひきこもり支援に関する厚生労働省・自治体の資料",
    "■ 高齢者虐待防止法、養護者支援関連の対応マニュアル(セルフネグレクトの記載)",
    "■ 各市町村の重層的支援・包括的相談支援の窓口情報",
    "■ ひきこもり地域支援センター・生活困窮者自立支援・精神保健福祉センター等の情報",
]
y = Cm(4.5)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=13, color=DARK_GRAY)
    y += Cm(1.25)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "ご注意 ・ 学び続けるために",
        "■ 重層的支援体制整備事業は市町村の任意事業です。実施状況・窓口は地域ごとに異なります\n"
        "■ 困難事例への対応は、知識だけでなく「事例を多機関で振り返る」ことを通じて深まります\n"
        "■ 本資料は 2026年5月時点の一般的な情報を基に構成しています",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 32: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, STEEL)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), CORAL)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 事例検討",
         size=18, color=CORAL, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "現場で出会った困難事例を持ち寄り、つなぎ方・支え方を一緒に考えましょう。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/困難事例・複合課題への対応_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
