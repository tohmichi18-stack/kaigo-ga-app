"""主任ケアマネ研修スライド生成スクリプト 第8弾

テーマ: 看取り期のケアマネジメントとACP(人生会議)の支援技術
重点: ターミナルケアマネジメント加算の算定要件 / 家族・多職種との合意形成
トーン: 実務型(支援技術・手順・場面対応)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(看取り・ACP:夕暮れの落ち着いた基調) =====
DUSK        = RGBColor(0x42, 0x3A, 0x5E)   # メイン(夕暮れの紫紺)
DUSK_MID    = RGBColor(0x60, 0x57, 0x80)
DUSK_LT     = RGBColor(0x8C, 0x83, 0xA6)
AMBER       = RGBColor(0xC4, 0x8C, 0x3A)   # アクセント(灯り)
SAGE        = RGBColor(0x52, 0x7C, 0x63)   # 穏やか・OK
ROSE        = RGBColor(0xA8, 0x5C, 0x66)   # 家族・感情
TEAL        = RGBColor(0x2E, 0x6B, 0x73)   # 多職種連携
RED         = RGBColor(0xB0, 0x44, 0x3C)   # 注意・NG
LIGHT_BG    = RGBColor(0xF2, 0xF1, 0xF5)
MID_GRAY    = RGBColor(0xD7, 0xD4, 0xDE)
DARK_GRAY   = RGBColor(0x32, 0x30, 0x38)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xEA, 0xE8, 0xEF)
RED_BG      = RGBColor(0xF7, 0xE6, 0xE4)
AMBER_BG    = RGBColor(0xF9, 0xF0, 0xDD)
SAGE_BG     = RGBColor(0xE5, 0xEC, 0xE7)
ROSE_BG     = RGBColor(0xF6, 0xE8, 0xEA)
TEAL_BG     = RGBColor(0xDF, 0xEA, 0xEA)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), DUSK)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=DUSK)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), AMBER)


def add_footer(slide, page_num, total=33):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(26), Cm(0.5),
             "主任CM研修 / 看取り期のケアマネジメントとACPの支援技術  ※算定要件は最新の告示・通知をご確認ください",
             size=9, color=DUSK_LT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=DUSK_LT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=DUSK_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=DUSK, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=AMBER):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, DUSK)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=accent_color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=33, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color, fontsize=15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(fontsize)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


def add_arrow(slide, left, top, width, height, color):
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    arr.shadow.inherit = False
    return arr


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, DUSK)
add_rect(s, 0, Cm(8), SW, Cm(0.15), AMBER)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), DUSK_LT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.4), SW - Cm(2), Cm(2.5),
         "看取り期のケアマネジメントと\nACP(人生会議)の支援技術",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
add_text(s, Cm(1.5), Cm(9.7), SW - Cm(3), Cm(2),
         "ー 本人の意思を支え、家族・多職種とともに最期の時を支える ー",
         size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(13), SW - Cm(3), Cm(2),
         "■ ACP(人生会議)の意味と、本人の意思を支える技術\n"
         "■ ターミナルケアマネジメント加算の算定要件と実務\n"
         "■ 家族の揺れに寄り添い、多職種と合意を形成する",
         size=14, color=AMBER_BG, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "看取り期のケアマネジメントに求められる役割と、その特有の難しさを理解する",
    "ACP(人生会議)の本来の意味を理解し、「一度決めて終わり」ではないと捉え直す",
    "本人の意思を聴き、支える具体的な技術(問いかけ・受けとめ・推定意思)を持ち帰る",
    "ターミナルケアマネジメント加算の算定要件と、算定漏れ・記録不備の防ぎ方を整理する",
    "家族の揺れ・家族間の相違に寄り添い、多職種と合意を形成する技術を学ぶ",
    "看取り後のグリーフケア・CM自身のケアまでを含めた支援を考える",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), AMBER)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "看取り期のケアマネジメントとは",            "10分"),
    ("2", "ACP(人生会議)の理解",                      "15分"),
    ("3", "ACPの支援技術 ― 意思を聴き、支える",        "20分"),
    ("4", "ターミナルケアマネジメント加算の実務",      "15分"),
    ("5", "家族・多職種との合意形成",                  "20分"),
    ("6", "看取り後の支援 ・ まとめ",                  "10分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, DUSK_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=DUSK)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=AMBER, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "看取り期の\nケアマネジメントとは",
              "Care Management at the End of Life")


# =============================================================
#  SLIDE 5: 在宅看取りの広がり
# =============================================================
s = add_slide()
add_header(s, "1. 看取り期のケアマネジメント", "在宅での看取りを支える時代へ")
add_footer(s, 5)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3),
        "「最期をどこで迎えるか」という問い",
        "「住み慣れた自宅で最期まで過ごしたい」と願う人は少なくない。\n"
        "国も、地域包括ケアの中で在宅での看取りを支える体制づくりを進めている。\n"
        "居宅CMは、その願いを生活の場で支える、数少ない伴走者である。",
        title_bg=DUSK, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(5),
        "看取り期に本人・家族が願うこと",
        "■ 痛み・苦痛のない穏やかな時間\n\n"
        "■ 住み慣れた場所・なじみの人の中で\n\n"
        "■ 自分らしさ・尊厳が保たれること\n\n"
        "■ 家族に過度な負担をかけたくない\n\n"
        "■ やり残したことへの思い・別れの時間",
        title_bg=DUSK_MID, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(5),
        "在宅看取りを可能にする条件",
        "■ 本人・家族の意思と心構え\n\n"
        "■ 24時間対応できる医療(訪問診療・看護)\n\n"
        "■ 介護サービスと家族介護の体制\n\n"
        "■ 急変時の方針の共有(救急搬送の是非)\n\n"
        "■ 多職種が連携し、情報を共有できる関係",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.5), SW - Cm(2.4), Cm(1.2),
         "▶ CMは、これらの条件を「つなぎ」「整え」「揺れを支える」役割を担う。",
         size=13, color=DUSK, bold=True)


# =============================================================
#  SLIDE 6: CMの役割と難しさ
# =============================================================
s = add_slide()
add_header(s, "1. 看取り期のケアマネジメント", "看取り期のCMの役割と、特有の難しさ")
add_footer(s, 6)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.3),
        "看取り期のCMの役割",
        "■ 本人の意思を聴き、ケアに反映する\n\n"
        "■ 状態変化に応じてプランを機動的に見直す\n\n"
        "■ 医療(訪問診療・看護)との連携をつなぐ\n\n"
        "■ 家族の不安・揺れに寄り添う\n\n"
        "■ 急変時・看取りの方針を多職種と共有\n\n"
        "■ 看取り後の家族支援(グリーフケア)",
        title_bg=DUSK_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.3),
        "看取り期ならではの難しさ",
        "■ 時間が限られ、状態が急速に変化する\n\n"
        "■ 本人の意思が確認しづらくなっていく\n\n"
        "■ 家族の感情が揺れ動き、葛藤も生じる\n\n"
        "■ 「正解」がない判断の連続\n\n"
        "■ 医療の比重が高まり、連携が複雑に\n\n"
        "■ CM自身も感情を揺さぶられる",
        title_bg=ROSE, body_bg=ROSE_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "看取り期のケアマネジメントで大切にしたい姿勢",
        "■ 「何をするか」より「誰のための時間か」 ― 主役は最期まで本人\n\n"
        "■ 答えを出すことより、本人・家族の「揺れ」に寄り添い続けること\n\n"
        "■ 一人で抱えない ― 医療・看護・介護のチームで支える\n\n"
        "■ 「できなかったこと」を悔やむより、「共に過ごせた時間」を大切にする",
        title_bg=DUSK, body_size=13)


# =============================================================
#  SLIDE 7: 第2部
# =============================================================
section_title(2, "ACP(人生会議)の理解",
              "Understanding Advance Care Planning")


# =============================================================
#  SLIDE 8: ACPとは
# =============================================================
s = add_slide()
add_header(s, "2. ACPの理解", "ACP(人生会議)とは")
add_footer(s, 8)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3.3),
        "ACP = アドバンス・ケア・プランニング",
        "■ もしものときに備え、本人が望む医療・ケアについて、本人を中心に、家族等や\n"
        "    医療・ケアチームが、繰り返し話し合い、共有する取組み\n"
        "■ 日本では「人生会議」という愛称で普及が図られている",
        title_bg=DUSK, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(5),
        "ACPで大切にすること",
        "■ 主役は「本人」 ― 本人の価値観・希望\n\n"
        "■ 家族等・医療ケアチームと「共に」話す\n\n"
        "■ 一度ではなく「繰り返し」話し合う\n\n"
        "■ 話し合った内容を「共有」する\n\n"
        "■ 「決めること」より「話し合う過程」",
        title_bg=DUSK_MID, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(5),
        "ACPは「書類づくり」ではない",
        "■ 事前指示書を書かせることが目的ではない\n\n"
        "■ 「対話のプロセス」そのものに意味がある\n\n"
        "■ 本人が自分の思いに気づき、\n   大切な人と分かち合うことが核心\n\n"
        "▶ 様式を埋めることをゴールにしない",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.5), SW - Cm(2.4), Cm(1.2),
         "▶ ACPは「最終段階だけの話」ではない。元気なうちから始められる、人生の対話。",
         size=13, color=DUSK, bold=True)


# =============================================================
#  SLIDE 9: ガイドライン
# =============================================================
s = add_slide()
add_header(s, "2. ACPの理解", "意思決定プロセスのガイドライン ― よりどころ")
add_footer(s, 9)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "厚労省「人生の最終段階における医療・ケアの決定プロセスに関するガイドライン」",
        "人生の最終段階の医療・ケアのあり方は、本人の意思決定を基本に進める ―\n"
        "その手順と考え方を示したもの。看取り期の判断に迷ったときの「よりどころ」になる。",
        title_bg=DUSK, body_size=12)

points = [
    ("本人の意思が\n確認できる場合",
     "・本人の意思決定を基本とする\n・本人と医療ケアチームが\n  話し合いを繰り返す\n・本人が意思を伝えられるうちに"),
    ("本人の意思が\n確認できない場合",
     "・家族等が本人の意思を推定\n  → 推定意思を尊重\n・推定できない場合は本人にとって\n  最善の方針を、家族等とチームで"),
    ("共通して\n大切なこと",
     "・複数の専門職からなるチームで\n・話し合いの内容は文書にまとめ\n  共有する\n・心身の状態変化に応じて繰り返す"),
]
col_w = Cm(10.2)
col_h = Cm(7)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(7.3)
colors = [DUSK_MID, ROSE, SAGE]
for i, (title, body) in enumerate(points):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.7), colors[i])
    add_text(s, x, y, col_w, Cm(1.7),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.7), col_w, col_h - Cm(1.7), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(2), col_w - Cm(0.8), col_h - Cm(2.3),
             body, size=11, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(14.7), Cm(31.4), Cm(2.3),
        "ガイドラインの根底にある考え方",
        "■ 「誰かが一人で決める」のではなく「本人を中心に、チームで、話し合って決める」\n"
        "■ 第7弾で学んだ「意思決定支援」が、人生の最終段階でも一貫した土台になる",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 10: ACPの誤解
# =============================================================
s = add_slide()
add_header(s, "2. ACPの理解", "ACPをめぐる「よくある誤解」")
add_footer(s, 10)

col_widths = [Cm(15.7), Cm(15.7)]
headers = ["よくある誤解", "本来の考え方"]
rows = [
    ["延命治療をするか/しないかを\n決めることだ",
     "治療の選択は一部。価値観・どう生きたいか\nを語り合う、もっと広い対話"],
    ["一度決めたら、それで終わり",
     "意思は変わってよい。状態や状況に応じて\n何度も話し合い、見直していく"],
    ["人生の最終段階になってから\n始めるもの",
     "元気なうちから始められる。早く始めるほど\n本人の言葉でじっくり語り合える"],
    ["書類(事前指示書)を\n書くことがゴール",
     "対話のプロセスそのものが大切。\n書類は話し合いの結果を残す手段"],
    ["医師が主導して進めるもの",
     "主役は本人。医療・介護・家族等の\nチームで支える。CMも担い手の一人"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=DUSK, first_col_color=RED,
          row_height=Cm(1.85), body_size=12, first_col_bold=False)

add_box(s, Cm(1.2), Cm(15.3), Cm(31.4), Cm(2.2),
        "主任CMが正したい誤解",
        "■ ACPを「重い手続き」と捉えると、本人も家族も身構えてしまう\n"
        "■ 「これからどう過ごしたい?」という、日々の暮らしの延長の対話として伝える",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)


# =============================================================
#  SLIDE 11: 第3部
# =============================================================
section_title(3, "ACPの支援技術",
              "Skills for Supporting ACP")


# =============================================================
#  SLIDE 12: いつ始めるか
# =============================================================
s = add_slide()
add_header(s, "3. ACPの支援技術", "いつ始めるか ― 「きっかけ」を逃さない")
add_footer(s, 12)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「いつか」ではなく「今日の暮らし」から",
        "改まって「人生会議をしましょう」と切り出すより、日々の関わりの中の\n"
        "自然な「きっかけ」を捉えて、少しずつ話を深めていくほうがうまくいく。",
        title_bg=DUSK, body_size=12)

triggers = [
    ("生活の節目",
     "・要介護度の変化\n・入退院のとき\n・サービスの見直しのとき"),
    ("本人の言葉",
     "・「もう長くないかも」\n・「迷惑をかけたくない」\n・「家にいたい」等のつぶやき"),
    ("身体の変化",
     "・病状の進行・診断\n・食べられなくなってきた\n・できないことが増えてきた"),
    ("家族の不安",
     "・「どうしたらいいか」\n・「もしものとき心配」\n・家族からの相談"),
]
col_w = Cm(7.7)
col_h = Cm(5.8)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.2)
colors = [DUSK_MID, AMBER, TEAL, ROSE]
for i, (title, body) in enumerate(triggers):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), colors[i])
    add_text(s, x, y, col_w, Cm(1.3),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(1.6), col_w - Cm(0.6), col_h - Cm(1.9),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "早く始めることの意味",
        "■ 本人が自分の言葉で語れるうちに始めれば、推定でなく「本人の意思」を残せる\n"
        "■ 「縁起でもない」と避けず、本人が語りたそうなサインを逃さない\n"
        "■ ただし、本人が話したくないときは無理強いしない ― タイミングと本人のペースを尊重",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=13)


# =============================================================
#  SLIDE 13: 何を話すか
# =============================================================
s = add_slide()
add_header(s, "3. ACPの支援技術", "何を話すか ― 対話のテーマ")
add_footer(s, 13)

themes = [
    ("価値観・大切なこと",
     "・どんなことを大切に生きてきたか\n・何があれば心穏やかでいられるか\n・してほしくないこと"),
    ("過ごし方・療養場所",
     "・最期までどこで過ごしたいか\n・家で過ごす上で大切にしたいこと\n・誰と、どんな時間を過ごしたいか"),
    ("医療・ケアの希望",
     "・受けたい/受けたくない医療\n・苦痛をどう和らげたいか\n・急変時にどうしてほしいか"),
    ("大切な人・気がかり",
     "・誰に意思を託したいか(代弁者)\n・伝えておきたいこと\n・やり残したこと・気がかり"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [DUSK_MID, SAGE, TEAL, ROSE]
for i, (title, body) in enumerate(themes):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "話す順番のコツ",
        "■ いきなり「延命治療どうしますか」から入らない ― 本人も家族も身構えてしまう\n"
        "■ まず「価値観・どう過ごしたいか」から。本人の人生・思いを聴くことから始める\n"
        "■ 医療の具体的な選択は、本人の価値観が見えてきた後に、医療職と一緒に",
        title_bg=DUSK, body_size=13)


# =============================================================
#  SLIDE 14: 聴く技術
# =============================================================
s = add_slide()
add_header(s, "3. ACPの支援技術", "聴く技術 ― 本人の思いを引き出す")
add_footer(s, 14)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.5),
        "本人が語りやすい関わり",
        "■ 落ち着いて話せる環境・時間をつくる\n\n"
        "■ 開かれた問いかけ\n   「これからどう過ごしたいですか」\n   「何が一番気がかりですか」\n\n"
        "■ 沈黙を待つ ― 答えを急かさない\n\n"
        "■ 言葉を受けとめ、返す(反映)\n\n"
        "■ 評価・否定をしない",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.5),
        "避けたい関わり",
        "■ 一般論・正解を押しつける\n\n"
        "■ 「そんな弱気にならないで」と励まし、\n   本人の不安にふたをする\n\n"
        "■ 質問攻めにする(尋問になる)\n\n"
        "■ 本人より家族にばかり話を向ける\n\n"
        "■ 「決めさせよう」と焦る",
        title_bg=RED, body_bg=RED_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "「答え」を引き出すより、「思い」に寄り添う",
        "■ ACPの対話は「結論を出す面接」ではない ― 本人の揺れ・迷いに付き合うことが支援\n\n"
        "■ 言葉にならない思いも大切に ― 表情・沈黙・ためらいにも、本人の気持ちが表れる\n\n"
        "■ 「つらい話をしてくれてありがとう」と、語ってくれたこと自体をねぎらう\n\n"
        "■ 第4弾のスーパービジョンの問いかけ技術が、ACPの対話でもそのまま生きる",
        title_bg=DUSK, body_size=13)


# =============================================================
#  SLIDE 15: 意思が確認できない場合
# =============================================================
s = add_slide()
add_header(s, "3. ACPの支援技術", "本人の意思が確認しづらいとき")
add_footer(s, 15)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "認知症の進行・状態の悪化で、本人が意思を表せないとき",
        "「意思が確認できない」=「意思がない」ではない。\n"
        "本人がこれまでどう生きてきたかから、その意思を「推定」して尊重する。",
        title_bg=DUSK, body_size=12)

steps = [
    ("1", "まず本人に向き合う",
     "意思を表せないように見えても、本人に語りかけ、表情・反応から思いを読み取る努力を続ける"),
    ("2", "推定意思を探る",
     "これまでの言動・価値観・人生から「本人ならどう考えるか」を、家族等と一緒に推定する"),
    ("3", "事前の意思表示を確認",
     "過去のACPの記録、本人が語っていた言葉、書き残したものがあれば、よりどころにする"),
    ("4", "最善の方針をチームで",
     "推定もできない場合は、本人にとって何が最善かを、家族等と医療ケアチームで話し合う"),
]
y = Cm(7)
colors = [DUSK_MID, ROSE, AMBER, SAGE]
for i, (num, title, body) in enumerate(steps):
    add_numbered_circle(s, Cm(1.5), y + Cm(0.3), Cm(1.2), num, colors[i])
    add_rect(s, Cm(3.2), y, Cm(7.5), Cm(1.8), colors[i])
    add_text(s, Cm(3.3), y, Cm(7.3), Cm(1.8),
             title, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(10.7), y, SW - Cm(11.9), Cm(1.8), TABLE_BG)
    add_text(s, Cm(11.1), y + Cm(0.15), SW - Cm(12.5), Cm(1.5),
             body, size=11, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    y += Cm(2.1)

add_box(s, Cm(1.2), Cm(15.6), Cm(31.4), Cm(1.6),
        "",
        "▶ 推定意思は「家族の希望」とは違う。「本人ならどうか」を、家族とともに考え続ける。",
        title_bg=WHITE, body_bg=AMBER_BG, title_size=1, body_size=12)


# =============================================================
#  SLIDE 16: 記録と共有
# =============================================================
s = add_slide()
add_header(s, "3. ACPの支援技術", "話し合った内容を「記録」し「共有」する")
add_footer(s, 16)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.3),
        "何を記録するか",
        "■ 話し合った日付・場面・同席者\n\n"
        "■ 本人が語った言葉(できるだけ本人の表現で)\n\n"
        "■ 本人の価値観・希望・気がかり\n\n"
        "■ 療養場所・医療ケアの希望\n\n"
        "■ 意思を託す人(代弁者)\n\n"
        "■ 「事実」と「CMの解釈」は分けて書く",
        title_bg=DUSK_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.3),
        "誰と共有するか",
        "■ 本人・家族等\n\n"
        "■ 主治医・訪問看護\n\n"
        "■ 訪問介護等のサービス事業所\n\n"
        "■ サービス担当者会議で共有\n\n"
        "▶ 共有されない記録は、いざという時に\n   活かされない",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "記録・共有のポイント",
        "■ 意思は変わる ― 「いつ時点の意思か」を明確にし、更新したら最新版を共有する\n\n"
        "■ 急変時に医療職がすぐ参照できる形にしておく(本人宅に情報を置く等の工夫)\n\n"
        "■ 救急隊・搬送先にも伝わるよう、関係者で「どこに・何を」残すかを取り決めておく\n\n"
        "■ 記録は「本人の意思を守る道具」 ― ターミナルCM加算の算定根拠にもなる(第4部)",
        title_bg=DUSK, body_size=13)


# =============================================================
#  SLIDE 17: 第4部
# =============================================================
section_title(4, "ターミナルケア\nマネジメント加算の実務",
              "Terminal Care Management Bonus")


# =============================================================
#  SLIDE 18: 加算の概要と要件
# =============================================================
s = add_slide()
add_header(s, "4. ターミナルCM加算", "算定要件の全体像")
add_footer(s, 18)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "加算の趣旨",
        "末期の利用者が在宅で最期を過ごせるよう、CMが頻回の訪問・多職種連携・\n"
        "看取りまでを一貫して支える業務を評価するもの。",
        title_bg=DUSK, body_size=12)

items = [
    ("①", "対象利用者",
     "末期の悪性腫瘍その他、\n医学的知見に基づき回復の\n見込みがないと診断"),
    ("②", "死亡日・死亡前\n14日以内の訪問",
     "本人・家族の同意のもと、\n2回以上の利用者宅訪問\n(状態変化等の理由が必要)"),
    ("③", "多職種との連携",
     "主治医・訪問看護等と連携し、\n利用者の状態等の情報を\n得ながら支援"),
    ("④", "記録・看取り対応",
     "訪問日・内容・連携状況等を\n記録。死亡日を含めた対応を\n行う"),
]
col_w = Cm(7.7)
col_h = Cm(6.8)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.2)
for i, (num, title, body) in enumerate(items):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2.1), DUSK_MID)
    add_text(s, x, y + Cm(0.15), col_w, Cm(0.8),
             num, size=20, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.15), y + Cm(0.9), col_w - Cm(0.3), Cm(1.1),
             title, size=12, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2.1), col_w, col_h - Cm(2.1), TABLE_BG)
    add_text(s, x + Cm(0.25), y + Cm(2.4), col_w - Cm(0.5), col_h - Cm(2.7),
             body, size=10.5, color=DARK_GRAY, line_spacing=1.35,
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "算定にあたっての注意",
        "■ 単位数・細かな要件は改定で変わりうる ― 最新の告示・解釈通知・Q&Aを必ず確認すること\n"
        "■ 「死亡日 + 死亡前14日以内」の業務が算定の核。この期間の記録が決め手になる",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 19: 実務フローと算定漏れ防止
# =============================================================
s = add_slide()
add_header(s, "4. ターミナルCM加算", "実務フローと、算定漏れ・記録不備の防止")
add_footer(s, 19)

steps = [
    ("末期の確認",
     "主治医意見書・診療情報提供書等で「末期である旨」を文書で確認。本人・家族の在宅看取りの意向を確認・記録"),
    ("ターミナル期の\nプラン再作成",
     "サービス担当者会議を開催。訪問診療・訪問看護との連携を強化し、看取りの方針を明文化"),
    ("頻回訪問と\n多職種共有",
     "状態変化に応じて訪問(2回以上必須)。訪問の都度「訪問理由・状態・対応」を記録。主治医・看護と情報共有"),
    ("死亡日対応・\n記録の確定",
     "死亡日の対応を記録。算定根拠となる書類(末期の確認・訪問記録・連携記録)を整理・確定"),
]
y = Cm(4.4)
colors = [DUSK_MID, TEAL, AMBER, SAGE]
for i, (title, body) in enumerate(steps):
    add_numbered_circle(s, Cm(1.5), y + Cm(0.3), Cm(1.2), i + 1, colors[i])
    add_rect(s, Cm(3.2), y, Cm(6.5), Cm(1.9), colors[i])
    add_text(s, Cm(3.3), y, Cm(6.3), Cm(1.9),
             title, size=12.5, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(9.7), y, SW - Cm(10.9), Cm(1.9), TABLE_BG)
    add_text(s, Cm(10.1), y + Cm(0.15), SW - Cm(11.5), Cm(1.6),
             body, size=11, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    y += Cm(2.15)

add_box(s, Cm(1.2), Cm(13.4), Cm(31.4), Cm(3.7),
        "算定漏れ・指摘を防ぐチェックポイント",
        "■ 末期である旨が「文書」で確認できるか ■ 14日以内に「2回以上」の訪問記録があるか\n"
        "■ 各訪問記録に「訪問理由・状態変化・対応」が明記されているか\n"
        "■ 多職種との連携(情報共有)の記録があるか ■ 死亡日の対応が記録されているか\n"
        "▶ 「やったのに算定していない」「記録が不十分で算定できない」を主任CMが点検する",
        title_bg=DUSK, body_size=12)


# =============================================================
#  SLIDE 20: 医療との連動
# =============================================================
s = add_slide()
add_header(s, "4. ターミナルCM加算", "医療の看取り評価との連動を理解する")
add_footer(s, 20)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "看取りは「介護」と「医療」が連動して支える",
        "在宅看取りでは、CM(介護)と訪問診療・訪問看護(医療)が、それぞれの役割で\n"
        "本人・家族を支える。互いの動きを理解することで、連携がスムーズになる。",
        title_bg=DUSK, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "CM(介護)が担うこと",
        "■ 生活全体のマネジメント\n\n"
        "■ 本人の意思・価値観を聴き、つなぐ\n\n"
        "■ 介護サービスの調整\n\n"
        "■ 家族の不安・揺れに寄り添う\n\n"
        "■ 多職種をつなぐハブ役",
        title_bg=DUSK_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "医療(訪問診療・看護)が担うこと",
        "■ 病状の管理・症状緩和(痛み等)\n\n"
        "■ 医学的な見通しの説明\n\n"
        "■ 死亡確認・看取りの医療対応\n\n"
        "■ 24時間の医療的バックアップ\n\n"
        "■ 急変時の判断・対応",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(2),
         "▶ 医療側にも在宅での看取りを評価するしくみがある ― 介護・医療が「同じ利用者を、\n"
         "▶ それぞれの立場で評価し、支えている」と理解し、対等な連携を心がける。",
         size=13, color=DUSK, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 21: 第5部
# =============================================================
section_title(5, "家族・多職種との\n合意形成",
              "Building Consensus")


# =============================================================
#  SLIDE 22: 家族の揺れ
# =============================================================
s = add_slide()
add_header(s, "5. 合意形成", "家族の「揺れ」を理解する")
add_footer(s, 22)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "揺れるのは当然 ― 家族を「決められない人」と見ない",
        "大切な人の最期を前に、家族の気持ちが揺れ動くのは自然なこと。\n"
        "「揺れ」を問題視せず、揺れに寄り添うことが、合意形成の出発点になる。",
        title_bg=DUSK, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "家族の中で起きていること",
        "■ 「家で看たい」と「自信がない」の間で\n\n"
        "■ 本人の意思と、家族自身の願いの違い\n\n"
        "■ 「何かしてあげたい」という焦り\n\n"
        "■ 予期悲嘆(失う前から始まる悲しみ)\n\n"
        "■ 介護疲れ・睡眠不足・孤立",
        title_bg=ROSE, body_bg=ROSE_BG, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "CMが家族にできること",
        "■ 揺れる気持ちをそのまま受けとめる\n\n"
        "■ 「迷って当然」と伝え、孤立させない\n\n"
        "■ 見通しを具体的に伝え、不安を和らげる\n\n"
        "■ 「いつでも方針は変えられる」と保証する\n\n"
        "■ 家族自身の休息・支援にも目を向ける",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ 「在宅か入院か」を急いで決めさせない。揺れながら進む過程に、CMが伴走する。",
         size=13, color=DUSK, bold=True)


# =============================================================
#  SLIDE 23: 家族間の意見の相違
# =============================================================
s = add_slide()
add_header(s, "5. 合意形成", "家族間で意見が分かれたとき")
add_footer(s, 23)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "よくある場面",
        "同居の家族は「家で看取りたい」、離れて暮らす家族は「入院を」 ―\n"
        "家族の中で意見が割れることは珍しくない。それぞれに事情と思いがある。",
        title_bg=DUSK, body_size=12)

col_widths = [Cm(10.5), Cm(10.5), Cm(10.4)]
headers = ["対応の原則", "具体的な関わり", "避けたいこと"]
rows = [
    ["本人の意思に\n立ち返る",
     "「ご本人はどう望んで\nおられたか」を中心に置く",
     "家族の声の大きさで\n方針が決まること"],
    ["それぞれの思いを\n聴く",
     "各家族の事情・不安・\n罪悪感を個別に受けとめる",
     "一方の家族だけと\n話を進めること"],
    ["情報を共有し\n見通しを揃える",
     "病状・見通し・選択肢を\n家族全員が同じく理解",
     "情報の差が\n対立を深めること"],
    ["話し合いの場を\nつくる",
     "多職種同席のカンファ等で\n冷静に話せる場を設定",
     "CM一人で家族の\n仲裁を抱え込むこと"],
]
add_table(s, Cm(1.2), Cm(7.3), col_widths, headers, rows,
          header_bg=DUSK, first_col_color=DUSK,
          row_height=Cm(1.85), body_size=11)

add_box(s, Cm(1.2), Cm(15.4), Cm(31.4), Cm(1.8),
        "",
        "▶ 家族間の対立の調整は重い。主任CM・多職種・必要なら医療職と分担し、一人で抱えない。",
        title_bg=WHITE, body_bg=AMBER_BG, title_size=1, body_size=12)


# =============================================================
#  SLIDE 24: 多職種との連携
# =============================================================
s = add_slide()
add_header(s, "5. 合意形成", "多職種との連携 ― 看取りを支えるチーム")
add_footer(s, 24)

members = [
    ("主治医・\n訪問診療医",
     "病状管理・症状緩和\n医学的見通し・死亡確認"),
    ("訪問看護師",
     "日々の状態観察・医療処置\n本人家族の心身の支え"),
    ("訪問介護員等",
     "日常生活の支援\n本人の様子の変化に気づく"),
    ("薬剤師",
     "疼痛緩和の薬剤管理\n服薬しやすい工夫"),
    ("CM",
     "全体のマネジメント\n意思をつなぎ、揺れを支える"),
]
col_w = Cm(6.2)
col_h = Cm(5.5)
gap = Cm(0.2)
start_x = (SW - (col_w * 5 + gap * 4)) // 2
y = Cm(4.5)
colors = [TEAL, DUSK_MID, SAGE, ROSE, AMBER]
for i, (title, body) in enumerate(members):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.8), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(1.8),
             title, size=12.5, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.8), col_w, col_h - Cm(1.8), TABLE_BG)
    add_text(s, x + Cm(0.2), y + Cm(2.1), col_w - Cm(0.4), col_h - Cm(2.4),
             body, size=10, color=DARK_GRAY, line_spacing=1.35,
             align=PP_ALIGN.CENTER)

add_box(s, Cm(1.2), Cm(10.5), Cm(31.4), Cm(6.5),
        "連携を機能させるために",
        "■ 平時から「顔の見える関係」を ― 看取り期になって初めて連絡、では間に合わない\n\n"
        "■ 情報共有の手段を決めておく(電話・FAX・ICT・連絡ノート等)― 急変時に迷わないように\n\n"
        "■ 「誰が・何を担うか」「急変時に誰に連絡するか」を、文書で明確にして共有する\n\n"
        "■ 状態は急変する ― 小さな変化も「些細かも」と思わず、チームで共有する\n\n"
        "■ CMは指示する立場ではなく「つなぐ」立場 ― 各職種の専門性を尊重し、対等に連携する",
        title_bg=DUSK, body_size=13)


# =============================================================
#  SLIDE 25: カンファレンスの活用
# =============================================================
s = add_slide()
add_header(s, "5. 合意形成", "サービス担当者会議・カンファレンスを合意の場に")
add_footer(s, 25)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「場」をつくることが合意形成の鍵",
        "本人・家族・多職種が同じ情報を共有し、一緒に方針を確認する「場」をつくる。\n"
        "CMは、その場のファシリテーター(第4弾の技術がここで生きる)。",
        title_bg=DUSK, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "会議で確認したいこと",
        "■ 本人の意思・希望(最優先で)\n\n"
        "■ 病状の見通しの共有\n\n"
        "■ 療養場所・看取りの方針\n\n"
        "■ 急変時の対応(救急搬送の是非)\n\n"
        "■ 各職種の役割分担・連絡体制",
        title_bg=DUSK_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "ファシリテーターとしてのCM",
        "■ 本人・家族が安心して話せる雰囲気を\n\n"
        "■ 専門用語をかみくだき、理解を揃える\n\n"
        "■ 立場の弱い声(本人・家族)を引き出す\n\n"
        "■ 「決まったこと」「保留」を明確にする\n\n"
        "■ 方針は変えてよいと、必ず添える",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "急変時・救急搬送をめぐる合意",
        "■ 「急に容体が変わったとき、どうするか」を、落ち着いているうちに話し合っておく\n"
        "■ 在宅看取りの方針でも、いざという時に家族が動揺し救急車を呼ぶことは自然に起こる\n"
        "■ 「呼んでもよい」と保証しつつ、本人の意思・連絡先・対応を関係者・救急隊にも伝わる形で共有",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 26: 第6部
# =============================================================
section_title(6, "看取り後の支援 ・ まとめ",
              "After the Passing & Wrap Up")


# =============================================================
#  SLIDE 27: グリーフケア
# =============================================================
s = add_slide()
add_header(s, "6. 看取り後の支援", "看取り後の支援 ― グリーフケア")
add_footer(s, 27)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "看取りは「死亡」で終わらない",
        "大切な人を亡くした家族の悲嘆(グリーフ)に寄り添うことも、看取り支援の一部。\n"
        "ターミナルケアマネジメント加算でも、死亡日を含めた対応が想定されている。",
        title_bg=DUSK, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "遺族(家族)への関わり",
        "■ ねぎらいの言葉をかける\n   「よく支えてこられましたね」\n\n"
        "■ 悲しみを否定せず、そのまま受けとめる\n\n"
        "■ 自責感(「もっとできたのでは」)に\n   寄り添う\n\n"
        "■ 必要に応じ、専門的な支援につなぐ",
        title_bg=ROSE, body_bg=ROSE_BG, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "デスカンファレンス(振り返り)",
        "■ 関わった多職種で看取りを振り返る\n\n"
        "■ 「何がよかったか」「何ができたか」\n   を中心に\n\n"
        "■ 反省点も、次への学びとして共有\n\n"
        "■ 関わった職員の気持ちもケアする場に",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ 振り返りは「評価・反省会」ではなく、チームが次の看取りに向かうための学びと癒やしの場。",
         size=13, color=DUSK, bold=True)


# =============================================================
#  SLIDE 28: CM自身のケア
# =============================================================
s = add_slide()
add_header(s, "6. 看取り後の支援", "支える人を支える ― CM自身のケア")
add_footer(s, 28)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "看取りは、CM自身の心も揺さぶる",
        "看取りに関わるCMもまた、悲しみ・無力感・自責の念を抱える。\n"
        "「支える人」自身がケアされなければ、支援を続けることはできない。",
        title_bg=DUSK, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "CMが抱えやすいもの",
        "■ 「もっとできたのでは」という自責\n\n"
        "■ 無力感・喪失感\n\n"
        "■ 感情を出せずに溜め込む\n\n"
        "■ 看取りが続くときの消耗\n\n"
        "■ 「専門職だから平気」という思い込み",
        title_bg=RED, body_bg=RED_BG, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "主任CMができること",
        "■ 看取り後、担当CMの気持ちに目を向ける\n\n"
        "■ 振り返りの場で、感情を語れるようにする\n\n"
        "■ 「よく支えた」とねぎらう\n\n"
        "■ 一人に看取りが偏らないよう配慮\n\n"
        "■ 支持的スーパービジョン(第4弾)を活かす",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ CM自身のケアは「甘え」ではない。支援を続けるための、専門職としての必要条件。",
         size=13, color=DUSK, bold=True)


# =============================================================
#  SLIDE 29: 主任CMの役割
# =============================================================
s = add_slide()
add_header(s, "6. まとめ", "看取り支援における主任CMの役割")
add_footer(s, 29)

roles = [
    ("実践を伝える",
     "・ACPの対話、看取りの関わりを\n  同行・SVを通じて後輩に伝える\n・経験の浅いCMを一人にしない"),
    ("仕組みを整える",
     "・加算の算定漏れ・記録不備を点検\n・多職種連携の体制を整える\n・看取り後の振り返りを定例化"),
    ("揺れを支える",
     "・本人・家族の揺れに伴走する\n・家族間の調整を分担して支える\n・担当CMの気持ちもケアする"),
    ("地域につなぐ",
     "・在宅看取りを支える地域資源を把握\n・医療・介護の顔の見える関係づくり\n・地域ケア会議で課題を共有"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [DUSK_MID, TEAL, ROSE, SAGE]
for i, (title, body) in enumerate(roles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "看取り支援は「チームの力」と「主任CMの支え」で成り立つ",
        "■ 看取りは、一人のCMが背負うには重すぎる ― チームで支える仕組みを主任CMがつくる\n"
        "■ ACP・合意形成・グリーフケアは「技術」 ― 経験を言葉にし、後輩に伝え、地域に広げる\n"
        "■ 本シリーズで学んだ意思決定支援・SV・ファシリテーションのすべてが、看取りの場で生きる",
        title_bg=DUSK, body_size=13)


# =============================================================
#  SLIDE 30: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 30)

actions = [
    ("Action 1", "「きっかけ」を捉えてACPの対話を始める",
     "・担当ケースの中で、価値観・これからの過ごし方を聴ける場面を意識する\n"
     "・改まった会議でなく、日々の対話の延長として始める"),
    ("Action 2", "ターミナルCM加算の記録を点検する",
     "・看取りに関わったケースの記録が、算定要件を満たす形で残っているか確認する\n"
     "・「末期の文書確認・14日以内2回訪問・連携・死亡日対応」をチェックする"),
    ("Action 3", "多職種との連携体制を見直す",
     "・看取りに関わる主治医・訪問看護等との連絡手段・役割分担を再確認する\n"
     "・急変時・救急搬送をめぐる方針を、関係者で共有できているか点検する"),
]
y = Cm(4.5)
colors = [DUSK_MID, AMBER, TEAL]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=24, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=16, bold=True, color=DUSK)
    add_text(s, Cm(6), y + Cm(1.4), SW - Cm(7.5), Cm(2.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 31: クロージング
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, DUSK)
add_rect(s, 0, Cm(7), SW, Cm(0.1), AMBER)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=AMBER, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "最期の時間をどう過ごすかは、\nその人の人生の、最後の自己決定。",
         size=25, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "答えのない問いに、本人と家族とともに向き合う日々。\n"
         "あなたが交わした対話、揺れに寄り添った時間、つないだ多職種の手。\n"
         "その一つひとつが、誰かの「その人らしい最期」を、確かに支えています。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 32: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 32)

refs = [
    "■ 厚生労働省「人生の最終段階における医療・ケアの決定プロセスに関するガイドライン」",
    "■ 厚生労働省「人生会議(ACP)」普及啓発関連資料",
    "■ 厚生労働省「認知症の人の日常生活・社会生活における意思決定支援ガイドライン」",
    "■ 厚生労働省「介護報酬改定」関係告示・通知(ターミナルケアマネジメント加算)",
    "■ 介護給付費単位数表・算定基準告示・解釈通知・Q&A",
    "■ 在宅看取り・グリーフケアに関する専門書・研修資料",
    "■ 日本介護支援専門員協会等による看取り支援関連の研修資料",
]
y = Cm(4.5)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=13, color=DARK_GRAY)
    y += Cm(1.25)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "ご注意 ・ 学び続けるために",
        "■ ターミナルケアマネジメント加算の単位数・要件は改定で変わります ―\n"
        "    算定にあたっては最新の告示・解釈通知・Q&A、保険者の指導内容を必ずご確認ください\n"
        "■ 看取り支援は、知識だけでなく「一つひとつの看取りの振り返り」を通じて深まります\n"
        "■ 本資料は 2026年5月時点の一般的な情報を基に構成しています",
        title_bg=AMBER, body_bg=AMBER_BG, body_size=12)


# =============================================================
#  SLIDE 33: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, DUSK)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), AMBER)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 事例検討",
         size=18, color=AMBER, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "現場で関わった看取りの経験を持ち寄り、ともに学びを深めましょう。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/看取り期のケアマネジメントとACP支援技術_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
