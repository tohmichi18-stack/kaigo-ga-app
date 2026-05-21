"""予習スライド用のデザイン基盤（python-pptx ヘルパー）。

プロフェッショナル基調・当日メモ用余白を意識したレイアウト関数群。
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- カラーパレット（落ち着いた紺×ティールの専門資料調）----
NAVY = RGBColor(0x1F, 0x3A, 0x5F)      # 主見出し
NAVY_DARK = RGBColor(0x15, 0x29, 0x44)
TEAL = RGBColor(0x2E, 0x8B, 0x8B)      # アクセント
TEAL_LIGHT = RGBColor(0xE3, 0xF0, 0xF0)
GOLD = RGBColor(0xC9, 0x9A, 0x3A)      # 強調（質問・注意）
INK = RGBColor(0x26, 0x2B, 0x33)       # 本文
GRAY = RGBColor(0x5C, 0x63, 0x6E)      # 補足
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)     # 背景パネル
LINE = RGBColor(0xD2, 0xD8, 0xDF)      # 罫線
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MEMO_BG = RGBColor(0xFC, 0xFB, 0xF4)   # メモ欄の地色
MEMO_LINE = RGBColor(0xE4, 0xDD, 0xC4)

JP_FONT = "Meiryo"
JP_FONT_BOLD = "Meiryo"

# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _set_font(run, size, color, bold=False, font=JP_FONT, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    # 日本語(EastAsia)フォントも明示
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', font)


def add_text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.15, wrap=True,
             space_after=2):
    """runs: [(text, size, color, bold), ...] を1段落として配置。
    runs が [[...],[...]] のリストのリストなら複数段落。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    if runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for seg in para:
            text, size, color, bold = seg[0], seg[1], seg[2], seg[3]
            font = seg[4] if len(seg) > 4 else JP_FONT
            r = p.add_run()
            r.text = text
            _set_font(r, size, color, bold, font)
    return box


def add_rect(slide, left, top, width, height, fill, line=None,
             line_w=0.75, shadow=False, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    if radius is not None:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def page_bg(slide, color=WHITE):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)


def header(slide, kicker, title, number):
    """標準コンテンツスライドの上部見出し。戻り値: 本文開始 top(Emu)。"""
    # 左の縦アクセントバー
    add_rect(slide, Inches(0.55), Inches(0.52), Inches(0.10), Inches(0.92), TEAL)
    add_text(slide, Inches(0.80), Inches(0.50), Inches(9.5), Inches(0.32),
             [(kicker, 13, TEAL, True)])
    add_text(slide, Inches(0.80), Inches(0.78), Inches(11.4), Inches(0.74),
             [(title, 27, NAVY, True)])
    # ヘッダ下の細い罫線
    add_rect(slide, Inches(0.80), Inches(1.62), Inches(11.73), Pt(1.4), LINE)
    # ページ番号
    add_text(slide, Inches(12.4), Inches(7.04), Inches(0.7), Inches(0.32),
             [(f"{number:02d}", 11, GRAY, True)], align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.80), Inches(7.04), Inches(7.0), Inches(0.32),
             [("令和8年度 第3回 介護支援専門員研修会｜予習資料", 9, GRAY, False)])
    return Inches(1.86)


def footer_note(slide, text):
    add_text(slide, Inches(0.80), Inches(7.04), Inches(9.5), Inches(0.32),
             [(text, 9, GRAY, False)])


def chip(slide, left, top, text, fill, fg=WHITE, width=Inches(1.6)):
    add_rect(slide, left, top, width, Inches(0.34), fill, radius=0.5)
    add_text(slide, left, top + Emu(10000), width, Inches(0.30),
             [(text, 11, fg, True)], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)


def memo_panel(slide, left, top, width, height, label="当日メモ"):
    """当日メモ用の余白パネル（罫線入り）。"""
    add_rect(slide, left, top, width, height, MEMO_BG, line=MEMO_LINE, line_w=1.0)
    add_text(slide, left + Inches(0.12), top + Inches(0.05),
             Inches(2.0), Inches(0.3),
             [("✎ " + label, 10, RGBColor(0x9A, 0x86, 0x3A), True)])
    # 横罫線
    n = max(1, int((height - Inches(0.45)) / Inches(0.42)))
    y = top + Inches(0.52)
    while y < top + height - Inches(0.10):
        add_rect(slide, left + Inches(0.15), y, width - Inches(0.30), Pt(0.75),
                 MEMO_LINE)
        y += Inches(0.42)


def bullet_block(slide, left, top, width, items, gap=0.06,
                 body_size=14, lead_color=TEAL):
    """items: [(見出し, 本文)] or [文字列]。簡易箇条書き。"""
    y = top
    for it in items:
        if isinstance(it, tuple):
            head, body = it
        else:
            head, body = None, it
        add_rect(slide, left, y + Inches(0.07), Inches(0.12), Inches(0.12),
                 lead_color)
        runs = []
        if head:
            runs.append((head + "　", body_size, NAVY, True))
        runs.append((body, body_size, INK, False))
        box = add_text(slide, left + Inches(0.26), y, width - Inches(0.26),
                       Inches(0.6), [runs], line_spacing=1.18)
        # 概算で次の y を決める（2行想定の余裕）
        approx_lines = max(1, int((len(head or '') + len(body)) /
                                  (width.inches * 3.2)) + 1)
        y += Inches(0.30 * approx_lines + gap)
    return y
