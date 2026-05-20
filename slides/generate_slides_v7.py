"""主任ケアマネ研修スライド生成スクリプト 第7弾

テーマ: 認知症基本法施行後のケアマネジメント
重点: 本人の意思決定支援 / 認知症の人と家族の地域共生に向けた実践
トーン: 実務型(理念の理解 + 実践の手順・視点)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(認知症シンボル=オレンジ基調 + 共生のグリーン) =====
ORANGE      = RGBColor(0xC2, 0x57, 0x1C)   # メイン(認知症シンボルカラー)
ORANGE_MID  = RGBColor(0xD9, 0x7B, 0x36)
ORANGE_LT   = RGBColor(0xE4, 0xA4, 0x6E)
GREEN       = RGBColor(0x3C, 0x7A, 0x4E)   # 共生・地域
TEAL        = RGBColor(0x2C, 0x6E, 0x73)   # 連携
GOLD        = RGBColor(0xC9, 0x9A, 0x2E)   # ヒント・希望
BROWN       = RGBColor(0x6B, 0x4A, 0x33)   # 落ち着き
RED         = RGBColor(0xB5, 0x3F, 0x33)   # 注意・NG
LIGHT_BG    = RGBColor(0xF6, 0xF1, 0xEB)
MID_GRAY    = RGBColor(0xDD, 0xD6, 0xCD)
DARK_GRAY   = RGBColor(0x35, 0x30, 0x2B)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xF0, 0xE9, 0xDF)
RED_BG      = RGBColor(0xF8, 0xE6, 0xE3)
GOLD_BG     = RGBColor(0xFA, 0xF1, 0xDA)
GREEN_BG    = RGBColor(0xE4, 0xED, 0xE6)
ORANGE_BG   = RGBColor(0xFB, 0xEC, 0xDC)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), ORANGE)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=ORANGE)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), GREEN)


def add_footer(slide, page_num, total=32):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(26), Cm(0.5),
             "主任CM研修 / 認知症基本法施行後のケアマネジメント",
             size=9, color=ORANGE_LT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=ORANGE_LT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=ORANGE_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=ORANGE, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=GREEN):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, ORANGE)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=33, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=GOLD_BG, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color, fontsize=15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(fontsize)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, ORANGE)
add_rect(s, 0, Cm(8), SW, Cm(0.15), GREEN)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), ORANGE_LT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=GOLD_BG, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.4), SW - Cm(2), Cm(2.5),
         "認知症基本法施行後の\nケアマネジメント",
         size=32, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
add_text(s, Cm(1.5), Cm(9.7), SW - Cm(3), Cm(2),
         "ー 本人の意思を中心に、認知症の人と家族の地域共生を支える ー",
         size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(13), SW - Cm(3), Cm(2),
         "■ 認知症基本法の理念とケアマネジメントへの影響\n"
         "■ 本人の意思決定支援 ― 形成・表明・実現を支える\n"
         "■ 認知症の人と家族の、地域共生に向けた実践",
         size=14, color=GOLD_BG, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=LIGHT_BG, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "認知症基本法の理念と、それがケアマネジメントに求めるものを理解する",
    "「支援される人」から「権利の主体」へ ― 認知症観の転換を捉える",
    "意思決定支援の3つのプロセス(形成・表明・実現)を実践に移せる",
    "診断直後から進行期まで、各段階に応じた支援の視点を持つ",
    "認知症の人と家族が地域で共生するための、地域資源と実践を知る",
    "主任CMとして、本人中心の実践を事業所・地域に広げる視点を持つ",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), GREEN)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "認知症基本法の理解 ― 共生社会への道",       "10分"),
    ("2", "ケアマネジメントへの影響 ― 認知症観の転換", "10分"),
    ("3", "本人の意思決定支援 ― 形成・表明・実現",      "25分"),
    ("4", "段階に応じた認知症の人と家族への実践",       "20分"),
    ("5", "地域共生に向けた実践と地域資源",             "20分"),
    ("6", "まとめ ― 主任CMが持ち帰るもの",              " 5分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, ORANGE_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=ORANGE)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=GREEN, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "認知症基本法の理解",
              "Understanding the Basic Act on Dementia")


# =============================================================
#  SLIDE 5: 認知症基本法とは
# =============================================================
s = add_slide()
add_header(s, "1. 認知症基本法の理解", "認知症基本法とは")
add_footer(s, 5)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3.3),
        "正式名称と施行",
        "■ 「共生社会の実現を推進するための認知症基本法」\n"
        "■ 令和6年(2024年)1月1日 施行\n"
        "■ 認知症の人が尊厳を保持し、希望を持って暮らせる「共生社会」の実現を目的とする",
        title_bg=ORANGE, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(4.5),
        "法律が定める主な事項",
        "■ 基本理念(7つ)\n\n"
        "■ 国・地方公共団体・国民・事業者の責務\n\n"
        "■ 認知症施策推進基本計画の策定\n\n"
        "■ 認知症の日(9月21日)・認知症月間(9月)",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(4.5),
        "「新しい認知症観」への転換",
        "■ 認知症は「誰もがなりうるもの」\n\n"
        "■ 認知症になっても「希望を持って\n   自分らしく暮らし続けられる」\n\n"
        "■ 認知症の人は「支えられるだけの存在」\n   ではなく、社会の対等な構成員",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13), SW - Cm(2.4), Cm(2),
         "▶ 「認知症の人のための法律」であると同時に、「社会全体のあり方を変える法律」。\n"
         "▶ ケアマネジメントの実践も、この理念を土台に問い直すことが求められる。",
         size=13, color=ORANGE, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 6: 7つの基本理念
# =============================================================
s = add_slide()
add_header(s, "1. 認知症基本法の理解", "7つの基本理念")
add_footer(s, 6)

principles = [
    ("1", "自らの意思による生活",
     "認知症の人が基本的人権を持つ個人として、\n自らの意思で日常生活・社会生活を営める"),
    ("2", "正しい知識と理解",
     "国民が認知症に関する正しい知識と理解を\n深められるようにする"),
    ("3", "社会的障壁の除去",
     "障壁を取り除き、対等な構成員として\n地域で安全・安心に自立して暮らせる"),
    ("4", "切れ目のないサービス",
     "本人の意向を尊重し、保健・医療・福祉の\nサービスが切れ目なく提供される"),
    ("5", "本人と家族への支援",
     "認知症の人と家族等が、地域で\n安心して暮らせるよう支援する"),
    ("6", "研究等の推進",
     "予防・診断・治療等の研究を推進し、\nその成果を広く活用する"),
    ("7", "分野を超えた総合的取組",
     "教育・地域づくり・雇用・保健・医療・\n福祉等が連携した総合的な取組"),
]
col_w = Cm(10.2)
gap_x = Cm(0.35)
y_start = Cm(4.4)
row_h = Cm(3.7)
for i, (num, title, body) in enumerate(principles):
    col = i % 3
    row = i // 3
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    if i == 6:
        x = Cm(1.2) + (col_w + gap_x)  # center the last one
    add_rect(s, x, y, col_w, row_h, TABLE_BG, line_color=MID_GRAY)
    add_numbered_circle(s, x + Cm(0.25), y + Cm(0.25), Cm(1), num, ORANGE_MID, fontsize=13)
    add_text(s, x + Cm(1.4), y + Cm(0.3), col_w - Cm(1.6), Cm(0.9),
             title, size=12, bold=True, color=ORANGE)
    add_text(s, x + Cm(0.4), y + Cm(1.3), col_w - Cm(0.8), row_h - Cm(1.5),
             body, size=10, color=DARK_GRAY, line_spacing=1.3)

add_text(s, Cm(1.2), Cm(16.5), SW - Cm(2.4), Cm(1),
         "▶ 共通して流れるのは「本人の意思の尊重」と「地域での共生」 ― 本日の研修の二本柱。",
         size=12, color=GREEN, bold=True)


# =============================================================
#  SLIDE 7: 共生とは
# =============================================================
s = add_slide()
add_header(s, "1. 認知症基本法の理解", "「共生社会」とは何か")
add_footer(s, 7)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "基本法が掲げる「共生社会」",
        "認知症の人を含めた国民一人ひとりが、互いに人格と個性を尊重し合いながら\n"
        "共生する活力ある社会 ― 「認知症の人を支える社会」から一歩進んだ考え方。",
        title_bg=ORANGE, body_size=13)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "これまで(支える側 / 支えられる側)",
        "■ 認知症の人 = ケアの「対象」\n\n"
        "■ 「できなくなったこと」に注目\n\n"
        "■ 周囲が「守る」「管理する」\n\n"
        "■ 本人不在で物事が決まりがち\n\n"
        "▶ 善意であっても、本人の力を奪うことに",
        title_bg=BROWN, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "これから(共に生きる)",
        "■ 認知症の人 = 社会の対等な構成員\n\n"
        "■ 「できること」「強み」に注目\n\n"
        "■ 本人が「役割」と「出番」を持つ\n\n"
        "■ 本人の声を中心に物事が決まる\n\n"
        "▶ 本人の力を活かし、共に社会をつくる",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(2),
         "▶ 共生とは「認知症の人がいてもいい社会」ではなく「認知症の人と共につくる社会」。\n"
         "▶ ケアマネジメントも「本人のために」から「本人と共に」へ。",
         size=13, color=ORANGE, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 8: 第2部
# =============================================================
section_title(2, "ケアマネジメントへの影響",
              "Impact on Care Management")


# =============================================================
#  SLIDE 9: 認知症観の転換
# =============================================================
s = add_slide()
add_header(s, "2. ケアマネジメントへの影響", "認知症観の転換 ― ケアマネジメントが変わる")
add_footer(s, 9)

col_widths = [Cm(8), Cm(11.7), Cm(11.7)]
headers = ["観点", "これまでの傾向", "基本法が求める方向"]
rows = [
    ["主体は誰か", "支援者・家族が決める", "本人が決める(支援を受けて)"],
    ["注目する点", "できないこと・リスク", "できること・強み・希望"],
    ["アセスメント", "心身機能・課題の把握中心", "本人の思い・歴史・価値観も"],
    ["目標設定", "安全・現状維持が中心", "本人が望む暮らし・役割・出番"],
    ["家族の位置", "「介護力」として捉える", "家族も支援を受ける対象"],
    ["地域の役割", "サービスの提供場所", "共に生きる場・つながりの場"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=ORANGE, first_col_color=ORANGE,
          row_height=Cm(1.5), body_size=12)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "主任CMが意識したいこと",
        "■ 「本人のため」と思った支援が、本人から決定権・役割・自信を奪っていないか\n"
        "■ アセスメント様式の「課題」欄だけでなく、「本人の願い・強み」を必ず聴き取り、書く\n"
        "■ ケアプランの主語を「本人」にする ― 「○○させる」でなく「○○したい/する」",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)


# =============================================================
#  SLIDE 10: 本人視点のアセスメント
# =============================================================
s = add_slide()
add_header(s, "2. ケアマネジメントへの影響", "本人を理解する ― パーソン・センタードな視点")
add_footer(s, 10)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「認知症の人」ではなく「認知症とともに生きる、その人」",
        "診断名や症状だけで人を見ない。その人の歴史・価値観・関係性・好み ―\n"
        "「その人らしさ」を理解することが、すべての支援の出発点になる。",
        title_bg=ORANGE, body_size=12)

aspects = [
    ("生活歴・人生史",
     "・どんな人生を歩んできたか\n・仕事・役割・誇りにしてきたこと\n・大切にしてきた習慣・場所"),
    ("価値観・好み",
     "・何を大切にしているか\n・好きなこと・苦手なこと\n・譲れないこだわり"),
    ("関係性",
     "・家族・友人・地域とのつながり\n・支えとなる人は誰か\n・本人が安心できる人・場"),
    ("今の思い・希望",
     "・これからどう暮らしたいか\n・不安に思っていること\n・やってみたいこと・出番"),
]
col_w = Cm(7.7)
col_h = Cm(6.3)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.2)
colors = [BROWN, ORANGE_MID, TEAL, GREEN]
for i, (title, body) in enumerate(aspects):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), colors[i])
    add_text(s, x, y, col_w, Cm(1.3),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(1.6), col_w - Cm(0.6), col_h - Cm(1.9),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "情報は「本人から」「継続的に」",
        "■ 家族や記録からの情報も大切だが、まず「本人自身」に聴く ― 本人が一番の情報源\n"
        "■ 一度のアセスメントで終わらせない ― 関係を重ねる中で、少しずつ語られることがある",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)


# =============================================================
#  SLIDE 11: 第3部
# =============================================================
section_title(3, "本人の意思決定支援",
              "Supporting Decision-Making")


# =============================================================
#  SLIDE 12: 意思決定支援とは
# =============================================================
s = add_slide()
add_header(s, "3. 意思決定支援", "意思決定支援とは ― 「代行」ではなく「支援」")
add_footer(s, 12)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(3),
        "意思決定支援とは",
        "認知症等により意思決定が困難に見える人に対し、本人が自ら意思決定できるよう\n"
        "支援する一連の行為。厚労省「認知症の人の日常生活・社会生活における\n"
        "意思決定支援ガイドライン」が、その考え方と実践を示している。",
        title_bg=ORANGE, body_size=13)

add_box(s, Cm(1.2), Cm(8), Cm(15.5), Cm(5),
        "大前提となる考え方",
        "■ 認知症があっても、本人には意思がある\n\n"
        "■ 「意思決定できない」と決めつけない\n\n"
        "■ 適切な支援があれば、多くの場面で\n   本人は自ら決められる\n\n"
        "■ 決めるのは本人、支えるのが支援者",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_box(s, Cm(17.5), Cm(8), Cm(15), Cm(5),
        "「代行決定」との違い",
        "■ 意思決定支援:本人が決められるよう支える\n\n"
        "■ 代行決定:本人に代わって他者が決める\n\n"
        "■ まず徹底して意思決定支援を尽くす\n\n"
        "▶ 代行決定は「最後の手段」。安易に\n   家族・支援者が決めてしまわない",
        title_bg=BROWN, body_size=12)

add_text(s, Cm(1.2), Cm(13.5), SW - Cm(2.4), Cm(1.5),
         "▶ 「どうせ分からない」「決められない」という思い込みが、本人から意思決定の機会を奪う。\n"
         "▶ 主任CMは、この思い込みに気づき、問い直す役割を持つ。",
         size=13, color=ORANGE, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 13: 意思決定支援の基本原則
# =============================================================
s = add_slide()
add_header(s, "3. 意思決定支援", "意思決定支援ガイドラインの基本原則")
add_footer(s, 13)

principles = [
    ("原則1",
     "本人の意思の尊重",
     "・本人の自己決定を尊重する\n・本人が自ら意思決定できるよう支援\n・決定を他者が安易に覆さない\n・時間をかけ、本人のペースを大切に"),
    ("原則2",
     "本人の意思決定能力への配慮",
     "・能力は「ある/ない」の二択ではない\n・説明の仕方・環境次第で発揮される\n・能力を固定的に判断しない\n・場面ごとに丁寧に見極める"),
    ("原則3",
     "チームによる早期からの\n継続的支援",
     "・多職種・家族等のチームで支える\n・診断後の早期から関わる\n・本人の状態変化に応じて継続的に\n・支援の経過を共有・記録する"),
]
col_w = Cm(10.2)
col_h = Cm(9.5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(4.5)
colors = [ORANGE_MID, TEAL, GREEN]
for i, (label, title, body) in enumerate(principles):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2.2), colors[i])
    add_text(s, x + Cm(0.2), y + Cm(0.2), col_w - Cm(0.4), Cm(0.7),
             label, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.2), y + Cm(0.85), col_w - Cm(0.4), Cm(1.2),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2.2), col_w, col_h - Cm(2.2), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(2.6), col_w - Cm(0.8), col_h - Cm(3),
             body, size=12, color=DARK_GRAY, line_spacing=1.6)

add_box(s, Cm(1.2), Cm(15), Cm(31.4), Cm(2.2),
        "意思決定能力の捉え方",
        "■ 「認知症だから決められない」ではない ― 説明をやさしく、選択肢を具体的に、環境を整えれば\n"
        "    決められることは多い ■ 「能力がない」のではなく「支援が足りていない」と考える",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)


# =============================================================
#  SLIDE 14: 意思決定支援の3プロセス
# =============================================================
s = add_slide()
add_header(s, "3. 意思決定支援", "意思決定支援の3つのプロセス")
add_footer(s, 14)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2),
        "「決める」は一瞬ではなく、3つの段階の積み重ね",
        "意思決定支援は「形成 → 表明 → 実現」の流れで本人を支えていく。",
        title_bg=ORANGE, body_size=12)

procs = [
    ("意思形成支援",
     "本人が「考え、思いを持つ」のを支える",
     "・分かりやすい情報提供\n  (やさしい言葉・図・実物)\n・選択肢を具体的に示す\n・落ち着いた環境・十分な時間\n・本人が信頼する人の同席"),
    ("意思表明支援",
     "本人が「思いを表す」のを支える",
     "・本人のペースを待つ\n・言葉以外の表現も受けとめる\n  (表情・態度・行動)\n・誘導・急かしをしない\n・複数回・場面を変えて確認"),
    ("意思実現支援",
     "本人の意思を「実現する」のを支える",
     "・本人の意思を生活・ケアに反映\n・本人の能力を活かす\n・チーム・社会資源で後押し\n・実現後も振り返り、見直す"),
]
col_w = Cm(10.2)
col_h = Cm(8.5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(6.7)
colors = [ORANGE_MID, TEAL, GREEN]
for i, (title, sub, body) in enumerate(procs):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2), colors[i])
    add_text(s, x + Cm(0.2), y + Cm(0.25), col_w - Cm(0.4), Cm(0.8),
             f"STEP {i+1}  {title}", size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.2), y + Cm(1.05), col_w - Cm(0.4), Cm(0.8),
             sub, size=10, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, x, y + Cm(2), col_w, col_h - Cm(2), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(2.3), col_w - Cm(0.8), col_h - Cm(2.6),
             body, size=11, color=DARK_GRAY, line_spacing=1.5)
    if i < 2:
        arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                 x + col_w + Cm(0.02), y + Cm(0.5),
                                 Cm(0.36), Cm(1))
        arr.fill.solid(); arr.fill.fore_color.rgb = GOLD
        arr.line.fill.background(); arr.shadow.inherit = False

add_text(s, Cm(1.2), Cm(16.5), SW - Cm(2.4), Cm(1),
         "▶ どの段階でつまずいているかを見極め、その段階に合った支援を行う。",
         size=12, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 15: 実践のポイント
# =============================================================
s = add_slide()
add_header(s, "3. 意思決定支援", "意思決定支援の実践 ― してよいこと/いけないこと")
add_footer(s, 15)

col_widths = [Cm(15.7), Cm(15.7)]
headers = ["意思決定を支える関わり", "意思決定を奪う関わり"]
rows = [
    ["本人に直接、ゆっくり問いかける",
     "本人を飛ばして家族にばかり聞く"],
    ["「どちらがいいですか」と選択肢を示す",
     "「これでいいですね?」と誘導する"],
    ["実物・写真・図など見て分かる形で示す",
     "口頭の説明だけで済ませる"],
    ["本人のペースで、答えを待つ",
     "急かす・沈黙を埋めてしまう"],
    ["表情・しぐさなど非言語のサインも読む",
     "「言葉で言えない=意思なし」とみなす"],
    ["決めた後も「これでよかったか」と確認",
     "一度決めたら見直さない"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=ORANGE, first_col_color=GREEN,
          row_height=Cm(1.55), body_size=12, first_col_bold=False)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "「本人の最善の利益」をめぐる注意",
        "■ 代行決定が避けられない場合も、まず「本人ならどう考えるか」(本人の推定意思)を最優先\n"
        "■ 支援者が考える『客観的な最善』を、本人の意思より優先させない ― 価値観は人それぞれ",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)


# =============================================================
#  SLIDE 16: ACPとの連続性
# =============================================================
s = add_slide()
add_header(s, "3. 意思決定支援", "日々の意思決定支援が、将来の備えにつながる")
add_footer(s, 16)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "「小さな意思決定」の積み重ねが土台になる",
        "今日着る服、昼食のメニュー、散歩に行くかどうか ― 日々の小さな選択を支えることが、\n"
        "やがて「これからの暮らし方」「医療・ケアの希望」といった大きな意思決定の支えになる。",
        title_bg=ORANGE, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "早期からの関わりが鍵",
        "■ 診断後の早い段階から、本人の思いを聴く\n\n"
        "■ 言葉にできるうちに、価値観・希望を\n   共有しておく\n\n"
        "■ 「人生会議(ACP)」も、本人が主体的に\n   参加できる段階から始める\n\n"
        "■ 一度きりでなく、折に触れて対話を重ねる",
        title_bg=TEAL, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "CMが大切にしたい姿勢",
        "■ 「決めさせる」のではなく「一緒に考える」\n\n"
        "■ 本人が語った言葉を、記録に残す\n\n"
        "■ 意思は変わってよい ― 変化を尊重する\n\n"
        "■ 本人の意思を、家族・多職種と共有し、\n   ケアプランに確実に反映する",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ 意思決定支援は「特別な場面の技術」ではなく、日々の関わりそのもの。",
         size=13, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 17: 第4部
# =============================================================
section_title(4, "段階に応じた\n認知症の人と家族への実践",
              "Practice Across the Stages")


# =============================================================
#  SLIDE 18: 段階に応じた支援
# =============================================================
s = add_slide()
add_header(s, "4. 段階に応じた実践", "進行に応じた支援の視点")
add_footer(s, 18)

col_widths = [Cm(6.5), Cm(9.5), Cm(8.4), Cm(7.6)]
headers = ["段階", "本人の状態(例)", "本人への支援", "家族への支援"]
rows = [
    ["診断前後\n・初期",
     "不安・とまどい\nできることは多い",
     "診断後の空白を埋める\n本人の思いを聴く\n本人同士の出会いの場",
     "正しい知識\n受けとめの支援"],
    ["中期",
     "生活に支障\nBPSDが出ることも",
     "できることを活かす\n環境を整える\n役割・出番を保つ",
     "介護負担の軽減\nレスパイト"],
    ["重度",
     "コミュニケーション\nが難しくなる",
     "非言語のサインを読む\n快・不快を大切に\n尊厳あるケア",
     "看取りへの心構え\n意思決定の支え"],
    ["人生の\n最終段階",
     "全面的なケアが\n必要に",
     "本人の推定意思を尊重\n苦痛のない穏やかさ",
     "グリーフケア\n振り返りの支援"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=ORANGE, first_col_color=ORANGE,
          row_height=Cm(2.4), body_size=10.5, header_size=12)

add_box(s, Cm(1.2), Cm(15.5), Cm(31.4), Cm(2.2),
        "段階で区切りすぎない",
        "■ 進行の速さ・現れ方は人により大きく異なる ― 「段階」はあくまで支援を考える目安\n"
        "■ どの段階でも一貫するのは「本人の意思を中心に」「できることを活かす」という姿勢",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)


# =============================================================
#  SLIDE 19: 診断直後の支援
# =============================================================
s = add_slide()
add_header(s, "4. 段階に応じた実践", "診断直後の支援 ―「空白の期間」をなくす")
add_footer(s, 19)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "「診断されたのに、つながる先がない」期間",
        "認知症と診断されても、介護保険サービスを使うほどではない時期がある。\n"
        "この「空白の期間」に本人・家族が孤立し、不安を抱え込んでしまうことが課題。",
        title_bg=ORANGE, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "診断直後の本人・家族の思い",
        "■ 「これからどうなるのか」という不安\n\n"
        "■ 「人に知られたくない」という思い\n\n"
        "■ ショック・否認・落ち込み\n\n"
        "■ 何を相談していいか分からない\n\n"
        "■ 「もう何もできない」という思い込み",
        title_bg=BROWN, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "CM・支援者ができること",
        "■ 早期からつながり、伴走する\n\n"
        "■ 正しい情報を、希望とともに伝える\n\n"
        "■ 「できること」に目を向けられるよう支える\n\n"
        "■ 認知症カフェ・本人ミーティング等\n   同じ立場の人との出会いにつなぐ\n\n"
        "■ 認知症地域支援推進員・包括と連携",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.5),
         "▶ 「サービスにつなぐ」だけが支援ではない。「人・場・情報につなぐ」ことが、\n"
         "▶ 診断直後の本人・家族には何より大切。",
         size=13, color=ORANGE, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 20: BPSDの理解
# =============================================================
s = add_slide()
add_header(s, "4. 段階に応じた実践", "BPSDの理解 ― 「困った行動」ではなく「サイン」")
add_footer(s, 20)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "BPSD(行動・心理症状)をどう捉えるか",
        "徘徊・興奮・不安・拒否などのBPSDは、「困らせる行動」ではなく、\n"
        "本人が何かを伝えようとしている「サイン」 ― その背景を読み解くことが支援の出発点。",
        title_bg=ORANGE, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "BPSDの背景にあるもの",
        "■ 身体の不調(痛み・便秘・発熱・脱水)\n\n"
        "■ 不安・孤独・退屈・居場所のなさ\n\n"
        "■ 環境(騒音・まぶしさ・なじみのなさ)\n\n"
        "■ 関わり方(急かす・否定する・命令)\n\n"
        "■ 薬の影響",
        title_bg=BROWN, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "支援の考え方",
        "■ 「症状を抑える」前に「理由を探す」\n\n"
        "■ 本人の立場で「なぜ?」を考える\n\n"
        "■ 背景(痛み・環境・関わり)に手を当てる\n\n"
        "■ 多職種で情報を持ち寄り検討する\n\n"
        "■ 薬は安易に頼らない(ドラッグロック)",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ 第6弾で学んだ「行動の背景を考える」視点が、ここでも生きる。拘束・抑制は最後まで避ける。",
         size=13, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 21: 家族への支援
# =============================================================
s = add_slide()
add_header(s, "4. 段階に応じた実践", "家族への支援 ― 家族も「支えられる対象」")
add_footer(s, 21)

supports = [
    ("知識・情報の支援",
     "・認知症の正しい理解を伝える\n・進行の見通し・関わり方のヒント\n・使える制度・サービスの情報"),
    ("感情面の支援",
     "・とまどい・怒り・悲しみを受けとめる\n・「介護のつらさ」を語れる場\n・家族会・ピアサポートにつなぐ"),
    ("負担の軽減",
     "・レスパイト(通所・短期入所)\n・介護負担の分散\n・家族の生活・仕事との両立支援"),
    ("多様な家族への配慮",
     "・遠距離介護・独居高齢者を支える家族\n・ヤングケアラー(子・孫世代)\n・働きながら介護する世代"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [TEAL, BROWN, GREEN, ORANGE_MID]
for i, (title, body) in enumerate(supports):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "家族支援は「本人支援」でもある",
        "■ 家族が安定して本人と向き合えることは、本人の暮らしの安定に直結する\n"
        "■ ただし「家族のため」が「本人の意思」を覆さないよう注意 ― 本人と家族、両方の声を聴く\n"
        "■ ヤングケアラーなど、見えにくい家族介護者の存在に気づき、教育・福祉と連携してつなぐ",
        title_bg=ORANGE, body_size=13)


# =============================================================
#  SLIDE 22: 若年性認知症
# =============================================================
s = add_slide()
add_header(s, "4. 段階に応じた実践", "若年性認知症 ― 現役世代ならではの課題")
add_footer(s, 22)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "65歳未満で発症する認知症",
        "働き盛り・子育て世代での発症。高齢者の認知症とは異なる課題があり、\n"
        "介護保険だけでなく、就労・経済・家族など多面的な支援が必要になる。",
        title_bg=ORANGE, body_size=12)

issues = [
    ("就労・経済",
     "・仕事の継続・退職の問題\n・収入の減少・住宅ローン\n・障害年金等の制度活用"),
    ("家族・子ども",
     "・配偶者の就労と介護の両立\n・子がヤングケアラーになりうる\n・親世代の介護と重なることも"),
    ("制度のはざま",
     "・介護保険サービスが\n  年齢層に合わないことも\n・障害福祉サービスの活用"),
    ("本人の思い",
     "・「まだ働きたい」「役割を持ちたい」\n・同世代との出会いの場が少ない\n・社会参加・就労継続の支援"),
]
col_w = Cm(7.7)
col_h = Cm(6.3)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.2)
colors = [BROWN, TEAL, ORANGE_MID, GREEN]
for i, (title, body) in enumerate(issues):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), colors[i])
    add_text(s, x, y, col_w, Cm(1.3),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(1.6), col_w - Cm(0.6), col_h - Cm(1.9),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "つなぐ先を知っておく",
        "■ 若年性認知症支援コーディネーター(都道府県に配置)が、就労・生活の相談に対応\n"
        "■ 介護・医療・障害福祉・就労支援・年金など、分野を超えた連携が不可欠",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)


# =============================================================
#  SLIDE 23: 第5部
# =============================================================
section_title(5, "地域共生に向けた実践",
              "Practice for an Inclusive Community")


# =============================================================
#  SLIDE 24: 本人の声を中心に
# =============================================================
s = add_slide()
add_header(s, "5. 地域共生に向けた実践", "本人の声を中心に ― 当事者発信とピアサポート")
add_footer(s, 24)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.7),
        "「認知症の人について」から「認知症の人と共に」",
        "認知症の本人が自ら語り、発信し、施策づくりにも参画する動きが広がっている。\n"
        "支援者は「代弁する」のではなく、「本人が語れる場・機会」をつくり、支える。",
        title_bg=ORANGE, body_size=12)

items = [
    ("本人ミーティング",
     "認知症の本人同士が集い、\n思い・体験・希望を語り合う場。\n本人にとっての安心と力の源に"),
    ("ピアサポート",
     "同じ立場の本人による支え合い。\n診断直後の人に、先を歩む本人が\n寄り添い、希望を伝える"),
    ("本人の社会参画",
     "施策づくり・地域づくりへの参画。\n講演・発信など、本人が役割を\n持って社会と関わる"),
]
col_w = Cm(10.2)
col_h = Cm(5.5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(7.5)
colors = [ORANGE_MID, TEAL, GREEN]
for i, (title, body) in enumerate(items):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), colors[i])
    add_text(s, x, y, col_w, Cm(1.3),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.6), col_w - Cm(0.8), col_h - Cm(1.9),
             body, size=12, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "CMにできること",
        "■ 「もう無理」と思い込む本人に、当事者発信の事例を伝え、出会いの場につなぐ\n"
        "■ ケアプランに「社会参加」「役割」「人とのつながり」を、本人の願いとして位置づける\n"
        "■ 本人の言葉を、支援チーム・地域に「翻訳」して届ける ― ただし代弁に終わらせない",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=13)


# =============================================================
#  SLIDE 25: 地域資源
# =============================================================
s = add_slide()
add_header(s, "5. 地域共生に向けた実践", "認知症にかかわる地域資源を知る")
add_footer(s, 25)

col_widths = [Cm(9), Cm(12.5), Cm(9.9)]
headers = ["地域資源", "役割・機能", "CMの活用の仕方"]
rows = [
    ["認知症地域支援\n推進員",
     "市町村に配置。医療・介護・地域の\n連携支援、本人・家族の相談",
     "困りごとの相談\n地域資源につなぐ"],
    ["認知症初期\n集中支援チーム",
     "複数専門職が、初期の段階で\n集中的に本人・家族を支援",
     "診断後の関わりに迷う\nケースで連携"],
    ["認知症カフェ\n(オレンジカフェ)",
     "本人・家族・地域住民・専門職が\n集い、つながり・情報交換する場",
     "本人・家族の\n孤立を防ぐ場として"],
    ["チームオレンジ",
     "認知症サポーター等が、本人・家族の\nニーズに合わせて支援をつなぐ仕組み",
     "見守り・外出支援等\n生活の支え合いに"],
    ["家族会・\nピアの集まり",
     "家族同士・本人同士が\n体験を分かち合い、支え合う",
     "本人・家族を\n継続的な支えにつなぐ"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=ORANGE, first_col_color=ORANGE,
          row_height=Cm(2.0), body_size=10.5, header_size=12)

add_box(s, Cm(1.2), Cm(15.5), Cm(31.4), Cm(2.2),
        "「フォーマル」だけでなく「インフォーマル」も",
        "■ 介護保険サービス以外の、地域のつながり・居場所・支え合いを知り、つなぐのもCMの役割\n"
        "■ 主任CMは、地域ケア会議等を通じて、足りない資源を地域とともに育てる視点を持つ",
        title_bg=GREEN, body_bg=GREEN_BG, body_size=12)


# =============================================================
#  SLIDE 26: 権利擁護
# =============================================================
s = add_slide()
add_header(s, "5. 地域共生に向けた実践", "権利擁護 ― 意思と暮らしを守るしくみ")
add_footer(s, 26)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "認知症の進行とともに高まる権利擁護のニーズ",
        "判断能力が低下しても、本人の意思と財産・暮らしが守られるよう、\n"
        "制度を「本人の意思を実現する道具」として活用する。",
        title_bg=ORANGE, body_size=12)

items = [
    ("成年後見制度",
     "・判断能力が不十分な人の\n  財産管理・身上保護を支える\n・本人の意思尊重が制度の基本\n・市民後見人の活用も広がる"),
    ("日常生活自立\n支援事業",
     "・福祉サービスの利用援助\n・日常的な金銭管理の支援\n・成年後見より軽度な段階で\n  社会福祉協議会が実施"),
    ("虐待防止・\n消費者被害対策",
     "・高齢者虐待への気づきと対応\n  (第6弾の学びと連動)\n・悪質商法・特殊詐欺からの保護\n・地域包括・行政と連携"),
]
col_w = Cm(10.2)
col_h = Cm(7)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(7.2)
colors = [TEAL, GREEN, BROWN]
for i, (title, body) in enumerate(items):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.5), colors[i])
    add_text(s, x, y, col_w, Cm(1.5),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.5), col_w, col_h - Cm(1.5), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.9), col_w - Cm(0.8), col_h - Cm(2.3),
             body, size=11, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(14.7), Cm(31.4), Cm(2.3),
        "権利擁護も「意思決定支援」が出発点",
        "■ 制度につなぐ前に、まず本人の意思決定支援を尽くす ■ 制度利用後も、本人の意思を中心に\n"
        "■ ニーズの早期発見はCMの役割 ― 「気づいたら、地域包括・行政につなぐ」",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)


# =============================================================
#  SLIDE 27: 第6部
# =============================================================
section_title(6, "まとめ", "Wrap Up")


# =============================================================
#  SLIDE 28: 主任CMの役割
# =============================================================
s = add_slide()
add_header(s, "Closing", "認知症ケアマネジメントにおける主任CMの役割")
add_footer(s, 28)

roles = [
    ("実践のモデルを示す",
     "・本人中心・意思決定支援の実践を、\n  自ら体現して見せる\n・後輩CMの同行・SVで伝える"),
    ("事業所の文化を育てる",
     "・「本人に聴く」を当たり前にする\n・ケアプランの主語を本人にする\n・事例検討で認知症観を問い直す"),
    ("地域とつなぐ",
     "・地域資源を把握し、つなぐ\n・足りない資源を地域と育てる\n・地域ケア会議で課題を発信"),
    ("学び続け、発信する",
     "・基本法・ガイドラインの理解を深める\n・当事者の声に学ぶ\n・共生社会づくりの担い手になる"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [ORANGE_MID, GREEN, TEAL, GOLD]
for i, (title, body) in enumerate(roles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "基本法の理念を、現場の実践に翻訳する",
        "■ 「共生社会」「意思の尊重」という理念は、現場の一つひとつの関わりの中で実現される\n"
        "■ 主任CMは、理念と現場をつなぐ「翻訳者」 ― 抽象的な理念を、明日の実践に変える\n"
        "■ その積み重ねが、認知症になっても希望を持って暮らせる地域をつくっていく",
        title_bg=ORANGE, body_size=13)


# =============================================================
#  SLIDE 29: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 29)

actions = [
    ("Action 1", "ケアプランの「主語」を点検する",
     "・担当ケースのケアプランを見直し、主語が「本人」になっているか確認する\n"
     "・「させる」「させられる」表現を、本人の願い・意思の言葉に書き換える"),
    ("Action 2", "「本人に直接聴く」を意識する",
     "・次の訪問・面談で、家族を介さず本人自身に問いかける場面を意識的につくる\n"
     "・選択肢を具体的に示し、本人のペースで答えを待つ"),
    ("Action 3", "地域資源を1つ、訪ねてみる",
     "・担当圏域の認知症カフェ・本人ミーティング等を調べ、一度足を運ぶ\n"
     "・認知症地域支援推進員と顔をつなぎ、つなぎ先を増やす"),
]
y = Cm(4.5)
colors = [ORANGE_MID, TEAL, GREEN]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=24, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=16, bold=True, color=ORANGE)
    add_text(s, Cm(6), y + Cm(1.4), SW - Cm(7.5), Cm(2.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 30: クロージング
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, ORANGE)
add_rect(s, 0, Cm(7), SW, Cm(0.1), GREEN)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=GOLD_BG, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "認知症になっても、その人の人生は\n続いていく ― 主役は、いつも本人。",
         size=25, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "「どうせ分からない」ではなく「どう伝えれば分かり合えるか」。\n"
         "「何ができないか」ではなく「何を一緒にできるか」。\n"
         "その問いの立て方を変えることが、共生社会への確かな一歩になります。",
         size=14, color=LIGHT_BG, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 31: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 31)

refs = [
    "■ 共生社会の実現を推進するための認知症基本法(令和6年1月施行)",
    "■ 認知症施策推進基本計画(政府)/ 各都道府県・市町村の認知症施策推進計画",
    "■ 厚生労働省「認知症の人の日常生活・社会生活における意思決定支援ガイドライン」",
    "■ 厚生労働省「人生の最終段階における医療・ケアの決定プロセスに関するガイドライン」",
    "■ 認知症の本人・家族による発信、本人ミーティング・家族会等の活動",
    "■ 認知症地域支援推進員・認知症初期集中支援チーム等の地域資源情報",
    "■ 日本介護支援専門員協会等による認知症ケアマネジメント関連の研修資料",
]
y = Cm(4.5)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=13, color=DARK_GRAY)
    y += Cm(1.25)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "学び続けるために",
        "■ 基本計画・自治体施策は更新されます。最新の法令・計画・ガイドラインをご確認ください\n"
        "■ 何より大切な学びの源は「認知症の本人の声」 ― 当事者の発信に直接ふれる機会を持つこと\n"
        "■ 本資料は 2026年5月時点の一般的な情報を基に構成しています",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)


# =============================================================
#  SLIDE 32: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, ORANGE)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), GREEN)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 意見交換",
         size=18, color=GOLD_BG, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "本人の意思を中心に置く実践を、事業所・地域に広げていきましょう。",
         size=14, color=LIGHT_BG, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/認知症基本法施行後のケアマネジメント_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
