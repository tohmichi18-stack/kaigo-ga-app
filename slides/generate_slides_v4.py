"""主任ケアマネ研修スライド生成スクリプト 第4弾

テーマ: 主任ケアマネの実務とスーパービジョン技術
重点: 後輩CM指導 / 地域ケア会議ファシリテーション / 地域包括支援センターとの協働
トーン: 実務型(技術・手順・場面対応)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(人材育成・協働トーン:ティール基調) =====
TEAL        = RGBColor(0x1B, 0x4D, 0x4A)   # メイン
TEAL_MID    = RGBColor(0x2E, 0x6E, 0x69)
TEAL_LIGHT  = RGBColor(0x4F, 0x8F, 0x89)
ORANGE      = RGBColor(0xD9, 0x6E, 0x2A)   # アクセント(行動・注意)
GOLD        = RGBColor(0xC9, 0x9A, 0x2E)   # ヒント
CORAL       = RGBColor(0xC0, 0x4A, 0x3E)   # NG・リスク
SAGE        = RGBColor(0x5C, 0x8A, 0x5A)   # OK・推奨
LIGHT_BG    = RGBColor(0xF0, 0xF3, 0xF2)
MID_GRAY    = RGBColor(0xD4, 0xDC, 0xDA)
DARK_GRAY   = RGBColor(0x32, 0x38, 0x37)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xE5, 0xEC, 0xEB)
CORAL_BG    = RGBColor(0xFA, 0xE8, 0xE6)
GOLD_BG     = RGBColor(0xFA, 0xF1, 0xDC)
SAGE_BG     = RGBColor(0xE6, 0xEE, 0xE5)
ORANGE_BG   = RGBColor(0xFB, 0xEC, 0xDD)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), TEAL)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=TEAL)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), ORANGE)


def add_footer(slide, page_num, total=37):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(24), Cm(0.5),
             "主任CM研修 / 主任ケアマネの実務とスーパービジョン技術",
             size=9, color=TEAL_LIGHT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=TEAL_LIGHT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=TEAL_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=TEAL, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=ORANGE):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, TEAL)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=accent_color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=34, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, TEAL)
add_rect(s, 0, Cm(8), SW, Cm(0.15), ORANGE)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), TEAL_LIGHT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.4), SW - Cm(2), Cm(2.5),
         "主任ケアマネの実務と\nスーパービジョン技術",
         size=32, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
add_text(s, Cm(1.5), Cm(9.7), SW - Cm(3), Cm(2),
         "ー 後輩ケアマネを育て、地域の議論を導き、地域包括と協働する ー",
         size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(13), SW - Cm(3), Cm(2),
         "■ スーパービジョンの 3 機能と実践技術\n"
         "■ 後輩CM への OJT・同行訪問・フィードバック\n"
         "■ 地域ケア会議のファシリテーション\n"
         "■ 地域包括支援センターとの協働",
         size=14, color=GOLD, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "主任CMに求められる「育てる・つなぐ・導く」役割を言語化できる",
    "スーパービジョンの3機能(管理的・教育的・支持的)を理解し、実践に移せる",
    "後輩CMの成長段階に応じたOJT・同行訪問・フィードバックの技術を持ち帰る",
    "地域ケア会議でのファシリテーション ― 発言を引き出し合意を形成する技術を学ぶ",
    "地域包括支援センターとの役割分担と協働のあり方を整理する",
    "「プレイヤー」から「育てる人・つなぐ人」への意識転換を図る",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), ORANGE)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "主任CMの役割 ― 「育てる・つなぐ・導く」",      "10分"),
    ("2", "スーパービジョン技術 ― 3機能と実践",          "25分"),
    ("3", "後輩CM指導 ― OJT・同行訪問・フィードバック",   "20分"),
    ("4", "地域ケア会議のファシリテーション",             "20分"),
    ("5", "地域包括支援センターとの協働",                 "10分"),
    ("6", "まとめ ― 主任CMが持ち帰るもの",                " 5分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, TEAL_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=TEAL)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=ORANGE, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "主任CMの役割", "The Role of the Chief Care Manager")


# =============================================================
#  SLIDE 5: 主任CMに求められる4機能
# =============================================================
s = add_slide()
add_header(s, "1. 主任CMの役割", "主任CMに求められる 4 つの機能")
add_footer(s, 5)

funcs = [
    ("育てる", "人材育成", "後輩CMのスーパービジョン\nOJT・研修の企画・運営",
     SAGE),
    ("つなぐ", "連携・協働", "地域包括・医療機関・他事業所\n多職種ネットワークの構築",
     TEAL_MID),
    ("導く",   "リーダーシップ", "地域ケア会議のファシリテート\n事業所内の方針づくり",
     ORANGE),
    ("支える", "相談・支持", "困難事例への助言\n後輩のメンタルサポート",
     GOLD),
]
col_w = Cm(7.7)
col_h = Cm(9.5)
gap = Cm(0.3)
total_w = col_w * 4 + gap * 3
start_x = (SW - total_w) // 2
y = Cm(4.5)
for i, (verb, label, body, color) in enumerate(funcs):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2.5), color)
    add_text(s, x + Cm(0.2), y + Cm(0.3), col_w - Cm(0.4), Cm(1.2),
             verb, size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.2), y + Cm(1.6), col_w - Cm(0.4), Cm(0.7),
             label, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, x, y + Cm(2.5), col_w, col_h - Cm(2.5), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(2.9), col_w - Cm(0.6), col_h - Cm(3.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "4機能に共通する土台",
        "■ 自分の実践を「言葉」にできること ― 暗黙知を形式知に変える力\n"
        "■ 「正解を教える」のではなく「考えるプロセスを支える」姿勢",
        title_bg=TEAL, body_size=13)


# =============================================================
#  SLIDE 6: プレイヤーから育てる人へ
# =============================================================
s = add_slide()
add_header(s, "1. 主任CMの役割", "「プレイヤー」から「育てる人」への転換")
add_footer(s, 6)

col_widths = [Cm(8), Cm(11.7), Cm(11.7)]
headers = ["観点", "プレイヤー(担当CM)", "主任CM(育てる人)"]
rows = [
    ["成果の単位", "自分の担当ケースの質", "チーム・事業所全体の質"],
    ["時間の使い方", "自分の支援に集中", "後輩に関わる時間を確保する"],
    ["困難事例", "自分で抱えて解決する", "後輩が対応できるよう支える"],
    ["評価のものさし", "利用者の満足・状態改善", "後輩の成長・自立"],
    ["失敗への向き合い", "自分のミスを正す", "後輩の失敗を学びに変える"],
    ["視野の範囲", "事業所内・担当圏域", "地域全体・制度・他機関"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=TEAL, first_col_color=TEAL,
          row_height=Cm(1.5), body_size=12)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "意識転換のポイント",
        "■ 「自分がやった方が早い」を手放す ― 後輩が経験する機会を奪わない\n"
        "■ プレイヤーをやめるのではなく、「プレイヤーの背中を見せながら育てる」二刀流へ\n"
        "■ 主任CMの仕事の成果は「すぐには見えない」 ― 半年・1年の時間軸で考える",
        title_bg=ORANGE, body_bg=ORANGE_BG, body_size=12)


# =============================================================
#  SLIDE 7: 第2部
# =============================================================
section_title(2, "スーパービジョン技術", "Supervision Skills")


# =============================================================
#  SLIDE 8: スーパービジョンとは
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "スーパービジョン(SV)とは何か")
add_footer(s, 8)

add_box(s, Cm(1.2), Cm(4.5), Cm(31.4), Cm(3.3),
        "定義",
        "経験ある専門職(スーパーバイザー)が、後進(スーパーバイジー)に対して、\n"
        "支持的な関係性のもとで行う、専門職としての成長を支える継続的な過程。\n"
        "「指導」「管理」とは異なり、バイジー自身の気づきと力を引き出すことが目的。",
        title_bg=TEAL, body_size=13)

add_box(s, Cm(1.2), Cm(8.3), Cm(15.5), Cm(8.5),
        "SVであるもの",
        "■ バイジーの実践を一緒に振り返る\n\n"
        "■ 問いかけを通じて気づきを促す\n\n"
        "■ バイジーの感情を受けとめる\n\n"
        "■ 専門職としての枠組みを共有する\n\n"
        "■ 継続的・定期的に行われる\n\n"
        "■ 守秘義務と安心感が前提",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=13)

add_box(s, Cm(17.5), Cm(8.3), Cm(15), Cm(8.5),
        "SVでないもの(混同しやすい)",
        "■ 業務指示・命令(=マネジメント)\n\n"
        "■ ケースの肩代わり(=代行)\n\n"
        "■ 一方的な助言・正解の提示(=ティーチング)\n\n"
        "■ 評価・査定のための面談(=人事考課)\n\n"
        "■ カウンセリング(=治療的関わり)\n\n"
        "▶ 重なる部分はあるが、目的が異なる",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=13)


# =============================================================
#  SLIDE 9: SVの3機能
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "スーパービジョンの 3 つの機能")
add_footer(s, 9)

functions = [
    ("管理的機能", "Administrative",
     "・業務が適切に行われているか確認\n"
     "・記録・手続きの質の担保\n"
     "・業務量・負荷の調整\n"
     "・組織のルールとの橋渡し",
     TEAL_MID),
    ("教育的機能", "Educational",
     "・専門知識・技術の伝達\n"
     "・アセスメント力・判断力の育成\n"
     "・実践の振り返りと意味づけ\n"
     "・「考え方」を育てる",
     ORANGE),
    ("支持的機能", "Supportive",
     "・感情面のサポート\n"
     "・バーンアウトの予防\n"
     "・孤立感の解消・安心感の提供\n"
     "・モチベーションの維持",
     SAGE),
]
col_w = Cm(10.2)
col_h = Cm(9)
gap = Cm(0.4)
total_w = col_w * 3 + gap * 2
start_x = (SW - total_w) // 2
y = Cm(4.5)
for i, (jp, en, body, color) in enumerate(functions):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.8), color)
    add_text(s, x + Cm(0.2), y + Cm(0.2), col_w - Cm(0.4), Cm(0.9),
             jp, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.2), y + Cm(1.1), col_w - Cm(0.4), Cm(0.6),
             en, size=11, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, x, y + Cm(1.8), col_w, col_h - Cm(1.8), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(2.2), col_w - Cm(0.8), col_h - Cm(2.6),
             body, size=13, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.7),
        "3機能のバランス",
        "■ どれか1つに偏らない ― 「管理」だけだと萎縮し、「支持」だけだと成長しない\n"
        "■ バイジーの状態に応じて重心を変える(疲弊時は支持を、停滞時は教育を厚く)\n"
        "■ 評価と支援を同じ人が担う難しさを自覚する ― 「査定」と「育成」の場面を分ける工夫を",
        title_bg=TEAL, body_size=12)


# =============================================================
#  SLIDE 10: SVの形態
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "SVの形態 ― 個人・グループ・ピア")
add_footer(s, 10)

col_widths = [Cm(6.5), Cm(8.7), Cm(8.7), Cm(7.5)]
headers = ["形態", "特徴", "向いている場面", "主任CMの留意点"]
rows = [
    ["個人SV",
     "1対1。深く掘り下げ\nられる。関係性が密",
     "新人の育成\n困難事例の振り返り\n感情面のケア",
     "1対1ゆえの圧を\n和らげる。安心の場に"],
    ["グループSV",
     "複数バイジー。多様な\n視点・相互学習が起きる",
     "事例検討会\nテーマ別の学び\n経験者の学び合い",
     "発言の偏りを防ぐ\n心理的安全性の確保"],
    ["ピアSV",
     "同僚同士で対等に。\nバイザーは固定しない",
     "中堅同士の学び\n孤立しがちな一人\nケアマネの支え合い",
     "脱線・愚痴大会に\nならない構造づくり"],
    ["ライブSV",
     "面接・訪問に同席し、\nその場で振り返る",
     "実践技術の伝達\n同行訪問\n初期の育成",
     "利用者への配慮\n後で必ず振り返る"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=TEAL, first_col_color=ORANGE,
          row_height=Cm(2.4), body_size=11, header_size=12)

add_text(s, Cm(1.2), Cm(16), SW - Cm(2.4), Cm(1.2),
         "▶ 形態は組み合わせて使う。事業所の規模・人員に応じて現実的な仕組みを設計する。",
         size=13, color=TEAL, bold=True)


# =============================================================
#  SLIDE 11: SVの基本プロセス
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "SVセッションの基本プロセス")
add_footer(s, 11)

steps = [
    ("1", "場をつくる",
     "・守秘義務と安心感を確認\n・今日扱うテーマを合意\n・評価の場ではないと伝える"),
    ("2", "語ってもらう",
     "・バイジーに事例・状況を語らせる\n・遮らず、最後まで聴く\n・事実と感情の両方を受けとめる"),
    ("3", "問いかけ・掘り下げる",
     "・「そのとき何を考えた?」\n・「他にどんな見方ができる?」\n・批判せず、一緒に考える"),
    ("4", "気づきを整理する",
     "・バイジー自身の言葉でまとめる\n・うまくいった点も必ず確認\n・次の一歩を具体化する"),
    ("5", "ふりかえる",
     "・セッション自体を振り返る\n・次回テーマを確認\n・ねぎらいで終える"),
]
y = Cm(4.5)
for i, (num, title, body) in enumerate(steps):
    color = [TEAL_MID, TEAL_LIGHT, ORANGE, SAGE, GOLD][i]
    add_numbered_circle(s, Cm(1.5), y + Cm(0.4), Cm(1.3), num, color)
    add_rect(s, Cm(3.4), y, Cm(6), Cm(2.1), color)
    add_text(s, Cm(3.4), y, Cm(6), Cm(2.1),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(9.4), y, SW - Cm(10.6), Cm(2.1), TABLE_BG)
    add_text(s, Cm(9.8), y + Cm(0.15), SW - Cm(11.2), Cm(1.9),
             body, size=12, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    y += Cm(2.4)

add_text(s, Cm(1.2), Cm(16.7), SW - Cm(2.4), Cm(0.8),
         "▶ 「教える」ステップがないことに注目。主役はバイジー、バイザーは伴走者。",
         size=12, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 12: 質問技法
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "気づきを促す問いかけの技術")
add_footer(s, 12)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.3),
        "開かれた質問(オープン・クエスチョン)",
        "■ 「はい/いいえ」で終わらない問い\n\n"
        "■ 例:\n"
        "  ・「そのとき、どう感じましたか?」\n"
        "  ・「何が一番気がかりですか?」\n"
        "  ・「他にどんな選択肢がありそう?」\n"
        "  ・「もし時間を戻せたら、どうしますか?」\n\n"
        "▶ 考えを広げ、深める",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.3),
        "閉じた質問(クローズド・クエスチョン)",
        "■ 事実確認・焦点化に有効\n\n"
        "■ 例:\n"
        "  ・「訪問は何回行いましたか?」\n"
        "  ・「家族には伝えましたか?」\n\n"
        "■ 使いどころ\n"
        "  ・話が拡散したとき焦点を絞る\n"
        "  ・事実関係を整理するとき\n\n"
        "▶ 多用すると「尋問」になるので注意",
        title_bg=TEAL_MID, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "そのほかの技法",
        "■ 反映(リフレクション):バイジーの言葉を返す 「『つらかった』と感じたのですね」\n"
        "■ 言い換え(パラフレーズ):要点を整理して返す 「つまり、家族との板挟みだった、と」\n"
        "■ 沈黙を待つ:すぐ答えを埋めない。考える時間を奪わない\n"
        "■ 承認・ねぎらい:「よく抱えましたね」「その判断は妥当だったと思います」\n"
        "■ 焦点化:「いま一番話したいのはどの点ですか?」と一緒に絞る\n"
        "■ 直面化(慎重に):矛盾や避けている点を、責めずに示す 「さっきと少し違って聞こえました」",
        title_bg=TEAL, body_size=12)


# =============================================================
#  SLIDE 13: SVでやってはいけないこと
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "SVでやってはいけないこと ― よくある失敗")
add_footer(s, 13)

ng_items = [
    ("答えをすぐ与える",
     "バイジーが考える前に「こうすればいい」と正解を出す。\n→ 自分で考える力が育たず、依存を生む"),
    ("詰問・尋問になる",
     "「なぜやらなかった?」と問い詰める。\n→ 萎縮し、本当のことを語らなくなる"),
    ("自分の経験談に終始",
     "「私のときはこうだった」と自分語りが中心に。\n→ バイジーの状況に合わず、置いていかれる"),
    ("評価とSVを混同",
     "支援の場のはずが査定のように感じさせる。\n→ 安心して弱みを出せなくなる"),
    ("感情を扱わない",
     "事実・手続きだけを確認し、つらさを素通り。\n→ 支持的機能が働かず、孤立感が残る"),
    ("時間・場の不在",
     "「いつでも相談して」で終わり、定例の場がない。\n→ SVが偶発的になり、機能しない"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
for i, (title, body) in enumerate(ng_items):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), CORAL)
    add_text(s, x + Cm(0.5), y + Cm(0.1), col_w - Cm(1), Cm(0.8),
             "NG  " + title, size=14, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), CORAL_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.35)


# =============================================================
#  SLIDE 14: バイザーの自己点検
# =============================================================
s = add_slide()
add_header(s, "2. スーパービジョン技術", "スーパーバイザー自身の自己点検")
add_footer(s, 14)

add_box(s, Cm(1.2), Cm(4.5), Cm(31.4), Cm(2.7),
        "バイザーも完璧ではない",
        "■ バイザー自身が悩み、迷い、感情に揺れる存在であることを認める\n"
        "■ 「自分も学び続ける一人」という姿勢が、バイジーの安心につながる",
        title_bg=TEAL, body_size=13)

checks = [
    "自分の価値観を、バイジーに押しつけていないか",
    "「教えたい」気持ちが先走り、聴くことを忘れていないか",
    "バイジーの成長を、自分の手柄にしようとしていないか",
    "苦手なバイジーに対して、関わりが薄くなっていないか",
    "自分自身がSV・相談を受けられる場を持っているか",
    "セッション後、自分の関わりを振り返っているか",
]
y = Cm(7.7)
for i, c in enumerate(checks):
    bg = TABLE_BG if i % 2 == 0 else WHITE
    add_rect(s, Cm(1.2), y, Cm(31.4), Cm(1.3), bg, line_color=MID_GRAY)
    add_rect(s, Cm(1.6), y + Cm(0.3), Cm(0.7), Cm(0.7), WHITE, line_color=TEAL)
    add_text(s, Cm(2.7), y + Cm(0.15), SW - Cm(4), Cm(1),
             c, size=13, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE)
    y += Cm(1.4)

add_text(s, Cm(1.2), Cm(16.5), SW - Cm(2.4), Cm(1),
         "▶ 「バイザーへのSV(=SVのSV)」を、地域の主任CM同士で持ち合う発想も大切。",
         size=13, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 15: 第3部
# =============================================================
section_title(3, "後輩CM指導", "Mentoring Junior Care Managers")


# =============================================================
#  SLIDE 16: 成長段階に応じた指導
# =============================================================
s = add_slide()
add_header(s, "3. 後輩CM指導", "成長段階に応じた関わり方")
add_footer(s, 16)

col_widths = [Cm(6), Cm(9), Cm(9), Cm(8.4)]
headers = ["段階", "状態・特徴", "つまずきやすい点", "主任CMの関わり"]
rows = [
    ["新人期\n(〜1年)",
     "制度・手続きの習得期。\n不安が大きい",
     "書類・期限の管理\n何が分からないか不明",
     "丁寧な同行・確認\n安心の土台づくり"],
    ["習熟期\n(1〜3年)",
     "一通りこなせる。\n自己流が出始める",
     "アセスメントの浅さ\n困難事例で行き詰まる",
     "問いかけ中心のSV\n視野を広げる"],
    ["中堅期\n(3〜7年)",
     "自立して動ける。\nマンネリ・燃え尽きも",
     "成長実感の停滞\n後輩指導の負担",
     "新たな役割の付与\nピアSVへの橋渡し"],
    ["指導者期\n(7年〜)",
     "後輩を育てる側へ。\n主任研修の対象",
     "育成スキルの不足\n自分の学びの停滞",
     "育成の伴走\n地域活動への誘い"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=TEAL, first_col_color=ORANGE,
          row_height=Cm(2.3), body_size=11, header_size=12)

add_box(s, Cm(1.2), Cm(15.3), Cm(31.4), Cm(2.2),
        "原則",
        "■ 同じ「指導」でも、新人には支持を厚く、中堅には問いを多く ― 段階で重心を変える\n"
        "■ 年数はあくまで目安。一人ひとりの実際の状態を見て関わる",
        title_bg=ORANGE, body_bg=ORANGE_BG, body_size=12)


# =============================================================
#  SLIDE 17: OJTの設計
# =============================================================
s = add_slide()
add_header(s, "3. 後輩CM指導", "OJTの設計 ― 「見て覚えろ」を卒業する")
add_footer(s, 17)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(7),
        "OJTが機能しない事業所の特徴",
        "■ 「忙しいから自分で覚えて」が口ぐせ\n\n"
        "■ 指導が場当たり的(計画がない)\n\n"
        "■ 教える人によって言うことが違う\n\n"
        "■ 振り返りの時間がとれていない\n\n"
        "■ できていない点ばかり指摘される\n\n"
        "■ 「育てる」が評価されない組織風土",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=13)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(7),
        "機能するOJTの4要素",
        "■ 計画:育成計画を文書化(到達目標・期間)\n\n"
        "■ 段階:Show(やってみせる)→ Tell(説明)\n   → Do(やらせる)→ Check(振り返る)\n\n"
        "■ 役割分担:指導担当を明確化\n   (複数で関わる場合は方針を統一)\n\n"
        "■ 振り返り:定例の振り返り時間を確保",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=13)

add_box(s, Cm(1.2), Cm(11.8), Cm(31.4), Cm(5.2),
        "育成計画に盛り込む項目(例)",
        "■ 到達目標:3か月後・6か月後・1年後に「何ができるようになるか」を具体的に\n"
        "■ 担当ケースの段階的付与:軽度・安定ケース → 徐々に複雑なケースへ\n"
        "■ 同行訪問の予定:初期は必ず同行、徐々に独り立ち\n"
        "■ 研修・事例検討への参加計画\n"
        "■ 振り返り面談:月1回など定例で設定 ― 「困ったときだけ」にしない",
        title_bg=TEAL, body_size=13)


# =============================================================
#  SLIDE 18: 同行訪問の活用
# =============================================================
s = add_slide()
add_header(s, "3. 後輩CM指導", "同行訪問の活用 ― 最良の学びの場")
add_footer(s, 18)

phases = [
    ("訪問前", TEAL_MID,
     "・後輩に「今日の訪問のねらい」を言わせる\n"
     "・想定される展開を一緒に確認\n"
     "・主任CMの役割(見守る/補助する)を決めておく"),
    ("訪問中", ORANGE,
     "・原則、後輩に任せる ― 口を出しすぎない\n"
     "・利用者の前で後輩を否定しない\n"
     "・後輩が詰まったときの合図・フォローを事前に決める"),
    ("訪問後", SAGE,
     "・必ずその日のうちに振り返る\n"
     "・後輩に「自己評価」を先に語らせる\n"
     "・できた点を具体的に承認 → 改善点は1〜2点に絞る"),
]
y = Cm(4.5)
for label, color, body in phases:
    add_rect(s, Cm(1.2), y, Cm(5), Cm(3.3), color)
    add_text(s, Cm(1.2), y, Cm(5), Cm(3.3),
             label, size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(6.2), y, SW - Cm(7.4), Cm(3.3), TABLE_BG)
    add_text(s, Cm(6.7), y + Cm(0.3), SW - Cm(8), Cm(2.7),
             body, size=13, color=DARK_GRAY, line_spacing=1.45)
    y += Cm(3.6)

add_box(s, Cm(1.2), Cm(15.4), Cm(31.4), Cm(2.1),
        "陥りやすい罠",
        "■ 主任CMが「主役」になってしまう ― 後輩の成長機会を奪う\n"
        "■ 振り返りを「次の訪問のついで」にして、結局しないまま流れてしまう",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)


# =============================================================
#  SLIDE 19: フィードバックの技術
# =============================================================
s = add_slide()
add_header(s, "3. 後輩CM指導", "フィードバックの技術 ― 伝わる伝え方")
add_footer(s, 19)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.5),
        "効果的なフィードバックの原則",
        "■ 具体的に:「良かった」でなく「○○の声かけが良かった」\n\n"
        "■ 行動に焦点:人格でなく「やったこと」に\n\n"
        "■ タイミング良く:できるだけ早く\n\n"
        "■ 量を絞る:改善点は一度に1〜2点まで\n\n"
        "■ 対話で:一方通行でなく、本人の考えも聞く",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.5),
        "「サンドイッチ」だけに頼らない",
        "■ 良い点→改善点→良い点 で挟む手法は有効だが…\n\n"
        "■ 多用すると「褒めは前置き」と見抜かれる\n\n"
        "■ 大切なのは:\n"
        "  ・改善点を「相手の成長のため」と心から伝える\n"
        "  ・本人が「次どうするか」を考えられるよう問う\n\n"
        "▶ 形式よりも、信頼関係と本気度",
        title_bg=GOLD, body_bg=GOLD_BG, body_size=12)

# 言い換え例
col_widths = [Cm(15.7), Cm(15.7)]
headers = ["伝わりにくい言い方", "伝わる言い方"]
rows = [
    ["「アセスメントが甘いね」",
     "「ご本人の意欲の部分、もう少し聞けると\n計画が変わるかも。どう思う?」"],
    ["「なんで連絡しなかったの」",
     "「あの場面で家族に連絡していたら、\nどう違ったと思う?」"],
    ["「もっと頑張って」",
     "「いま一番大変なのはどこ?\n一緒に整理してみよう」"],
]
add_table(s, Cm(1.2), Cm(11.3), col_widths, headers, rows,
          header_bg=TEAL, first_col_color=CORAL,
          row_height=Cm(1.7), body_size=11, first_col_bold=False)


# =============================================================
#  SLIDE 20: 困難事例での指導 + 燃え尽き予防
# =============================================================
s = add_slide()
add_header(s, "3. 後輩CM指導", "困難事例での指導と、燃え尽きの予防")
add_footer(s, 20)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(6.5),
        "困難事例 ― 抱え込ませない",
        "■ 「一人で背負わせない」を最優先に\n\n"
        "■ 後輩が抱える事例を主任CMが把握する仕組み\n  (定例の事例共有・進捗確認)\n\n"
        "■ チームで対応する事例を切り分ける\n\n"
        "■ 「相談していい」を言葉と態度で示す\n\n"
        "■ 虐待・希死念慮等は即・組織対応\n  (後輩個人の判断に委ねない)",
        title_bg=TEAL_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(6.5),
        "燃え尽き(バーンアウト)のサイン",
        "■ 情緒的消耗感:「もう疲れた」が口ぐせに\n\n"
        "■ 脱人格化:利用者を「事務的」に扱い始める\n\n"
        "■ 達成感の低下:「自分は役に立っていない」\n\n"
        "■ 行動の変化:遅刻・欠勤・記録の遅れ・口数減\n\n"
        "▶ 「真面目で熱心な人」ほど陥りやすい",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11.5), Cm(31.4), Cm(5.5),
        "主任CMができる予防的関わり",
        "■ 業務量の偏りを把握し、調整する(管理的機能)\n"
        "■ 「うまくいったケース」を一緒に振り返り、達成感を言語化する\n"
        "■ 感情を吐き出せる場(支持的SV)を定期的に持つ\n"
        "■ 一人で完結する仕事を減らし、チームで支える構造をつくる\n"
        "■ 不調のサインに気づいたら、早めに声をかける ― 「最近どう?」の一言から",
        title_bg=TEAL, body_size=13)


# =============================================================
#  SLIDE 21: 第4部
# =============================================================
section_title(4, "地域ケア会議の\nファシリテーション",
              "Facilitating Community Care Meetings")


# =============================================================
#  SLIDE 22: 地域ケア会議とは
# =============================================================
s = add_slide()
add_header(s, "4. 地域ケア会議", "地域ケア会議とは ― 5 つの機能")
add_footer(s, 22)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.5),
        "位置づけ",
        "介護保険法に基づき市町村等が設置。多職種が協働し、\n"
        "個別課題の解決から地域課題の発見・政策提言までを担う場。",
        title_bg=TEAL, body_size=13)

functions = [
    ("個別課題\n解決", "支援困難事例の検討\n多職種で支援策を練る"),
    ("ネットワーク\n構築", "関係機関の顔の見える\n関係づくり"),
    ("地域課題\n発見", "個別事例の積み重ねから\n地域の課題を見出す"),
    ("地域づくり\n資源開発", "不足する社会資源を\n地域で創り出す"),
    ("政策形成", "地域課題を介護保険\n事業計画等に反映"),
]
col_w = Cm(6.2)
col_h = Cm(6.5)
gap = Cm(0.2)
total_w = col_w * 5 + gap * 4
start_x = (SW - total_w) // 2
y = Cm(7.3)
colors = [TEAL_MID, TEAL_LIGHT, ORANGE, SAGE, GOLD]
for i, (title, body) in enumerate(functions):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2.3), colors[i])
    add_text(s, x + Cm(0.1), y + Cm(0.1), col_w - Cm(0.2), Cm(2.1),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2.3), col_w, col_h - Cm(2.3), TABLE_BG)
    add_text(s, x + Cm(0.2), y + Cm(2.5), col_w - Cm(0.4), col_h - Cm(2.7),
             body, size=11, color=DARK_GRAY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "「個別」から「地域」へ ― 会議の積み上げ構造",
        "■ 1件の困難事例の検討が、同じ課題を抱える他のケースの発見につながる\n"
        "■ 「この地域に○○が足りない」という気づきが、資源開発・政策提言へとつながっていく",
        title_bg=ORANGE, body_bg=ORANGE_BG, body_size=12)


# =============================================================
#  SLIDE 23: 主任CMの関わり方
# =============================================================
s = add_slide()
add_header(s, "4. 地域ケア会議", "地域ケア会議における主任CMの立場")
add_footer(s, 23)

roles = [
    ("事例提供者として",
     "・自分(事業所)の困難事例を持ち込む\n"
     "・「うまくいかない」を開示する勇気\n"
     "・事例の匿名化・同意取得を徹底\n"
     "・検討で得た助言を実践に戻す"),
    ("助言者・参加者として",
     "・他のCMの事例に専門的助言\n"
     "・経験を押しつけず、選択肢を示す\n"
     "・多職種の視点を引き出す\n"
     "・若手CMの発言を後押しする"),
    ("ファシリテーターとして",
     "・会議の進行・時間管理\n"
     "・発言を引き出し、論点を整理\n"
     "・合意形成と次のアクション化\n"
     "・個別課題から地域課題への橋渡し"),
]
col_w = Cm(10.2)
col_h = Cm(8.5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(4.5)
colors = [TEAL_MID, SAGE, ORANGE]
for i, (title, body) in enumerate(roles):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.5), colors[i])
    add_text(s, x + Cm(0.2), y, col_w - Cm(0.4), Cm(1.5),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.5), col_w, col_h - Cm(1.5), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.9), col_w - Cm(0.8), col_h - Cm(2.3),
             body, size=12, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "本日の重点 ― ファシリテーター役",
        "■ 主任研修・更新研修でも、地域ケア会議の運営力が期待されている\n"
        "■ 「会議を仕切る」のではなく「参加者の力を引き出して、結論を一緒に作る」のがファシリテーション\n"
        "■ 司会(進行)・記録・タイムキープなど、役割を分担する設計も大切",
        title_bg=TEAL, body_size=12)


# =============================================================
#  SLIDE 24: ファシリテーションの基本
# =============================================================
s = add_slide()
add_header(s, "4. 地域ケア会議", "ファシリテーションの基本姿勢")
add_footer(s, 24)

principles = [
    ("中立を保つ",
     "特定の意見・職種に肩入れしない。\n自分の結論に誘導しない"),
    ("全員を参加させる",
     "発言の偏りをなくす。\n声の小さい人・若手にも振る"),
    ("否定しない場をつくる",
     "「それは違う」を言わせない。\nどんな意見もまず受けとめる"),
    ("論点を可視化する",
     "ホワイトボード等で見える化。\n「いま何を話しているか」を共有"),
    ("時間を守る",
     "終了時刻を意識した進行。\n1事例の検討時間を区切る"),
    ("結論を出す",
     "「で、誰が何をいつまでに」を\n必ず明確にして終える"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.6)
for i, (title, body) in enumerate(principles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_numbered_circle(s, x + Cm(0.3), y + Cm(0.3), Cm(1.1), i + 1, ORANGE)
    add_rect(s, x, y, col_w, Cm(1.7), TEAL_MID)
    add_text(s, x + Cm(1.7), y, col_w - Cm(2), Cm(1.7),
             title, size=15, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.7), col_w, row_h - Cm(1.7), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.9), col_w - Cm(1), row_h - Cm(2),
             body, size=12, color=DARK_GRAY, line_spacing=1.35)


# =============================================================
#  SLIDE 25: 会議の設計
# =============================================================
s = add_slide()
add_header(s, "4. 地域ケア会議", "会議の設計 ― 準備で 8 割が決まる")
add_footer(s, 25)

steps = [
    ("事前準備",
     "・会議の目的・到達点を明確化\n・事例の選定と提供者との打合せ\n・参加メンバーの選定(必要な職種)\n・資料の事前配布・匿名化の確認"),
    ("会議の冒頭",
     "・目的とゴール、終了時刻を共有\n・グランドルールの確認(否定しない 等)\n・参加者の自己紹介(関係づくり)"),
    ("検討の進行",
     "・事例提供 → 質問・情報整理 → 検討\n・論点を絞り、可視化しながら進める\n・発言の偏りを調整、多職種の視点を引き出す"),
    ("まとめ",
     "・支援方針・役割分担を具体化\n・「誰が・何を・いつまでに」を確認\n・地域課題として残す点を整理\n・振り返り(会議自体の改善点)"),
]
y = Cm(4.5)
colors = [TEAL_MID, TEAL_LIGHT, ORANGE, SAGE]
for i, (title, body) in enumerate(steps):
    add_rect(s, Cm(1.2), y, Cm(5.5), Cm(2.8), colors[i])
    add_numbered_circle(s, Cm(1.5), y + Cm(0.3), Cm(1), i + 1, WHITE if False else TEAL)
    add_text(s, Cm(1.2), y + Cm(0.1), Cm(5.5), Cm(2.6),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(6.9), y, SW - Cm(8.1), Cm(2.8), TABLE_BG)
    add_text(s, Cm(7.3), y + Cm(0.25), SW - Cm(8.9), Cm(2.3),
             body, size=12, color=DARK_GRAY, line_spacing=1.35)
    y += Cm(3.1)

add_text(s, Cm(1.2), Cm(16.7), SW - Cm(2.4), Cm(0.8),
         "▶ 「事例提供者が安心して持ち込める」準備が、会議の質を決める。",
         size=12, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 26: 困った場面への対応
# =============================================================
s = add_slide()
add_header(s, "4. 地域ケア会議", "困った場面への対応 ― ファシリテーターの引き出し")
add_footer(s, 26)

col_widths = [Cm(9), Cm(11.5), Cm(10.9)]
headers = ["困った場面", "なぜ起こるか", "対応の引き出し"]
rows = [
    ["特定の人が\n話し続ける",
     "熱心さ・経験の豊富さ\n沈黙への不安",
     "「ありがとうございます。\n他の方はいかがですか?」と振る"],
    ["誰も発言\nしない",
     "心理的安全性の不足\n論点が不明確",
     "問いを具体化・小さくする\n「まず○○さんから一言」と指名"],
    ["批判・説教\nが始まる",
     "事例提供者への\n評価的なまなざし",
     "「ここは責める場でなく\n一緒に考える場」と立て直す"],
    ["話が\n脱線する",
     "論点が共有\nされていない",
     "可視化した論点に立ち返る\n「今日の検討点に戻すと…」"],
    ["結論が\nまとまらない",
     "情報過多\n時間配分のミス",
     "「今日決められること」と\n「持ち帰り」を切り分ける"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=TEAL, first_col_color=CORAL,
          row_height=Cm(2.0), body_size=11, header_size=12)

add_text(s, Cm(1.2), Cm(16), SW - Cm(2.4), Cm(1.2),
         "▶ 困った場面の多くは「準備」と「グランドルール」で予防できる。起きてからの対応は最終手段。",
         size=12, color=TEAL, bold=True)


# =============================================================
#  SLIDE 27: 第5部
# =============================================================
section_title(5, "地域包括支援センター\nとの協働",
              "Collaboration with Community Support Centers")


# =============================================================
#  SLIDE 28: 地域包括の機能
# =============================================================
s = add_slide()
add_header(s, "5. 地域包括との協働", "地域包括支援センターの機能と居宅CMの関係")
add_footer(s, 28)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "地域包括支援センターとは",
        "市町村が設置する地域包括ケアの中核機関。保健師・社会福祉士・主任CM等が配置され、\n"
        "高齢者の総合相談・権利擁護・地域のケアマネジメント支援などを担う。",
        title_bg=TEAL, body_size=12)

funcs = [
    ("総合相談支援",
     "高齢者・家族の\nあらゆる相談の窓口",
     "居宅CM:つなぎ先\nとして活用"),
    ("権利擁護",
     "高齢者虐待対応\n成年後見・消費者被害",
     "居宅CM:虐待の\n通報・連携"),
    ("包括的・継続的\nケアマネジメント",
     "地域のCMへの\n支援・後方支援",
     "居宅CM:困難事例\nの相談・助言を得る"),
    ("介護予防\nケアマネジメント",
     "要支援者等への\n介護予防支援",
     "居宅CM:委託\n要支援→要介護の連続"),
]
col_w = Cm(7.7)
col_h = Cm(6.8)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7)
colors = [TEAL_MID, ORANGE, TEAL_LIGHT, SAGE]
for i, (title, body, rel) in enumerate(funcs):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.7), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(1.7),
             title, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.7), col_w, col_h - Cm(1.7), TABLE_BG)
    add_text(s, x + Cm(0.25), y + Cm(2), col_w - Cm(0.5), Cm(2.2),
             body, size=11, color=DARK_GRAY,
             align=PP_ALIGN.CENTER, line_spacing=1.3)
    add_rect(s, x + Cm(0.25), y + Cm(4.2), col_w - Cm(0.5), Cm(0.05), MID_GRAY)
    add_text(s, x + Cm(0.25), y + Cm(4.4), col_w - Cm(0.5), Cm(2),
             rel, size=10, color=TEAL, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.3)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "主任CMの視点",
        "■ 地域包括は「上位機関」ではなく「協働のパートナー」 ― 対等な専門職同士\n"
        "■ 包括の主任CMと居宅の主任CMが、地域のケアマネジメントの質を共に支える",
        title_bg=ORANGE, body_bg=ORANGE_BG, body_size=12)


# =============================================================
#  SLIDE 29: 連携の実際
# =============================================================
s = add_slide()
add_header(s, "5. 地域包括との協働", "連携の実際 ― 場面別の役割分担")
add_footer(s, 29)

col_widths = [Cm(8), Cm(12), Cm(11.4)]
headers = ["場面", "居宅CMの役割", "地域包括の役割"]
rows = [
    ["要支援→要介護\nの移行",
     "要介護のケアマネジメント\nを引き継ぐ。情報を確実に受領",
     "介護予防支援からの\n引き継ぎ。情報提供"],
    ["支援困難事例",
     "事例を抱え込まず相談\n日々の支援を担う",
     "後方支援・助言\n多機関調整のハブ"],
    ["高齢者虐待\nの疑い",
     "サインの発見・通報\n事実の記録",
     "通報の受理・事実確認\n市町村との対応"],
    ["権利擁護\n(成年後見等)",
     "ニーズの発見\n本人・家族への説明",
     "申立て支援・制度説明\n関係機関へのつなぎ"],
    ["地域づくり",
     "個別事例から見える\n地域課題を持ち寄る",
     "地域ケア会議の開催\n資源開発・政策提言"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=TEAL, first_col_color=ORANGE,
          row_height=Cm(2.0), body_size=11, header_size=12)

add_text(s, Cm(1.2), Cm(16), SW - Cm(2.4), Cm(1.2),
         "▶ 「どこからが包括の仕事か」で迷うより、「利用者にとって最善の動き」を一緒に考える。",
         size=12, color=TEAL, bold=True)


# =============================================================
#  SLIDE 30: 困難事例での協働
# =============================================================
s = add_slide()
add_header(s, "5. 地域包括との協働", "困難事例での協働 ― 早めに・具体的に・記録して")
add_footer(s, 30)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(7),
        "相談をためらわせるもの",
        "■ 「自分の力不足と思われたくない」\n\n"
        "■ 「こんなことで相談していいのか」\n\n"
        "■ 「忙しそうで声をかけづらい」\n\n"
        "■ 「何をどう伝えればいいか分からない」\n\n"
        "▶ 主任CMが後輩に「早めの相談」を\n   モデルとして示すことが大切",
        title_bg=CORAL, body_bg=CORAL_BG, body_size=13)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(7),
        "相談を活かすコツ",
        "■ 早めに:こじれる前の段階で\n\n"
        "■ 具体的に:「何に困り、何を求めるか」を整理\n\n"
        "■ 経過を記録:支援経過・関わった機関を時系列で\n\n"
        "■ ゴールを共有:相談で「何を一緒に決めたいか」\n\n"
        "■ 平時から関係づくり:困ってから初めて\n   連絡、にしない",
        title_bg=SAGE, body_bg=SAGE_BG, body_size=13)

add_box(s, Cm(1.2), Cm(12), Cm(31.4), Cm(5),
        "主任CMが事業所内でつくる仕組み",
        "■ 「相談していい」を文化にする ― 抱え込みを評価しない、相談を歓迎する\n"
        "■ 困難事例リストを事業所で共有し、包括への相談タイミングを逃さない\n"
        "■ 包括の担当者と「顔の見える関係」を平時から維持する(合同事例検討・地域の集まり)\n"
        "■ 多機関(医療・障害・生活困窮・権利擁護)との連携窓口を整理しておく",
        title_bg=TEAL, body_size=13)


# =============================================================
#  SLIDE 31: 地域づくりへの参画
# =============================================================
s = add_slide()
add_header(s, "5. 地域包括との協働", "地域づくりへの参画 ― 一事業所を超えて")
add_footer(s, 31)

add_box(s, Cm(1.2), Cm(4.5), Cm(31.4), Cm(2.7),
        "なぜ主任CMが地域づくりに関わるのか",
        "■ 目の前の利用者の困りごとの背景には、地域の資源不足・制度の隙間がある\n"
        "■ 個別支援を地域の力に変えていくのが、主任CMに期待される役割",
        title_bg=TEAL, body_size=13)

activities = [
    ("個別の気づきを\n持ち寄る",
     "日々の支援で見えた\n「地域に足りないもの」を\n地域ケア会議等で発信"),
    ("ネットワークに\n参加する",
     "ケアマネ連絡会・職能団体\n多職種連携の会・\n地域の協議体に顔を出す"),
    ("資源開発に\n関わる",
     "インフォーマルな支え合い\n通いの場・生活支援の\n仕組みづくりに協力"),
    ("次世代を\n育てる",
     "実習生・新人の受け入れ\n地域の研修への講師協力\nCM全体の底上げ"),
]
col_w = Cm(7.7)
col_h = Cm(6.8)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(8)
colors = [TEAL_MID, ORANGE, SAGE, GOLD]
for i, (title, body) in enumerate(activities):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2), colors[i])
    add_text(s, x + Cm(0.1), y, col_w - Cm(0.2), Cm(2),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(2), col_w, col_h - Cm(2), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(2.3), col_w - Cm(0.6), col_h - Cm(2.6),
             body, size=11, color=DARK_GRAY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.35)

add_text(s, Cm(1.2), Cm(16.3), SW - Cm(2.4), Cm(1),
         "▶ 「忙しくてそれどころではない」のが現実。だからこそ、できる範囲で一歩ずつ。",
         size=13, color=ORANGE, bold=True)


# =============================================================
#  SLIDE 32: 第6部
# =============================================================
section_title(6, "まとめ", "Wrap Up")


# =============================================================
#  SLIDE 33: 主任CMの3つの軸
# =============================================================
s = add_slide()
add_header(s, "Closing", "主任CMの実務を支える 3 つの軸")
add_footer(s, 33)

axes = [
    ("人を育てる軸",
     "スーパービジョン / OJT",
     "・SVの3機能を意識して関わる\n・問いかけで気づきを引き出す\n・成長段階に応じて重心を変える\n・燃え尽きを予防し、支える",
     SAGE),
    ("場を導く軸",
     "ファシリテーション",
     "・中立を保ち、全員を参加させる\n・準備で会議の質を決める\n・論点を可視化し、結論を出す\n・個別課題を地域課題へつなぐ",
     ORANGE),
    ("地域とつなぐ軸",
     "多機関協働",
     "・地域包括と対等に協働する\n・早めに・具体的に相談する\n・平時から顔の見える関係を保つ\n・地域づくりに一歩ずつ関わる",
     TEAL_MID),
]
col_w = Cm(10.2)
col_h = Cm(10)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(4.5)
for i, (title, sub, body, color) in enumerate(axes):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(2), color)
    add_text(s, x + Cm(0.2), y + Cm(0.25), col_w - Cm(0.4), Cm(0.9),
             title, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.2), y + Cm(1.15), col_w - Cm(0.4), Cm(0.7),
             sub, size=12, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, x, y + Cm(2), col_w, col_h - Cm(2), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(2.4), col_w - Cm(1), col_h - Cm(2.8),
             body, size=12, color=DARK_GRAY, line_spacing=1.6)

add_box(s, Cm(1.2), Cm(15), Cm(31.4), Cm(2.2),
        "3つの軸に共通する原点",
        "■ 「正解を与える人」ではなく「考えるプロセスを支える人」\n"
        "■ 主任CMの成果は、後輩・地域・利用者を通じて、時間をかけて表れる",
        title_bg=TEAL, body_size=13)


# =============================================================
#  SLIDE 34: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 34)

actions = [
    ("Action 1", "後輩CMとの「定例の振り返り」を設ける",
     "・「困ったときだけ」をやめ、月1回など定例のSVの時間をつくる\n"
     "・1回目は「最近どう?」から。評価でなく支える場であることを伝える"),
    ("Action 2", "次の会議を「準備」から設計する",
     "・次に関わる地域ケア会議・事例検討会で、目的・ゴール・進行を事前設計\n"
     "・グランドルール(否定しない 等)を冒頭で共有してみる"),
    ("Action 3", "地域包括の担当者と一度話す",
     "・困難事例がなくても、平時に顔をつなぐ\n"
     "・自分の事業所の困難事例リストを整理し、相談のタイミングを見直す"),
]
y = Cm(4.5)
colors = [SAGE, ORANGE, TEAL_MID]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=22, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=16, bold=True, color=TEAL)
    add_text(s, Cm(6), y + Cm(1.4), SW - Cm(7.5), Cm(2.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 35: クロージングメッセージ
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, TEAL)
add_rect(s, 0, Cm(7), SW, Cm(0.1), ORANGE)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=ORANGE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "ケアマネジャーを育てることは、\nその先にいる利用者を支えること。",
         size=26, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "主任CMの一日は地味で、成果はすぐには見えません。\n"
         "けれど、あなたが後輩に向けた問いかけ、会議で引き出した一言、\n"
         "地域包括とつないだ一本の連絡が、地域のケアの質を少しずつ変えていきます。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 36: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 36)

refs = [
    "■ 厚生労働省「地域ケア会議運営マニュアル」",
    "■ 厚生労働省「主任介護支援専門員研修」「主任介護支援専門員更新研修」関連資料",
    "■ 日本介護支援専門員協会「スーパービジョン」関連研修・テキスト",
    "■ 厚生労働省「適切なケアマネジメント手法の手引き」",
    "■ 各都道府県・職能団体のスーパービジョン研修・ファシリテーション研修",
    "■ 地域包括支援センター運営に関する手引き・通知",
    "■ 対人援助職向けのスーパービジョン・バーンアウト関連の専門書",
]
y = Cm(4.5)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=14, color=DARK_GRAY)
    y += Cm(1.3)

add_box(s, Cm(1.2), Cm(14.3), Cm(31.4), Cm(2.7),
        "学び続けるために",
        "■ スーパービジョンもファシリテーションも「技術」 ― 知識だけでなく、実践と振り返りで磨かれる\n"
        "■ 主任CM自身も、地域の主任CM同士で学び合い、支え合う場を持つことが大切\n"
        "■ 本資料は 2026年5月時点の情報を基に構成しています。",
        title_bg=ORANGE, body_bg=ORANGE_BG, body_size=12)


# =============================================================
#  SLIDE 37: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, TEAL)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), ORANGE)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 意見交換",
         size=18, color=GOLD, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "本日の学びを、ぜひ事業所・地域の中で共有してください。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/主任CM実務とスーパービジョン技術_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
