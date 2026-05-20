"""主任ケアマネ研修スライド生成スクリプト 第10弾

テーマ: 法定研修制度改正への対応と主任ケアマネ更新研修の戦略
重点: 研修体系の見直しを踏まえた、事業所内での計画的な人材育成
トーン: 実務型(戦略・計画・管理の視点)
時間: 90分
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===== カラーパレット(人材育成・研修:フォレストグリーン基調 + 銅) =====
FOREST      = RGBColor(0x29, 0x47, 0x37)   # メイン
FOREST_MID  = RGBColor(0x40, 0x63, 0x50)
FOREST_LT   = RGBColor(0x6E, 0x8C, 0x7C)
COPPER      = RGBColor(0xB3, 0x77, 0x39)   # アクセント
TEAL        = RGBColor(0x2C, 0x6E, 0x73)   # 連携
GOLD        = RGBColor(0xC4, 0x9A, 0x36)   # ヒント
SLATE       = RGBColor(0x44, 0x53, 0x60)   # 制度・落ち着き
RED         = RGBColor(0xB2, 0x44, 0x39)   # 注意・リスク
LIGHT_BG    = RGBColor(0xF1, 0xF3, 0xF1)
MID_GRAY    = RGBColor(0xD5, 0xDB, 0xD6)
DARK_GRAY   = RGBColor(0x2E, 0x33, 0x30)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_BG    = RGBColor(0xE7, 0xEC, 0xE8)
RED_BG      = RGBColor(0xF8, 0xE6, 0xE4)
GOLD_BG     = RGBColor(0xFA, 0xF1, 0xDA)
COPPER_BG   = RGBColor(0xF7, 0xEC, 0xDC)
TEAL_BG     = RGBColor(0xDF, 0xEA, 0xEA)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Cm(33.867)
prs.slide_height = Cm(19.05)
SW = prs.slide_width
SH = prs.slide_height
blank_layout = prs.slide_layouts[6]


# ===== ヘルパー =====
def add_slide():
    return prs.slides.add_slide(blank_layout)


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *,
             size=14, bold=False, color=DARK_GRAY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=None, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def add_header(slide, section_label, title):
    add_rect(slide, 0, 0, SW, Cm(1.6), FOREST)
    add_text(slide, Cm(0.8), Cm(0.25), Cm(22), Cm(0.6),
             section_label, size=11, color=WHITE, bold=True)
    add_text(slide, Cm(0.8), Cm(1.9), SW - Cm(1.6), Cm(1.3),
             title, size=23, bold=True, color=FOREST)
    add_rect(slide, Cm(0.8), Cm(3.2), Cm(2.5), Cm(0.12), COPPER)


def add_footer(slide, page_num, total=27):
    add_text(slide, Cm(0.8), SH - Cm(0.8), Cm(27), Cm(0.5),
             "主任CM研修 / 法定研修制度改正への対応と主任CM更新研修の戦略  ※研修制度は最新情報をご確認ください",
             size=9, color=FOREST_LT)
    add_text(slide, SW - Cm(3), SH - Cm(0.8), Cm(2.2), Cm(0.5),
             f"{page_num} / {total}", size=9, color=FOREST_LT, align=PP_ALIGN.RIGHT)


def add_box(slide, left, top, width, height, title, body, *,
            title_bg=FOREST_MID, body_bg=TABLE_BG, title_color=WHITE,
            body_color=DARK_GRAY, title_size=14, body_size=12,
            body_line_spacing=None):
    add_rect(slide, left, top, width, Cm(1.0), title_bg)
    add_text(slide, left + Cm(0.2), top + Cm(0.1), width - Cm(0.4), Cm(0.8),
             title, size=title_size, bold=True, color=title_color,
             anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, left, top + Cm(1.0), width, height - Cm(1.0), body_bg)
    add_text(slide, left + Cm(0.3), top + Cm(1.1), width - Cm(0.6),
             height - Cm(1.1), body, size=body_size, color=body_color,
             line_spacing=body_line_spacing)


def add_table(slide, left, top, col_widths, headers, rows, *,
              header_bg=FOREST, header_color=WHITE,
              row_bg_a=TABLE_BG, row_bg_b=WHITE,
              header_size=12, body_size=11, row_height=Cm(1.4),
              first_col_color=None, first_col_bold=True):
    total_w = sum(col_widths, Cm(0))
    add_rect(slide, left, top, total_w, Cm(1.0), header_bg)
    x = left
    for i, h in enumerate(headers):
        add_text(slide, x + Cm(0.15), top + Cm(0.1), col_widths[i] - Cm(0.3), Cm(0.8),
                 h, size=header_size, bold=True, color=header_color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += col_widths[i]

    y = top + Cm(1.0)
    for r_i, row in enumerate(rows):
        bg = row_bg_a if r_i % 2 == 0 else row_bg_b
        add_rect(slide, left, y, total_w, row_height, bg, line_color=MID_GRAY)
        x = left
        for i, c in enumerate(row):
            color = first_col_color if (i == 0 and first_col_color) else DARK_GRAY
            bold = first_col_bold and (i == 0)
            align = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            add_text(slide, x + Cm(0.15), y + Cm(0.15),
                     col_widths[i] - Cm(0.3), row_height - Cm(0.3),
                     c, size=body_size, bold=bold, color=color,
                     align=align, anchor=MSO_ANCHOR.MIDDLE)
            x += col_widths[i]
        y += row_height
    return y


def section_title(num, jp, en, *, accent_color=COPPER):
    s = add_slide()
    add_rect(s, 0, 0, SW, SH, FOREST)
    add_rect(s, 0, Cm(8.5), SW, Cm(0.1), accent_color)
    add_text(s, Cm(2), Cm(6.5), SW - Cm(4), Cm(1.5),
             f"PART {num}", size=22, color=accent_color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(9.5), SW - Cm(2), Cm(2),
             jp, size=33, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1), Cm(13), SW - Cm(2), Cm(1),
             en, size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)
    return s


def add_numbered_circle(slide, left, top, diameter, num, color, fontsize=15):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_top = Cm(0)
    tf.margin_bottom = Cm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(num)
    r.font.name = FONT
    r.font.size = Pt(fontsize)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return circle


def add_arrow(slide, left, top, width, height, color):
    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    arr.shadow.inherit = False
    return arr


# =============================================================
#  SLIDE 1: 表紙
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, FOREST)
add_rect(s, 0, Cm(8), SW, Cm(0.15), COPPER)
add_rect(s, 0, Cm(12.3), SW, Cm(0.05), FOREST_LT)

add_text(s, Cm(2), Cm(4.0), SW - Cm(4), Cm(1),
         "主任介護支援専門員 研修会",
         size=18, color=COPPER, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(1), Cm(5.4), SW - Cm(2), Cm(2.5),
         "法定研修制度改正への対応と\n主任ケアマネ更新研修の戦略",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.25)
add_text(s, Cm(1.5), Cm(9.7), SW - Cm(3), Cm(2),
         "ー 研修体系の見直しを踏まえた、計画的な人材育成 ー",
         size=16, color=WHITE, align=PP_ALIGN.CENTER)

add_text(s, Cm(1.5), Cm(13), SW - Cm(3), Cm(2),
         "■ CM法定研修体系の全体像と、制度改正の動向\n"
         "■ 主任ケアマネ更新研修を「逃さない」戦略\n"
         "■ 事業所内での計画的な研修管理・人材育成",
         size=14, color=GOLD_BG, align=PP_ALIGN.CENTER, line_spacing=1.5)

add_text(s, Cm(2), Cm(16.3), SW - Cm(4), Cm(1),
         "主任ケアマネジャーの会  主催  ・  研修時間 90 分",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =============================================================
#  SLIDE 2: 本日の目標
# =============================================================
s = add_slide()
add_header(s, "Introduction", "本日の研修目標")
add_footer(s, 2)

goals = [
    "介護支援専門員の法定研修体系の全体像を、改めて整理する",
    "研修制度の見直し・改正の動向と、注視すべき論点を理解する",
    "主任ケアマネ更新研修の位置づけと、「更新を逃す」リスクを把握する",
    "更新研修を計画的に受講するための「逆算の戦略」を持ち帰る",
    "事業所内で研修受講を管理し、人材育成計画に組み込む方法を学ぶ",
    "研修を「受けっぱなし」にせず、事業所に還元する仕組みを考える",
]
y = Cm(4.5)
for g in goals:
    add_rect(s, Cm(1.2), y, Cm(0.5), Cm(1.5), COPPER)
    add_text(s, Cm(2), y + Cm(0.15), SW - Cm(3.5), Cm(1.2),
             g, size=15, color=DARK_GRAY, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.3)
    y += Cm(1.8)


# =============================================================
#  SLIDE 3: アジェンダ
# =============================================================
s = add_slide()
add_header(s, "Agenda", "本日の流れ(90分)")
add_footer(s, 3)

agenda = [
    ("1", "CM法定研修体系の全体像",                "10分"),
    ("2", "研修制度改正の動向 ― 何が論点か",        "15分"),
    ("3", "主任ケアマネ更新研修の戦略",            "20分"),
    ("4", "事業所内での計画的な人材育成",          "20分"),
    ("5", "研修を「活かす」 ― 学びの還元",         "15分"),
    ("6", "まとめ ― 主任CMが持ち帰るもの",         "10分"),
]
y = Cm(4.7)
for num, title, mins in agenda:
    add_numbered_circle(s, Cm(1.5), y, Cm(1.3), num, FOREST_MID)
    add_text(s, Cm(3.3), y + Cm(0.15), Cm(22), Cm(1),
             title, size=17, bold=True, color=FOREST)
    add_text(s, SW - Cm(4), y + Cm(0.2), Cm(3), Cm(1),
             mins, size=14, color=COPPER, bold=True, align=PP_ALIGN.RIGHT)
    y += Cm(1.85)


# =============================================================
#  SLIDE 4: 第1部
# =============================================================
section_title(1, "CM法定研修体系の全体像",
              "The Statutory Training System")


# =============================================================
#  SLIDE 5: 研修体系マップ
# =============================================================
s = add_slide()
add_header(s, "1. 法定研修体系の全体像", "介護支援専門員の法定研修体系マップ")
add_footer(s, 5)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2),
        "資格取得から、主任CM、その更新まで ― 研修は続いていく",
        "CMの研修は「一度受けて終わり」ではなく、資格の取得・更新と結びついた継続的なもの。",
        title_bg=FOREST, body_size=12)

# 研修の流れ(横フロー)
stages = [
    ("実務研修", "資格取得時", FOREST_MID),
    ("専門研修\nⅠ・Ⅱ", "更新に必要", TEAL),
    ("主任CM\n研修", "主任になる", COPPER),
    ("主任CM\n更新研修", "主任を続ける", SLATE),
]
col_w = Cm(7.2)
gap = Cm(0.9)
total_w = col_w * 4 + gap * 3
start_x = (SW - total_w) // 2
y = Cm(7)
for i, (title, sub, color) in enumerate(stages):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(3), color)
    add_text(s, x + Cm(0.2), y + Cm(0.4), col_w - Cm(0.4), Cm(1.5),
             title, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Cm(0.2), y + Cm(1.9), col_w - Cm(0.4), Cm(0.8),
             sub, size=11, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 3:
        add_arrow(s, x + col_w + Cm(0.1), y + Cm(1), gap - Cm(0.2), Cm(1), COPPER)

add_box(s, Cm(1.2), Cm(11.5), Cm(31.4), Cm(5.5),
        "押さえておきたい基本構造",
        "■ 介護支援専門員証は 5年ごとに更新 ― 更新には所定の研修(専門研修等)の修了が必要\n\n"
        "■ 主任介護支援専門員も、5年ごとに 主任介護支援専門員更新研修 を受ける必要がある\n\n"
        "■ 実務未経験・ブランクのある人向けに、再研修・更新研修(実務未経験者向け)もある\n\n"
        "■ 研修の名称・時間数・要件の詳細や運用は 都道府県により異なる ― 必ず自県の情報を確認",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 6: なぜ法定研修か
# =============================================================
s = add_slide()
add_header(s, "1. 法定研修体系の全体像", "なぜ法定研修があるのか ― その意味")
add_footer(s, 6)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(6),
        "法定研修の役割",
        "■ ケアマネジメントの質を、全国で一定に保つ\n\n"
        "■ 制度改正・最新の知見を確実に届ける\n\n"
        "■ 専門職としての継続的な学びを担保する\n\n"
        "■ 資格・更新と結びつけ、受講を確実に\n\n"
        "■ 主任CMなど、段階的な力量形成を支える",
        title_bg=FOREST_MID, body_size=13)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(6),
        "一方で指摘される課題",
        "■ 研修時間が長く、負担が大きい\n\n"
        "■ 業務を離れて受講する間の体制確保\n\n"
        "■ 受講費用・移動の負担\n\n"
        "■ 内容の重複感・「やらされ感」\n\n"
        "▶ これらが、制度見直しの背景になっている",
        title_bg=RED, body_bg=RED_BG, body_size=13)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "主任CMとしての捉え方",
        "■ 法定研修を「やらされる負担」と見るか、「質を支えるしくみ」と見るかで、関わり方が変わる\n\n"
        "■ 制度への不満を言うだけでなく、「どうすれば学びを実りあるものにできるか」を考える\n\n"
        "■ 主任CMは、後輩が研修を前向きに受けられるよう、意味づけを支える立場にある\n\n"
        "■ 同時に、現場の負担の実態を、職能団体等を通じて制度に伝えていく役割も持つ",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 7: 第2部
# =============================================================
section_title(2, "研修制度改正の動向",
              "Trends in Training System Reform")


# =============================================================
#  SLIDE 8: 見直しの背景
# =============================================================
s = add_slide()
add_header(s, "2. 研修制度改正の動向", "なぜ研修体系が見直されているのか")
add_footer(s, 8)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "研修体系の見直しが議論されている",
        "国において、介護支援専門員の法定研修のあり方について見直しの議論が進められている。\n"
        "その背景には、現場が抱える複数の課題がある。",
        title_bg=FOREST, body_size=12)

reasons = [
    ("研修負担の重さ",
     "・研修時間が長い\n・受講中の業務代替が困難\n・費用・移動の負担"),
    ("ケアマネ不足",
     "・有資格者の減少・高齢化\n・受験者数の減少傾向\n・人材確保が急務"),
    ("内容への課題",
     "・科目間・段階間の重複感\n・実務との結びつきの弱さ\n・「学びの実感」が薄い場合"),
    ("学び方の変化",
     "・オンライン研修の活用\n・ICTを前提とした見直し\n・柔軟な受講方法の要請"),
]
col_w = Cm(7.7)
col_h = Cm(5.8)
gap = Cm(0.3)
start_x = (SW - (col_w * 4 + gap * 3)) // 2
y = Cm(7.2)
colors = [RED, COPPER, SLATE, TEAL]
for i, (title, body) in enumerate(reasons):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), colors[i])
    add_text(s, x, y, col_w, Cm(1.3),
             title, size=14, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.3), y + Cm(1.6), col_w - Cm(0.6), col_h - Cm(1.9),
             body, size=11, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "見直しの方向性(議論されている論点)",
        "■ 研修の負担軽減・効率化(時間数・受講方法の見直し)\n"
        "■ オンライン・ICTの活用 ■ 科目内容の整理・実務に即した見直し\n"
        "▶ ただし、具体的な改正内容・施行時期は流動的 ― 最新の国・都道府県の情報を必ず確認",
        title_bg=COPPER, body_bg=COPPER_BG, body_size=12)


# =============================================================
#  SLIDE 9: 注視すべきこと
# =============================================================
s = add_slide()
add_header(s, "2. 研修制度改正の動向", "主任CMとして注視すべきこと")
add_footer(s, 9)

col_widths = [Cm(10), Cm(12), Cm(9.4)]
headers = ["注視するポイント", "なぜ重要か", "確認の仕方"]
rows = [
    ["研修時間・\nカリキュラムの変更",
     "受講計画・業務調整に\n直接影響する",
     "都道府県・実施機関の\n案内を定期的に確認"],
    ["受講方法\n(オンライン等)",
     "受講のしやすさ・\n業務との両立に関わる",
     "実施要綱・募集要項を\n毎年チェック"],
    ["更新の要件・\n経過措置",
     "「更新を逃す」リスクに\n直結する",
     "自分・職員の更新時期と\nあわせて確認"],
    ["施行時期・\n移行措置",
     "「いつから変わるか」で\n計画が変わる",
     "国・職能団体の\n発信を追う"],
]
add_table(s, Cm(1.2), Cm(4.5), col_widths, headers, rows,
          header_bg=FOREST, first_col_color=FOREST,
          row_height=Cm(2.0), body_size=11)

add_box(s, Cm(1.2), Cm(13.5), Cm(31.4), Cm(3.7),
        "情報源を「複数」持つ",
        "■ 都道府県(研修の実施主体)の案内 ― 受講の実務情報はここが基本\n"
        "■ 厚生労働省の検討会・通知 ― 制度改正の方向性\n"
        "■ 職能団体(日本介護支援専門員協会等)の発信 ― 現場目線の解説・動向\n"
        "▶ 主任CMは「最新の確かな情報源」を複数持ち、事業所内に正確に伝える役割を持つ",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 10: 第3部
# =============================================================
section_title(3, "主任ケアマネ\n更新研修の戦略",
              "Strategy for Chief CM Renewal Training")


# =============================================================
#  SLIDE 11: 主任研修・更新研修の位置づけ
# =============================================================
s = add_slide()
add_header(s, "3. 更新研修の戦略", "主任CM研修・更新研修の位置づけ")
add_footer(s, 11)

add_box(s, Cm(1.2), Cm(4.3), Cm(15.5), Cm(6.3),
        "主任介護支援専門員研修",
        "■ 主任CMとなるための研修\n\n"
        "■ 一定の実務経験等の受講要件がある\n\n"
        "■ ケアマネジメントの指導・人材育成・\n   地域づくり等を学ぶ\n\n"
        "■ 修了が、主任CMとしての出発点\n\n"
        "▶ 受講要件は都道府県により設定",
        title_bg=FOREST_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.3), Cm(15), Cm(6.3),
        "主任介護支援専門員更新研修",
        "■ 主任CMの資格的な立場を維持するための研修\n\n"
        "■ おおむね5年ごとに受講が必要\n\n"
        "■ 受講要件・実務要件が定められている\n\n"
        "■ 修了しないと、主任CMとしての\n   位置づけを維持できないことがある",
        title_bg=COPPER, body_bg=COPPER_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "なぜ「戦略」が必要なのか",
        "■ 主任CM更新研修は、受講時期・要件が決まっており、「うっかり」では済まされない\n\n"
        "■ 居宅介護支援事業所の管理者要件として主任CMが求められる ― 更新を逃すと事業所運営に影響\n\n"
        "■ 研修は数日間にわたり、業務を離れる ― 計画なしでは「受けたくても受けられない」\n\n"
        "■ だからこそ「いつ・誰が・どう受けるか」を逆算して計画する戦略が要る",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 12: 更新を逃すリスク
# =============================================================
s = add_slide()
add_header(s, "3. 更新研修の戦略", "「更新を逃す」と何が起きるか")
add_footer(s, 12)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "更新は「個人の問題」では済まない",
        "更新研修を受け損ねることは、本人だけでなく、事業所の運営にも影響しうる。\n"
        "「気づいたら期限が過ぎていた」を、組織として防ぐ必要がある。",
        title_bg=FOREST, body_size=12)

risks = [
    ("本人への影響",
     "・主任CMとしての立場の\n  維持に影響\n・再度の研修受講等が\n  必要になりうる"),
    ("事業所への影響",
     "・管理者要件を満たせなく\n  なるおそれ\n・特定事業所加算の要件\n  にも関わる"),
    ("利用者・地域への影響",
     "・指導・育成体制の弱体化\n・地域での主任CMの\n  役割が果たせなくなる"),
]
col_w = Cm(10.2)
col_h = Cm(5.5)
gap = Cm(0.4)
start_x = (SW - (col_w * 3 + gap * 2)) // 2
y = Cm(7.2)
colors = [SLATE, RED, COPPER]
for i, (title, body) in enumerate(risks):
    x = start_x + (col_w + gap) * i
    add_rect(s, x, y, col_w, Cm(1.3), colors[i])
    add_text(s, x, y, col_w, Cm(1.3),
             title, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1.3), col_w, col_h - Cm(1.3), TABLE_BG)
    add_text(s, x + Cm(0.4), y + Cm(1.7), col_w - Cm(0.8), col_h - Cm(2.1),
             body, size=12, color=DARK_GRAY, line_spacing=1.5)

add_box(s, Cm(1.2), Cm(14), Cm(31.4), Cm(3),
        "重要 ― 要件・経過措置は必ず確認を",
        "■ 管理者要件としての主任CMの扱いや、更新研修の要件・経過措置は、制度改正で変わりうる\n"
        "■ 「いつまでに・何を満たす必要があるか」は、都道府県・国の最新情報で必ず確認すること\n"
        "■ 本資料の記載は一般的な整理 ― 個別の判断は、必ず一次情報にあたる",
        title_bg=COPPER, body_bg=COPPER_BG, body_size=12)


# =============================================================
#  SLIDE 13: 逆算の計画
# =============================================================
s = add_slide()
add_header(s, "3. 更新研修の戦略", "「逆算」で計画する ― 更新研修受講のロードマップ")
add_footer(s, 13)

steps = [
    ("更新期限を\n確認する",
     "・自分・職員それぞれの主任CM更新の\n  期限を正確に把握する\n・介護支援専門員証の更新時期も併せて"),
    ("受講要件を\n満たすか確認",
     "・実務要件等を満たしているか確認\n・足りない場合は早めに手を打つ\n・都道府県の要件を確認"),
    ("受講時期を\n決め、申し込む",
     "・募集時期・定員を踏まえ、いつの回で\n  受けるかを決める\n・「期限ぎりぎり」を狙わない"),
    ("業務体制を\n整える",
     "・受講中の担当業務の代替を準備\n・チームで分担、利用者への説明\n・受講後の振り返りまで計画に含める"),
]
y = Cm(4.4)
colors = [FOREST_MID, COPPER, TEAL, SLATE]
for i, (title, body) in enumerate(steps):
    add_numbered_circle(s, Cm(1.5), y + Cm(0.4), Cm(1.2), i + 1, colors[i])
    add_rect(s, Cm(3.2), y, Cm(6.5), Cm(2.2), colors[i])
    add_text(s, Cm(3.3), y, Cm(6.3), Cm(2.2),
             title, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(9.7), y, SW - Cm(10.9), Cm(2.2), TABLE_BG)
    add_text(s, Cm(10.1), y + Cm(0.2), SW - Cm(11.5), Cm(1.9),
             body, size=11.5, color=DARK_GRAY,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
    y += Cm(2.45)

add_box(s, Cm(1.2), Cm(14.6), Cm(31.4), Cm(2.5),
        "「ぎりぎり」を狙わないこと",
        "■ 募集時期・定員・地域開催の有無により、希望の回で受けられないこともある\n"
        "■ 期限の1年以上前から動き出す ― 「受けられる回」に余裕を持って受講するのが鉄則",
        title_bg=FOREST, body_size=12)


# =============================================================
#  SLIDE 14: 受講中の業務
# =============================================================
s = add_slide()
add_header(s, "3. 更新研修の戦略", "受講中の業務をどう回すか")
add_footer(s, 14)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「研修に行けない」最大の理由は、業務がまわらないこと",
        "研修の数日間、担当CMが不在になる ― この体制を、事前に組織として準備しておく。",
        title_bg=FOREST, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "事前に準備すること",
        "■ 受講日程を早めに共有し、予定を空ける\n\n"
        "■ 担当ケースの引き継ぎ・代替担当の決定\n\n"
        "■ 緊急時の連絡・対応体制を決める\n\n"
        "■ 利用者・家族への事前の説明\n\n"
        "■ 訪問・モニタリング等の日程調整",
        title_bg=FOREST_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "組織で支える発想",
        "■ 「個人の都合」でなく「事業所の予定」に\n\n"
        "■ 研修受講を、業務計画に正式に組み込む\n\n"
        "■ 互いの研修を、チームで支え合う文化\n\n"
        "■ 1人に研修が重ならないよう年間で調整\n\n"
        "■ 主任CMが、この調整の旗振り役に",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.2),
         "▶ 「研修に行ける事業所」は、計画的に体制を整えている事業所。偶然ではなく仕組み。",
         size=13, color=FOREST, bold=True)


# =============================================================
#  SLIDE 15: 第4部
# =============================================================
section_title(4, "事業所内での\n計画的な人材育成",
              "Systematic Workforce Development")


# =============================================================
#  SLIDE 16: 研修受講管理
# =============================================================
s = add_slide()
add_header(s, "4. 計画的な人材育成", "研修受講を「管理」する ― 一覧化のすすめ")
add_footer(s, 16)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "「誰が・いつ・何を」を、見える化する",
        "職員一人ひとりの資格・研修の状況を一覧にしておくことが、計画的な育成の出発点。\n"
        "頭の中・個人任せにせず、事業所として把握・管理する。",
        title_bg=FOREST, body_size=12)

col_widths = [Cm(6), Cm(6.8), Cm(6.8), Cm(6.4), Cm(5.4)]
headers = ["職員", "資格・研修状況", "次の更新・研修", "受講予定", "受講要件"]
rows = [
    ["A(主任CM)", "主任CM", "主任CM更新研修", "○年度予定", "要件確認済"],
    ["B(CM)", "専門研修Ⅱ修了", "CM証 更新", "○年度", "—"],
    ["C(CM)", "専門研修Ⅰ修了", "専門研修Ⅱ", "次年度", "—"],
    ["D(CM)", "実務研修修了", "専門研修Ⅰ", "○年度", "実務経験を確認"],
]
add_table(s, Cm(1.7), Cm(7.3), col_widths, headers, rows,
          header_bg=FOREST, first_col_color=FOREST,
          row_height=Cm(1.5), body_size=11)

add_box(s, Cm(1.2), Cm(14.5), Cm(31.4), Cm(2.5),
        "受講管理表のポイント",
        "■ 資格・研修の修了状況、次の更新・研修の時期、受講要件を一覧に(様式は簡単なもので可)\n"
        "■ 年度初めに全員分を点検・更新 ― 「気づいたら期限切れ」を組織として防ぐ",
        title_bg=COPPER, body_bg=COPPER_BG, body_size=12)


# =============================================================
#  SLIDE 17: 育成計画への組み込み
# =============================================================
s = add_slide()
add_header(s, "4. 計画的な人材育成", "法定研修を「人材育成計画」に組み込む")
add_footer(s, 17)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "法定研修は、人材育成計画の「背骨」",
        "法定研修の受講時期は決まっている ― だからこそ、それを軸に\n"
        "事業所の人材育成計画を組み立てると、計画が現実的で続けやすくなる。",
        title_bg=FOREST, body_size=12)

add_box(s, Cm(1.2), Cm(7.3), Cm(15.5), Cm(5.5),
        "育成計画に盛り込む要素",
        "■ 各職員の法定研修の受講予定\n\n"
        "■ 主任CMをめざす職員の育成ステップ\n\n"
        "■ 事業所内研修・OJTの年間計画\n\n"
        "■ 外部研修・自己研鑽の機会\n\n"
        "■ 研修費用・受講中の体制の見込み",
        title_bg=FOREST_MID, body_size=12)

add_box(s, Cm(17.5), Cm(7.3), Cm(15), Cm(5.5),
        "「次の主任CM」を計画的に育てる",
        "■ 主任研修の受講要件(実務経験等)を\n   逆算し、対象者を早期に見極める\n\n"
        "■ 要件を満たせるよう、経験の機会を配分\n\n"
        "■ 主任研修の前から、指導・育成の経験を\n\n"
        "■ 主任CMの「複数配置」「後継」を見据える",
        title_bg=COPPER, body_bg=COPPER_BG, body_size=12)

add_text(s, Cm(1.2), Cm(13.3), SW - Cm(2.4), Cm(1.5),
         "▶ 人材育成計画は「立派な書類」より「毎年見直し、実際に使われる計画」であることが大切。\n"
         "▶ 法定研修の予定を軸にすれば、絵に描いた餅になりにくい。",
         size=13, color=FOREST, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 18: 費用と時間の確保
# =============================================================
s = add_slide()
add_header(s, "4. 計画的な人材育成", "研修の「費用」と「時間」をどう確保するか")
add_footer(s, 18)

add_box(s, Cm(1.2), Cm(4.5), Cm(15.5), Cm(6),
        "費用の確保",
        "■ 受講料・テキスト代・交通費等を把握\n\n"
        "■ 年間の研修費用を予算に組み込む\n\n"
        "■ 「誰が負担するか」を明確にする\n   (事業所負担/個人負担/折半 等)\n\n"
        "■ 活用できる助成・支援制度がないか確認\n\n"
        "■ 費用負担のルールを文書化しておく",
        title_bg=FOREST_MID, body_size=12)

add_box(s, Cm(17.5), Cm(4.5), Cm(15), Cm(6),
        "時間の確保",
        "■ 研修日を「業務予定」として確保する\n\n"
        "■ 受講中の業務代替を組織で準備\n\n"
        "■ 年間で受講が重ならないよう調整\n\n"
        "■ オンライン研修等、柔軟な受講方法も活用\n\n"
        "■ 「行きたくても行けない」をなくす",
        title_bg=TEAL, body_bg=TEAL_BG, body_size=12)

add_box(s, Cm(1.2), Cm(11), Cm(31.4), Cm(6),
        "「研修に投資する」という経営判断",
        "■ 研修の費用・時間は「コスト」ではなく、人材と事業所の質への「投資」\n\n"
        "■ 研修を受けやすい事業所は、職員の定着・採用の面でも強みになる\n\n"
        "■ 主任CMは、研修の意義を管理者・経営層に伝え、必要な資源の確保を後押しする\n\n"
        "■ 「育てる事業所」であることが、これからの人材獲得競争での生き残りにつながる",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 19: 第5部
# =============================================================
section_title(5, "研修を「活かす」",
              "Making Training Count")


# =============================================================
#  SLIDE 20: 受けっぱなしにしない
# =============================================================
s = add_slide()
add_header(s, "5. 研修を活かす", "「受けっぱなし」にしない ― 学びを定着させる")
add_footer(s, 20)

add_box(s, Cm(1.2), Cm(4.3), Cm(31.4), Cm(2.3),
        "研修の価値は「受けた後」に決まる",
        "時間と費用をかけて受けた研修も、職場に持ち帰って活かされなければ、\n"
        "「修了証」が残るだけ。学びを定着させ、還元する仕組みが要る。",
        title_bg=FOREST, body_size=12)

col_widths = [Cm(15.7), Cm(15.7)]
headers = ["受けっぱなしの研修", "活かされる研修"]
rows = [
    ["修了したら、それで終わり",
     "受講後に学びを振り返り、言葉にする"],
    ["資料がファイルにしまわれるだけ",
     "学びを事業所内で共有する場がある"],
    ["本人だけの学びにとどまる",
     "学びがチーム・他の職員に広がる"],
    ["日々の実務と結びつかない",
     "学びを具体的な実践・改善につなげる"],
    ["「行ってきました」報告で完了",
     "「何を実務に活かすか」まで確認する"],
]
add_table(s, Cm(1.2), Cm(7.3), col_widths, headers, rows,
          header_bg=FOREST, first_col_color=RED,
          row_height=Cm(1.5), body_size=12, first_col_bold=False)

add_text(s, Cm(1.2), Cm(15.3), SW - Cm(2.4), Cm(1.5),
         "▶ 「研修に行かせる」だけでなく「研修を活かす場をつくる」までが、育成のマネジメント。\n"
         "▶ ここを担うのが主任CMの役割。",
         size=13, color=FOREST, bold=True, line_spacing=1.4)


# =============================================================
#  SLIDE 21: 学びを還元する仕組み
# =============================================================
s = add_slide()
add_header(s, "5. 研修を活かす", "学びを事業所に還元する仕組み")
add_footer(s, 21)

ways = [
    ("報告・共有の場",
     "・研修内容を朝礼・会議で共有\n・「持ち帰った1つ」を必ず話す\n・資料を共有フォルダに整理"),
    ("事例検討との連動",
     "・研修で学んだ視点を事例検討に活かす\n・「研修で学んだ○○で考えると」\n・学びを実際のケースで試す"),
    ("OJT・SVとの連動",
     "・研修の学びを後輩指導に活かす\n・同行訪問で実践に落とし込む\n・スーパービジョンの中で振り返る"),
    ("実践・改善への反映",
     "・学びを業務の見直しにつなげる\n・「1つ変えてみる」を促す\n・後日、変化を振り返る"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [FOREST_MID, TEAL, COPPER, SLATE]
for i, (title, body) in enumerate(ways):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "法定研修と事業所内研修(OJT)をつなぐ",
        "■ 法定研修(外部)で得た「考え方・知識」を、事業所内研修・OJT(内部)で「実践」に変える\n"
        "■ 法定研修だけ、事業所内研修だけ、では育たない ― 両輪で回す\n"
        "■ 主任CMは、外部の学びと日々の実務を結びつける「橋渡し役」",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 22: 第6部
# =============================================================
section_title(6, "まとめ", "Wrap Up")


# =============================================================
#  SLIDE 23: 主任CMの役割
# =============================================================
s = add_slide()
add_header(s, "Closing", "研修・人材育成における主任CMの役割")
add_footer(s, 23)

roles = [
    ("情報を捉える",
     "・研修制度改正の動向を注視する\n・複数の確かな情報源を持つ\n・変化を事業所に正確に伝える"),
    ("計画をつくる",
     "・職員の研修受講を一覧で管理\n・人材育成計画に組み込む\n・「次の主任CM」を計画的に育てる"),
    ("体制を整える",
     "・受講中の業務体制を準備する\n・費用・時間の確保を後押しする\n・研修を受けやすい職場にする"),
    ("学びを活かす",
     "・「受けっぱなし」にしない\n・学びを共有・実践につなげる\n・法定研修とOJTを橋渡しする"),
]
col_w = Cm(15.5)
gap_x = Cm(0.4)
y_start = Cm(4.5)
row_h = Cm(3.7)
colors = [FOREST_MID, COPPER, TEAL, SLATE]
for i, (title, body) in enumerate(roles):
    col = i % 2
    row = i // 2
    x = Cm(1.2) + col * (col_w + gap_x)
    y = y_start + row * (row_h + Cm(0.3))
    add_rect(s, x, y, col_w, Cm(1), colors[i])
    add_text(s, x + Cm(0.5), y, col_w - Cm(1), Cm(1),
             title, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, y + Cm(1), col_w, row_h - Cm(1), TABLE_BG)
    add_text(s, x + Cm(0.5), y + Cm(1.2), col_w - Cm(1), row_h - Cm(1.4),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)

add_box(s, Cm(1.2), Cm(13.3), Cm(31.4), Cm(3.7),
        "「研修を受ける人」から「育成を設計する人」へ",
        "■ 主任CM自身も研修を受け続ける一人 ― だが、それだけにとどまらない\n"
        "■ 事業所全体の研修・育成を「設計し、まわす」のが主任CMの役割\n"
        "■ 計画的な人材育成こそ、ケアの質と、事業所の持続性を支える土台になる",
        title_bg=FOREST, body_size=13)


# =============================================================
#  SLIDE 24: 持ち帰りアクション
# =============================================================
s = add_slide()
add_header(s, "Closing", "明日からの 3 つのアクション")
add_footer(s, 24)

actions = [
    ("Action 1", "自分と職員の「更新期限」を確認する",
     "・主任CM更新研修・介護支援専門員証の更新時期を、全員分洗い出す\n"
     "・受講要件を満たしているかも併せて確認する"),
    ("Action 2", "研修受講管理表をつくる/見直す",
     "・「誰が・いつ・何の研修を」受けるかを一覧化する\n"
     "・年度初めに点検する習慣を、事業所のルールにする"),
    ("Action 3", "次の研修の「学びの還元」を設計する",
     "・次に職員が研修に行くとき、受講後に学びを共有する場をあらかじめ決めておく\n"
     "・「持ち帰った1つを実務でどう活かすか」まで確認する"),
]
y = Cm(4.5)
colors = [FOREST_MID, COPPER, TEAL]
for i, (label, title, body) in enumerate(actions):
    color = colors[i]
    add_rect(s, Cm(1.2), y, Cm(4.5), Cm(3.7), color)
    add_text(s, Cm(1.2), y + Cm(0.4), Cm(4.5), Cm(1),
             label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Cm(1.2), y + Cm(1.5), Cm(4.5), Cm(1.8),
             label.split(" ")[1], size=24, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Cm(5.7), y, SW - Cm(6.9), Cm(3.7), TABLE_BG)
    add_text(s, Cm(6), y + Cm(0.3), SW - Cm(7.5), Cm(1),
             title, size=16, bold=True, color=FOREST)
    add_text(s, Cm(6), y + Cm(1.4), SW - Cm(7.5), Cm(2.2),
             body, size=12, color=DARK_GRAY, line_spacing=1.4)
    y += Cm(4)


# =============================================================
#  SLIDE 25: クロージング
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, FOREST)
add_rect(s, 0, Cm(7), SW, Cm(0.1), COPPER)

add_text(s, Cm(2), Cm(3), SW - Cm(4), Cm(1),
         "Closing Message", size=14, color=COPPER, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(8), SW - Cm(4), Cm(3),
         "研修は「義務」ではなく、\n人と事業所の未来への投資。",
         size=26, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         line_spacing=1.4)
add_text(s, Cm(3), Cm(12.5), SW - Cm(6), Cm(3),
         "制度は変わり続け、研修の負担に追われる日々かもしれません。\n"
         "それでも、計画的に学びを設計し、一人ひとりの成長を支えること。\n"
         "その積み重ねが、ケアの質と、あなたの事業所の明日を確かに支えます。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER, line_spacing=1.6)


# =============================================================
#  SLIDE 26: 参考資料
# =============================================================
s = add_slide()
add_header(s, "Reference", "参考資料 ・ 情報源")
add_footer(s, 26)

refs = [
    "■ 厚生労働省「介護支援専門員(ケアマネジャー)」関連 ― 研修制度・検討会資料",
    "■ 各都道府県「介護支援専門員研修」実施要綱・募集要項(研修の実務情報)",
    "■ 各都道府県「主任介護支援専門員更新研修」の受講要件・日程情報",
    "■ 日本介護支援専門員協会等による研修制度・改正動向の発信",
    "■ 介護保険法・関係省令(研修・更新の根拠)",
    "■ 居宅介護支援事業所の運営基準(管理者要件に関する規定)",
]
y = Cm(4.5)
for r in refs:
    add_text(s, Cm(2), y, SW - Cm(4), Cm(1),
             r, size=13, color=DARK_GRAY)
    y += Cm(1.3)

add_box(s, Cm(1.2), Cm(13), Cm(31.4), Cm(4),
        "ご注意 ・ 情報の確認について",
        "■ 法定研修の体系・時間数・受講方法・更新要件、管理者要件等は、制度改正により変わります\n"
        "■ 研修の名称・要件・運用は都道府県ごとに異なります ― 必ず自県・国の最新情報をご確認ください\n"
        "■ 個別の受講判断・更新の可否は、都道府県の研修担当・実施機関に直接ご確認ください\n"
        "■ 本資料は 2026年5月時点の一般的な情報を基に構成しています",
        title_bg=COPPER, body_bg=COPPER_BG, body_size=12)


# =============================================================
#  SLIDE 27: 質疑応答
# =============================================================
s = add_slide()
add_rect(s, 0, 0, SW, SH, FOREST)
add_rect(s, Cm(11), Cm(7.5), Cm(11.8), Cm(0.12), COPPER)
add_text(s, Cm(2), Cm(7.8), SW - Cm(4), Cm(2),
         "ご清聴ありがとうございました",
         size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(11), SW - Cm(4), Cm(1.5),
         "質疑応答 ・ 情報交換",
         size=18, color=COPPER, align=PP_ALIGN.CENTER)
add_text(s, Cm(2), Cm(13), SW - Cm(4), Cm(1.5),
         "各事業所の研修管理・人材育成の工夫を持ち寄り、学び合いましょう。",
         size=14, color=MID_GRAY, align=PP_ALIGN.CENTER)


# ===== 保存 =====
output = "/home/user/kaigo-ga-app/slides/法定研修制度改正と主任CM更新研修の戦略_研修スライド.pptx"
prs.save(output)
print(f"Saved: {output}")
print(f"Total slides: {len(prs.slides)}")
