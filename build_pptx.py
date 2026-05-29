# -*- coding: utf-8 -*-
"""
特定事業所加算II 内部勉強会スライド（困難ケースの受け入れ体制）
PowerPoint(.pptx) 生成スクリプト
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- カラーパレット ----
NAVY   = RGBColor(0x0F, 0x30, 0x57)
BLUE   = RGBColor(0x1E, 0x60, 0x91)
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)
WARM   = RGBColor(0xE7, 0x6F, 0x51)
GOLD   = RGBColor(0xE9, 0xC4, 0x6A)
BG     = RGBColor(0xF4, 0xF7, 0xFA)
INK    = RGBColor(0x23, 0x30, 0x3A)
MUTED  = RGBColor(0x5A, 0x6B, 0x76)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xD7, 0xE1, 0xEA)
CARDBG = RGBColor(0xFF, 0xFF, 0xFF)
GOODBG = RGBColor(0xE9, 0xF7, 0xF4)
BADBG  = RGBColor(0xFD, 0xEE, 0xE9)
NOTEBG = RGBColor(0xFF, 0xF7, 0xE6)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    _set_fill(bg, BG)
    bg.shadow.inherit = False
    return s


def grad_bg(s, c1, c2):
    """対角グラデーション背景"""
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background()
    bg.shadow.inherit = False
    sp = bg.fill._xPr  # spPr
    # remove existing fill
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    grad = sp.makeelement(qn('a:gradFill'), {})
    gsLst = grad.makeelement(qn('a:gsLst'), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = grad.makeelement(qn('a:gs'), {'pos': str(pos)})
        srgb = grad.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (col[0], col[1], col[2])})
        gs.append(srgb)
        gsLst.append(gs)
    grad.append(gsLst)
    lin = grad.makeelement(qn('a:lin'), {'ang': '2700000', 'scaled': '1'})
    grad.append(lin)
    # insert grad fill before line props
    ln = sp.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp.append(grad)
    return bg


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    return tb, tf


def para(tf, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=4, bullet=False, level=0, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.level = level
    runs = text if isinstance(text, list) else [(text, color, bold)]
    for seg in runs:
        t, c, b = seg
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.color.rgb = c
        r.font.name = font
    if bullet:
        _bullet(p)
    else:
        _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def _bullet(p, char="・", color=ACCENT):
    pPr = p._p.get_or_add_pPr()
    pPr.set('indent', '-228600')
    pPr.set('marL', '228600')
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buf = pPr.makeelement(qn('a:buFont'), {'typeface': FONT})
    pPr.append(buf)
    bu = pPr.makeelement(qn('a:buChar'), {'char': char})
    pPr.append(bu)


def rect(s, x, y, w, h, fill=CARDBG, line=None, rounded=True, line_w=1.0):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def leftbar(s, x, y, w, h, color):
    """左に縦バーのある見出し用"""
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(8), h)
    _set_fill(bar, color)
    bar.shadow.inherit = False


def kicker(s, text, color=ACCENT):
    tb, tf = textbox(s, Inches(0.7), Inches(0.4), Inches(8), Inches(0.4))
    para(tf, text, size=13, color=color, bold=True, first=True)
    return tb


def heading(s, text):
    leftbar(s, Inches(0.7), Inches(0.78), Inches(0.12), Inches(0.62), ACCENT)
    tb, tf = textbox(s, Inches(0.95), Inches(0.72), Inches(11.6), Inches(0.8))
    para(tf, text, size=30, color=NAVY, bold=True, first=True)


def pill(s, x, y, text, color=ACCENT, tcolor=WHITE, w=Inches(2.0)):
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
    _set_fill(p, color)
    p.shadow.inherit = False
    try:
        p.adjustments[0] = 0.5
    except Exception:
        pass
    tf = p.text_frame; tf.word_wrap = True
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    para(tf, text, size=12, color=tcolor, bold=True, align=PP_ALIGN.CENTER, first=True)
    return p


def card(s, x, y, w, h, title=None, bullets=None, fill=CARDBG, line=LINE,
         title_color=BLUE, body_size=15, title_size=18):
    rect(s, x, y, w, h, fill=fill, line=line)
    pad = Inches(0.22)
    tb, tf = textbox(s, x + pad, y + Inches(0.12), w - pad * 2, h - Inches(0.2))
    first = True
    if title:
        para(tf, title, size=title_size, color=title_color, bold=True, first=True, space_after=5)
        first = False
    if bullets:
        for b in bullets:
            para(tf, b, size=body_size, color=INK, first=first, bullet=True, space_after=3)
            first = False
    return tf


def note(s, x, y, w, h, runs, bar_color=GOLD, fill=NOTEBG):
    rect(s, x, y, w, h, fill=fill, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(6), h)
    _set_fill(bar, bar_color); bar.shadow.inherit = False
    tb, tf = textbox(s, x + Inches(0.2), y, w - Inches(0.35), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, runs, size=14, first=True)


def footer(s, n):
    tb, tf = textbox(s, Inches(11.6), Inches(7.05), Inches(1.5), Inches(0.35))
    para(tf, "%d / %d" % (n, TOTAL), size=11, color=MUTED, align=PP_ALIGN.RIGHT, first=True)
    tb2, tf2 = textbox(s, Inches(0.7), Inches(7.05), Inches(9), Inches(0.35))
    para(tf2, "特定事業所加算II 内部勉強会｜困難ケース受け入れ体制", size=10, color=MUTED, first=True)


TOTAL = 22
_n = [0]


def newslide(kick=None, title=None, grad=None):
    _n[0] += 1
    if grad:
        s = prs.slides.add_slide(BLANK)
        grad_bg(s, grad[0], grad[1])
    else:
        s = slide()
        if kick:
            kicker(s, kick)
        if title:
            heading(s, title)
        footer(s, _n[0])
    return s


# ======================================================================
# Slide 1 タイトル
# ======================================================================
s = newslide(grad=(NAVY, ACCENT))
# badge
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.4), Inches(5.6), Inches(0.5))
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x2B, 0x55, 0x7A)
b.line.color.rgb = RGBColor(0x8F, 0xB6, 0xD4); b.line.width = Pt(1); b.shadow.inherit = False
btf = b.text_frame; para(btf, "特定事業所加算II 算定事業所　内部勉強会資料", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.4))
para(tf, "困難ケースの受け入れ体制で実現する", size=44, color=WHITE, bold=True, first=True, space_after=2)
para(tf, "確実性のある業務推進", size=44, color=WHITE, bold=True, space_after=10)
para(tf, "～ 加算IIの責任を、現場の力に変える ～", size=22, color=GOLD, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.5))
para(tf2, "対象：サービス提供責任者・訪問介護員・管理者　／　形式：事例検討型グループワーク",
     size=15, color=RGBColor(0xCF, 0xE0, 0xEE), first=True, space_after=3)
para(tf2, "開催日：2026年　　月　　日　　　会場：", size=15, color=RGBColor(0xCF, 0xE0, 0xEE))

# ======================================================================
# Slide 2 ねらい
# ======================================================================
s = newslide("PURPOSE", "本日のねらい")
tb, tf = textbox(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.6))
para(tf, [("「断らない事業所」ではなく、", MUTED, False),
          ("「組織として安全に受けられる事業所」", WARM, True),
          ("を目指す。", MUTED, False)], size=20, first=True)
cw, gap, x0, y0, ch = Inches(3.85), Inches(0.28), Inches(0.7), Inches(2.35), Inches(2.0)
data = [("1. 共通理解", BLUE, "特定事業所加算IIの趣旨と、私たちに求められる体制を再確認する。"),
        ("2. 体制の見える化", ACCENT, "困難ケースを「個人」でなく「仕組み」で受ける流れを理解する。"),
        ("3. 事例で実践", WARM, "4つの事例で、判断・連携・記録の勘所を体得する。")]
for i, (t, c, body) in enumerate(data):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    pill(s, x + Inches(0.22), y0 + Inches(0.2), t, color=c, w=Inches(1.9))
    tb, tf = textbox(s, x + Inches(0.22), y0 + Inches(0.75), cw - Inches(0.44), ch - Inches(0.9))
    para(tf, body, size=15, first=True)
note(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(0.95),
     [("合言葉：", WARM, True), ("属人化を減らし、誰が対応しても一定の質を保てる仕組みをつくる。", INK, False)])

# ======================================================================
# Slide 3 アジェンダ
# ======================================================================
s = newslide("AGENDA", "本日の流れ")
ag = [("第1部　理解編（約20分）", ["特定事業所加算IIの位置づけ", "なぜ困難ケース受け入れ体制が要なのか", "「困難ケース」の定義と類型"]),
      ("第2部　体制編（約20分）", ["受け入れの基本フロー（5ステップ）", "5つの体制づくり（情報・チーム・緊急時・記録・育成）"]),
      ("第3部　事例編（約30分）", ["事例1〜4のグループ検討", "NG対応とGOOD対応の比較"]),
      ("第4部　まとめ（約10分）", ["自事業所の体制セルフチェック", "明日からのアクション宣言"])]
cw, ch, gx, gy = Inches(5.85), Inches(2.1), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.75)
for i, (t, items) in enumerate(ag):
    x = x0 + (i % 2) * (cw + gx)
    y = y0 + (i // 2) * (ch + gy)
    card(s, x, y, cw, ch, title=t, bullets=items, title_color=BLUE, body_size=15, title_size=18)

# ======================================================================
# Slide 4 加算IIの位置づけ（表）
# ======================================================================
s = newslide("BASIC", "特定事業所加算IIの位置づけ（再確認）")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, [("特定事業所加算は、", INK, False), ("質の高いサービスを安定的に提供する体制", WARM, True),
          ("を評価する加算です。区分により求める要件が異なります。", INK, False)], size=16, first=True)
# table
rows, cols = 4, 4
tbl_shape = s.shapes.add_table(rows, cols, Inches(0.7), Inches(2.1), Inches(11.9), Inches(2.0))
table = tbl_shape.table
table.columns[0].width = Inches(3.5)
for c in range(1, 4):
    table.columns[c].width = Inches(2.8)
hdr = ["区分", "体制要件", "人材要件", "重度者等対応要件"]
body = [["加算I", "○", "○", "○"],
        ["加算II（当事業所）", "○", "○", "－"],
        ["加算III", "○", "－", "○"]]
for c, h in enumerate(hdr):
    cell = table.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = h; r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
for ri, row in enumerate(body, start=1):
    highlight = (ri == 2)
    for ci, val in enumerate(row):
        cell = table.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xFF, 0xF3, 0xCD) if highlight else (RGBColor(0xEE, 0xF4, 0xF9) if ri % 2 == 0 else WHITE)
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = val; r.font.size = Pt(14)
        r.font.bold = (ci == 0) or highlight; r.font.color.rgb = INK; r.font.name = FONT
note(s, Inches(0.7), Inches(4.35), Inches(11.9), Inches(1.55),
     [("加算IIは", INK, False), ("「体制要件」＋「人材要件（有資格者の割合・勤続年数等）」", WARM, True),
      ("で算定。重度者対応は必須ではないが、体制要件である研修・会議・緊急時対応・健康管理の質こそが、困難ケースを受けられる組織力の土台となる。", INK, False)])
tb, tf = textbox(s, Inches(0.9), Inches(6.0), Inches(11.6), Inches(0.5))
para(tf, "※要件・割合の詳細は最新の介護報酬告示・解釈通知（直近改定版）で必ずご確認ください。本資料は体制づくりの考え方を共有するものです。",
     size=11, color=MUTED, first=True)

# ======================================================================
# Slide 5 なぜ要なのか
# ======================================================================
s = newslide("WHY", "なぜ「困難ケース受け入れ体制」が要なのか")
items = [("① 加算の趣旨に直結する", "加算IIの体制要件（個別研修計画／定期会議／文書による指示と報告／緊急時対応／健康管理）は、そのまま困難ケースを安全に受けるための装置。"),
         ("② 加算の「説明責任」になる", "加算算定＝対外的に質を約束すること。実地指導の場でも、体制が機能していることを記録で示せることが重要。"),
         ("③ 地域からの信頼につながる", "「あの事業所なら相談できる」という評価が、安定した依頼・稼働率・人材定着を生む。"),
         ("④ 職員を守る", "体制があれば、難しいケースを個人の頑張りに依存させない。バーンアウト・離職・事故を防ぐ。")]
cw, ch, gx, gy = Inches(5.85), Inches(2.05), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.7)
for i, (t, b) in enumerate(items):
    x = x0 + (i % 2) * (cw + gx)
    y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    tb, tf = textbox(s, x + Inches(0.25), y + Inches(0.18), cw - Inches(0.5), ch - Inches(0.3))
    para(tf, t, size=18, color=BLUE, bold=True, first=True, space_after=5)
    para(tf, b, size=15, color=INK)

# ======================================================================
# Slide 6 困難ケースとは
# ======================================================================
s = newslide("DEFINE", "「困難ケース」とは何か")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, "「対応が難しい」を分解すると、原因ごとに必要な備えが見えてくる。", size=18, color=MUTED, first=True)
items = [("医療・身体面", BLUE, "重度（要介護4・5）、喀痰吸引・経管栄養等の医療的ケア、ターミナル、急変リスク。"),
         ("認知・精神面", ACCENT, "認知症によるBPSD（拒否・暴言・徘徊）、精神疾患、意思疎通の困難。"),
         ("家族・環境面", WARM, "家族関係の葛藤、介護力不足、ハラスメント、不衛生な住環境、孤立。"),
         ("社会・制度面", GOLD, "経済的困窮、サービス拒否、虐待の疑い、多問題を抱える世帯。")]
cw, ch, gx, gy = Inches(5.85), Inches(1.55), Inches(0.3), Inches(0.28)
x0, y0 = Inches(0.7), Inches(2.1)
for i, (t, c, b) in enumerate(items):
    x = x0 + (i % 2) * (cw + gx)
    y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    tcol = RGBColor(0x5A, 0x45, 0x00) if c == GOLD else WHITE
    pill(s, x + Inches(0.22), y + Inches(0.18), t, color=c, tcolor=tcol, w=Inches(1.9))
    tb, tf = textbox(s, x + Inches(0.22), y + Inches(0.66), cw - Inches(0.44), ch - Inches(0.8))
    para(tf, b, size=14, first=True)
note(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.1),
     [("視点：", ACCENT, True), ("困難さは利用者の「属性」ではなく、支援する側の体制の不足から生まれることもある。体制を整えれば「困難」は「対応可能」に変わる。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# Slide 7 受け入れフロー
# ======================================================================
s = newslide("FLOW", "困難ケース受け入れの基本フロー")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, [("「受けるか／断るか」を個人が即答しない。", INK, False), ("5ステップで組織判断する。", WARM, True)], size=18, first=True)
steps = [("1", "相談受付", "情報を一旦すべて受け止める"),
         ("2", "情報収集・\nアセスメント", "リスクと必要体制を見極める"),
         ("3", "受入判定会議", "組織で可否・条件を判断"),
         ("4", "体制構築・開始", "役割分担・手順・緊急時を明示"),
         ("5", "モニタリング", "定期会議で改善・撤退判断")]
n = len(steps)
total_w = Inches(11.9)
sw = Inches(2.15); gap = (total_w - sw * n) / (n - 1)
x0, y0, h = Inches(0.7), Inches(2.2), Inches(1.7)
for i, (num, t, sub) in enumerate(steps):
    x = x0 + i * (sw + gap)
    box = rect(s, x, y0, sw, h, fill=WHITE, line=BLUE, line_w=2)
    # number circle
    cr = s.shapes.add_shape(MSO_SHAPE.OVAL, x + sw/2 - Inches(0.22), y0 + Inches(0.15), Inches(0.44), Inches(0.44))
    _set_fill(cr, BLUE); cr.shadow.inherit = False
    para(cr.text_frame, num, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, x + Inches(0.08), y0 + Inches(0.62), sw - Inches(0.16), h - Inches(0.7), anchor=MSO_ANCHOR.TOP)
    para(tf, t, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=3)
    para(tf, sub, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    if i < n - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + sw - Inches(0.02), y0 + h/2 - Inches(0.12), gap + Inches(0.04), Inches(0.24))
        _set_fill(ar, BLUE); ar.shadow.inherit = False
note(s, Inches(0.7), Inches(4.3), Inches(11.9), Inches(0.85),
     [("NG：", WARM, True), ("電話口でサ責が一人で「うちでは無理です」／逆に「大丈夫です受けます」と即断する。", INK, False)],
     bar_color=WARM, fill=BADBG)
note(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(0.85),
     [("GOOD：", ACCENT, True), ("「持ち帰ってチームで検討し、◯日までにご連絡します」と、判断を仕組みに乗せる。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 体制づくり 8-12 共通関数
# ======================================================================
def system_slide(kick, title, left_t, left_items, right_t, right_items, note_runs, note_kind="note"):
    s = newslide(kick, title)
    cw, ch, gx = Inches(5.85), Inches(3.05), Inches(0.3)
    x0, y0 = Inches(0.7), Inches(1.7)
    card(s, x0, y0, cw, ch, title=left_t, bullets=left_items, title_color=BLUE, body_size=15)
    card(s, x0 + cw + gx, y0, cw, ch, title=right_t, bullets=right_items, title_color=BLUE, body_size=15)
    if note_kind == "good":
        note(s, Inches(0.7), Inches(5.05), Inches(11.9), Inches(1.15), note_runs, bar_color=ACCENT, fill=GOODBG)
    else:
        note(s, Inches(0.7), Inches(5.05), Inches(11.9), Inches(1.15), note_runs)
    return s

# Slide 8
system_slide("SYSTEM 1/5", "体制づくり①　情報収集とアセスメント",
    "受付時に必ず押さえる項目",
    ["心身状態・医療的ケアの有無", "認知症・精神症状とその現れ方", "家族構成・介護力・キーパーソン",
     "住環境・動物・衛生・段差等", "過去のサービス利用歴と中止理由", "これまで何が「困難」とされたか"],
    "リスクアセスメントの観点",
    ["利用者の安全（急変・転倒・服薬）", "職員の安全（感染・腰痛・ハラスメント）", "必要なスキル・資格は揃うか",
     "必要な人員・時間を確保できるか", "連携先（医療・ケアマネ・地域包括）"],
    [("ポイント：", GOLD, True), ("「受けられない理由」だけでなく「何が揃えば受けられるか（条件）」をセットで整理する。", INK, False)])

# Slide 9
system_slide("SYSTEM 2/5", "体制づくり②　チーム体制・役割分担",
    "「一人で抱えない」設計",
    ["主担当＋副担当の複数担当制", "困難ケースは経験者とペアで導入", "サ責が司令塔として手順書を整備",
     "相談・SOSを上げてよい文化を明示"],
    "サービス提供責任者の役割",
    ["受入判定会議の招集・進行", "個別の手順書・留意点の作成と周知", "ヘルパーへの文書による指示",
     "提供後の報告の収集とフィードバック"],
    [("加算要件との接続：", ACCENT, True), ("「文書等による指示及び提供後の報告」「定期的な会議」は加算IIの体制要件そのもの。困難ケースほど、この運用が品質を左右する。", INK, False)],
    note_kind="good")

# Slide 10
system_slide("SYSTEM 3/5", "体制づくり③　緊急時対応・リスク管理",
    "緊急時対応の「明示」",
    ["急変・事故・連絡不能時の連絡フローを1枚に", "連絡先（家族／主治医／ケアマネ／管理者／119）の優先順位",
     "各利用者宅に応じた個別の留意事項", "夜間・休日の連絡当番"],
    "職員を守るリスク管理",
    ["ハラスメント発生時の手順（中断・退避・報告）", "感染症・身体的負担への備え", "ヒヤリ・ハットの収集と共有",
     "「撤退・中止」も正当な選択肢と明文化"],
    [("カスタマーハラスメント：", GOLD, True), ("「利用者・家族だから我慢」ではなく、組織として線引きと対応手順を持つことが、結果的にサービス継続を可能にする。", INK, False)])

# Slide 11
system_slide("SYSTEM 4/5", "体制づくり④　記録・報告・情報共有",
    "記録は「確実性」の証拠",
    ["受入判定会議の議事録（誰が・何を根拠に判断したか）", "個別手順書・指示書とその更新履歴",
     "日々のサービス提供記録・報告", "ヒヤリ・ハット／事故報告／苦情対応"],
    "情報共有の仕組み",
    ["申し送りノート・チャット等の一元化", "変化があればその日のうちに共有", "定期会議で全員が同じ理解を持つ",
     "ケアマネ・多職種への報告ルート"],
    [("実地指導の視点：", ACCENT, True), ("「やっている」を「記録で示せる」へ。困難ケースほど、判断の根拠と経過の記録が事業所を守る。", INK, False)],
    note_kind="good")

# Slide 12
system_slide("SYSTEM 5/5", "体制づくり⑤　研修・人材育成",
    "個別研修計画（加算要件）",
    ["一人ひとりの課題に応じた個別研修計画の作成・実施", "困難ケース対応スキル（認知症ケア・医療的ケア・接遇）を計画に位置づけ",
     "OJT＝経験者同行による実地学習"],
    "本日のような事例検討会",
    ["うまくいった／いかなかった事例を共有", "「自分ならどうするか」を言語化", "暗黙知を手順・ナレッジに変える",
     "新人も同じ判断軸を持てるようにする"],
    [("狙い：", GOLD, True), ("研修を「加算のための義務」で終わらせず、困難ケースを受けられる現場力に転換する。", INK, False)])

# ======================================================================
# Slide 13 事例セクション扉
# ======================================================================
s = newslide(grad=(WARM, GOLD))
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.5), Inches(2.4), Inches(0.5))
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xC8, 0x55, 0x3A); b.line.fill.background(); b.shadow.inherit = False
para(b.text_frame, "CASE STUDY", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.2))
para(tf, "第3部　事例で学ぶ", size=42, color=WHITE, bold=True, first=True, space_after=2)
para(tf, "困難ケースの受け入れ", size=42, color=WHITE, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(1.4))
para(tf2, "各事例：① 状況把握 → ② NG対応 → ③ GOOD対応 → ④ 体制のポイント", size=18, color=WHITE, bold=True, first=True, space_after=6)
para(tf2, "グループで「自分の事業所ならどう動くか」を3分間ディスカッションしてください。", size=15, color=WHITE)

# ======================================================================
# 事例 共通関数
# ======================================================================
def case_slide(emoji_title, situation, ng_items, good_items, point_runs):
    s = newslide()
    # header bar
    hd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.75))
    grad_fill_shape(hd, WARM, GOLD)
    hd.line.fill.background(); hd.shadow.inherit = False
    tf = hd.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, emoji_title, size=22, color=WHITE, bold=True, first=True)
    # situation
    rect(s, Inches(0.7), Inches(1.4), Inches(11.9), Inches(1.15), fill=WHITE, line=LINE)
    tb, tf = textbox(s, Inches(0.95), Inches(1.5), Inches(11.4), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [("状況：", NAVY, True)] + situation, size=15, first=True)
    # NG / GOOD
    cw, gx, y0, ch = Inches(5.85), Inches(0.3), Inches(2.75), Inches(2.35)
    x0 = Inches(0.7)
    rect(s, x0, y0, cw, ch, fill=BADBG, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, y0, Pt(6), ch); _set_fill(bar, WARM); bar.shadow.inherit = False
    tb, tf = textbox(s, x0 + Inches(0.25), y0 + Inches(0.12), cw - Inches(0.4), ch - Inches(0.2))
    para(tf, "NG対応", size=17, color=WARM, bold=True, first=True, space_after=4)
    for it in ng_items:
        para(tf, it, size=14, bullet=True, space_after=3)
    x1 = x0 + cw + gx
    rect(s, x1, y0, cw, ch, fill=GOODBG, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x1, y0, Pt(6), ch); _set_fill(bar, ACCENT); bar.shadow.inherit = False
    tb, tf = textbox(s, x1 + Inches(0.25), y0 + Inches(0.12), cw - Inches(0.4), ch - Inches(0.2))
    para(tf, "GOOD対応", size=17, color=ACCENT, bold=True, first=True, space_after=4)
    for it in good_items:
        para(tf, it, size=14, bullet=True, space_after=3)
    note(s, Inches(0.7), Inches(5.35), Inches(11.9), Inches(1.1), [("体制のポイント：", GOLD, True)] + point_runs)
    return s


def grad_fill_shape(shape, c1, c2):
    sp = shape.fill._xPr
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    grad = sp.makeelement(qn('a:gradFill'), {})
    gsLst = grad.makeelement(qn('a:gsLst'), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = grad.makeelement(qn('a:gs'), {'pos': str(pos)})
        srgb = grad.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (col[0], col[1], col[2])})
        gs.append(srgb); gsLst.append(gs)
    grad.append(gsLst)
    lin = grad.makeelement(qn('a:lin'), {'ang': '0', 'scaled': '1'})
    grad.append(lin)
    ln = sp.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp.append(grad)


# Slide 14 事例1
case_slide("🧩 事例1：訪問拒否のある認知症（BPSD）のケース",
    [("独居・要介護3、アルツハイマー型認知症。ヘルパーを「泥棒」と呼び玄関で拒否。日によって態度が大きく変わり、担当が定まらず複数のヘルパーが交代で入っていた。", INK, False)],
    ["毎回ちがうヘルパーが訪問し関係が築けない", "「拒否されたので帰りました」を個人判断で繰り返す", "記録が共有されず原因が分析されない"],
    ["主担当を固定し、声かけ・訪問時間を統一", "拒否のパターンを記録・分析し手順書化", "ケアマネ・主治医と共有し服薬・受診も検討"],
    [("BPSDは「対応の一貫性」が鍵。担当固定＋記録の蓄積＋チームでの分析を仕組みにする（＝情報共有・会議・研修の体制が効く）。", INK, False)])

# Slide 15 事例2
case_slide("🩺 事例2：医療依存度が高いケース（喀痰吸引等）",
    [("要介護5、神経難病で喀痰吸引・経管栄養が必要。退院に伴い在宅へ。家族から「夜間も含めて対応してほしい」と依頼。", INK, False)],
    ["資格・体制を確認せず「受けます」と即答", "研修未修了のヘルパーに任せてしまう", "急変時の連絡先・手順が未整備"],
    ["喀痰吸引等研修修了者と登録状況を確認のうえ受入判定", "主治医・訪問看護と役割分担と指示書を明確化", "急変時フロー・緊急連絡網を1枚にして全員に配布"],
    [("「できる人がいるか」「指示系統が明確か」「緊急時が定まっているか」の3点を会議で確認してから受ける。医療連携は文書で残す。", INK, False)])

# Slide 16 事例3
case_slide("⚠️ 事例3：家族関係が複雑／ハラスメントリスク",
    [("同居の家族がサービス内容に細かく介入し、ヘルパーに強い口調で要求を繰り返す。担当ヘルパーが訪問のたびに精神的に疲弊。担当を外してほしいと相談があった。", INK, False)],
    ["「利用者のため」と我慢を強い、個人で抱えさせる", "担当を黙って交代させ根本対応をしない", "記録に残さず管理者まで情報が上がらない"],
    ["事実を記録し、管理者・ケアマネと共有", "契約・重要事項に基づく対応範囲を家族へ説明", "ハラスメント手順に沿い、必要なら複数名訪問・条件提示"],
    [("職員を守ることがサービス継続の前提。「我慢」ではなく「手順」で対応。組織として線引きを示すことが信頼にもつながる。", INK, False)])

# Slide 17 事例4
case_slide("🏠 事例4：セルフネグレクト／不衛生な住環境",
    [("独居・要介護2。室内にゴミが堆積、入浴・受診を拒否。近隣からの苦情もあり、地域包括支援センターから相談が入った。", INK, False)],
    ["「不衛生だから」と訪問介護単独で抱え込む", "本人の意思を無視して一気に片付けようとする", "虐待・権利擁護の視点が抜ける"],
    ["地域包括・行政・ケアマネと役割分担して関わる", "本人との関係構築を優先し段階的に支援", "セルフネグレクト＝権利擁護課題として記録・共有"],
    [("多問題ケースは1事業所で抱えない。地域のネットワークに「つなぐ」ことも、困難ケース受け入れ体制の一部。", INK, False)])

# ======================================================================
# Slide 18 組織で受ける3原則
# ======================================================================
s = newslide("ORGANIZATION", "困難ケースを「組織で受ける」3原則")
items = [("原則1", BLUE, "一人で判断しない", "受入可否・撤退は会議で決める。個人を矢面に立たせない。"),
         ("原則2", ACCENT, "条件で受ける", "「Yes/No」でなく「何が揃えば受けられるか」で考える。"),
         ("原則3", WARM, "記録で残す", "判断根拠・経過・連携を残し、誰でも引き継げる状態にする。")]
cw, gap, x0, y0, ch = Inches(3.85), Inches(0.28), Inches(0.7), Inches(1.7), Inches(2.5)
for i, (p, c, t, b) in enumerate(items):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    pill(s, x + Inches(0.22), y0 + Inches(0.2), p, color=c, w=Inches(1.5))
    tb, tf = textbox(s, x + Inches(0.25), y0 + Inches(0.75), cw - Inches(0.5), ch - Inches(0.9))
    para(tf, t, size=20, color=NAVY, bold=True, first=True, space_after=6)
    para(tf, b, size=15, color=INK)
note(s, Inches(0.7), Inches(4.5), Inches(11.9), Inches(1.4),
     [("これが「確実性のある業務推進」：", ACCENT, True),
      ("属人的な勘や善意ではなく、仕組みと記録で再現性を担保する。これこそ加算IIを算定する事業所の責任の果たし方。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# Slide 19 連携マップ
# ======================================================================
s = newslide("NETWORK", "多職種・地域との連携マップ")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, "困難ケースほど、訪問介護「だけ」では完結しない。つなぐ先を全員が知っておく。", size=17, color=MUTED, first=True)
items = [("医療", "主治医／訪問看護／薬剤師。急変・医療的ケア・服薬管理の連携。指示は文書で。"),
         ("介護・相談", "ケアマネ（中心）／福祉用具／通所。情報共有と担当者会議の活用。"),
         ("地域・行政", "地域包括支援センター／市区町村／民生委員。虐待・権利擁護・経済問題。")]
cw, gap, x0, y0, ch = Inches(3.85), Inches(0.28), Inches(0.7), Inches(2.2), Inches(2.1)
for i, (t, b) in enumerate(items):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    hd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.22), y0 + Inches(0.2), cw - Inches(0.44), Inches(0.55))
    _set_fill(hd, BLUE); hd.shadow.inherit = False
    para(hd.text_frame, t, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, x + Inches(0.25), y0 + Inches(0.95), cw - Inches(0.5), ch - Inches(1.1))
    para(tf, b, size=15, first=True)
note(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(1.2),
     [("サ責の動き：", GOLD, True), ("変化を察知したら、まずケアマネへ報告・相談。一事業所で抱えず、サービス担当者会議で体制を組み直す。", INK, False)])

# ======================================================================
# Slide 20 セルフチェック
# ======================================================================
s = newslide("CHECK", "自事業所の体制セルフチェック")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, "「はい」と言える項目はいくつ？　△の項目を今日の宿題に。", size=17, color=MUTED, first=True)
left = ["困難ケースの受入を会議で判断する仕組みがある", "受付時の情報収集項目が標準化されている",
        "個別の手順書・指示書を作成している", "サービス提供後の報告が回収・共有される"]
right = ["緊急時連絡フローが1枚で全員に共有されている", "ハラスメント対応手順が明文化されている",
         "個別研修計画に困難ケース対応が含まれる", "ケアマネ・地域包括への連携ルートが明確"]
cw, gx, x0, y0, ch = Inches(5.85), Inches(0.3), Inches(0.7), Inches(2.15), Inches(2.6)
for col, items in enumerate((left, right)):
    x = x0 + col * (cw + gx)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    tb, tf = textbox(s, x + Inches(0.3), y0 + Inches(0.2), cw - Inches(0.55), ch - Inches(0.3))
    first = True
    for it in items:
        p = para(tf, [("✔ ", ACCENT, True), (it, INK, False)], size=15, first=first, space_after=10)
        first = False
note(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(1.15),
     [("使い方：", GOLD, True), ("各自で採点 → グループで「△を1つ、来月までにどう改善するか」を1つ決める。", INK, False)])

# ======================================================================
# Slide 21 まとめ
# ======================================================================
s = newslide("SUMMARY", "まとめ：加算IIの責任を現場の力に")
cw, gx, x0, y0, ch = Inches(5.85), Inches(0.3), Inches(0.7), Inches(1.7), Inches(3.4)
card(s, x0, y0, cw, ch, title="今日のキーメッセージ",
     bullets=["困難さは「属性」でなく「体制不足」から生まれる", "受け入れは5ステップで組織判断",
              "「一人で判断しない・条件で受ける・記録で残す」", "加算の体制要件＝困難ケースを受ける装置"],
     title_color=BLUE, body_size=16)
x1 = x0 + cw + gx
rect(s, x1, y0, cw, ch, fill=CARDBG, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), y0 + Inches(0.18), cw - Inches(0.5), ch - Inches(0.3))
para(tf, "明日からのアクション宣言", size=18, color=BLUE, bold=True, first=True, space_after=4)
para(tf, "グループで1つずつ発表してください。", size=13, color=MUTED, space_after=6)
para(tf, "・私の事業所が次に整えるのは「　　　　　」", size=16, space_after=4)
para(tf, "・私が個人として変えるのは「　　　　　」", size=16, space_after=8)
rect(s, x1 + Inches(0.25), y0 + Inches(2.5), cw - Inches(0.5), Inches(0.75), fill=GOODBG, line=None)
tb, tf = textbox(s, x1 + Inches(0.4), y0 + Inches(2.55), cw - Inches(0.8), Inches(0.65), anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("小さな一歩を、必ず", INK, False), ("仕組み（手順・記録）", ACCENT, True), ("に落とす。", INK, False)], size=15, first=True)

# ======================================================================
# Slide 22 クロージング
# ======================================================================
s = newslide(grad=(NAVY, ACCENT))
tb, tf = textbox(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.5))
para(tf, "ご清聴ありがとうございました", size=40, color=WHITE, bold=True, first=True, space_after=10)
para(tf, "「断らない」のではなく、「組織で安全に受けられる」事業所へ。", size=20, color=GOLD, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.6))
para(tf2, "質疑・ディスカッションタイム", size=18, color=RGBColor(0xCF, 0xE0, 0xEE), first=True, space_after=8)
para(tf2, "※本資料の要件・割合等は、運用前に最新の介護報酬告示・解釈通知でご確認ください。",
     size=12, color=RGBColor(0xCF, 0xE0, 0xEE))

# ----------------------------------------------------------------------
out = "特定事業所加算II_困難ケース受け入れ体制_勉強会.pptx"
prs.save(out)
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst))
