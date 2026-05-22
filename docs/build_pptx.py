# -*- coding: utf-8 -*-
"""中央区主任ケアマネの会 研修スライドの PowerPoint 生成スクリプト。
docs/chuo-shunin-caremgr-training.html と同じ内容・配色で .pptx を出力する。
実行: python3 docs/build_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- 配色 ----
MAIN      = RGBColor(0x1F, 0x7A, 0x6B)
MAIN_DARK = RGBColor(0x15, 0x5B, 0x50)
ACCENT    = RGBColor(0xE8, 0xA3, 0x3D)
INK       = RGBColor(0x23, 0x30, 0x2E)
SOFT      = RGBColor(0xF4, 0xF7, 0xF6)
LINE      = RGBColor(0xD6, 0xE3, 0xE0)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY      = RGBColor(0x8A, 0xA1, 0x9C)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

TOTAL = 15
_page = [0]


def _set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn("a:" + tag))
        if el is None:
            el = rPr.makeelement(qn("a:" + tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold) または list of runs。"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = ln if isinstance(ln, list) else [ln]
        for (text, size, color, bold) in runs:
            r = p.add_run()
            r.text = text
            _set_font(r, size, color, bold)
        p.space_after = Pt(4)
    return tb


def rect(slide, x, y, w, h, fill, line=None, rounded=False, line_w=1):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def base(cover=False):
    s = prs.slides.add_slide(BLANK)
    _page[0] += 1
    if cover:
        bg = rect(s, 0, 0, SW, SH, MAIN_DARK)
        rect(s, 0, 0, SW, Inches(0.18), ACCENT)
    else:
        rect(s, 0, 0, SW, SH, WHITE)
    return s


def footer(slide, cover=False):
    col = RGBColor(0xC9, 0xD9, 0xD5) if cover else GRAY
    y = SH - Inches(0.55)
    line_col = RGBColor(0x3A, 0x52, 0x4D) if cover else LINE
    rect(slide, Inches(0.9), y - Inches(0.08), SW - Inches(1.8),
         Emu(9525), line_col)
    add_text(slide, Inches(0.9), y, Inches(9), Inches(0.4),
             [[("中央区 主任ケアマネの会に参加する意義", 10, col, False)]])
    add_text(slide, SW - Inches(1.8), y, Inches(0.9), Inches(0.4),
             [[("%d / %d" % (_page[0], TOTAL), 10, col, False)]],
             align=PP_ALIGN.RIGHT)


def header(slide, kicker, title):
    add_text(slide, Inches(0.95), Inches(0.55), Inches(11), Inches(0.5),
             [[(kicker, 13, MAIN, True)]])
    rect(slide, Inches(0.95), Inches(1.05), Inches(0.14), Inches(0.7), ACCENT)
    add_text(slide, Inches(1.25), Inches(1.0), Inches(11.2), Inches(0.95),
             [[(title, 30, MAIN_DARK, True)]], anchor=MSO_ANCHOR.MIDDLE)


def lead(slide, y, runs):
    return add_text(slide, Inches(0.95), y, Inches(11.4), Inches(0.9),
                    [runs])


def bullets(slide, x, y, w, items, size=18, gap=0.62):
    """items: list of list-of-runs。"""
    for i, runs in enumerate(items):
        cy = y + Inches(gap) * i
        rect(slide, x, cy + Inches(0.10), Inches(0.16), Inches(0.16), MAIN)
        add_text(slide, x + Inches(0.34), cy, w - Inches(0.34),
                 Inches(gap), [runs])


def num_bullets(slide, x, y, w, items, size=18, gap=0.78):
    for i, runs in enumerate(items):
        cy = y + Inches(gap) * i
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, cy,
                                      Inches(0.42), Inches(0.42))
        circ.fill.solid()
        circ.fill.fore_color.rgb = ACCENT
        circ.line.fill.background()
        circ.shadow.inherit = False
        tf = circ.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        _set_font(r, 16, WHITE, True)
        add_text(slide, x + Inches(0.62), cy, w - Inches(0.62),
                 Inches(gap), [runs], anchor=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, pill, title, body):
    box = rect(slide, x, y, w, h, SOFT, line=LINE, rounded=True)
    pad = Inches(0.22)
    # pill
    pw = Inches(1.35)
    pill_box = rect(slide, x + pad, y + pad, pw, Inches(0.36), MAIN,
                    rounded=True)
    tf = pill_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = pill
    _set_font(r, 11, WHITE, True)
    add_text(slide, x + pad, y + pad + Inches(0.5), w - pad * 2,
             Inches(0.9), [[(title, 16, MAIN_DARK, True)]])
    add_text(slide, x + pad, y + h - Inches(1.45), w - pad * 2,
             Inches(1.3), [[(body, 12.5, INK, False)]])


def panel(slide, x, y, w, h, title, items, accent=False):
    box = rect(slide, x, y, w, h, SOFT, rounded=True)
    rect(slide, x, y, w, Inches(0.12), ACCENT if accent else MAIN)
    pad = Inches(0.3)
    add_text(slide, x + pad, y + pad, w - pad * 2, Inches(0.6),
             [[(title, 19, MAIN_DARK, True)]])
    bullets(slide, x + pad, y + Inches(1.0), w - pad * 2, items,
            size=15, gap=0.62)


# ============================================================
# 1. 表紙
# ============================================================
s = base(cover=True)
add_text(s, Inches(1.1), Inches(1.2), Inches(11), Inches(0.5),
         [[("居宅介護支援事業所 ｜ 事業所内研修", 15, ACCENT, True)]])
add_text(s, Inches(1.1), Inches(1.85), Inches(11.2), Inches(2.8),
         [[("東京都中央区", 44, WHITE, True)],
          [("「主任ケアマネの会」に参加する意義", 44, WHITE, True)]])
add_text(s, Inches(1.1), Inches(4.35), Inches(11), Inches(0.6),
         [[("― 地域連携と人材育成を、事業所の力に変える ―", 20, WHITE, False)]])
rect(s, Inches(1.1), Inches(5.2), Inches(6.4), Emu(9525),
     RGBColor(0x3A, 0x52, 0x4D))
add_text(s, Inches(1.1), Inches(5.35), Inches(8), Inches(1.3),
         [[("対象：管理者・主任介護支援専門員", 14,
            RGBColor(0xC9, 0xD9, 0xD5), False)],
          [("ねらい：会への参加を「個人の活動」から「事業所の戦略」へ",
            14, RGBColor(0xC9, 0xD9, 0xD5), False)]])
footer(s, cover=True)

# ============================================================
# 2. アジェンダ
# ============================================================
s = base()
header(s, "AGENDA", "本日の研修のねらいと流れ")
lead(s, Inches(2.05),
     [("主任ケアマネの会への参加を、単なる「外部の集まり」ではなく、", 17, INK, False),
      ("事業所の運営基盤を強くする取り組み", 17, MAIN_DARK, True),
      ("として捉え直します。", 17, INK, False)])
num_bullets(s, Inches(1.05), Inches(2.95), Inches(11.2), [
    [("主任ケアマネに求められる役割の再確認", 17, INK, False)],
    [("「中央区主任ケアマネの会」とは何か", 17, INK, False)],
    [("参加する6つの意義（連携・情報・困難事例・研修要件・育成・地域貢献）",
      17, INK, False)],
    [("事業所経営から見たメリットと留意点", 17, INK, False)],
    [("当事業所としての参加方針・今後の取り組み", 17, INK, False)],
])
footer(s)

# ============================================================
# 3. 主任ケアマネの役割
# ============================================================
s = base()
header(s, "前提の確認", "主任ケアマネに求められる役割")
lead(s, Inches(2.0),
     [("主任介護支援専門員は、自らの担当業務にとどまらず", 16, INK, False),
      ("「地域」と「事業所内」をつなぐ中核", 16, MAIN_DARK, True),
      ("としての機能が期待されています。", 16, INK, False)])
cw, ch, gap = Inches(3.75), Inches(2.5), Inches(0.3)
x0 = Inches(0.95)
cy = Inches(2.75)
cards3 = [
    ("役割 1", "ケアマネジメントの質の管理",
     "事業所内の支援困難事例への助言、適切なケアプラン作成の指導。"),
    ("役割 2", "人材育成・スーパーバイズ",
     "後進のケアマネへのOJT、相談対応、実習指導者としての関わり。"),
    ("役割 3", "地域づくり・多職種連携",
     "地域包括支援センターや他職種と連携し、地域課題に取り組む。"),
]
for i, (p, t, b) in enumerate(cards3):
    card(s, x0 + (cw + gap) * i, cy, cw, ch, p, t, b)
add_text(s, Inches(0.95), Inches(5.5), Inches(11.4), Inches(0.6),
         [[("→ これらの役割を一人で抱えず磨くための「場」が、", 16, INK, False),
           ("主任ケアマネの会", 16, MAIN_DARK, True),
           ("です。", 16, INK, False)]])
footer(s)

# ============================================================
# 4. 会とは
# ============================================================
s = base()
header(s, "WHAT", "「中央区 主任ケアマネの会」とは")
pw, ph = Inches(5.65), Inches(3.0)
panel(s, Inches(0.95), Inches(2.05), pw, ph, "概要", [
    [("中央区内で活動する", 14, INK, False),
     ("主任ケアマネが集う自主的な研鑽の場", 14, MAIN_DARK, True)],
    [("地域包括支援センターや行政と連携した運営", 14, INK, False)],
    [("定例会・事例検討会・研修会などを実施", 14, INK, False)],
])
panel(s, Inches(6.85), Inches(2.05), pw, ph, "主な活動内容", [
    [("支援困難事例の検討・情報交換", 14, INK, False)],
    [("制度改正・行政情報の共有", 14, INK, False)],
    [("多職種・地域資源とのネットワークづくり", 14, INK, False)],
    [("主任ケアマネ更新研修につながる学びの機会", 14, INK, False)],
], accent=True)
note = rect(s, Inches(0.95), Inches(5.35), Inches(11.4), Inches(0.85),
            RGBColor(0xFD, 0xF3, 0xDF), rounded=True)
add_text(s, Inches(1.25), Inches(5.5), Inches(10.8), Inches(0.6),
         [[("※ 開催頻度・申込方法など最新の運営状況は、担当の地域包括支援センターにご確認ください。",
            14, RGBColor(0x8A, 0x66, 0x1F), True)]])
footer(s)

# ============================================================
# 5. 6つの意義 全体像
# ============================================================
s = base()
header(s, "OVERVIEW", "参加する「6つの意義」")
six = [
    ("意義 1", "地域ネットワークの構築", "顔の見える関係で、連携先・相談先を確保。"),
    ("意義 2", "最新情報のキャッチアップ", "制度改正・報酬改定・行政情報をいち早く入手。"),
    ("意義 3", "困難事例への対応力", "他事業所の知見を借り、抱え込みを防ぐ。"),
    ("意義 4", "研修要件への対応", "主任更新研修・法定研修につながる学び。"),
    ("意義 5", "後進育成への還元", "学びを事業所内OJTに展開し、組織力UP。"),
    ("意義 6", "地域包括ケアへの貢献", "地域課題の解決に関わり、事業所の信頼向上。"),
]
cw, ch = Inches(3.75), Inches(2.05)
gx, gy = Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.95), Inches(2.1)
for i, (p, t, b) in enumerate(six):
    col, row = i % 3, i // 3
    card(s, x0 + (cw + gx) * col, y0 + (ch + gy) * row, cw, ch, p, t, b)
footer(s)

# ============================================================
# 意義スライド共通ヘルパ
# ============================================================
def gi_slide(kicker, title, lead_runs, items, tail_runs):
    s = base()
    header(s, kicker, title)
    lead(s, Inches(2.0), lead_runs)
    bullets(s, Inches(1.05), Inches(2.85), Inches(11.2), items,
            size=17, gap=0.74)
    if tail_runs:
        add_text(s, Inches(0.95), Inches(5.85), Inches(11.4), Inches(0.6),
                 [tail_runs])
    footer(s)
    return s

# 6. 意義1
gi_slide("意義 1", "地域ネットワークの構築・連携先の確保",
    [("中央区という", 16, INK, False),
     ("同じ地域", 16, MAIN_DARK, True),
     ("で働く主任ケアマネと「顔の見える関係」を築けます。", 16, INK, False)],
    [
     [("連携先の確保", 16, MAIN_DARK, True),
      ("：医療機関・他事業所・インフォーマル資源への“つなぎ”がスムーズに", 16, INK, False)],
     [("緊急時の相談先", 16, MAIN_DARK, True),
      ("：困ったとき、すぐに相談できる同職種の仲間ができる", 16, INK, False)],
     [("地域資源の把握", 16, MAIN_DARK, True),
      ("：区内のサービス事情・空き状況などの生きた情報が得られる", 16, INK, False)],
     [("担当者交代時のリスク低減", 16, MAIN_DARK, True),
      ("：個人ではなく事業所として地域につながる", 16, INK, False)],
    ],
    [("→ ネットワークは", 16, INK, False),
     ("事業所の“見えない資産”", 16, MAIN_DARK, True),
     ("。利用者支援の選択肢が広がります。", 16, INK, False)])

# 7. 意義2
gi_slide("意義 2", "制度改正・最新情報のキャッチアップ",
    [("介護保険制度は改正が頻繁です。", 16, INK, False),
     ("情報の遅れは、運営リスクに直結", 16, MAIN_DARK, True),
     ("します。", 16, INK, False)],
    [
     [("制度改正・報酬改定", 16, MAIN_DARK, True),
      ("の解釈や実務対応をいち早く共有できる", 16, INK, False)],
     [("区独自の事業・行政情報", 16, MAIN_DARK, True),
      ("（中央区の施策や様式変更など）が得られる", 16, INK, False)],
     [("運営指導（実地指導）対策", 16, MAIN_DARK, True),
      ("：他事業所の指摘事例から学べる", 16, INK, False)],
     [("会で得た情報を持ち帰り、", 16, INK, False),
      ("事業所内で正確に共有", 16, MAIN_DARK, True),
      ("できる", 16, INK, False)],
    ],
    [("→ 「知らなかった」による減算・返戻を防ぎ、", 16, INK, False),
     ("適正な運営", 16, MAIN_DARK, True),
     ("を支えます。", 16, INK, False)])

# 8. 意義3（2パネル構成）
s = base()
header(s, "意義 3", "困難事例への対応力強化")
pw, ph = Inches(5.65), Inches(2.95)
panel(s, Inches(0.95), Inches(2.05), pw, ph, "事業所内だけでは限界がある", [
    [("虐待・セルフネグレクト・8050問題などの複合課題", 14, INK, False)],
    [("クレーム対応・支援拒否ケース", 14, INK, False)],
    [("一人で抱え込むと", 14, INK, False),
     ("燃え尽き・離職", 14, MAIN_DARK, True),
     ("のリスク", 14, INK, False)],
])
panel(s, Inches(6.85), Inches(2.05), pw, ph, "会の事例検討で得られるもの", [
    [("多様な視点・経験からの", 14, INK, False),
     ("具体的な助言", 14, MAIN_DARK, True)],
    [("「他もやっている」という", 14, INK, False),
     ("心理的な支え", 14, MAIN_DARK, True)],
    [("地域包括・行政との", 14, INK, False),
     ("連携の糸口", 14, MAIN_DARK, True)],
], accent=True)
add_text(s, Inches(0.95), Inches(5.35), Inches(11.4), Inches(0.6),
         [[("→ 対応の質が上がり、担当ケアマネの", 16, INK, False),
           ("精神的負担の軽減", 16, MAIN_DARK, True),
           ("にもつながります。", 16, INK, False)]])
footer(s)

# 9. 意義4
gi_slide("意義 4", "研修要件への対応・専門性の維持",
    [("主任介護支援専門員は", 16, INK, False),
     ("更新（5年ごと）", 16, MAIN_DARK, True),
     ("が必要であり、継続的な研鑽が前提です。", 16, INK, False)],
    [
     [("会での学びは", 16, INK, False),
      ("主任更新研修・法定研修", 16, MAIN_DARK, True),
      ("の内容理解を深める土台になる", 16, INK, False)],
     [("最新の実務知識に触れ続けることで", 16, INK, False),
      ("専門性が陳腐化しない", 16, MAIN_DARK, True)],
     [("事例提供・発表の経験が、", 16, INK, False),
      ("指導者・実習指導者としての力", 16, MAIN_DARK, True),
      ("を養う", 16, INK, False)],
     [("「資格を持っているだけ」から「機能している主任ケアマネ」へ",
       16, INK, False)],
    ],
    [("→ 特定事業所加算など、", 16, INK, False),
     ("主任ケアマネの実働が事業所評価に直結", 16, MAIN_DARK, True),
     ("します。", 16, INK, False)])

# 10. 意義5（3カード）
s = base()
header(s, "意義 5", "後進育成への還元 ― 事業所内OJTの強化")
lead(s, Inches(2.0),
     [("会への参加は", 16, INK, False),
      ("個人で完結させず、事業所に“持ち帰る”", 16, MAIN_DARK, True),
      ("ことで価値が何倍にもなります。", 16, INK, False)])
cw, ch, gap = Inches(3.75), Inches(2.4), Inches(0.3)
cy = Inches(2.75)
cards5 = [
    ("持ち帰り 1", "朝礼・内部研修で共有",
     "会で得た制度情報・事例を、定例の場で全員に展開する。"),
    ("持ち帰り 2", "事例検討の手法を導入",
     "会で学んだ進め方を、事業所内の事例検討会に応用する。"),
    ("持ち帰り 3", "若手のロールモデルに",
     "主任ケアマネの学ぶ姿勢が、後進の目標になる。"),
]
for i, (p, t, b) in enumerate(cards5):
    card(s, Inches(0.95) + (cw + gap) * i, cy, cw, ch, p, t, b)
add_text(s, Inches(0.95), Inches(5.4), Inches(11.4), Inches(0.6),
         [[("→ 一人の学びが", 16, INK, False),
           ("チーム全体の底上げ", 16, MAIN_DARK, True),
           ("に。育成は採用難への有効な対策です。", 16, INK, False)]])
footer(s)

# 11. 意義6
gi_slide("意義 6", "地域包括ケアシステムへの貢献と事業所の信頼",
    [("主任ケアマネの会は、", 16, INK, False),
     ("中央区の地域包括ケアを支える担い手のひとつ", 16, MAIN_DARK, True),
     ("です。", 16, INK, False)],
    [
     [("地域ケア会議や地域課題の検討に関わり、", 16, INK, False),
      ("地域づくりに参画", 16, MAIN_DARK, True),
      ("できる", 16, INK, False)],
     [("会での積極的な活動が、地域包括・行政・他事業所からの", 16, INK, False),
      ("信頼につながる", 16, MAIN_DARK, True)],
     [("「地域に開かれ、地域に貢献している事業所」という", 16, INK, False),
      ("評判・ブランド", 16, MAIN_DARK, True),
      ("に", 16, INK, False)],
     [("その評判は、利用者紹介・採用・連携のしやすさに還元される",
       16, INK, False)],
    ],
    [("→ 地域への貢献は、めぐりめぐって", 16, INK, False),
     ("事業所自身の安定経営", 16, MAIN_DARK, True),
     ("を支えます。", 16, INK, False)])

# ============================================================
# 12. 経営メリット表
# ============================================================
s = base()
header(s, "FOR MANAGEMENT", "事業所経営から見たメリット整理")
rows = [
    ("観点", "会への参加がもたらすもの"),
    ("運営の適正化", "制度改正・指導事例の把握による減算・返戻リスクの低減"),
    ("支援の質", "困難事例対応力の向上、ケアプラン点検の視点獲得"),
    ("人材育成", "主任ケアマネを軸としたOJT機能・指導体制の強化"),
    ("定着・離職防止", "抱え込みの解消、相談先の確保による心理的負担の軽減"),
    ("営業・信頼", "地域ネットワーク拡大、紹介・連携のしやすさ、事業所の評判向上"),
]
tx, ty = Inches(0.95), Inches(2.0)
tw = Inches(11.45)
tbl = s.shapes.add_table(len(rows), 2, tx, ty, tw, Inches(3.0)).table
tbl.columns[0].width = Inches(3.2)
tbl.columns[1].width = tw - Inches(3.2)
for ri, (c0, c1) in enumerate(rows):
    for ci, txt in enumerate((c0, c1)):
        cell = tbl.cell(ri, ci)
        cell.margin_left = Inches(0.18)
        cell.margin_right = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = txt
        if ri == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = MAIN
            _set_font(r, 15, WHITE, True)
        elif ci == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = SOFT
            _set_font(r, 14, MAIN_DARK, True)
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            _set_font(r, 14, INK, False)
add_text(s, Inches(0.95), Inches(5.55), Inches(11.4), Inches(0.6),
         [[("→ 参加時間は「コスト」ではなく", 16, INK, False),
           ("事業所の基盤への投資", 16, MAIN_DARK, True),
           ("と捉えられます。", 16, INK, False)]])
footer(s)

# ============================================================
# 13. 留意点
# ============================================================
s = base()
header(s, "CAUTION", "参加にあたっての留意点")
lead(s, Inches(2.0),
     [("効果を最大化し、形骸化させないために、次の点を意識します。",
       17, INK, False)])
bullets(s, Inches(1.05), Inches(2.8), Inches(11.2), [
    [("業務調整", 16, MAIN_DARK, True),
     ("：参加日は訪問・面談スケジュールを事前に調整し、無理なく出席する",
      16, INK, False)],
    [("守秘義務", 16, MAIN_DARK, True),
     ("：事例検討では個人が特定されないよう配慮し、情報管理を徹底する",
      16, INK, False)],
    [("“持ち帰り”を仕組み化", 16, MAIN_DARK, True),
     ("：参加報告を定例化し、学びを必ず事業所内で共有する", 16, INK, False)],
    [("属人化の回避", 16, MAIN_DARK, True),
     ("：可能な範囲で複数名・交代制での参加も検討する", 16, INK, False)],
    [("受け身にならない", 16, MAIN_DARK, True),
     ("：情報を得るだけでなく、事例提供など発信する側にも回る",
      16, INK, False)],
], gap=0.74)
footer(s)

# ============================================================
# 14. 当事業所の方針
# ============================================================
s = base()
header(s, "ACTION PLAN", "当事業所としての方針・今後の取り組み")
lead(s, Inches(2.0),
     [("本研修を踏まえ、会への参加を", 17, INK, False),
      ("事業所の正式な取り組み", 17, MAIN_DARK, True),
      ("として位置づけます。", 17, INK, False)])
num_bullets(s, Inches(1.05), Inches(2.85), Inches(11.2), [
    [("主任ケアマネの会への", 16, INK, False),
     ("継続的な参加", 16, MAIN_DARK, True),
     ("を業務として明確化する", 16, INK, False)],
    [("参加後は", 16, INK, False),
     ("定例ミーティングで報告", 16, MAIN_DARK, True),
     ("し、学びを全員で共有する", 16, INK, False)],
    [("会で得た事例検討の手法を、", 16, INK, False),
     ("事業所内の事例検討会", 16, MAIN_DARK, True),
     ("に取り入れる", 16, INK, False)],
    [("制度改正情報は", 16, INK, False),
     ("速やかに業務手順へ反映", 16, MAIN_DARK, True),
     ("する", 16, INK, False)],
    [("次回更新時期・研修計画と", 16, INK, False),
     ("連動させて参加計画", 16, MAIN_DARK, True),
     ("を立てる", 16, INK, False)],
], gap=0.72)
footer(s)

# ============================================================
# 15. まとめ
# ============================================================
s = base(cover=True)
add_text(s, Inches(1.1), Inches(1.0), Inches(11), Inches(0.5),
         [[("CONCLUSION", 14, ACCENT, True)]])
add_text(s, Inches(1.1), Inches(1.9), Inches(11.2), Inches(3.0),
         [[("主任ケアマネの会への参加は、", 30, WHITE, True)],
          [("個人の研鑽", 30, ACCENT, True),
           ("であると同時に、", 30, WHITE, True)],
          [("事業所の連携力・育成力・信頼", 30, ACCENT, True),
           ("を", 30, WHITE, True)],
          [("高める「投資」である。", 30, WHITE, True)]],
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.1), Inches(5.0), Inches(11.2), Inches(1.2),
         [[("地域とつながり、学びを持ち帰り、チームで活かす。",
            16, RGBColor(0xC9, 0xD9, 0xD5), False)],
          [("その積み重ねが、利用者支援の質と事業所の未来を支えます。",
            16, RGBColor(0xC9, 0xD9, 0xD5), False)]],
         align=PP_ALIGN.CENTER)
footer(s, cover=True)

OUT = "docs/中央区主任ケアマネの会_研修スライド.pptx"
prs.save(OUT)
print("saved:", OUT, "/ slides:", len(prs.slides._sldIdLst))
