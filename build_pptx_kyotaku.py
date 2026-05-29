# -*- coding: utf-8 -*-
"""
特定事業所加算II 内部勉強会スライド（居宅介護支援事業所版／困難事例の受け入れ体制）
PowerPoint(.pptx) 生成スクリプト
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- カラーパレット（居宅版は青緑×藍系で差別化）----
NAVY   = RGBColor(0x10, 0x33, 0x5C)
BLUE   = RGBColor(0x21, 0x6A, 0x8C)
ACCENT = RGBColor(0x1B, 0x9A, 0x8B)
WARM   = RGBColor(0xD9, 0x6A, 0x4A)
GOLD   = RGBColor(0xE3, 0xB5, 0x4A)
BG     = RGBColor(0xF3, 0xF7, 0xF9)
INK    = RGBColor(0x21, 0x2E, 0x38)
MUTED  = RGBColor(0x57, 0x68, 0x73)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xD4, 0xE0, 0xE6)
CARDBG = RGBColor(0xFF, 0xFF, 0xFF)
GOODBG = RGBColor(0xE6, 0xF5, 0xF2)
BADBG  = RGBColor(0xFB, 0xEC, 0xE6)
NOTEBG = RGBColor(0xFD, 0xF4, 0xE2)

FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    _set_fill(bg, BG)
    bg.shadow.inherit = False
    return s


def _grad(sp, c1, c2, ang):
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    grad = sp.makeelement(qn('a:gradFill'), {})
    gsLst = grad.makeelement(qn('a:gsLst'), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = grad.makeelement(qn('a:gs'), {'pos': str(pos)})
        srgb = grad.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (col[0], col[1], col[2])})
        gs.append(srgb); gsLst.append(gs)
    grad.append(gsLst)
    lin = grad.makeelement(qn('a:lin'), {'ang': str(ang), 'scaled': '1'})
    grad.append(lin)
    ln = sp.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp.append(grad)


def grad_bg(s, c1, c2):
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background(); bg.shadow.inherit = False
    _grad(bg.fill._xPr, c1, c2, 2700000)
    return bg


def grad_fill_shape(shape, c1, c2, ang=0):
    _grad(shape.fill._xPr, c1, c2, ang)


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    return tb, tf


def para(tf, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=4, bullet=False, level=0, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.level = level
    runs = text if isinstance(text, list) else [(text, color, bold)]
    for seg in runs:
        t, c, b = seg
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.color.rgb = c
        r.font.name = font
    if bullet:
        _bullet(p)
    else:
        _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def _bullet(p, char="・"):
    pPr = p._p.get_or_add_pPr()
    pPr.set('indent', '-228600')
    pPr.set('marL', '228600')
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': FONT}))
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': char}))


def rect(s, x, y, w, h, fill=CARDBG, line=None, rounded=True, line_w=1.0):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def leftbar(s, x, y, w, h, color):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(8), h)
    _set_fill(bar, color); bar.shadow.inherit = False


def kicker(s, text, color=ACCENT):
    tb, tf = textbox(s, Inches(0.7), Inches(0.4), Inches(8), Inches(0.4))
    para(tf, text, size=13, color=color, bold=True, first=True)


def heading(s, text):
    leftbar(s, Inches(0.7), Inches(0.78), Inches(0.12), Inches(0.62), ACCENT)
    tb, tf = textbox(s, Inches(0.95), Inches(0.72), Inches(11.6), Inches(0.8))
    para(tf, text, size=29, color=NAVY, bold=True, first=True)


def pill(s, x, y, text, color=ACCENT, tcolor=WHITE, w=Inches(2.0)):
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
    _set_fill(p, color); p.shadow.inherit = False
    try:
        p.adjustments[0] = 0.5
    except Exception:
        pass
    tf = p.text_frame; tf.word_wrap = True
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    para(tf, text, size=12, color=tcolor, bold=True, align=PP_ALIGN.CENTER, first=True)
    return p


def card(s, x, y, w, h, title=None, bullets=None, fill=CARDBG, line=LINE,
         title_color=BLUE, body_size=15, title_size=18):
    rect(s, x, y, w, h, fill=fill, line=line)
    pad = Inches(0.22)
    tb, tf = textbox(s, x + pad, y + Inches(0.12), w - pad * 2, h - Inches(0.2))
    first = True
    if title:
        para(tf, title, size=title_size, color=title_color, bold=True, first=True, space_after=5)
        first = False
    if bullets:
        for b in bullets:
            para(tf, b, size=body_size, color=INK, first=first, bullet=True, space_after=3)
            first = False
    return tf


def note(s, x, y, w, h, runs, bar_color=GOLD, fill=NOTEBG):
    rect(s, x, y, w, h, fill=fill, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(6), h)
    _set_fill(bar, bar_color); bar.shadow.inherit = False
    tb, tf = textbox(s, x + Inches(0.2), y, w - Inches(0.35), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, runs, size=14, first=True)


TOTAL = 22
_n = [0]


def footer(s, n):
    tb, tf = textbox(s, Inches(11.6), Inches(7.05), Inches(1.5), Inches(0.35))
    para(tf, "%d / %d" % (n, TOTAL), size=11, color=MUTED, align=PP_ALIGN.RIGHT, first=True)
    tb2, tf2 = textbox(s, Inches(0.7), Inches(7.05), Inches(9), Inches(0.35))
    para(tf2, "特定事業所加算II 内部勉強会｜居宅介護支援・困難事例受け入れ体制", size=10, color=MUTED, first=True)


def newslide(kick=None, title=None, grad=None):
    _n[0] += 1
    if grad:
        s = prs.slides.add_slide(BLANK)
        grad_bg(s, grad[0], grad[1])
    else:
        s = slide()
        if kick:
            kicker(s, kick)
        if title:
            heading(s, title)
        footer(s, _n[0])
    return s


def system_slide(kick, title, left_t, left_items, right_t, right_items, note_runs, note_kind="note"):
    s = newslide(kick, title)
    cw, ch, gx = Inches(5.85), Inches(3.05), Inches(0.3)
    x0, y0 = Inches(0.7), Inches(1.7)
    card(s, x0, y0, cw, ch, title=left_t, bullets=left_items, title_color=BLUE, body_size=15)
    card(s, x0 + cw + gx, y0, cw, ch, title=right_t, bullets=right_items, title_color=BLUE, body_size=15)
    if note_kind == "good":
        note(s, Inches(0.7), Inches(5.05), Inches(11.9), Inches(1.15), note_runs, bar_color=ACCENT, fill=GOODBG)
    else:
        note(s, Inches(0.7), Inches(5.05), Inches(11.9), Inches(1.15), note_runs)
    return s


def case_slide(title_text, situation, ng_items, good_items, point_runs):
    s = newslide()
    hd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.75))
    grad_fill_shape(hd, WARM, GOLD, 0)
    hd.line.fill.background(); hd.shadow.inherit = False
    tf = hd.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, title_text, size=21, color=WHITE, bold=True, first=True)
    rect(s, Inches(0.7), Inches(1.4), Inches(11.9), Inches(1.15), fill=WHITE, line=LINE)
    tb, tf = textbox(s, Inches(0.95), Inches(1.5), Inches(11.4), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [("状況：", NAVY, True)] + situation, size=15, first=True)
    cw, gx, y0, ch = Inches(5.85), Inches(0.3), Inches(2.75), Inches(2.35)
    x0 = Inches(0.7)
    rect(s, x0, y0, cw, ch, fill=BADBG, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, y0, Pt(6), ch); _set_fill(bar, WARM); bar.shadow.inherit = False
    tb, tf = textbox(s, x0 + Inches(0.25), y0 + Inches(0.12), cw - Inches(0.4), ch - Inches(0.2))
    para(tf, "NG対応", size=17, color=WARM, bold=True, first=True, space_after=4)
    for it in ng_items:
        para(tf, it, size=14, bullet=True, space_after=3)
    x1 = x0 + cw + gx
    rect(s, x1, y0, cw, ch, fill=GOODBG, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x1, y0, Pt(6), ch); _set_fill(bar, ACCENT); bar.shadow.inherit = False
    tb, tf = textbox(s, x1 + Inches(0.25), y0 + Inches(0.12), cw - Inches(0.4), ch - Inches(0.2))
    para(tf, "GOOD対応", size=17, color=ACCENT, bold=True, first=True, space_after=4)
    for it in good_items:
        para(tf, it, size=14, bullet=True, space_after=3)
    note(s, Inches(0.7), Inches(5.35), Inches(11.9), Inches(1.1), [("体制のポイント：", GOLD, True)] + point_runs)
    return s


# ======================================================================
# Slide 1 タイトル
# ======================================================================
s = newslide(grad=(NAVY, ACCENT))
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.4), Inches(6.6), Inches(0.5))
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x21, 0x55, 0x76)
b.line.color.rgb = RGBColor(0x8F, 0xB6, 0xD4); b.line.width = Pt(1); b.shadow.inherit = False
para(b.text_frame, "居宅介護支援事業所｜特定事業所加算II 算定　内部勉強会資料", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.4))
para(tf, "困難事例の受け入れ体制で実現する", size=42, color=WHITE, bold=True, first=True, space_after=2)
para(tf, "確実性のあるケアマネジメント", size=42, color=WHITE, bold=True, space_after=10)
para(tf, "～ 加算IIの責任を、事業所の支援力に変える ～", size=21, color=GOLD, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.5))
para(tf2, "対象：管理者・主任介護支援専門員・介護支援専門員　／　形式：事例検討型グループワーク",
     size=15, color=RGBColor(0xCF, 0xE0, 0xEE), first=True, space_after=3)
para(tf2, "開催日：2026年　　月　　日　　　会場：", size=15, color=RGBColor(0xCF, 0xE0, 0xEE))

# ======================================================================
# Slide 2 ねらい
# ======================================================================
s = newslide("PURPOSE", "本日のねらい")
tb, tf = textbox(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.6))
para(tf, [("「断れない事業所」ではなく、", MUTED, False),
          ("「組織として困難事例を受けきれる事業所」", WARM, True),
          ("を目指す。", MUTED, False)], size=20, first=True)
cw, gap, x0, y0, ch = Inches(3.85), Inches(0.28), Inches(0.7), Inches(2.35), Inches(2.0)
data = [("1. 共通理解", BLUE, "居宅介護支援における加算IIの趣旨と、求められる体制を再確認する。"),
        ("2. 体制の見える化", ACCENT, "困難事例を主任ケアマネ中心に「チームで受ける」流れを理解する。"),
        ("3. 事例で実践", WARM, "4つの事例で、受任判断・連携・記録の勘所を体得する。")]
for i, (t, c, body) in enumerate(data):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    pill(s, x + Inches(0.22), y0 + Inches(0.2), t, color=c, w=Inches(1.9))
    tb, tf = textbox(s, x + Inches(0.22), y0 + Inches(0.75), cw - Inches(0.44), ch - Inches(0.9))
    para(tf, body, size=15, first=True)
note(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(0.95),
     [("合言葉：", WARM, True), ("ケアマネ個人の力量に依存せず、誰が担当しても一定の質を保てる仕組みをつくる。", INK, False)])

# ======================================================================
# Slide 3 アジェンダ
# ======================================================================
s = newslide("AGENDA", "本日の流れ")
ag = [("第1部　理解編（約20分）", ["特定事業所加算IIの位置づけ（居宅介護支援）", "なぜ困難事例の受け入れ体制が要なのか", "「困難事例」の定義と類型"]),
      ("第2部　体制編（約20分）", ["受任〜ケアマネジメントの基本フロー", "5つの体制づくり（受任・SV・24時間・記録・研修）"]),
      ("第3部　事例編（約30分）", ["事例1〜4のグループ検討", "NG対応とGOOD対応の比較"]),
      ("第4部　まとめ（約10分）", ["自事業所の体制セルフチェック", "明日からのアクション宣言"])]
cw, ch, gx, gy = Inches(5.85), Inches(2.1), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.75)
for i, (t, items) in enumerate(ag):
    x = x0 + (i % 2) * (cw + gx)
    y = y0 + (i // 2) * (ch + gy)
    card(s, x, y, cw, ch, title=t, bullets=items, title_color=BLUE, body_size=15, title_size=18)

# ======================================================================
# Slide 4 加算IIの位置づけ（表・居宅版）
# ======================================================================
s = newslide("BASIC", "特定事業所加算IIの位置づけ（居宅介護支援）")
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5))
para(tf, [("居宅介護支援の特定事業所加算は、", INK, False), ("中重度・支援困難ケースにも対応できる質の高い事業所", WARM, True),
          ("を評価する加算です。", INK, False)], size=16, first=True)
rows, cols = 5, 3
tbl_shape = s.shapes.add_table(rows, cols, Inches(0.7), Inches(2.0), Inches(11.9), Inches(2.3))
table = tbl_shape.table
table.columns[0].width = Inches(2.6)
table.columns[1].width = Inches(5.0)
table.columns[2].width = Inches(4.3)
hdr = ["区分", "主な人員要件（主任ケアマネ・常勤ケアマネ）", "中重度要件（要介護3〜5）"]
body = [["加算I", "主任CM 2名以上＋常勤CM 3名以上 等", "40%以上 など"],
        ["加算II（当事業所）", "主任CM 1名以上＋常勤CM 3名以上", "要件なし"],
        ["加算III", "主任CM 1名以上＋常勤CM 2名以上", "要件なし"],
        ["加算A", "主任CM＋常勤CM＋非常勤（他事業所と連携可）", "要件なし"]]
for c, h in enumerate(hdr):
    cell = table.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = h; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
for ri, row in enumerate(body, start=1):
    highlight = (ri == 2)
    for ci, val in enumerate(row):
        cell = table.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xFF, 0xF3, 0xCD) if highlight else (RGBColor(0xEE, 0xF4, 0xF6) if ri % 2 == 0 else WHITE)
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = val; r.font.size = Pt(12.5)
        r.font.bold = (ci == 0) or highlight; r.font.color.rgb = INK; r.font.name = FONT
note(s, Inches(0.7), Inches(4.5), Inches(11.9), Inches(1.45),
     [("全区分共通の要件として、", INK, False),
      ("24時間連絡体制・定期的な会議・研修計画と実施・運営基準減算/集中減算なし・他法人との事例検討会、そして『支援困難事例の受け入れ』", WARM, True),
      ("が求められます（＝困難事例対応は加算の根幹）。", INK, False)])
tb, tf = textbox(s, Inches(0.9), Inches(6.05), Inches(11.6), Inches(0.5))
para(tf, "※CM＝介護支援専門員。人員数・割合・要件の詳細は最新の介護報酬告示・解釈通知（直近改定版）で必ずご確認ください。",
     size=11, color=MUTED, first=True)

# ======================================================================
# Slide 5 なぜ要なのか（居宅版・加算要件への明記を強調）
# ======================================================================
s = newslide("WHY", "なぜ「困難事例の受け入れ体制」が要なのか")
note(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(1.0),
     [("加算要件そのもの：", ACCENT, True),
      ("居宅介護支援の特定事業所加算では「地域包括支援センター等から支援困難な事例を紹介された場合でも、担当する介護支援専門員を配置していること」が要件。困難事例を受けられる体制＝加算の前提です。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)
items = [("① 加算の趣旨に直結する", "24時間連絡体制・会議・研修・SVといった加算要件は、そのまま困難事例を支えるための装置になる。"),
         ("② 地域からの信頼につながる", "地域包括や行政から「あの事業所なら任せられる」と頼られ、安定した紹介・受任につながる。"),
         ("③ ケアマネを守る", "体制があれば困難事例を一人で抱えさせない。バーンアウト・離職・抱え込みによる事故を防ぐ。"),
         ("④ 利用者の権利を守る", "支援拒否・虐待・権利擁護が絡む事例こそ、組織と多職種で受けることが利用者保護になる。")]
cw, ch, gx, gy = Inches(5.85), Inches(1.6), Inches(0.3), Inches(0.25)
x0, y0 = Inches(0.7), Inches(2.75)
for i, (t, b) in enumerate(items):
    x = x0 + (i % 2) * (cw + gx)
    y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    tb, tf = textbox(s, x + Inches(0.25), y + Inches(0.15), cw - Inches(0.5), ch - Inches(0.3))
    para(tf, t, size=17, color=BLUE, bold=True, first=True, space_after=4)
    para(tf, b, size=14, color=INK)

# ======================================================================
# Slide 6 困難事例とは（居宅版類型）
# ======================================================================
s = newslide("DEFINE", "「困難事例」とは何か")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, "ケアマネジメントが難航する要因を分解すると、必要な備えが見えてくる。", size=18, color=MUTED, first=True)
items = [("本人の状態・拒否", BLUE, "認知症・精神疾患、サービス拒否・支援拒否、意思決定支援が必要なケース。"),
         ("医療・退院支援", ACCENT, "医療ニーズが高い、退院直後で在宅移行が急、ターミナル・看取り。"),
         ("家族・世帯", WARM, "介護負担過多、虐待（疑い）、8050・ヤングケアラー、複合課題世帯。"),
         ("社会・制度・権利擁護", GOLD, "経済的困窮、独居・身寄りなし、成年後見、虐待対応・権利擁護。")]
cw, ch, gx, gy = Inches(5.85), Inches(1.55), Inches(0.3), Inches(0.28)
x0, y0 = Inches(0.7), Inches(2.1)
for i, (t, c, b) in enumerate(items):
    x = x0 + (i % 2) * (cw + gx)
    y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    tcol = RGBColor(0x5A, 0x45, 0x00) if c == GOLD else WHITE
    pill(s, x + Inches(0.22), y + Inches(0.18), t, color=c, tcolor=tcol, w=Inches(2.4))
    tb, tf = textbox(s, x + Inches(0.22), y + Inches(0.66), cw - Inches(0.44), ch - Inches(0.8))
    para(tf, b, size=14, first=True)
note(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.1),
     [("視点：", ACCENT, True), ("困難さは利用者の「属性」ではなく、抱え込みや連携不足など支援する側の体制から生まれることも多い。体制を整えれば「困難」は「対応可能」に変わる。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# Slide 7 受任〜ケアマネジメント基本フロー
# ======================================================================
s = newslide("FLOW", "受任〜ケアマネジメントの基本フロー")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, [("困難事例は担当一人で抱えない。", INK, False), ("受任判断から組織・チームで動かす。", WARM, True)], size=18, first=True)
steps = [("1", "インテーク／\n受任判断", "情報を受け止め組織で受任を判断"),
         ("2", "アセスメント", "課題分析・リスクと支援ニーズ把握"),
         ("3", "ケアプラン作成", "本人意向と多職種の視点を反映"),
         ("4", "サービス担当者\n会議", "多職種で役割分担・緊急時を共有"),
         ("5", "モニタリング", "定期訪問・再評価・見直し")]
n = len(steps)
total_w = Inches(11.9)
sw = Inches(2.15); gap = (total_w - sw * n) / (n - 1)
x0, y0, h = Inches(0.7), Inches(2.2), Inches(1.7)
for i, (num, t, sub) in enumerate(steps):
    x = x0 + i * (sw + gap)
    rect(s, x, y0, sw, h, fill=WHITE, line=BLUE, line_w=2)
    cr = s.shapes.add_shape(MSO_SHAPE.OVAL, x + sw/2 - Inches(0.22), y0 + Inches(0.15), Inches(0.44), Inches(0.44))
    _set_fill(cr, BLUE); cr.shadow.inherit = False
    para(cr.text_frame, num, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, x + Inches(0.08), y0 + Inches(0.62), sw - Inches(0.16), h - Inches(0.7))
    para(tf, t, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=3)
    para(tf, sub, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    if i < n - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + sw - Inches(0.02), y0 + h/2 - Inches(0.12), gap + Inches(0.04), Inches(0.24))
        _set_fill(ar, BLUE); ar.shadow.inherit = False
note(s, Inches(0.7), Inches(4.3), Inches(11.9), Inches(0.85),
     [("NG：", WARM, True), ("紹介の電話を受けた担当が「手一杯なので無理です」と即断する／逆に抱え込んで一人で背負う。", INK, False)],
     bar_color=WARM, fill=BADBG)
note(s, Inches(0.7), Inches(5.3), Inches(11.9), Inches(0.85),
     [("GOOD：", ACCENT, True), ("「管理者・主任ケアマネと検討し、◯日までに回答します」と、受任判断を会議に乗せる。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 体制 8-12
# ======================================================================
system_slide("SYSTEM 1/5", "体制づくり①　インテーク・受任判断",
    "受任判断で押さえる項目",
    ["本人・家族の状況、支援拒否の有無", "医療ニーズ・退院時期・緊急度", "担当できる人員・件数の余力（標準担当件数）",
     "必要な専門性（認知症・医療・権利擁護）", "地域包括・行政からの紹介背景"],
    "「断らない」を支える仕組み",
    ["受任は担当個人でなく管理者・主任CMが判断", "件数・難易度の偏りをチームで調整", "受けられない場合も理由と代替案を記録・連絡",
     "「何が揃えば受けられるか」を整理"],
    [("加算要件との接続：", ACCENT, True), ("地域包括等からの支援困難事例の紹介に応じ、担当ケアマネを配置できる体制が加算の要件。受任判断の仕組み化がその裏づけになる。", INK, False)],
    note_kind="good")

system_slide("SYSTEM 2/5", "体制づくり②　主任ケアマネによるスーパービジョン",
    "主任ケアマネ（管理者）の役割",
    ["困難事例の方針をSV（助言・同行・面談）で支える", "担当者会議・事例検討会の招集と進行", "ケアプラン点検と質の担保",
     "新人・経験の浅いCMの育成（OJT）"],
    "「一人で抱えない」チーム支援",
    ["主担当＋副担当／複数名訪問の体制", "困難事例は経験者とペアで担当", "相談・SOSを上げてよい文化を明示",
     "ケース進行を会議で共有し方針を統一"],
    [("ポイント：", ACCENT, True), ("主任ケアマネの配置は加算の人員要件であると同時に、困難事例を組織で受けるための要。SV機能を形だけにしない。", INK, False)],
    note_kind="good")

system_slide("SYSTEM 3/5", "体制づくり③　24時間連絡体制・緊急時対応",
    "24時間連絡体制の確保（加算要件）",
    ["利用者・家族からの夜間休日の連絡先を明示", "連絡当番・転送・記録のルール化", "緊急時の対応手順（救急・受診・家族連絡）を整備",
     "事業所として体制を負担（個人携帯任せにしない）"],
    "リスク・権利擁護への備え",
    ["虐待（疑い）発見時の通報・連携手順", "ハラスメント・暴力リスクへの対応手順", "災害・急変時のBCP的な備え",
     "成年後見・権利擁護の相談ルート"],
    [("カスタマーハラスメント等：", GOLD, True), ("「ケアマネだから我慢」ではなく、組織として線引きと手順を持つことが、結果的に支援継続を可能にする。", INK, False)])

system_slide("SYSTEM 4/5", "体制づくり④　記録・ケアプラン・担当者会議の質",
    "記録は「確実性」の証拠",
    ["受任判断・SVの記録（誰が何を根拠に判断したか）", "アセスメント・ケアプランとその更新履歴", "サービス担当者会議の記録",
     "モニタリング記録（月次の訪問・面接）"],
    "ケアマネジメントの質",
    ["本人の意向・自立支援に基づくケアプラン", "インフォーマル資源も組み込んだ総合的プラン", "多職種の意見を反映した担当者会議",
     "運営基準減算・集中減算に該当しない運用"],
    [("実地指導の視点：", ACCENT, True), ("「やっている」を「記録で示せる」へ。困難事例ほど、判断根拠と経過の記録が事業所とケアマネを守る。", INK, False)],
    note_kind="good")

system_slide("SYSTEM 5/5", "体制づくり⑤　研修計画・事例検討会",
    "研修計画に基づく実施（加算要件）",
    ["個々のCMの課題に応じた研修計画の作成・実施", "認知症・医療連携・権利擁護・意思決定支援を計画化", "主任CMによるOJT・同行訪問"],
    "他法人との事例検討会（加算要件）",
    ["他法人の居宅介護支援事業所と共同で事例検討会・研修", "本日のような内部事例検討会の定例化", "暗黙知を手順・ナレッジに変える",
     "新人も同じ判断軸を持てるようにする"],
    [("狙い：", GOLD, True), ("研修・事例検討会を「加算のための義務」で終わらせず、困難事例を受けきる事業所の支援力に転換する。", INK, False)])

# ======================================================================
# Slide 13 事例扉
# ======================================================================
s = newslide(grad=(WARM, GOLD))
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.5), Inches(2.4), Inches(0.5))
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xC2, 0x55, 0x38); b.line.fill.background(); b.shadow.inherit = False
para(b.text_frame, "CASE STUDY", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.2))
para(tf, "第3部　事例で学ぶ", size=42, color=WHITE, bold=True, first=True, space_after=2)
para(tf, "困難事例の受け入れと支援", size=42, color=WHITE, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(1.4))
para(tf2, "各事例：① 状況把握 → ② NG対応 → ③ GOOD対応 → ④ 体制のポイント", size=18, color=WHITE, bold=True, first=True, space_after=6)
para(tf2, "グループで「自分の事業所ならどう受任し、どう動くか」を3分間ディスカッション。", size=15, color=WHITE)

# ======================================================================
# 事例 14-17（居宅版）
# ======================================================================
case_slide("🏠 事例1：支援拒否のある独居高齢者（地域包括から紹介）",
    [("要介護2・独居、認知症の疑い。サービス利用を強く拒否し近隣トラブルも。地域包括支援センターから「担当してほしい」と紹介が入った。", INK, False)],
    ["「拒否が強く対応困難」と受任を断る", "一度の訪問拒否で「支援不可」と判断", "担当ケアマネが一人で説得を抱え込む"],
    ["管理者・主任CMと協議し受任を組織判断", "地域包括と役割分担し関係構築を優先・段階的に", "本人の意向・意思決定支援を軸にプラン化し記録"],
    [("地域包括からの困難事例の受け入れは加算要件の核心。断らずに『どう受けるか』を組織で設計し、関係構築から始める。", INK, False)])

case_slide("💰 事例2：認知症で金銭管理が困難・権利擁護が絡むケース",
    [("要介護3・独居、認知症が進行し金銭管理や契約が困難。親族とは疎遠。サービス費の滞納や悪質商法の被害も心配される。", INK, False)],
    ["介護サービスの調整だけで完結させる", "金銭・契約の問題を「家族の問題」と放置", "成年後見等の制度につなげず抱え込む"],
    ["地域包括・行政・社協と連携し権利擁護へ", "成年後見制度・日常生活自立支援事業を検討", "意思決定支援の経過を記録し担当者会議で共有"],
    [("権利擁護が絡む事例は1事業所で抱えない。多職種・行政につなぐこと自体が困難事例受け入れ体制の一部。", INK, False)])

case_slide("⚠️ 事例3：家族の介護負担・虐待(疑い)が絡む複合課題世帯",
    [("要介護4の母を息子が在宅介護。息子は仕事を失い経済的に困窮、母への不適切な対応（虐待の疑い）も見られる。8050的な複合課題。", INK, False)],
    ["「家庭の事情」として踏み込まず様子見", "虐待の疑いを記録・通報せず抱える", "介護サービスの追加調整だけで対応"],
    ["虐待の疑いは速やかに市町村・地域包括へ通報・連携", "息子（家族）の生活課題も含め多機関で支援", "事実経過を客観的に記録し担当者会議で共有"],
    [("虐待対応・複合課題は通報と多機関連携が原則。ケアマネが一人で判断せず、組織と地域のネットワークで動く。", INK, False)])

case_slide("🩺 事例4：医療ニーズの高い退院支援・看取り",
    [("末期がんで余命が限られ、本人は在宅での看取りを希望。退院日が迫り、医療・介護・家族の調整を短期間で整える必要がある。", INK, False)],
    ["退院日直前まで調整に着手しない", "医療側との情報共有が不十分なまま受任", "急変時の方針・連絡体制を決めずに開始"],
    ["退院前カンファレンスに参加し医療と方針共有", "訪問看護・主治医・薬剤師と役割分担を明確化", "本人の意向（ACP）と急変時対応を担当者会議で確認"],
    [("医療ニーズの高い事例は時間との勝負。24時間連絡体制と医療介護連携の段取り力が、在宅生活の質を左右する。", INK, False)])

# ======================================================================
# Slide 18 組織で受ける3原則
# ======================================================================
s = newslide("ORGANIZATION", "困難事例を「組織で受ける」3原則")
items = [("原則1", BLUE, "一人で判断しない", "受任・方針・撤退は管理者・主任CMと会議で決める。"),
         ("原則2", ACCENT, "つないで受ける", "1事業所で抱えず、多職種・地域・行政と役割分担する。"),
         ("原則3", WARM, "記録で残す", "判断根拠・経過・連携を残し、誰でも引き継げる状態に。")]
cw, gap, x0, y0, ch = Inches(3.85), Inches(0.28), Inches(0.7), Inches(1.7), Inches(2.5)
for i, (p, c, t, b) in enumerate(items):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    pill(s, x + Inches(0.22), y0 + Inches(0.2), p, color=c, w=Inches(1.5))
    tb, tf = textbox(s, x + Inches(0.25), y0 + Inches(0.75), cw - Inches(0.5), ch - Inches(0.9))
    para(tf, t, size=20, color=NAVY, bold=True, first=True, space_after=6)
    para(tf, b, size=15, color=INK)
note(s, Inches(0.7), Inches(4.5), Inches(11.9), Inches(1.4),
     [("これが「確実性のあるケアマネジメント」：", ACCENT, True),
      ("個人の力量や善意でなく、主任CMによるSV・会議・記録という仕組みで再現性を担保する。これが加算IIを算定する事業所の責任の果たし方。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# Slide 19 連携マップ（居宅版）
# ======================================================================
s = newslide("NETWORK", "多職種・地域との連携マップ")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, "困難事例ほど、居宅介護支援「だけ」では完結しない。つなぐ先を全員が把握する。", size=17, color=MUTED, first=True)
items = [("医療", "主治医／訪問看護／病院MSW／薬剤師。退院支援・看取り・服薬の連携。"),
         ("介護・インフォーマル", "各サービス事業所／福祉用具／家族・近隣／ボランティア・通いの場。"),
         ("地域・行政・権利擁護", "地域包括支援センター／市町村／社協／成年後見・虐待対応窓口。")]
cw, gap, x0, y0, ch = Inches(3.85), Inches(0.28), Inches(0.7), Inches(2.2), Inches(2.1)
for i, (t, b) in enumerate(items):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    hd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.22), y0 + Inches(0.2), cw - Inches(0.44), Inches(0.55))
    _set_fill(hd, BLUE); hd.shadow.inherit = False
    para(hd.text_frame, t, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, x + Inches(0.25), y0 + Inches(0.95), cw - Inches(0.5), ch - Inches(1.1))
    para(tf, b, size=14, first=True)
note(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(1.2),
     [("ケアマネの動き：", GOLD, True), ("ケアマネは地域のハブ。困難事例ほどサービス担当者会議・地域ケア会議を活用し、抱え込まず多機関で支える。", INK, False)])

# ======================================================================
# Slide 20 セルフチェック（居宅版）
# ======================================================================
s = newslide("CHECK", "自事業所の体制セルフチェック")
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5))
para(tf, "「はい」と言える項目はいくつ？　△の項目を今日の宿題に。", size=17, color=MUTED, first=True)
left = ["困難事例の受任を管理者・主任CMが判断する仕組みがある", "担当件数・難易度の偏りをチームで調整している",
        "主任CMによるSV（助言・同行）が機能している", "24時間連絡体制が事業所として整っている"]
right = ["虐待・権利擁護の通報・連携ルートが明確", "研修計画に困難事例対応が位置づけられている",
         "他法人との事例検討会を実施している", "受任判断・SV・会議の記録が残っている"]
cw, gx, x0, y0, ch = Inches(5.85), Inches(0.3), Inches(0.7), Inches(2.15), Inches(2.6)
for col, items in enumerate((left, right)):
    x = x0 + col * (cw + gx)
    rect(s, x, y0, cw, ch, fill=CARDBG, line=LINE)
    tb, tf = textbox(s, x + Inches(0.3), y0 + Inches(0.2), cw - Inches(0.55), ch - Inches(0.3))
    first = True
    for it in items:
        para(tf, [("✔ ", ACCENT, True), (it, INK, False)], size=15, first=first, space_after=10)
        first = False
note(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(1.15),
     [("使い方：", GOLD, True), ("各自で採点 → グループで「△を1つ、来月までにどう改善するか」を1つ決める。", INK, False)])

# ======================================================================
# Slide 21 まとめ
# ======================================================================
s = newslide("SUMMARY", "まとめ：加算IIの責任を事業所の支援力に")
cw, gx, x0, y0, ch = Inches(5.85), Inches(0.3), Inches(0.7), Inches(1.7), Inches(3.4)
card(s, x0, y0, cw, ch, title="今日のキーメッセージ",
     bullets=["困難事例の受け入れは加算IIの根幹（要件）", "受任は5ステップで組織判断、担当一人で抱えない",
              "「一人で判断しない・つないで受ける・記録で残す」", "主任CMのSV・会議・研修・記録が再現性を生む"],
     title_color=BLUE, body_size=16)
x1 = x0 + cw + gx
rect(s, x1, y0, cw, ch, fill=CARDBG, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), y0 + Inches(0.18), cw - Inches(0.5), ch - Inches(0.3))
para(tf, "明日からのアクション宣言", size=18, color=BLUE, bold=True, first=True, space_after=4)
para(tf, "グループで1つずつ発表してください。", size=13, color=MUTED, space_after=6)
para(tf, "・私の事業所が次に整えるのは「　　　　　」", size=16, space_after=4)
para(tf, "・私が個人として変えるのは「　　　　　」", size=16, space_after=8)
rect(s, x1 + Inches(0.25), y0 + Inches(2.5), cw - Inches(0.5), Inches(0.75), fill=GOODBG, line=None)
tb, tf = textbox(s, x1 + Inches(0.4), y0 + Inches(2.55), cw - Inches(0.8), Inches(0.65), anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("小さな一歩を、必ず", INK, False), ("仕組み（手順・記録）", ACCENT, True), ("に落とす。", INK, False)], size=15, first=True)

# ======================================================================
# Slide 22 クロージング
# ======================================================================
s = newslide(grad=(NAVY, ACCENT))
tb, tf = textbox(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.5))
para(tf, "ご清聴ありがとうございました", size=40, color=WHITE, bold=True, first=True, space_after=10)
para(tf, "「断れない」のではなく、「組織で困難事例を受けきれる」事業所へ。", size=20, color=GOLD, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.6))
para(tf2, "質疑・ディスカッションタイム", size=18, color=RGBColor(0xCF, 0xE0, 0xEE), first=True, space_after=8)
para(tf2, "※本資料の要件・人員数・割合等は、運用前に最新の介護報酬告示・解釈通知でご確認ください。",
     size=12, color=RGBColor(0xCF, 0xE0, 0xEE))

# ----------------------------------------------------------------------
out = "特定事業所加算II_居宅介護支援_困難事例受け入れ体制_勉強会.pptx"
prs.save(out)
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst))
