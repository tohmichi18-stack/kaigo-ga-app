# -*- coding: utf-8 -*-
"""
モニタリング記録の書き方（居宅介護支援 勉強会 補足資料）
PowerPoint(.pptx) 生成スクリプト
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY   = RGBColor(0x10, 0x33, 0x5C)
BLUE   = RGBColor(0x21, 0x6A, 0x8C)
ACCENT = RGBColor(0x1B, 0x9A, 0x8B)
WARM   = RGBColor(0xD9, 0x6A, 0x4A)
GOLD   = RGBColor(0xE3, 0xB5, 0x4A)
BG     = RGBColor(0xF3, 0xF7, 0xF9)
INK    = RGBColor(0x21, 0x2E, 0x38)
MUTED  = RGBColor(0x57, 0x68, 0x73)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xD4, 0xE0, 0xE6)
CARDBG = RGBColor(0xFF, 0xFF, 0xFF)
GOODBG = RGBColor(0xE6, 0xF5, 0xF2)
BADBG  = RGBColor(0xFB, 0xEC, 0xE6)
NOTEBG = RGBColor(0xFD, 0xF4, 0xE2)
GREYBG = RGBColor(0xEE, 0xF1, 0xF3)
FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    _set_fill(bg, BG); bg.shadow.inherit = False
    return s


def _grad(sp, c1, c2, ang):
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    grad = sp.makeelement(qn('a:gradFill'), {})
    gsLst = grad.makeelement(qn('a:gsLst'), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = grad.makeelement(qn('a:gs'), {'pos': str(pos)})
        srgb = grad.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (col[0], col[1], col[2])})
        gs.append(srgb); gsLst.append(gs)
    grad.append(gsLst)
    grad.append(grad.makeelement(qn('a:lin'), {'ang': str(ang), 'scaled': '1'}))
    ln = sp.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp.append(grad)


def grad_bg(s, c1, c2):
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background(); bg.shadow.inherit = False
    _grad(bg.fill._xPr, c1, c2, 2700000)
    return bg


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    return tb, tf


def para(tf, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=4, bullet=False, level=0, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after); p.level = level
    runs = text if isinstance(text, list) else [(text, color, bold)]
    for t, c, b in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = c; r.font.name = font
    _bullet(p) if bullet else _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def _bullet(p, char="・"):
    pPr = p._p.get_or_add_pPr()
    pPr.set('indent', '-228600'); pPr.set('marL', '228600')
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': FONT}))
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': char}))


def rect(s, x, y, w, h, fill=CARDBG, line=None, rounded=True, line_w=1.0):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def leftbar(s, x, y, h, color):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(8), h)
    _set_fill(bar, color); bar.shadow.inherit = False


def kicker(s, text, color=ACCENT):
    tb, tf = textbox(s, Inches(0.7), Inches(0.4), Inches(10), Inches(0.4))
    para(tf, text, size=13, color=color, bold=True, first=True)


def heading(s, text):
    leftbar(s, Inches(0.7), Inches(0.78), Inches(0.62), ACCENT)
    tb, tf = textbox(s, Inches(0.95), Inches(0.72), Inches(11.7), Inches(0.8))
    para(tf, text, size=28, color=NAVY, bold=True, first=True)


def pill(s, x, y, text, color=ACCENT, tcolor=WHITE, w=Inches(2.0)):
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
    _set_fill(p, color); p.shadow.inherit = False
    try:
        p.adjustments[0] = 0.5
    except Exception:
        pass
    tf = p.text_frame; tf.word_wrap = True; tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    para(tf, text, size=12, color=tcolor, bold=True, align=PP_ALIGN.CENTER, first=True)
    return p


def card(s, x, y, w, h, title=None, bullets=None, fill=CARDBG, line=LINE,
         title_color=BLUE, body_size=15, title_size=18):
    rect(s, x, y, w, h, fill=fill, line=line)
    pad = Inches(0.22)
    tb, tf = textbox(s, x + pad, y + Inches(0.12), w - pad * 2, h - Inches(0.2))
    first = True
    if title:
        para(tf, title, size=title_size, color=title_color, bold=True, first=True, space_after=5); first = False
    if bullets:
        for b in bullets:
            para(tf, b, size=body_size, color=INK, first=first, bullet=True, space_after=3); first = False
    return tf


def note(s, x, y, w, h, runs, bar_color=GOLD, fill=NOTEBG):
    rect(s, x, y, w, h, fill=fill, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(6), h)
    _set_fill(bar, bar_color); bar.shadow.inherit = False
    tb, tf = textbox(s, x + Inches(0.2), y, w - Inches(0.35), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, runs, size=14, first=True)


TOTAL = 13
_n = [0]


def footer(s, n):
    tb, tf = textbox(s, Inches(11.6), Inches(7.05), Inches(1.5), Inches(0.35))
    para(tf, "%d / %d" % (n, TOTAL), size=11, color=MUTED, align=PP_ALIGN.RIGHT, first=True)
    tb2, tf2 = textbox(s, Inches(0.7), Inches(7.05), Inches(9), Inches(0.35))
    para(tf2, "居宅介護支援 勉強会 補足｜モニタリング記録の書き方", size=10, color=MUTED, first=True)


def newslide(kick=None, title=None, grad=None):
    _n[0] += 1
    if grad:
        s = prs.slides.add_slide(BLANK); grad_bg(s, grad[0], grad[1])
    else:
        s = slide()
        if kick:
            kicker(s, kick)
        if title:
            heading(s, title)
        footer(s, _n[0])
    return s


# ======================================================================
# 1 タイトル
# ======================================================================
s = newslide(grad=(NAVY, ACCENT))
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.5), Inches(6.4), Inches(0.5))
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x21, 0x55, 0x76)
b.line.color.rgb = RGBColor(0x8F, 0xB6, 0xD4); b.line.width = Pt(1); b.shadow.inherit = False
para(b.text_frame, "居宅介護支援 勉強会 補足資料（ケアマネジメント深掘り編）", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.0))
para(tf, "モニタリング記録の書き方", size=46, color=WHITE, bold=True, first=True, space_after=8)
para(tf, "～ 「経過のメモ」から「目標の達成度評価」へ ～", size=20, color=GOLD, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.0))
para(tf2, "対象：介護支援専門員・主任介護支援専門員　／　ケアプランと連動した記録の型を身につける",
     size=15, color=RGBColor(0xCF, 0xE0, 0xEE), first=True)

# ======================================================================
# 2 モニタリングとは／基準
# ======================================================================
s = newslide("BASIS", "モニタリングとは何か（運営基準のおさらい）")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(2.7), title="モニタリングの目的",
     bullets=["ケアプランが適切に実施されているか確認", "目標の達成状況を評価する", "新たな課題・変化を把握する",
              "→ ケアプランの継続・修正・終了の判断につなぐ"], title_color=BLUE, body_size=15)
card(s, Inches(0.7) + Inches(6.15), Inches(1.7), Inches(5.85), Inches(2.7), title="運営基準上の原則",
     bullets=["少なくとも月1回、利用者宅を訪問し面接", "少なくとも月1回、モニタリング結果を記録", "特段の事情がない限り継続的に実施",
              "一定要件下のオンラインモニタリングは別途規定"], title_color=ACCENT, body_size=15)
note(s, Inches(0.7), Inches(4.6), Inches(11.9), Inches(1.3),
     [("ここがポイント：", GOLD, True),
      ("モニタリング記録は「訪問したことの証明」ではなく「ケアプランの目標がどこまで達成されたかの評価」。だからケアプラン（第2表）と必ずひも付けて書く。", INK, False)])
tb, tf = textbox(s, Inches(0.9), Inches(6.0), Inches(11.6), Inches(0.5))
para(tf, "※訪問頻度・記録頻度・オンライン要件の詳細は最新の運営基準・解釈通知でご確認ください。", size=11, color=MUTED, first=True)

# ======================================================================
# 3 ありがちなNG
# ======================================================================
s = newslide("PROBLEM", "ありがちな『書けていない』モニタリング記録")
items = [("出来事の羅列だけ", "「変わりなし」「お元気でした」だけ。何を評価したのか不明。"),
         ("ケアプランと無関係", "目標に触れず、世間話や天気のメモで終わっている。"),
         ("主観・印象のみ", "「良くなった気がする」など、根拠（事実）が示されない。"),
         ("課題を見落とす", "変化やリスクの芽に気づかず、見直しにつながらない。")]
cw, ch, gx, gy = Inches(5.85), Inches(1.7), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.75)
for i, (t, b) in enumerate(items):
    x = x0 + (i % 2) * (cw + gx); y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=BADBG, line=None)
    leftbar(s, x, y, ch, WARM)
    tb, tf = textbox(s, x + Inches(0.3), y + Inches(0.18), cw - Inches(0.55), ch - Inches(0.3))
    para(tf, [("✕ ", WARM, True), (t, NAVY, True)], size=17, first=True, space_after=4)
    para(tf, b, size=14, color=INK)
note(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.95),
     [("共通の原因：", GOLD, True), ("「訪問記録」と「モニタリング（評価）」を混同している。モニタリングは“目標に対する評価”が主役。", INK, False)])

# ======================================================================
# 4 書く前提：何を見るか（視点）
# ======================================================================
s = newslide("VIEWPOINT", "モニタリングで確認する5つの視点")
items = [("1. 目標の達成度", ACCENT, "短期目標は達成できたか・近づいているか"),
         ("2. サービスの実施・効果", BLUE, "計画どおり提供され、効果が出ているか"),
         ("3. 本人・家族の満足/意向", BLUE, "本人はどう感じているか、意向に変化は"),
         ("4. 新たな課題・変化", WARM, "心身・環境・家族状況の変化やリスク"),
         ("5. 計画の妥当性", GOLD, "今のケアプランは継続でよいか・見直すか")]
y = Inches(1.7); rh = Inches(0.92)
for i, (t, col, b) in enumerate(items):
    yy = y + i * (rh + Inches(0.08))
    rect(s, Inches(0.7), yy, Inches(11.9), rh, fill=WHITE, line=LINE)
    leftbar(s, Inches(0.7), yy, rh, col)
    tb, tf = textbox(s, Inches(1.0), yy + Inches(0.1), Inches(4.2), rh - Inches(0.2), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=17, color=NAVY, bold=True, first=True)
    tb2, tf2 = textbox(s, Inches(5.3), yy + Inches(0.1), Inches(7.1), rh - Inches(0.2), anchor=MSO_ANCHOR.MIDDLE)
    para(tf2, b, size=14, color=INK, first=True)

# ======================================================================
# 5 書き方の型（事実→評価→判断）
# ======================================================================
s = newslide("FORMAT", "書き方の型：事実 → 評価 → 判断")
steps = [("① 事実", "観察・聞き取った客観的事実", "本人・家族の発言、できていること、サービス実施状況、心身・生活の様子", BLUE),
         ("② 評価", "目標に照らした分析", "短期目標に対しどこまで達成できたか。なぜそうなったか（要因）", ACCENT),
         ("③ 判断", "今後どうするか", "計画は継続/修正/終了か。次の課題・対応・再アセスメントの要否", GOLD)]
cw = Inches(3.85); gap = Inches(0.275); x0 = Inches(0.7); y0 = Inches(1.8); ch = Inches(2.9)
for i, (t, sub, b, col) in enumerate(steps):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, fill=WHITE, line=col, line_w=2)
    hd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y0 + Inches(0.2), cw - Inches(0.4), Inches(0.55))
    _set_fill(hd, col); hd.shadow.inherit = False
    para(hd.text_frame, t, size=18, color=(RGBColor(0x5A, 0x45, 0x00) if col == GOLD else WHITE), bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, x + Inches(0.25), y0 + Inches(0.95), cw - Inches(0.5), ch - Inches(1.1))
    para(tf, sub, size=15, color=NAVY, bold=True, first=True, space_after=6)
    para(tf, b, size=13.5, color=INK)
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + cw - Inches(0.02), y0 + ch/2 - Inches(0.12), gap + Inches(0.04), Inches(0.24))
        _set_fill(ar, GOLD); ar.shadow.inherit = False
note(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(0.95),
     [("コツ：", ACCENT, True), ("「事実」と「ケアマネの判断」を混ぜない。事実を書いてから、それをどう評価し、どう判断したかを分けて書く。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 6 SOAPの活用
# ======================================================================
s = newslide("OPTION", "書きやすくする型：SOAPの活用")
items = [("S（主観的情報）", "本人・家族が語った言葉・訴え", "「最近よく眠れる」「外に出るのが億劫」", BLUE),
         ("O（客観的情報）", "観察・測定した事実", "歩行の様子、サービス実施状況、体重・服薬", BLUE),
         ("A（評価・分析）", "S/Oをもとに目標に照らした判断", "短期目標『伝い歩き』はおおむね達成", ACCENT),
         ("P（計画）", "今後の対応・方針", "現計画継続。次回は買い物の外出を評価", GOLD)]
y = Inches(1.7); rh = Inches(1.05)
for i, (t, role, ex, col) in enumerate(items):
    yy = y + i * (rh + Inches(0.06))
    rect(s, Inches(0.7), yy, Inches(11.9), rh, fill=WHITE, line=LINE)
    tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), yy + rh/2 - Inches(0.26), Inches(2.5), Inches(0.52))
    _set_fill(tag, col); tag.shadow.inherit = False
    para(tag.text_frame, t, size=14, color=(RGBColor(0x5A, 0x45, 0x00) if col == GOLD else WHITE), bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, Inches(3.6), yy + Inches(0.12), Inches(3.6), rh - Inches(0.24), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, role, size=14, color=NAVY, bold=True, first=True)
    tb2, tf2 = textbox(s, Inches(7.3), yy + Inches(0.12), Inches(5.1), rh - Inches(0.24), anchor=MSO_ANCHOR.MIDDLE)
    para(tf2, [("例：", MUTED, True), (ex, INK, False)], size=13.5, first=True)
note(s, Inches(0.7), Inches(6.3), Inches(11.9), Inches(0.0) + Inches(0.7),
     [("ヒント：", GOLD, True), ("様式が自由記載でも、頭の中でS・O・A・Pに分けると「事実」と「評価」が自然に整理できる。", INK, False)])

# ======================================================================
# 7 Before / After 1
# ======================================================================
s = newslide("EXAMPLE", "記載例①　Before / After（達成度評価）")
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5))
para(tf, "短期目標：「見守りがあれば家の中を伝い歩きできる」（Aさん・要介護2）", size=15, color=MUTED, first=True)
# before
rect(s, Inches(0.7), Inches(2.05), Inches(5.85), Inches(3.4), fill=BADBG, line=None)
leftbar(s, Inches(0.7), Inches(2.05), Inches(3.4), WARM)
tb, tf = textbox(s, Inches(1.0), Inches(2.2), Inches(5.4), Inches(3.1))
para(tf, "✕ Before", size=17, color=WARM, bold=True, first=True, space_after=6)
para(tf, "「お変わりなくお過ごしです。デイにも通えています。引き続き様子を見ます。」", size=14.5, space_after=8)
para(tf, "→ 目標に触れていない", size=13, color=MUTED, space_after=2)
para(tf, "→ 事実が曖昧で評価がない", size=13, color=MUTED, space_after=2)
para(tf, "→ 次の判断につながらない", size=13, color=MUTED)
# after
x1 = Inches(0.7) + Inches(6.15)
rect(s, x1, Inches(2.05), Inches(5.85), Inches(3.4), fill=GOODBG, line=None)
leftbar(s, x1, Inches(2.05), Inches(3.4), ACCENT)
tb, tf = textbox(s, x1 + Inches(0.3), Inches(2.2), Inches(5.4), Inches(3.1))
para(tf, "○ After", size=17, color=ACCENT, bold=True, first=True, space_after=6)
para(tf, [("【事実】", BLUE, True), ("居室から台所まで壁を伝い自力で移動。ふらつきは減少。本人「前より怖くない」。", INK, False)], size=13.5, space_after=4)
para(tf, [("【評価】", BLUE, True), ("短期目標はおおむね達成。下肢筋力向上が要因と考えられる。", INK, False)], size=13.5, space_after=4)
para(tf, [("【判断】", BLUE, True), ("現計画を継続。次回は屋外歩行・買い物の可否を評価し、長期目標の前進を確認する。", INK, False)], size=13.5)

# ======================================================================
# 8 Before / After 2（変化・課題の発見）
# ======================================================================
s = newslide("EXAMPLE", "記載例②　変化・新たな課題を捉える")
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5))
para(tf, "モニタリングは「うまくいっている確認」だけでなく「変化の早期発見」も役割。", size=15, color=MUTED, first=True)
rect(s, Inches(0.7), Inches(2.05), Inches(5.85), Inches(3.4), fill=BADBG, line=None)
leftbar(s, Inches(0.7), Inches(2.05), Inches(3.4), WARM)
tb, tf = textbox(s, Inches(1.0), Inches(2.2), Inches(5.4), Inches(3.1))
para(tf, "✕ Before", size=17, color=WARM, bold=True, first=True, space_after=6)
para(tf, "「最近少し元気がない様子。家族も大変そう。見守り継続。」", size=14.5, space_after=8)
para(tf, "→ 印象だけで事実がない", size=13, color=MUTED, space_after=2)
para(tf, "→ リスクが具体化されない", size=13, color=MUTED, space_after=2)
para(tf, "→ 対応（誰が何を）がない", size=13, color=MUTED)
x1 = Inches(0.7) + Inches(6.15)
rect(s, x1, Inches(2.05), Inches(5.85), Inches(3.4), fill=GOODBG, line=None)
leftbar(s, x1, Inches(2.05), Inches(3.4), ACCENT)
tb, tf = textbox(s, x1 + Inches(0.3), Inches(2.2), Inches(5.4), Inches(3.1))
para(tf, "○ After", size=17, color=ACCENT, bold=True, first=True, space_after=6)
para(tf, [("【事実】", BLUE, True), ("食事量が2週間で減少、体重1.5kg減。妻「夜間の介護で眠れない」と疲労を訴え。", INK, False)], size=13.5, space_after=4)
para(tf, [("【評価】", BLUE, True), ("低栄養リスクと介護者の負担増。現計画では対応が不十分と判断。", INK, False)], size=13.5, space_after=4)
para(tf, [("【判断】", BLUE, True), ("主治医へ受診勧奨、ショート利用を検討。サービス担当者会議を開催し計画見直し（再アセスメント）へ。", INK, False)], size=13.5)
note(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.95),
     [("重要：", GOLD, True), ("「変化あり・新たな課題」を捉えたら、モニタリングはケアプラン見直し（再アセスメント・担当者会議）に確実につなぐ。", INK, False)])

# ======================================================================
# 9 困難事例でのモニタリング記録
# ======================================================================
s = newslide("DIFFICULT CASE", "困難事例におけるモニタリング記録の留意点")
cards = [("支援拒否・関係構築途上", "拒否の事実と背景、関わりの工夫と本人の反応を具体的に。小さな変化も評価として残す。", WARM),
         ("認知症・意思決定支援", "本人の言葉・表情・行動の事実を丁寧に。意思の推定根拠、家族・多職種からの情報も記録。", ACCENT),
         ("虐待・複合課題世帯", "客観的事実を時系列で（憶測と区別）。リスク評価と、通報・多機関連携の対応を明記。", WARM),
         ("看取り・状態変化", "状態変化のスピードに合わせ頻回に。本人の意向（ACP）と急変時対応の確認を記録。", ACCENT)]
cw, ch, gx, gy = Inches(5.85), Inches(1.85), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.7)
for i, (t, b, col) in enumerate(cards):
    x = x0 + (i % 2) * (cw + gx); y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    leftbar(s, x, y, ch, col)
    tb, tf = textbox(s, x + Inches(0.3), y + Inches(0.18), cw - Inches(0.55), ch - Inches(0.3))
    para(tf, t, size=16, color=NAVY, bold=True, first=True, space_after=4)
    para(tf, b, size=13.5, color=INK)
note(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.95),
     [("共通：", GOLD, True), ("困難事例ほど「事実」と「ケアマネの判断・推測」を明確に区別して書く。記録が事業所と利用者の双方を守る。", INK, False)])

# ======================================================================
# 10 書くときの注意点
# ======================================================================
s = newslide("TIPS", "記録の質を上げる実務の注意点")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(3.5), title="表現・内容",
     bullets=["5W1Hで具体的に（いつ・誰が・何を）", "事実と評価・推測を書き分ける", "本人・家族の言葉はそのまま「　」で残す",
              "略語・身内表現を避け、第三者が読んで分かる", "断定できないことは「〜と思われる」で根拠とともに"],
     title_color=BLUE, body_size=14)
card(s, Inches(0.7) + Inches(6.15), Inches(1.7), Inches(5.85), Inches(3.5), title="運用・コンプライアンス",
     bullets=["月1回の訪問・面接と記録を確実に（基準）", "訪問日・記録日・記録者を明確に", "目標の評価期間が来たら必ず達成度を判定",
              "見直しが必要なら担当者会議・再アセスメントへ", "改ざんと誤解されぬよう追記・修正の運用を統一"],
     title_color=ACCENT, body_size=14)
note(s, Inches(0.7), Inches(5.4), Inches(11.9), Inches(0.95),
     [("合言葉：", GOLD, True), ("「あとで自分以外の誰かが読んでも、利用者の状況と判断の根拠が分かる」記録を目指す。", INK, False)])

# ======================================================================
# 11 演習
# ======================================================================
s = newslide("WORKSHOP", "演習：このモニタリング記録を直してみよう")
rect(s, Inches(0.7), Inches(1.6), Inches(11.9), Inches(1.7), fill=GREYBG, line=LINE)
tb, tf = textbox(s, Inches(1.0), Inches(1.75), Inches(11.3), Inches(1.4))
para(tf, "【お題】次の記録を「事実→評価→判断」で書き直してください。", size=15, color=NAVY, bold=True, first=True, space_after=6)
para(tf, "「Bさん、今日もデイに行けていて安心。家族も特に困っていない様子。このまま継続でよいと思う。」",
     size=15, color=INK)
card(s, Inches(0.7), Inches(3.5), Inches(11.9), Inches(2.6), title="グループで話し合う問い",
     bullets=["短期目標は何だった？ その達成度は記録から判断できる？",
              "「安心」「様子」は事実？ 観察した具体的な事実は何？",
              "家族の状況は本当に確認した？（聞き取りの事実は？）",
              "継続でよい根拠は？ 見直すべき変化は隠れていない？"],
     title_color=BLUE, body_size=15)

# ======================================================================
# 12 全体の流れの中の位置づけ
# ======================================================================
s = newslide("CYCLE", "モニタリングはケアマネジメントの「回す力」")
steps = [("アセスメント", BLUE), ("ケアプラン", BLUE), ("担当者会議", BLUE), ("サービス提供", BLUE), ("モニタリング", ACCENT)]
n = len(steps); total_w = Inches(11.9)
sw = Inches(2.0); gap = (total_w - sw * n) / (n - 1)
x0, y0, h = Inches(0.7), Inches(2.3), Inches(1.1)
for i, (t, col) in enumerate(steps):
    x = x0 + i * (sw + gap)
    r = rect(s, x, y0, sw, h, fill=(GOODBG if col == ACCENT else WHITE), line=col, line_w=2)
    tb, tf = textbox(s, x + Inches(0.1), y0, sw - Inches(0.2), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER, first=True)
    if i < n - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + sw - Inches(0.02), y0 + h/2 - Inches(0.11), gap + Inches(0.04), Inches(0.22))
        _set_fill(ar, GOLD); ar.shadow.inherit = False
# feedback arrow
fb = s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, x0 + Inches(1.0), y0 + h + Inches(0.45), total_w - Inches(2.0), Inches(0.4))
_set_fill(fb, ACCENT); fb.shadow.inherit = False
para(fb.text_frame, "モニタリングの評価・判断が次のアセスメント／計画見直しに戻る", size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
note(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(1.2),
     [("つまり：", ACCENT, True),
      ("モニタリング記録は過程の終点であり次の起点。ここで目標を評価し変化を捉えるから、ケアマネジメントのサイクルが回り、質が上がり続ける。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 13 まとめ
# ======================================================================
s = newslide("SUMMARY", "まとめ：モニタリング記録の要点")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(3.6), title="チェックリスト",
     bullets=["月1回の訪問・面接と記録ができているか", "短期目標の達成度を評価しているか",
              "事実と評価・判断を書き分けているか", "本人・家族の言葉を具体的に残しているか",
              "変化・新たな課題を捉えているか", "見直しが必要なら会議・再アセスにつないだか"],
     title_color=BLUE, body_size=14)
x1 = Inches(0.7) + Inches(6.15)
rect(s, x1, Inches(1.7), Inches(5.85), Inches(3.6), fill=CARDBG, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), Inches(1.9), Inches(5.35), Inches(3.3))
para(tf, "キーメッセージ", size=18, color=BLUE, bold=True, first=True, space_after=8)
para(tf, [("モニタリングは", INK, False), ("「目標の達成度評価」", WARM, True), ("。", INK, False)], size=16, space_after=10)
para(tf, "ケアプラン（第2表）の目標とひも付け、事実→評価→判断で書く。", size=15, space_after=10)
para(tf, "変化を捉えたら見直しへ。記録が事業所とケアマネ、そして利用者を守る。", size=15)
note(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.85),
     [("※", MUTED, True), ("訪問・記録の頻度やオンライン要件、様式の詳細は最新の運営基準・解釈通知でご確認ください。", MUTED, False)],
     bar_color=MUTED, fill=GREYBG)

# ----------------------------------------------------------------------
out = "モニタリング記録の書き方_勉強会補足.pptx"
prs.save(out)
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst))
