"""ヤングケアラーについてのスライド生成スクリプト"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# カラーパレット
COLOR_PRIMARY = RGBColor(0x2E, 0x5E, 0x8A)      # ディープブルー
COLOR_ACCENT = RGBColor(0xF2, 0x8C, 0x28)       # オレンジ
COLOR_LIGHT = RGBColor(0xF5, 0xF7, 0xFA)        # 薄いグレー
COLOR_DARK = RGBColor(0x22, 0x2B, 0x3A)         # 濃紺
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY = RGBColor(0x66, 0x70, 0x80)
COLOR_GREEN = RGBColor(0x2E, 0x8B, 0x57)

FONT_JP = "Yu Gothic"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]


def add_bg(slide, color):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def add_text(slide, left, top, width, height, text, size=18, bold=False,
             color=COLOR_DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font=FONT_JP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_rect(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def add_round_rect(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


def page_header(slide, title, page_num, total):
    # 上部の帯
    add_rect(slide, 0, 0, prs.slide_width, Inches(1.05), COLOR_PRIMARY)
    # タイトル
    add_text(slide, Inches(0.6), Inches(0.22), Inches(11), Inches(0.7),
             title, size=28, bold=True, color=COLOR_WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    # ページ番号
    add_text(slide, Inches(11.8), Inches(0.32), Inches(1.3), Inches(0.5),
             f"{page_num} / {total}", size=12, color=COLOR_WHITE,
             align=PP_ALIGN.RIGHT)
    # 下のアクセント線
    add_rect(slide, 0, Inches(1.05), prs.slide_width, Inches(0.08), COLOR_ACCENT)
    # フッター
    add_text(slide, Inches(0.6), Inches(7.05), Inches(11), Inches(0.35),
             "ヤングケアラーを知る・支える", size=10, color=COLOR_GRAY)


TOTAL = 12

# --------- スライド1: タイトル ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_PRIMARY)
# 装飾
add_rect(s, 0, Inches(5.6), prs.slide_width, Inches(0.15), COLOR_ACCENT)
add_rect(s, 0, Inches(5.85), prs.slide_width, Inches(0.05), COLOR_WHITE)

add_text(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.0),
         "ヤングケアラー", size=66, bold=True, color=COLOR_WHITE)
add_text(s, Inches(0.8), Inches(2.6), Inches(11.7), Inches(0.7),
         "Young Carers / Young Caregivers", size=24, color=COLOR_ACCENT)
add_text(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(1.4),
         "家族の世話を担う子どもたちを\nどう理解し、どう支えるか",
         size=28, color=COLOR_WHITE)
add_text(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4),
         "2026年5月", size=14, color=COLOR_LIGHT)

# --------- スライド2: 目次 ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "目次", 2, TOTAL)

items = [
    ("01", "ヤングケアラーとは"),
    ("02", "どんなケアをしているのか"),
    ("03", "日本における現状（統計）"),
    ("04", "背景にある社会的要因"),
    ("05", "本人が直面する課題"),
    ("06", "周囲が気づくサイン"),
    ("07", "支援の仕組みと相談窓口"),
    ("08", "私たちにできること"),
    ("09", "まとめ"),
]
col_w = Inches(6.0)
row_h = Inches(0.55)
for i, (num, label) in enumerate(items):
    col = i // 5
    row = i % 5
    left = Inches(0.7) + col * col_w
    top = Inches(1.6) + row * Inches(0.95)
    add_round_rect(s, left, top, Inches(0.7), Inches(0.7), COLOR_ACCENT)
    add_text(s, left, top, Inches(0.7), Inches(0.7), num,
             size=18, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.9), top + Inches(0.05),
             Inches(5.0), Inches(0.6), label,
             size=18, bold=True, color=COLOR_DARK,
             anchor=MSO_ANCHOR.MIDDLE)

# --------- スライド3: ヤングケアラーとは ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "01. ヤングケアラーとは", 3, TOTAL)

# 定義ボックス
add_round_rect(s, Inches(0.7), Inches(1.5), Inches(12), Inches(1.6),
               COLOR_LIGHT)
add_text(s, Inches(1.0), Inches(1.65), Inches(0.5), Inches(0.5),
         "“", size=48, bold=True, color=COLOR_ACCENT)
add_text(s, Inches(1.5), Inches(1.75), Inches(10.8), Inches(1.3),
         "本来大人が担うと想定されている家事や家族の世話などを\n"
         "日常的に行っている、おおむね18歳未満の子ども",
         size=22, bold=True, color=COLOR_DARK)
add_text(s, Inches(1.5), Inches(2.7), Inches(10.8), Inches(0.4),
         "（厚生労働省・こども家庭庁による定義をもとに）",
         size=12, color=COLOR_GRAY)

# ポイント3つ
points = [
    ("家族のケア", "親・きょうだい・祖父母など\n身近な家族をケアしている"),
    ("日常的・継続的", "一時的な手伝いではなく\n日々の生活に組み込まれている"),
    ("年齢に見合わない責任", "学業や友人関係に影響が出る\nほどの負担を担っている"),
]
for i, (title, desc) in enumerate(points):
    left = Inches(0.7) + i * Inches(4.13)
    top = Inches(3.5)
    add_round_rect(s, left, top, Inches(3.93), Inches(2.8), COLOR_WHITE,
                   line=COLOR_PRIMARY)
    add_rect(s, left, top, Inches(3.93), Inches(0.7), COLOR_PRIMARY)
    add_text(s, left, top, Inches(3.93), Inches(0.7), title,
             size=18, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.2), top + Inches(0.9),
             Inches(3.5), Inches(1.8), desc,
             size=15, color=COLOR_DARK, align=PP_ALIGN.CENTER)

# --------- スライド4: どんなケアをしているのか ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "02. どんなケアをしているのか", 4, TOTAL)

add_text(s, Inches(0.7), Inches(1.35), Inches(12), Inches(0.5),
         "ヤングケアラーが担うケアは、多岐にわたります。",
         size=16, color=COLOR_GRAY)

cares = [
    ("家事", "食事の準備・洗濯・\n掃除・買い物"),
    ("身体的ケア", "入浴・トイレ・\n着替えの介助"),
    ("感情的サポート", "病気の家族の\n話し相手・見守り"),
    ("きょうだいの世話", "幼い弟妹の世話・\n保育園送り迎え"),
    ("通訳", "日本語が苦手な\n家族の通訳・代筆"),
    ("金銭管理", "家計のやりくり・\n手続きの代行"),
    ("見守り", "目を離せない家族と\n常に一緒にいる"),
    ("通院付き添い", "病院への同行・\n薬の管理"),
]
cols, rows = 4, 2
cell_w = Inches(3.0)
cell_h = Inches(2.1)
start_left = Inches(0.7)
start_top = Inches(2.0)
for i, (t, d) in enumerate(cares):
    r = i // cols
    c = i % cols
    left = start_left + c * Inches(3.1)
    top = start_top + r * Inches(2.25)
    add_round_rect(s, left, top, cell_w, cell_h, COLOR_LIGHT)
    add_rect(s, left + Inches(0.15), top + Inches(0.25),
             Inches(0.12), Inches(0.45), COLOR_ACCENT)
    add_text(s, left + Inches(0.35), top + Inches(0.2),
             cell_w - Inches(0.4), Inches(0.5), t,
             size=18, bold=True, color=COLOR_PRIMARY)
    add_text(s, left + Inches(0.35), top + Inches(0.85),
             cell_w - Inches(0.4), Inches(1.1), d,
             size=14, color=COLOR_DARK)

# --------- スライド5: 日本における現状 ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "03. 日本における現状", 5, TOTAL)

add_text(s, Inches(0.7), Inches(1.3), Inches(12), Inches(0.5),
         "国の実態調査では、決して少なくない子どもがケアを担っていることが分かっています。",
         size=14, color=COLOR_GRAY)

# 大きな数字ブロック
stats = [
    ("中学2年生", "約 17 人に 1 人", "クラスに1〜2人いる計算"),
    ("全日制高校2年生", "約 24 人に 1 人", "見えにくいが確実に存在"),
    ("小学6年生", "約 15 人に 1 人", "低年齢化も指摘される"),
]
for i, (label, num, desc) in enumerate(stats):
    left = Inches(0.7) + i * Inches(4.13)
    top = Inches(2.0)
    add_round_rect(s, left, top, Inches(3.93), Inches(3.0),
                   COLOR_PRIMARY)
    add_text(s, left, top + Inches(0.25), Inches(3.93), Inches(0.5),
             label, size=16, bold=True, color=COLOR_LIGHT,
             align=PP_ALIGN.CENTER)
    add_text(s, left, top + Inches(0.9), Inches(3.93), Inches(1.2),
             num, size=28, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, left + Inches(1.2), top + Inches(2.15),
             Inches(1.53), Inches(0.04), COLOR_ACCENT)
    add_text(s, left, top + Inches(2.3), Inches(3.93), Inches(0.6),
             desc, size=12, color=COLOR_LIGHT,
             align=PP_ALIGN.CENTER)

add_round_rect(s, Inches(0.7), Inches(5.3), Inches(12), Inches(1.4),
               COLOR_LIGHT)
add_text(s, Inches(1.0), Inches(5.45), Inches(11.5), Inches(0.5),
         "■ 気づかれにくい問題", size=15, bold=True, color=COLOR_ACCENT)
add_text(s, Inches(1.0), Inches(5.9), Inches(11.5), Inches(0.8),
         "・自分が「ヤングケアラー」だと自覚していない子どもが多い\n"
         "・「家のことだから」と外に相談できないまま我慢してしまうケースが目立つ",
         size=13, color=COLOR_DARK)

# --------- スライド6: 背景にある社会的要因 ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "04. 背景にある社会的要因", 6, TOTAL)

factors = [
    ("少子高齢化", "祖父母世代の介護を\n孫が担う家庭が増加"),
    ("核家族化・ひとり親", "家族内に\n大人の手が少ない"),
    ("共働きの増加", "保護者が日中不在の\n時間が長くなる"),
    ("精神疾患・障害", "親やきょうだいに\nケアが必要な場合"),
    ("経済的困窮", "外部サービスを\n利用しづらい"),
    ("外国ルーツの家庭", "言語の壁を\n子どもが埋めている"),
]
for i, (t, d) in enumerate(factors):
    r = i // 3
    c = i % 3
    left = Inches(0.7) + c * Inches(4.13)
    top = Inches(1.6) + r * Inches(2.7)
    add_round_rect(s, left, top, Inches(3.93), Inches(2.5),
                   COLOR_WHITE, line=COLOR_ACCENT)
    add_round_rect(s, left + Inches(0.3), top - Inches(0.2),
                   Inches(0.7), Inches(0.7), COLOR_ACCENT)
    add_text(s, left + Inches(0.3), top - Inches(0.2),
             Inches(0.7), Inches(0.7), str(i + 1),
             size=18, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(1.2), top + Inches(0.05),
             Inches(2.6), Inches(0.6), t,
             size=17, bold=True, color=COLOR_PRIMARY,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.3), top + Inches(1.0),
             Inches(3.3), Inches(1.3), d,
             size=14, color=COLOR_DARK)

# --------- スライド7: 本人が直面する課題 ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "05. 本人が直面する課題", 7, TOTAL)

# 4象限
sections = [
    ("学業面", COLOR_PRIMARY, [
        "・遅刻、欠席が増える",
        "・宿題ができない",
        "・進学・就職をあきらめる",
        "・授業に集中できない",
    ]),
    ("健康面", COLOR_ACCENT, [
        "・睡眠不足、慢性的な疲労",
        "・食事が不規則になる",
        "・心身の不調・うつ症状",
        "・通院・受診の時間がない",
    ]),
    ("人間関係", COLOR_GREEN, [
        "・友人と遊ぶ時間がない",
        "・部活動に参加できない",
        "・孤立感、疎外感",
        "・恋愛・結婚をあきらめる",
    ]),
    ("将来", COLOR_DARK, [
        "・キャリア形成の機会損失",
        "・経済的自立の遅れ",
        "・自己肯定感の低下",
        "・「自分」を持てなくなる",
    ]),
]
positions = [
    (Inches(0.7), Inches(1.5)),
    (Inches(6.93), Inches(1.5)),
    (Inches(0.7), Inches(4.4)),
    (Inches(6.93), Inches(4.4)),
]
for (t, color, items), (left, top) in zip(sections, positions):
    add_round_rect(s, left, top, Inches(5.7), Inches(2.7),
                   COLOR_LIGHT)
    add_rect(s, left, top, Inches(0.18), Inches(2.7), color)
    add_text(s, left + Inches(0.4), top + Inches(0.15),
             Inches(5.2), Inches(0.5), t,
             size=20, bold=True, color=color)
    add_text(s, left + Inches(0.4), top + Inches(0.75),
             Inches(5.2), Inches(1.9), "\n".join(items),
             size=14, color=COLOR_DARK)

# --------- スライド8: 周囲が気づくサイン ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "06. 周囲が気づくサイン", 8, TOTAL)

add_text(s, Inches(0.7), Inches(1.3), Inches(12), Inches(0.5),
         "学校・地域・職場で、こんな様子に気づいたら声をかけてみてください。",
         size=14, color=COLOR_GRAY)

signs = [
    "遅刻や欠席、早退が目立つ",
    "忘れ物・宿題未提出が多い",
    "授業中に眠そうにしている",
    "服装や持ち物の手入れが行き届かない",
    "保護者と連絡が取りづらい",
    "急いで帰宅する、放課後に活動しない",
    "年齢に比べて大人びた発言・態度",
    "家庭の話題を避ける、または極端に詳しい",
]
for i, sign in enumerate(signs):
    r = i // 2
    c = i % 2
    left = Inches(0.7) + c * Inches(6.13)
    top = Inches(2.0) + r * Inches(1.05)
    add_round_rect(s, left, top, Inches(5.93), Inches(0.85),
                   COLOR_LIGHT)
    add_round_rect(s, left + Inches(0.25), top + Inches(0.18),
                   Inches(0.5), Inches(0.5), COLOR_ACCENT)
    add_text(s, left + Inches(0.25), top + Inches(0.18),
             Inches(0.5), Inches(0.5), "✓",
             size=18, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.95), top, Inches(4.9), Inches(0.85),
             sign, size=15, color=COLOR_DARK,
             anchor=MSO_ANCHOR.MIDDLE)

add_round_rect(s, Inches(0.7), Inches(6.45), Inches(12), Inches(0.5),
               COLOR_PRIMARY)
add_text(s, Inches(0.7), Inches(6.45), Inches(12), Inches(0.5),
         "ポイント：「がんばっているね」より「困っていない？」の一声を。",
         size=14, bold=True, color=COLOR_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# --------- スライド9: 支援の仕組みと相談窓口 ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "07. 支援の仕組みと相談窓口", 9, TOTAL)

# 左：支援の仕組み
add_round_rect(s, Inches(0.7), Inches(1.5), Inches(6), Inches(5.2),
               COLOR_LIGHT)
add_rect(s, Inches(0.7), Inches(1.5), Inches(6), Inches(0.6),
         COLOR_PRIMARY)
add_text(s, Inches(0.7), Inches(1.5), Inches(6), Inches(0.6),
         "公的な支援の枠組み",
         size=18, bold=True, color=COLOR_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
support = [
    ("こども家庭庁", "実態把握・支援強化の旗振り役"),
    ("スクールソーシャルワーカー", "学校での発見・関係機関への橋渡し"),
    ("市区町村のこども家庭センター", "家庭全体の包括的な相談支援"),
    ("地域包括支援センター", "高齢者ケアに関する家族の相談"),
    ("障害者相談支援事業所", "障害のある家族のケアプラン作成"),
    ("ヤングケアラー・コーディネーター", "本人と支援を結ぶ専門人材"),
]
for i, (name, desc) in enumerate(support):
    top = Inches(2.3) + i * Inches(0.7)
    add_text(s, Inches(1.0), top, Inches(5.5), Inches(0.3), "● " + name,
             size=13, bold=True, color=COLOR_PRIMARY)
    add_text(s, Inches(1.2), top + Inches(0.3), Inches(5.3), Inches(0.3),
             desc, size=11, color=COLOR_GRAY)

# 右：相談窓口
add_round_rect(s, Inches(6.93), Inches(1.5), Inches(5.8), Inches(5.2),
               COLOR_LIGHT)
add_rect(s, Inches(6.93), Inches(1.5), Inches(5.8), Inches(0.6),
         COLOR_ACCENT)
add_text(s, Inches(6.93), Inches(1.5), Inches(5.8), Inches(0.6),
         "相談できる窓口（一例）",
         size=18, bold=True, color=COLOR_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
contacts = [
    ("24時間子供SOSダイヤル", "0120-0-78310"),
    ("チャイルドライン", "0120-99-7777"),
    ("よりそいホットライン", "0120-279-338"),
    ("いのちの電話", "0570-783-556"),
    ("各自治体の児童相談所", "189（いちはやく）"),
]
for i, (name, tel) in enumerate(contacts):
    top = Inches(2.3) + i * Inches(0.85)
    add_text(s, Inches(7.2), top, Inches(5.3), Inches(0.4), name,
             size=14, bold=True, color=COLOR_DARK)
    add_text(s, Inches(7.2), top + Inches(0.38), Inches(5.3), Inches(0.4),
             "☎ " + tel, size=13, color=COLOR_PRIMARY, bold=True)

# --------- スライド10: 私たちにできること ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "08. 私たちにできること", 10, TOTAL)

roles = [
    ("学校・教員", COLOR_PRIMARY, [
        "・サインに気づく目を持つ",
        "・SSWや養護教諭と連携",
        "・「家庭の事情」で片づけない",
        "・本人の声を否定せず聴く",
    ]),
    ("地域・近所", COLOR_ACCENT, [
        "・家事や送迎を手伝う",
        "・「いつでも頼って」と伝える",
        "・行政につなぐ役を担う",
        "・家族全体を孤立させない",
    ]),
    ("行政・専門職", COLOR_GREEN, [
        "・横断的なケース会議の実施",
        "・家族全体のアセスメント",
        "・レスパイト支援の充実",
        "・継続的なフォローアップ",
    ]),
    ("私たち一人ひとり", COLOR_DARK, [
        "・正しく知り、偏見を持たない",
        "・身近な子に関心を向ける",
        "・「がんばれ」より「休んでいい」",
        "・SOSを受けとめる準備を持つ",
    ]),
]
positions = [
    (Inches(0.7), Inches(1.5)),
    (Inches(6.93), Inches(1.5)),
    (Inches(0.7), Inches(4.4)),
    (Inches(6.93), Inches(4.4)),
]
for (t, color, items), (left, top) in zip(roles, positions):
    add_round_rect(s, left, top, Inches(5.7), Inches(2.7),
                   COLOR_WHITE, line=color)
    add_rect(s, left, top, Inches(5.7), Inches(0.55), color)
    add_text(s, left, top, Inches(5.7), Inches(0.55), t,
             size=18, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, left + Inches(0.3), top + Inches(0.75),
             Inches(5.2), Inches(1.9), "\n".join(items),
             size=14, color=COLOR_DARK)

# --------- スライド11: まとめ ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_WHITE)
page_header(s, "09. まとめ", 11, TOTAL)

messages = [
    ("家族を支える子どもは「えらい子」ではなく、\n支えられるべき子ども。",
     "ケアを担うことそのものを否定するのではなく、\n過度な負担から守る視点が必要。"),
    ("見つけにくいから「見つける仕組み」を。",
     "学校・地域・医療・福祉が連携し、\n大人の側からつながりに行くことが大切。"),
    ("「子どもが子どもでいられる時間」を取り戻す。",
     "勉強・遊び・友人関係・将来の選択肢——\nそのすべてを諦めずに済む社会へ。"),
]
for i, (title, desc) in enumerate(messages):
    top = Inches(1.55) + i * Inches(1.75)
    add_round_rect(s, Inches(0.7), top, Inches(12), Inches(1.55),
                   COLOR_LIGHT)
    add_rect(s, Inches(0.7), top, Inches(0.2), Inches(1.55), COLOR_ACCENT)
    add_round_rect(s, Inches(1.1), top + Inches(0.3),
                   Inches(0.95), Inches(0.95), COLOR_PRIMARY)
    add_text(s, Inches(1.1), top + Inches(0.3),
             Inches(0.95), Inches(0.95), str(i + 1),
             size=28, bold=True, color=COLOR_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(2.3), top + Inches(0.2),
             Inches(10.2), Inches(0.7), title,
             size=18, bold=True, color=COLOR_PRIMARY)
    add_text(s, Inches(2.3), top + Inches(0.85),
             Inches(10.2), Inches(0.7), desc,
             size=13, color=COLOR_DARK)

# --------- スライド12: 終わりに ---------
s = prs.slides.add_slide(blank_layout)
add_bg(s, COLOR_PRIMARY)
add_rect(s, 0, Inches(5.6), prs.slide_width, Inches(0.15), COLOR_ACCENT)

add_text(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.0),
         "Thank you", size=56, bold=True, color=COLOR_WHITE)
add_text(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(0.7),
         "ご清聴ありがとうございました。",
         size=22, color=COLOR_ACCENT)
add_text(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.5),
         "「気づくこと」から、すべてが始まります。\n"
         "あなたのまわりの子どもたちに、もう一度目を向けてみませんか。",
         size=18, color=COLOR_WHITE)
add_text(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.4),
         "References: 厚生労働省 / こども家庭庁 ヤングケアラー実態調査ほか",
         size=11, color=COLOR_LIGHT)

out = "/home/user/kaigo-ga-app/young_caregivers.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
