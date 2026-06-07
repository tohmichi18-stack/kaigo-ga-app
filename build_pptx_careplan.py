# -*- coding: utf-8 -*-
"""
ケアプラン第1表・第2表の考え方（居宅介護支援 勉強会 補足資料）
PowerPoint(.pptx) 生成スクリプト
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY   = RGBColor(0x10, 0x33, 0x5C)
BLUE   = RGBColor(0x21, 0x6A, 0x8C)
ACCENT = RGBColor(0x1B, 0x9A, 0x8B)
WARM   = RGBColor(0xD9, 0x6A, 0x4A)
GOLD   = RGBColor(0xE3, 0xB5, 0x4A)
BG     = RGBColor(0xF3, 0xF7, 0xF9)
INK    = RGBColor(0x21, 0x2E, 0x38)
MUTED  = RGBColor(0x57, 0x68, 0x73)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINE   = RGBColor(0xD4, 0xE0, 0xE6)
CARDBG = RGBColor(0xFF, 0xFF, 0xFF)
GOODBG = RGBColor(0xE6, 0xF5, 0xF2)
BADBG  = RGBColor(0xFB, 0xEC, 0xE6)
NOTEBG = RGBColor(0xFD, 0xF4, 0xE2)
FONT = "Yu Gothic"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    _set_fill(bg, BG); bg.shadow.inherit = False
    return s


def _grad(sp, c1, c2, ang):
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill'):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    grad = sp.makeelement(qn('a:gradFill'), {})
    gsLst = grad.makeelement(qn('a:gsLst'), {})
    for pos, col in ((0, c1), (100000, c2)):
        gs = grad.makeelement(qn('a:gs'), {'pos': str(pos)})
        srgb = grad.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (col[0], col[1], col[2])})
        gs.append(srgb); gsLst.append(gs)
    grad.append(gsLst)
    grad.append(grad.makeelement(qn('a:lin'), {'ang': str(ang), 'scaled': '1'}))
    ln = sp.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp.append(grad)


def grad_bg(s, c1, c2):
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background(); bg.shadow.inherit = False
    _grad(bg.fill._xPr, c1, c2, 2700000)
    return bg


def grad_fill_shape(shape, c1, c2, ang=0):
    _grad(shape.fill._xPr, c1, c2, ang)


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    return tb, tf


def para(tf, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         first=False, space_after=4, bullet=False, level=0, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after); p.level = level
    runs = text if isinstance(text, list) else [(text, color, bold)]
    for t, c, b in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = c; r.font.name = font
    _bullet(p) if bullet else _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def _bullet(p, char="・"):
    pPr = p._p.get_or_add_pPr()
    pPr.set('indent', '-228600'); pPr.set('marL', '228600')
    for tag in ('a:buChar', 'a:buAutoNum', 'a:buNone'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': FONT}))
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': char}))


def rect(s, x, y, w, h, fill=CARDBG, line=None, rounded=True, line_w=1.0):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def leftbar(s, x, y, h, color):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(8), h)
    _set_fill(bar, color); bar.shadow.inherit = False


def kicker(s, text, color=ACCENT):
    tb, tf = textbox(s, Inches(0.7), Inches(0.4), Inches(10), Inches(0.4))
    para(tf, text, size=13, color=color, bold=True, first=True)


def heading(s, text):
    leftbar(s, Inches(0.7), Inches(0.78), Inches(0.62), ACCENT)
    tb, tf = textbox(s, Inches(0.95), Inches(0.72), Inches(11.7), Inches(0.8))
    para(tf, text, size=28, color=NAVY, bold=True, first=True)


def pill(s, x, y, text, color=ACCENT, tcolor=WHITE, w=Inches(2.0)):
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.34))
    _set_fill(p, color); p.shadow.inherit = False
    try:
        p.adjustments[0] = 0.5
    except Exception:
        pass
    tf = p.text_frame; tf.word_wrap = True; tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    para(tf, text, size=12, color=tcolor, bold=True, align=PP_ALIGN.CENTER, first=True)
    return p


def card(s, x, y, w, h, title=None, bullets=None, fill=CARDBG, line=LINE,
         title_color=BLUE, body_size=15, title_size=18):
    rect(s, x, y, w, h, fill=fill, line=line)
    pad = Inches(0.22)
    tb, tf = textbox(s, x + pad, y + Inches(0.12), w - pad * 2, h - Inches(0.2))
    first = True
    if title:
        para(tf, title, size=title_size, color=title_color, bold=True, first=True, space_after=5); first = False
    if bullets:
        for b in bullets:
            para(tf, b, size=body_size, color=INK, first=first, bullet=True, space_after=3); first = False
    return tf


def note(s, x, y, w, h, runs, bar_color=GOLD, fill=NOTEBG):
    rect(s, x, y, w, h, fill=fill, line=None)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Pt(6), h)
    _set_fill(bar, bar_color); bar.shadow.inherit = False
    tb, tf = textbox(s, x + Inches(0.2), y, w - Inches(0.35), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, runs, size=14, first=True)


TOTAL = 13
_n = [0]


def footer(s, n):
    tb, tf = textbox(s, Inches(11.6), Inches(7.05), Inches(1.5), Inches(0.35))
    para(tf, "%d / %d" % (n, TOTAL), size=11, color=MUTED, align=PP_ALIGN.RIGHT, first=True)
    tb2, tf2 = textbox(s, Inches(0.7), Inches(7.05), Inches(9), Inches(0.35))
    para(tf2, "居宅介護支援 勉強会 補足｜ケアプラン第1表・第2表の考え方", size=10, color=MUTED, first=True)


def newslide(kick=None, title=None, grad=None):
    _n[0] += 1
    if grad:
        s = prs.slides.add_slide(BLANK); grad_bg(s, grad[0], grad[1])
    else:
        s = slide()
        if kick:
            kicker(s, kick)
        if title:
            heading(s, title)
        footer(s, _n[0])
    return s


# ======================================================================
# 1 タイトル
# ======================================================================
s = newslide(grad=(NAVY, ACCENT))
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.5), Inches(6.4), Inches(0.5))
b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x21, 0x55, 0x76)
b.line.color.rgb = RGBColor(0x8F, 0xB6, 0xD4); b.line.width = Pt(1); b.shadow.inherit = False
para(b.text_frame, "居宅介護支援 勉強会 補足資料（ケアマネジメント深掘り編）", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.0))
para(tf, "ケアプラン第1表・第2表の考え方", size=44, color=WHITE, bold=True, first=True, space_after=8)
para(tf, "～ 意向から始まり、ニーズ・目標・サービスへ一本の線でつなぐ ～", size=20, color=GOLD, bold=True)
tb2, tf2 = textbox(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.0))
para(tf2, "対象：介護支援専門員・主任介護支援専門員　／　困難事例にも通じるケアプランの基本を再確認",
     size=15, color=RGBColor(0xCF, 0xE0, 0xEE), first=True)

# ======================================================================
# 2 ケアプラン全体像
# ======================================================================
s = newslide("OVERVIEW", "居宅サービス計画書の全体像と第1表・第2表の役割")
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5))
para(tf, "居宅サービス計画書は複数の帳票で構成される。中でも第1表・第2表が計画の「背骨」。", size=16, color=MUTED, first=True)
rows = [("第1表", "居宅サービス計画書(1)", "意向・方針・全体像", "本人/家族の意向、総合的な援助の方針、認定情報、同意"),
        ("第2表", "居宅サービス計画書(2)", "課題・目標・サービス", "生活全般の解決すべき課題（ニーズ）、長期/短期目標、援助内容"),
        ("第3表", "週間サービス計画表", "1週間の組み立て", "曜日・時間ごとのサービスと主な日常生活上の活動"),
        ("他", "第4〜7表ほか", "会議・経過・給付管理", "担当者会議記録、支援経過、利用票・別表 等")]
y = Inches(2.05); rh = Inches(1.0)
for i, (no, name, role, detail) in enumerate(rows):
    yy = y + i * (rh + Inches(0.08))
    highlight = i < 2
    rect(s, Inches(0.7), yy, Inches(11.9), rh, fill=(GOODBG if highlight else WHITE), line=LINE)
    tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), yy + Inches(0.27), Inches(1.2), Inches(0.46))
    _set_fill(tag, ACCENT if highlight else MUTED); tag.shadow.inherit = False
    para(tag.text_frame, no, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, Inches(2.3), yy + Inches(0.1), Inches(3.0), rh - Inches(0.2), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, name, size=15, color=NAVY, bold=True, first=True, space_after=1)
    para(tf, role, size=12, color=BLUE, bold=True)
    tb2, tf2 = textbox(s, Inches(5.4), yy + Inches(0.1), Inches(7.0), rh - Inches(0.2), anchor=MSO_ANCHOR.MIDDLE)
    para(tf2, detail, size=13, color=INK, first=True)

# ======================================================================
# 3 一貫性の原則（流れ図）
# ======================================================================
s = newslide("PRINCIPLE", "最重要原則：第1表から第2表へ「一本の線」でつなぐ")
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5))
para(tf, [("良いケアプランは各欄がバラバラでなく、", INK, False), ("意向 → 方針 → ニーズ → 目標 → サービスが論理的につながっている。", WARM, True)], size=16, first=True)
steps = [("意向", "第1表\n本人・家族の望む暮らし", BLUE),
         ("方針", "第1表\n総合的な援助の方針", BLUE),
         ("ニーズ", "第2表\n解決すべき課題", ACCENT),
         ("目標", "第2表\n長期・短期目標", ACCENT),
         ("サービス", "第2表\n援助内容", ACCENT)]
n = len(steps); total_w = Inches(11.9)
sw = Inches(2.0); gap = (total_w - sw * n) / (n - 1)
x0, y0, h = Inches(0.7), Inches(2.4), Inches(1.7)
for i, (t, sub, col) in enumerate(steps):
    x = x0 + i * (sw + gap)
    rect(s, x, y0, sw, h, fill=WHITE, line=col, line_w=2)
    hd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.15), y0 + Inches(0.15), sw - Inches(0.3), Inches(0.5))
    _set_fill(hd, col); hd.shadow.inherit = False
    para(hd.text_frame, t, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, x + Inches(0.1), y0 + Inches(0.75), sw - Inches(0.2), h - Inches(0.85))
    para(tf, sub, size=12, color=INK, align=PP_ALIGN.CENTER, first=True)
    if i < n - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + sw - Inches(0.02), y0 + h/2 - Inches(0.11), gap + Inches(0.04), Inches(0.22))
        _set_fill(ar, GOLD); ar.shadow.inherit = False
note(s, Inches(0.7), Inches(4.5), Inches(11.9), Inches(1.4),
     [("点検の問い：", ACCENT, True),
      ("「このサービスは、どの目標のため？ その目標はどのニーズを解決する？ そのニーズは本人の意向とつながっている？」を逆向きにたどれるか。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 4 第1表の構成
# ======================================================================
s = newslide("FORM 1", "第1表の構成と各欄の考え方")
cards = [("利用者・家族の生活に対する意向", "本人と家族、それぞれの「こうありたい暮らし」。サービス名でなく生活の言葉で。"),
         ("総合的な援助の方針", "チーム全員の羅針盤。誰が・何を目指し・緊急時はどうするかを共有。"),
         ("介護認定審査会の意見／\nサービスの種類の指定", "意見が付されていれば計画に反映。指定がある場合はその範囲で。"),
         ("課題分析の結果・同意等", "アセスメント結果の総括。利用者の同意・説明日・交付を確実に。")]
cw, ch, gx, gy = Inches(5.85), Inches(2.0), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.7)
for i, (t, b) in enumerate(cards):
    x = x0 + (i % 2) * (cw + gx); y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    tb, tf = textbox(s, x + Inches(0.25), y + Inches(0.18), cw - Inches(0.5), ch - Inches(0.3))
    para(tf, t, size=17, color=BLUE, bold=True, first=True, space_after=5)
    para(tf, b, size=14, color=INK)

# ======================================================================
# 5 第1表① 意向の書き方
# ======================================================================
s = newslide("FORM 1", "第1表①　利用者・家族の意向の考え方")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(2.9), title="押さえる視点",
     bullets=["本人の言葉・思いを主語にする（本人と家族は分けて）", "「〜したい」という生活・人生レベルの願い",
              "サービス利用が前提の表現にしない", "意思表示が難しい人は意思決定支援の視点で代弁・推定",
              "ここがケアプラン全体の出発点になる"], title_color=BLUE, body_size=14)
x1 = Inches(0.7) + Inches(5.85) + Inches(0.3)
rect(s, x1, Inches(1.7), Inches(5.85), Inches(2.9), fill=WHITE, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), Inches(1.85), Inches(5.35), Inches(2.6))
para(tf, "書きぶりの例", size=17, color=BLUE, bold=True, first=True, space_after=6)
para(tf, [("△ ", WARM, True), ("「デイサービスを週2回利用したい」", INK, False)], size=14, space_after=2)
para(tf, "（＝サービスありき。意向ではなく手段）", size=12, color=MUTED, space_after=8)
para(tf, [("○ ", ACCENT, True), ("本人：「転ばないように歩けるようになり、また畑仕事をしたい」", INK, False)], size=14, space_after=2)
para(tf, [("○ ", ACCENT, True), ("家族：「日中も安心して仕事に行けるようにしたい」", INK, False)], size=14)
note(s, Inches(0.7), Inches(4.85), Inches(11.9), Inches(1.0),
     [("ポイント：", GOLD, True), ("意向は第2表のニーズ・目標へ直結する。ここがサービス名で埋まると、計画全体が「サービス調整表」になってしまう。", INK, False)])

# ======================================================================
# 6 第1表② 総合的な援助の方針
# ======================================================================
s = newslide("FORM 1", "第1表②　総合的な援助の方針の考え方")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(2.9), title="この欄の役割＝チームの羅針盤",
     bullets=["利用者像と支援の方向性をチームで共有", "誰が中心に、何を目指して関わるか", "本人の強み・本人にできることも明記",
              "困難事例ほど方針の言語化が支援をそろえる"], title_color=BLUE, body_size=14)
x1 = Inches(0.7) + Inches(5.85) + Inches(0.3)
card(s, x1, Inches(1.7), Inches(5.85), Inches(2.9), title="必ず盛り込みたい要素",
     bullets=["緊急時・急変時の連絡体制と対応方針", "医療との連携（主治医・訪問看護の役割）",
              "想定されるリスクと予防の方針", "権利擁護・意思決定支援の視点（必要時）"], title_color=ACCENT, body_size=14)
note(s, Inches(0.7), Inches(4.85), Inches(11.9), Inches(1.0),
     [("困難事例での要点：", ACCENT, True),
      ("支援拒否・虐待・看取り等では、緊急時対応と多職種の役割をこの欄に明記。方針が共有されていれば、担当が代わっても支援がぶれない。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 7 第2表の構成
# ======================================================================
s = newslide("FORM 2", "第2表の構成と各欄の考え方")
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5))
para(tf, "第2表は「課題（ニーズ）→ 目標 →（援助内容）」の対応関係が一目で読めることが命。", size=16, color=MUTED, first=True)
rows = [("生活全般の解決すべき課題（ニーズ）", "本人が望む生活と現状のギャップ。本人主体・ICFの視点で。", ACCENT),
        ("長期目標／短期目標（と期間）", "ニーズが解決・改善した状態。評価できる具体性と期間を。", BLUE),
        ("援助内容（サービス内容・種別・頻度・期間）", "目標達成の手段。保険給付外・インフォーマルも記載。", BLUE)]
y = Inches(2.05); rh = Inches(1.25)
for i, (t, b, col) in enumerate(rows):
    yy = y + i * (rh + Inches(0.12))
    rect(s, Inches(0.7), yy, Inches(11.9), rh, fill=WHITE, line=LINE)
    leftbar(s, Inches(0.7), yy, rh, col)
    tag = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), yy + rh/2 - Inches(0.25), Inches(0.5), Inches(0.5))
    _set_fill(tag, col); tag.shadow.inherit = False
    para(tag.text_frame, str(i + 1), size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, Inches(1.7), yy + Inches(0.12), Inches(10.6), rh - Inches(0.24), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, t, size=17, color=NAVY, bold=True, first=True, space_after=3)
    para(tf, b, size=14, color=INK)

# ======================================================================
# 8 第2表① ニーズ
# ======================================================================
s = newslide("FORM 2", "第2表①　生活全般の解決すべき課題（ニーズ）")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(2.95), title="ニーズの捉え方",
     bullets=["「望む生活」と「現状」のギャップが課題", "心身機能だけでなく活動・参加・環境（ICF）で捉える",
              "本人の困りごと・本人の言葉を起点に", "できないこと探しでなく、強み・できることも見る",
              "原因→結果の構造で整理（なぜ困っているか）"], title_color=ACCENT, body_size=14)
x1 = Inches(0.7) + Inches(5.85) + Inches(0.3)
rect(s, x1, Inches(1.7), Inches(5.85), Inches(2.95), fill=WHITE, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), Inches(1.85), Inches(5.35), Inches(2.7))
para(tf, "書きぶりの例", size=17, color=BLUE, bold=True, first=True, space_after=6)
para(tf, [("△ ", WARM, True), ("「歩行が不安定」（＝状態・課題分析の断片）", INK, False)], size=14, space_after=8)
para(tf, [("○ ", ACCENT, True), ("「転倒の不安なく歩けるようになり、自分で買い物に行けるようになりたい」", INK, False)], size=14, space_after=6)
para(tf, "→ 望む生活が見え、目標・サービスが導ける", size=12, color=MUTED)
note(s, Inches(0.7), Inches(4.9), Inches(11.9), Inches(0.95),
     [("注意：", GOLD, True), ("ニーズを「サービスが必要」と書かない（例：×訪問介護が必要）。それは手段。課題は本人の生活の言葉で。", INK, False)])

# ======================================================================
# 9 第2表② 目標
# ======================================================================
s = newslide("FORM 2", "第2表②　長期目標・短期目標と期間")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(2.95), title="目標設定の考え方",
     bullets=["長期目標＝ニーズが解決・改善した状態（ゴール）", "短期目標＝長期目標への段階・足がかり",
              "本人が主語／達成度を評価できる表現で", "「期間」は認定有効期間等もふまえ現実的に",
              "本人が見て「やってみよう」と思える言葉に"], title_color=BLUE, body_size=14)
x1 = Inches(0.7) + Inches(5.85) + Inches(0.3)
rect(s, x1, Inches(1.7), Inches(5.85), Inches(2.95), fill=WHITE, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), Inches(1.85), Inches(5.35), Inches(2.7))
para(tf, "長期 → 短期の例", size=17, color=BLUE, bold=True, first=True, space_after=6)
para(tf, [("長期：", BLUE, True), ("「一人で近所のスーパーまで買い物に行ける」（6か月）", INK, False)], size=14, space_after=8)
para(tf, [("短期：", ACCENT, True), ("「見守りがあれば家の中を伝い歩きできる」（2か月）", INK, False)], size=14, space_after=4)
para(tf, [("短期：", ACCENT, True), ("「下肢筋力が向上し、玄関の段差をまたげる」（3か月）", INK, False)], size=14)
note(s, Inches(0.7), Inches(4.9), Inches(11.9), Inches(0.95),
     [("ポイント：", GOLD, True), ("「〜の支援を受ける」は目標でなくサービス。目標は“本人がどうなるか”。評価日に達成を判定できる粒度にする。", INK, False)])

# ======================================================================
# 10 第2表③ 援助内容
# ======================================================================
s = newslide("FORM 2", "第2表③　援助内容（サービス内容・種別・頻度・期間）")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(2.95), title="記載の考え方",
     bullets=["短期目標を達成するための具体的な手段", "「サービス内容」は何をするかを具体的に",
              "保険給付対象か（※印）を区別", "種別・事業所・頻度・期間を明確に",
              "家族・近隣・インフォーマル資源も書く"], title_color=BLUE, body_size=14)
x1 = Inches(0.7) + Inches(5.85) + Inches(0.3)
rect(s, x1, Inches(1.7), Inches(5.85), Inches(2.95), fill=WHITE, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), Inches(1.85), Inches(5.35), Inches(2.7))
para(tf, "考え方のコツ", size=17, color=BLUE, bold=True, first=True, space_after=6)
para(tf, [("・", ACCENT, True), ("「短期目標」と援助内容が1対1で対応しているか", INK, False)], size=14, space_after=4)
para(tf, [("・", ACCENT, True), ("フォーマルサービスだけで埋めない（自助・互助も）", INK, False)], size=14, space_after=4)
para(tf, [("・", ACCENT, True), ("本人・家族が担う役割も明記し、自立を支える", INK, False)], size=14, space_after=4)
para(tf, [("・", ACCENT, True), ("頻度・期間はモニタリングで見直す前提で設定", INK, False)], size=14)
note(s, Inches(0.7), Inches(4.9), Inches(11.9), Inches(0.95),
     [("一貫性の確認：", ACCENT, True), ("援助内容 →（達成する）短期目標 →（解決する）ニーズ →（実現する）意向、と下から上へ説明できればOK。", INK, False)],
     bar_color=ACCENT, fill=GOODBG)

# ======================================================================
# 11 Before/After 通し例
# ======================================================================
s = newslide("EXAMPLE", "通し事例：第1表・第2表を一本の線でつなぐ")
tb, tf = textbox(s, Inches(0.7), Inches(1.4), Inches(12), Inches(0.5))
para(tf, "Aさん（要介護2・脳梗塞後の右麻痺・独居）。「また自分で料理を作りたい」という思い。", size=15, color=MUTED, first=True)
data = [("第1表 意向", "本人：また自分で簡単な料理を作り、自分の生活を続けたい", BLUE),
        ("第1表 方針", "残存機能を活かし、できることを増やしながら独居生活を支える。緊急時は長女と訪問看護に連絡", BLUE),
        ("第2表 ニーズ", "右手が使いにくく台所に立つのが不安。安全に調理を再開して自分らしく暮らしたい", ACCENT),
        ("第2表 目標", "短期：手すりと自助具を使い座って調理ができる（3か月）／長期：見守りで一食分を作れる（6か月）", ACCENT),
        ("第2表 援助内容", "訪問リハ（調理動作訓練・週1）／福祉用具（自助具）／訪問介護（共に調理・見守り）／長女の声かけ", GOLD)]
y = Inches(1.95); rh = Inches(0.86)
for i, (label, body, col) in enumerate(data):
    yy = y + i * (rh + Inches(0.06))
    rect(s, Inches(0.7), yy, Inches(11.9), rh, fill=WHITE, line=LINE)
    tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), yy + rh/2 - Inches(0.22), Inches(2.3), Inches(0.44))
    _set_fill(tag, col); tag.shadow.inherit = False
    para(tag.text_frame, label, size=13, color=(RGBColor(0x5A, 0x45, 0x00) if col == GOLD else WHITE), bold=True, align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(s, Inches(3.4), yy + Inches(0.05), Inches(9.0), rh - Inches(0.1), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, body, size=13.5, color=INK, first=True)
    if i < len(data) - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(1.95), yy + rh - Inches(0.02), Inches(0.2), Inches(0.1))
        _set_fill(ar, GOLD); ar.shadow.inherit = False

# ======================================================================
# 12 困難事例での第1表・第2表
# ======================================================================
s = newslide("DIFFICULT CASE", "困難事例における第1表・第2表のポイント")
cards = [("支援拒否のあるケース", "意向は「拒否」でなく背景にある思いを探る。方針に関係構築の段階と緊急時を明記。目標は小さく刻む。", WARM),
         ("認知症・意思決定支援", "意思の推定・代弁の根拠を残す。本人の言葉を可能な限り記載。方針に権利擁護の視点を。", ACCENT),
         ("虐待・複合課題世帯", "本人と家族の意向を分けて記載。方針に多機関連携・通報体制。家族支援もニーズに位置づける。", WARM),
         ("看取り・医療依存", "本人の意向（ACP）を中心に。方針に急変時対応・医療連携を明確化。目標・期間は状態変化に応じ柔軟に。", ACCENT)]
cw, ch, gx, gy = Inches(5.85), Inches(1.85), Inches(0.3), Inches(0.3)
x0, y0 = Inches(0.7), Inches(1.7)
for i, (t, b, col) in enumerate(cards):
    x = x0 + (i % 2) * (cw + gx); y = y0 + (i // 2) * (ch + gy)
    rect(s, x, y, cw, ch, fill=CARDBG, line=LINE)
    leftbar(s, x, y, ch, col)
    tb, tf = textbox(s, x + Inches(0.3), y + Inches(0.18), cw - Inches(0.55), ch - Inches(0.3))
    para(tf, t, size=16, color=NAVY, bold=True, first=True, space_after=4)
    para(tf, b, size=13.5, color=INK)
note(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.95),
     [("共通：", GOLD, True), ("困難事例ほど、第1表「総合的な援助の方針」に緊急時・連携・権利擁護を書き、チームと多職種で方針を共有することが要。", INK, False)])

# ======================================================================
# 13 まとめ・チェック
# ======================================================================
s = newslide("SUMMARY", "まとめ：ケアプラン点検の視点")
card(s, Inches(0.7), Inches(1.7), Inches(5.85), Inches(3.6), title="セルフチェック",
     bullets=["意向がサービス名でなく生活の言葉になっているか", "意向→方針→ニーズ→目標→サービスが一本でつながるか",
              "ニーズが本人主体・ICFの視点で書けているか", "目標は本人が主語で、評価できる具体性があるか",
              "援助内容にインフォーマル資源も入っているか", "困難事例で方針に緊急時・連携・権利擁護があるか"],
     title_color=BLUE, body_size=14)
x1 = Inches(0.7) + Inches(5.85) + Inches(0.3)
rect(s, x1, Inches(1.7), Inches(5.85), Inches(3.6), fill=CARDBG, line=LINE)
tb, tf = textbox(s, x1 + Inches(0.25), Inches(1.9), Inches(5.35), Inches(3.3))
para(tf, "キーメッセージ", size=18, color=BLUE, bold=True, first=True, space_after=8)
para(tf, [("ケアプランは「サービス調整表」ではなく", INK, False), ("本人の望む暮らしの設計図", WARM, True), ("。", INK, False)], size=15, space_after=10)
para(tf, "第1表で“どこを目指すか”を共有し、第2表で“どう実現するか”を具体化する。", size=15, space_after=10)
para(tf, "この一貫性こそ、困難事例でもチームがぶれずに支援できる土台になる。", size=15)
note(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.85),
     [("※", MUTED, True), ("帳票様式・記載要領の詳細は、最新の「介護サービス計画書の様式及び課題分析標準項目」等の通知でご確認ください。", MUTED, False)],
     bar_color=MUTED, fill=RGBColor(0xEE, 0xF1, 0xF3))

# ----------------------------------------------------------------------
out = "ケアプラン第1表・第2表の考え方_勉強会補足.pptx"
prs.save(out)
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst))
